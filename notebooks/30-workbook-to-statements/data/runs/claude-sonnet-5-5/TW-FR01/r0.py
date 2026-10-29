import openpyxl


def _norm(s):
    return " ".join(str(s).replace("\u2019", "'").replace("\xa0", " ").split()).strip().lower()


def _num(v):
    if isinstance(v, bool) or v is None:
        raise ValueError("not a number")
    if isinstance(v, (int, float)):
        return v
    t = str(v)
    for ch in (" ", "\xa0", "\u202f", "\u2009"):
        t = t.replace(ch, "")
    t = t.replace("\u2212", "-")
    x = float(t)
    return int(x) if x == int(x) else x


def extract(path):
    wb = openpyxl.load_workbook(path, data_only=True)

    CR, ACT, PAS, TFT = "CR", "Bilan actif", "Bilan passif", "TFT"
    cols = {CR: ("B", "C"), ACT: ("D", "E"), PAS: ("B", "C"), TFT: ("B", "C")}

    rowcache = {}

    def row_of(sheet, label):
        if sheet not in rowcache:
            m = {}
            ws = wb[sheet]
            for r in range(1, ws.max_row + 1):
                v = ws["A%d" % r].value
                if v is not None:
                    m.setdefault(_norm(v), r)
            rowcache[sheet] = m
        return rowcache[sheet][_norm(label)]

    def cell_ref(sheet, label, year, factor=1):
        col = cols[sheet][0 if year == "cy" else 1]
        return {"sheet": sheet, "cell": "%s%d" % (col, row_of(sheet, label)), "factor": factor}

    def line(parts, year):
        cells = [cell_ref(s, l, year, f) for (s, l, f) in parts]
        total = 0
        for c in cells:
            total += c["factor"] * _num(wb[c["sheet"]][c["cell"]].value)
        return {"value": total, "cells": cells}

    spec = {
        "revenue": [(CR, "Chiffre d'affaires net", 1)],
        "cost_of_sales": [(CR, "Coût des ventes", 1)],
        "distribution_costs": [(CR, "Frais commerciaux", 1)],
        "administrative_expenses": [(CR, "Frais administratifs", 1)],
        "other_operating_income": [(CR, "Autres produits d'exploitation", 1),
                                   (CR, "Autres charges d'exploitation", -1)],
        "finance_income": [(CR, "Produits financiers", 1)],
        "finance_costs": [(CR, "Charges financières", 1)],
        "income_tax_expense": [(CR, "Impôts sur les bénéfices", 1)],
        "property_plant_equipment": [(ACT, "Immobilisations corporelles", 1)],
        "right_of_use_assets": [(ACT, "Droits d'utilisation (contrats de location)", 1)],
        "intangible_assets": [(ACT, "Immobilisations incorporelles", 1)],
        "investments": [],
        "deferred_tax_assets": [(ACT, "Impôts différés actifs", 1)],
        "inventories": [(ACT, "Stocks de marchandises", 1)],
        "trade_receivables": [(ACT, "Créances clients et comptes rattachés", 1)],
        "other_current_assets": [(ACT, "Autres créances et charges constatées d'avance", 1)],
        "cash_and_cash_equivalents": [(ACT, "Disponibilités", 1)],
        "trade_payables": [(PAS, "Dettes fournisseurs et comptes rattachés", 1)],
        "other_current_liabilities": [(PAS, "Dettes fiscales et sociales (hors IS)", 1),
                                      (PAS, "Autres dettes", 1),
                                      (PAS, "Produits constatés d'avance", 1)],
        "current_borrowings": [(PAS, "Emprunts auprès des établissements de crédit (à moins d'un an)", 1),
                               (PAS, "Dettes locatives (à moins d'un an)", 1)],
        "current_tax_payable": [(PAS, "Impôt sur les sociétés à payer", 1)],
        "non_current_borrowings": [(PAS, "Emprunts auprès des établissements de crédit (à plus d'un an)", 1),
                                   (PAS, "Dettes locatives (à plus d'un an)", 1)],
        "provisions": [(PAS, "Provisions pour risques et charges", 1)],
        "deferred_tax_liabilities": [(PAS, "Impôts différés passifs", 1)],
        "share_capital": [(PAS, "Capital social", 1)],
        "retained_earnings": [(PAS, "Report à nouveau", 1), (PAS, "Résultat de l'exercice", 1)],
        "other_reserves": [(PAS, "Autres réserves", 1)],
        "net_cash_operating": [(TFT, "Flux net de trésorerie généré par l'activité", 1)],
        "net_cash_investing": [(TFT, "Flux net de trésorerie lié aux opérations d'investissement", 1)],
        "net_cash_financing": [(TFT, "Flux net de trésorerie lié aux opérations de financement", 1)],
        "fx_effect_on_cash": [(TFT, "Incidence des variations de cours des devises", 1)],
        "cash_at_beginning": [(TFT, "Trésorerie d'ouverture", 1)],
        "dividends": [(TFT, "Dividendes versés", -1)],
    }

    fields = {}
    for key, parts in spec.items():
        if parts:
            fields[key] = {"cy": line(parts, "cy"), "py": line(parts, "py")}
        else:
            fields[key] = {"cy": {"value": 0, "cells": []}, "py": {"value": 0, "cells": []}}

    source_totals = []

    def tot(sheet, label, formula=None, sum_of=None):
        for year in ("cy", "py"):
            c = cell_ref(sheet, label, year)
            d = {"sheet": c["sheet"], "cell": c["cell"], "year": year, "factor": 1}
            if formula is not None:
                d["formula"] = formula
            else:
                d["sum_of"] = [cell_ref(s, l, year, 1) for (s, l) in sum_of]
            source_totals.append(d)

    tot(CR, "Marge brute", "gross_profit")
    tot(CR, "Résultat d'exploitation", "operating_profit")
    tot(CR, "Résultat financier", "finance_income - finance_costs")
    tot(CR, "Résultat courant avant impôts", "profit_before_tax")
    tot(CR, "Résultat net de l'exercice", "profit_for_year")
    tot(ACT, "Total actif immobilisé", "total_non_current_assets")
    tot(ACT, "Total actif circulant", "total_current_assets")
    tot(ACT, "TOTAL ACTIF", "total_assets")
    tot(PAS, "Total capitaux propres", "total_equity")
    tot(PAS, "Total dettes", "total_liabilities - provisions")
    tot(PAS, "TOTAL PASSIF", "total_equity_and_liabilities")
    tot(TFT, "Variation de trésorerie", "net_change_in_cash")
    tot(TFT, "Trésorerie de clôture", "cash_at_end")
    tot(TFT, "Report à nouveau à la clôture", sum_of=[(PAS, "Report à nouveau")])

    r_first = row_of(ACT, "Immobilisations incorporelles")
    r_last = row_of(ACT, "Disponibilités")
    r_op = row_of(TFT, "Report à nouveau à l'ouverture")
    r_pr = row_of(TFT, "Résultat de l'exercice précédent affecté")

    ignored = [
        {"sheet": ACT, "range": "B%d:C%d" % (r_first, r_last),
         "reason": "Gross value and accumulated depreciation columns for 2025; only the Net columns feed the template."},
        {"sheet": TFT, "range": "B%d:C%d" % (r_op, r_pr),
         "reason": "Appropriation-of-result roll-forward (opening retained earnings and prior-year result allocated); "
                   "already captured by the balance sheet retained earnings and profit, only the dividends line and closing balance are used."},
    ]

    name = wb[CR]["A1"].value

    return {
        "entity_code": "TW-FR01",
        "source_entity_name": str(name).strip(),
        "currency": "EUR",
        "unit": 1,
        "period_end": {"cy": "2025-12-31", "py": "2024-12-31"},
        "fields": fields,
        "source_totals": source_totals,
        "ignored": ignored,
        "source_exceptions": [],
        "notes": (
            "Amounts are stored as text with spaces as thousands separators and are parsed in code. "
            "Asset lines use the Net columns (D for 2025, E for 2024). Other operating income is other operating income "
            "less other operating charges. Other current liabilities combine taxes and social charges (excl. income tax), "
            "other payables and deferred income. Borrowings and lease liabilities are split current / non-current. "
            "Retained earnings are carried forward plus profit for the year. Dividends come from the TFT appropriation table "
            "with the sign reversed. The workbook's 'Total dettes' excludes provisions, so it is declared as total_liabilities - provisions. "
            "All printed totals tie; no source exceptions."
        ),
    }
