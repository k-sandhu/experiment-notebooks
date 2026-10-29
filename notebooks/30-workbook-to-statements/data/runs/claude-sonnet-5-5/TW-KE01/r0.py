import re
from datetime import datetime
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


def _norm(s):
    return re.sub(r"\s+", " ", str(s)).strip().lower()


def extract(path):
    wb = load_workbook(path, data_only=True)

    def rows_of(ws):
        d = {}
        for r in range(1, ws.max_row + 1):
            v = ws.cell(r, 1).value
            if isinstance(v, str):
                d.setdefault(_norm(v), r)
        return d

    rows = {n: rows_of(wb[n]) for n in ("IS", "SOFP", "SCF")}

    def row(sheet, label):
        return rows[sheet][_norm(label)]

    # Year columns from the header row of the IS sheet
    ws_is = wb["IS"]
    hdr = {}
    for c in range(2, ws_is.max_column + 1):
        v = ws_is.cell(5, c).value
        m = re.search(r"(20\d\d)", str(v)) if v is not None else None
        if m:
            hdr[int(m.group(1))] = c
    cy_year = max(hdr)
    col = {"cy": get_column_letter(hdr[cy_year]), "py": get_column_letter(hdr[cy_year - 1])}

    head = str(ws_is.cell(5, hdr[cy_year]).value)
    unit = 1000 if "000" in head else (1000000 if re.search(r"m(n|illion)?\b", head, re.I) else 1)

    # Period end from the title line of the balance sheet
    title = str(wb["SOFP"]["A3"].value)
    m = re.search(r"(\d{1,2}) ([A-Za-z]+) (\d{4})", title)
    d_cy = datetime.strptime(" ".join(m.groups()), "%d %B %Y")
    d_py = d_cy.replace(year=d_cy.year - 1)
    period_end = {"cy": d_cy.strftime("%Y-%m-%d"), "py": d_py.strftime("%Y-%m-%d")}

    def val(sheet, cell):
        v = wb[sheet][cell].value
        return v

    def cellref(sheet, label, yr):
        return "%s%d" % (col[yr], row(sheet, label))

    # key -> list of (sheet, label, factor)
    spec = {
        "revenue": [("IS", "Sales", 1)],
        "cost_of_sales": [("IS", "Cost of goods sold", 1)],
        "distribution_costs": [("IS", "Selling and distribution expenses", 1)],
        "administrative_expenses": [
            ("IS", "Administrative salaries", 1),
            ("IS", "Office premises costs", 1),
            ("IS", "Other administrative expenses", 1),
        ],
        "other_operating_income": [("IS", "Other gains and losses (net)", 1)],
        "finance_income": [("IS", "Interest received", 1)],
        "finance_costs": [("IS", "Finance costs", 1)],
        "income_tax_expense": [("IS", "Tax on profit", 1)],
        "property_plant_equipment": [("SOFP", "Tangible assets", 1)],
        "right_of_use_assets": [("SOFP", "Leased assets (right of use)", 1)],
        "intangible_assets": [("SOFP", "Intangible assets", 1)],
        "investments": [],
        "deferred_tax_assets": [("SOFP", "Deferred tax assets", 1)],
        "inventories": [("SOFP", "Inventories", 1)],
        "trade_receivables": [("SOFP", "Trade debtors", 1)],
        "other_current_assets": [("SOFP", "Other receivables and prepayments", 1)],
        "cash_and_cash_equivalents": [("SOFP", "Cash", 1)],
        "trade_payables": [("SOFP", "Accounts payable", 1)],
        "other_current_liabilities": [("SOFP", "Accruals and other payables", 1)],
        "current_borrowings": [("SOFP", "Short-term borrowings", 1)],
        "current_tax_payable": [("SOFP", "Income tax payable", 1)],
        "non_current_borrowings": [("SOFP", "Loans due after one year", 1)],
        "provisions": [("SOFP", "Provisions", 1)],
        "deferred_tax_liabilities": [("SOFP", "Deferred tax liabilities", 1)],
        "share_capital": [("SOFP", "Share capital", 1)],
        "retained_earnings": [("SOFP", "Accumulated profits", 1)],
        "other_reserves": [("SOFP", "Other reserves", 1)],
        "net_cash_operating": [("SCF", "Net cash from operating activities", 1)],
        "net_cash_investing": [("SCF", "Net cash used in investing activities", 1)],
        "net_cash_financing": [("SCF", "Net cash from financing activities", 1)],
        "fx_effect_on_cash": [("SCF", "Foreign exchange differences on cash", 1)],
        "cash_at_beginning": [("SCF", "Cash at start of year", 1)],
        "dividends": [("SOFP", "Dividends paid to shareholders", -1)],
    }

    fields = {}
    for key, parts in spec.items():
        fields[key] = {}
        for yr in ("cy", "py"):
            cells = []
            total = 0
            for sheet, label, f in parts:
                ref = cellref(sheet, label, yr)
                cells.append({"sheet": sheet, "cell": ref, "factor": f})
                total += f * val(sheet, ref)
            fields[key][yr] = {"value": total, "cells": cells}

    # Printed totals: (sheet, label, formula)
    tot_spec = [
        ("IS", "Gross profit", "gross_profit"),
        ("IS", "Total administrative expenses", "administrative_expenses"),
        ("IS", "Profit from operations", "operating_profit"),
        ("IS", "Profit before tax", "profit_before_tax"),
        ("IS", "Net profit", "profit_for_year"),
        ("SOFP", "Total non-current assets", "total_non_current_assets"),
        ("SOFP", "Total current assets", "total_current_assets"),
        ("SOFP", "TOTAL ASSETS", "total_assets"),
        ("SOFP", "Total equity", "total_equity"),
        ("SOFP", "Total non-current liabilities", "total_non_current_liabilities"),
        ("SOFP", "Total current liabilities", "total_current_liabilities"),
        ("SOFP", "Total liabilities", "total_liabilities"),
        ("SOFP", "TOTAL EQUITY AND LIABILITIES", "total_equity_and_liabilities"),
        ("SOFP", "Profit for the year", "profit_for_year"),
        ("SOFP", "Closing retained earnings", "retained_earnings"),
        ("SCF", "Net change in cash and cash equivalents", "net_change_in_cash"),
        ("SCF", "Cash at end of year", "cash_at_end"),
    ]
    source_totals = []
    for sheet, label, formula in tot_spec:
        for yr in ("cy", "py"):
            source_totals.append({
                "sheet": sheet,
                "cell": cellref(sheet, label, yr),
                "year": yr,
                "factor": 1,
                "formula": formula,
            })

    # Opening retained earnings is last year's closing balance, already used from the balance sheet
    ignored = []
    for yr in ("cy", "py"):
        ignored.append({
            "sheet": "SOFP",
            "cell": cellref("SOFP", "Opening retained earnings", yr),
            "reason": "Opening retained earnings in the reconciliation; equals the prior-year closing "
                      "balance already taken from the balance sheet (and not a template input).",
        })

    # Detect printed totals that do not agree with their components
    source_exceptions = []
    for yr in ("cy", "py"):
        comp = sum(fields["administrative_expenses"][yr]["value"] for _ in [0])
        ref = cellref("IS", "Total administrative expenses", yr)
        printed = val("IS", ref)
        if printed != comp:
            source_exceptions.append({
                "check": "T IS!%s %s" % (ref, yr),
                "difference": printed - comp,
                "explanation": (
                    "Printed 'Total administrative expenses' differs from the sum of the three "
                    "administrative lines above it. The component lines are consistent with the printed "
                    "'Profit from operations', so the components were used and the printed subtotal "
                    "left as it is."
                ),
            })

    return {
        "entity_code": "TW-KE01",
        "source_entity_name": str(ws_is["A1"].value),
        "currency": "KES",
        "unit": unit,
        "period_end": period_end,
        "fields": fields,
        "source_totals": source_totals,
        "ignored": ignored,
        "source_exceptions": source_exceptions,
        "notes": (
            "The workbook is for Tarnwell E.A. Ltd, in KShs'000, matched to TW-KE01 (Tarnwell East Africa "
            "Limited, Kenya, KES). Administrative expenses are the sum of the three detail lines; the printed "
            "2025 subtotal is 2,700 higher than those lines, while operating profit ties to the detail, so it is "
            "declared as a source exception. Dividends come from 'Dividends paid to shareholders' in the "
            "retained earnings reconciliation (sign flipped); no separate declared figure is given. "
            "Short-term borrowings and loans due after one year are mapped to the borrowings lines, with no "
            "separate lease liabilities shown. Investments are absent from the workbook and are set to 0. "
            "Other gains and losses (net) is mapped to other operating income (expense), net."
        ),
    }
