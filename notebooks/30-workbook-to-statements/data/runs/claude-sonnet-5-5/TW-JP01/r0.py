import re
import openpyxl


FIELD_KEYS = [
    "revenue", "cost_of_sales", "distribution_costs", "administrative_expenses",
    "other_operating_income", "finance_income", "finance_costs", "income_tax_expense",
    "property_plant_equipment", "right_of_use_assets", "intangible_assets", "investments",
    "deferred_tax_assets", "inventories", "trade_receivables", "other_current_assets",
    "cash_and_cash_equivalents", "trade_payables", "other_current_liabilities",
    "current_borrowings", "current_tax_payable", "non_current_borrowings", "provisions",
    "deferred_tax_liabilities", "share_capital", "retained_earnings", "other_reserves",
    "net_cash_operating", "net_cash_investing", "net_cash_financing", "fx_effect_on_cash",
    "cash_at_beginning", "dividends",
]


def _eng(label):
    """English part of a bilingual label (from the first ASCII letter on)."""
    if not isinstance(label, str):
        return None
    m = re.search(r"[A-Za-z].*$", label.strip())
    return m.group(0).strip().lower() if m else None


def _find_row(ws, eng, after=None, before=None):
    target = eng.lower()
    lo = 0
    hi = ws.max_row + 1
    if after:
        lo = _find_row(ws, after)
    if before:
        hi = _find_row(ws, before, after=after)
    for r in range(lo + 1, hi):
        if _eng(ws.cell(row=r, column=1).value) == target:
            return r
    raise KeyError("%s: label %r not found" % (ws.title, eng))


def _year_cols(ws):
    cols = {}
    for r in range(1, 12):
        for c in range(2, ws.max_column + 1):
            v = ws.cell(row=r, column=c).value
            if isinstance(v, str):
                if "前期" in v:
                    cols["py"] = c
                elif "当期" in v:
                    cols["cy"] = c
        if len(cols) == 2:
            return cols
    raise KeyError("%s: year columns not found" % ws.title)


def _col_letter(c):
    return openpyxl.utils.get_column_letter(c)


def _unit(ws):
    for r in range(1, 8):
        v = ws.cell(row=r, column=1).value
        if isinstance(v, str) and "単位" in v:
            if "百万円" in v:
                return 1000000
            if "千円" in v:
                return 1000
            return 1
    return 1000000


def extract(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    names = wb.sheetnames

    def sheet(name, idx):
        return wb[name] if name in names else wb.worksheets[idx]

    IS = sheet("損益計算書", 0)
    BS = sheet("貸借対照表", 1)
    CF = sheet("CF計算書", 2)
    EQ = sheet("株主資本等変動", 3)
    colmap = {ws.title: _year_cols(ws) for ws in (IS, BS, CF, EQ)}

    def ref(ws, eng, year, factor=1, after=None, before=None):
        r = _find_row(ws, eng, after, before)
        c = colmap[ws.title][year]
        return {"sheet": ws.title, "cell": "%s%d" % (_col_letter(c), r), "factor": factor}

    def value_of(cells):
        total = 0
        for ce in cells:
            ws = wb[ce["sheet"]]
            total += ce["factor"] * ws[ce["cell"]].value
        return total

    CA = dict(after="Current assets", before="Total current assets")
    CL = dict(after="Current liabilities", before="Total current liabilities")

    # key -> list of (sheet, english label, factor, bounds)
    spec = {
        "revenue": [(IS, "Net sales", 1, {})],
        "cost_of_sales": [(IS, "Cost of sales", 1, {})],
        "distribution_costs": [(IS, "Selling expenses", 1, {})],
        "administrative_expenses": [(IS, "General and administrative", 1, {})],
        "other_operating_income": [(IS, "Miscellaneous income", 1, {})],
        "finance_income": [(IS, "Interest income", 1, {})],
        "finance_costs": [(IS, "Interest expense", 1, {})],
        "income_tax_expense": [(IS, "Current income taxes", 1, {}),
                               (IS, "Deferred income taxes", 1, {})],
        "property_plant_equipment": [(BS, "Property and equipment", 1, {})],
        "right_of_use_assets": [(BS, "Right-of-use assets", 1, {})],
        "intangible_assets": [(BS, "Intangible assets", 1, {})],
        "investments": [],
        "deferred_tax_assets": [(BS, "Deferred tax assets", 1, {})],
        "inventories": [(BS, "Inventories", 1, {})],
        "trade_receivables": [(BS, "Accounts receivable", 1, {})],
        "other_current_assets": [(BS, "Other", 1, CA)],
        "cash_and_cash_equivalents": [(BS, "Cash and deposits", 1, {})],
        "trade_payables": [(BS, "Accounts payable", 1, {})],
        "other_current_liabilities": [(BS, "Other", 1, CL)],
        "current_borrowings": [(BS, "Short-term loans", 1, {})],
        "current_tax_payable": [(BS, "Income taxes payable", 1, {})],
        "non_current_borrowings": [(BS, "Long-term loans", 1, {})],
        "provisions": [(BS, "Retirement benefit liability", 1, {})],
        "deferred_tax_liabilities": [(BS, "Deferred tax liabilities", 1, {})],
        "share_capital": [(BS, "Capital stock", 1, {}),
                          (BS, "Capital surplus", 1, {})],
        "retained_earnings": [(BS, "Retained earnings", 1, {})],
        "other_reserves": [(BS, "Accumulated other comprehensive income", 1, {})],
        "net_cash_operating": [(CF, "Operating activities", 1, {})],
        "net_cash_investing": [(CF, "Investing activities", 1, {})],
        "net_cash_financing": [(CF, "Financing activities", 1, {})],
        "fx_effect_on_cash": [(CF, "FX effect", 1, {})],
        "cash_at_beginning": [(CF, "Beginning balance", 1, {})],
        "dividends": [(EQ, "Dividends", -1, {})],
    }

    fields = {}
    for key in FIELD_KEYS:
        fields[key] = {}
        for year in ("cy", "py"):
            cells = [ref(ws, eng, year, f, **b) for ws, eng, f, b in spec[key]]
            fields[key][year] = {"value": value_of(cells), "cells": cells}

    # Printed totals and subtotals
    totals_spec = [
        (IS, "Gross profit", {}, "gross_profit"),
        (IS, "Total SG&A", {}, "distribution_costs + administrative_expenses"),
        (IS, "Operating income", {}, "gross_profit - distribution_costs - administrative_expenses"),
        (IS, "Total non-operating income", {}, "finance_income + other_operating_income"),
        (IS, "Ordinary income", {}, "profit_before_tax"),
        (IS, "Income before income taxes", {}, "profit_before_tax"),
        (IS, "Net income", {}, "profit_for_year"),
        (BS, "Total current assets", {}, "total_current_assets"),
        (BS, "Total non-current assets", {}, "total_non_current_assets"),
        (BS, "Total assets", {}, "total_assets"),
        (BS, "Total current liabilities", {}, "total_current_liabilities"),
        (BS, "Total liabilities", {}, "total_liabilities"),
        (BS, "Total net assets", {}, "total_equity"),
        (BS, "Total liabilities and net assets", {}, "total_equity_and_liabilities"),
        (CF, "Net change", {}, "net_change_in_cash"),
        (CF, "Ending balance", {}, "cash_at_end"),
        (EQ, "Net income", {}, "profit_for_year"),
        (EQ, "Ending balance", {}, "retained_earnings"),
    ]
    source_totals = []
    for ws, eng, bounds, formula in totals_spec:
        for year in ("cy", "py"):
            ce = ref(ws, eng, year, 1, **bounds)
            source_totals.append({
                "sheet": ce["sheet"], "cell": ce["cell"], "year": year,
                "factor": 1, "formula": formula,
            })

    # Period ends from the balance sheet date line
    cy_end, py_end = None, None
    for r in range(1, 6):
        v = BS.cell(row=r, column=1).value
        if isinstance(v, str):
            m = re.search(r"(\d{4})年(\d{1,2})月(\d{1,2})日", v)
            if m:
                y, mo, d = int(m.group(1)), int(m.group(2)), int(m.group(3))
                cy_end = "%04d-%02d-%02d" % (y, mo, d)
                py_end = "%04d-%02d-%02d" % (y - 1, mo, d)
                break

    eq_open_row = _find_row(EQ, "Beginning balance")
    eq_range = "%s%d:%s%d" % (
        _col_letter(min(colmap[EQ.title].values())), eq_open_row,
        _col_letter(max(colmap[EQ.title].values())), eq_open_row,
    )
    ignored = [{
        "sheet": EQ.title,
        "range": eq_range,
        "reason": "Opening retained earnings in the equity movement table: the current-year "
                  "figure repeats last year's closing balance and the prior-year figure is "
                  "not a template line; the template performs the roll-forward itself.",
    }]

    return {
        "entity_code": "TW-JP01",
        "source_entity_name": str(IS["A1"].value),
        "currency": "JPY",
        "unit": _unit(IS),
        "period_end": {"cy": cy_end, "py": py_end},
        "fields": fields,
        "source_totals": source_totals,
        "ignored": ignored,
        "source_exceptions": [],
        "notes": (
            "Workbook is in millions of yen (unit 1000000); column B is the prior year and "
            "column C the current year, detected from the headers. Miscellaneous income is "
            "mapped to other operating income and interest income to finance income; ordinary "
            "income therefore equals the template's profit before tax. Income tax expense is "
            "current plus deferred tax (the deferred line is a charge, as net income only "
            "reconciles that way). Capital stock and capital surplus are combined in share "
            "capital; accumulated other comprehensive income goes to other reserves. Dividends "
            "are shown negative in the equity table and are flipped with factor -1. 'Cash and "
            "deposits' is treated entirely as cash equivalents because it agrees with the cash "
            "flow ending balance. There are no investments or loans to group companies in the "
            "workbook. All printed totals tie and no source exceptions were found."
        ),
    }
