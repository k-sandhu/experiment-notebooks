import re
import openpyxl
from openpyxl.utils import get_column_letter

TB = "TB"
CF = "CF summary"

# line key -> (natural side, account codes). Natural side is the side that gives a positive number.
LINES = {
    "revenue": ("cr", [4000, 4010]),
    "cost_of_sales": ("dr", [5000, 5010]),
    "distribution_costs": ("dr", [6000]),
    "administrative_expenses": ("dr", [6100, 6110, 6120, 6130]),
    "other_operating_income": ("cr", [6500]),
    "finance_income": ("cr", [7000]),
    "finance_costs": ("dr", [7100, 7110]),
    "income_tax_expense": ("dr", [8000, 8010]),
    "property_plant_equipment": ("dr", [1500, 1510]),
    "right_of_use_assets": ("dr", [1550, 1560]),
    "intangible_assets": ("dr", [1600]),
    "investments": ("dr", []),
    "deferred_tax_assets": ("dr", [1700]),
    "inventories": ("dr", []),
    "trade_receivables": ("dr", [1100, 1110]),
    "other_current_assets": ("dr", [1200, 1210, 1220]),
    "cash_and_cash_equivalents": ("dr", [1000, 1010]),
    "trade_payables": ("cr", [2000]),
    "other_current_liabilities": ("cr", [2100, 2110, 2120]),
    "current_borrowings": ("cr", [2300]),
    "current_tax_payable": ("cr", [2200]),
    "non_current_borrowings": ("cr", [2400, 2410]),
    "provisions": ("cr", [2500]),
    "deferred_tax_liabilities": ("cr", [2600]),
    "share_capital": ("cr", [3000]),
    # b/f + current-year result (TB is pre-closing) - dividends
    "retained_earnings": ("cr", [3100, 3200, 4000, 4010, 5000, 5010, 6000, 6100, 6110,
                                 6120, 6130, 6500, 7000, 7100, 7110, 8000, 8010]),
    "other_reserves": ("cr", [3300]),
    "dividends": ("dr", [3200]),
}

CF_LINES = {
    "net_cash_operating": "net cash from operating",
    "net_cash_investing": "net cash from investing",
    "net_cash_financing": "net cash used in financing",
    "fx_effect_on_cash": "foreign exchange differences",
    "cash_at_beginning": "cash at start",
}


def _num(v):
    return isinstance(v, (int, float)) and not isinstance(v, bool)


def _date(text, fallback):
    m = re.search(r"(\d{1,2})/(\d{1,2})/(\d{4})", str(text or ""))
    if not m:
        return fallback
    return "%s-%02d-%02d" % (m.group(3), int(m.group(2)), int(m.group(1)))


def extract(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    tb = wb[TB]
    cf = wb[CF]

    # --- locate TB header, columns, account rows, total row
    hdr = None
    for r in range(1, tb.max_row + 1):
        if str(tb.cell(r, 1).value).strip().lower() == "account":
            hdr = r
            break
    dr_cols, cr_cols = [], []
    for c in range(1, tb.max_column + 1):
        t = str(tb.cell(hdr, c).value or "").strip().lower()
        if t == "debit":
            dr_cols.append(c)
        elif t == "credit":
            cr_cols.append(c)
    cols = {"cy": (dr_cols[0], cr_cols[0]), "py": (dr_cols[1], cr_cols[1])}

    rows = {}
    total_row = None
    last_row = hdr
    for r in range(hdr + 1, tb.max_row + 1):
        a = tb.cell(r, 1).value
        b = str(tb.cell(r, 2).value or "").strip().lower()
        if _num(a):
            rows[int(a)] = r
            last_row = r
        elif b == "total":
            total_row = r

    def tb_cells(codes, natural, year):
        out = []
        dcol, ccol = cols[year]
        for code in codes:
            r = rows.get(code)
            if r is None:
                continue
            for col, side in ((dcol, "dr"), (ccol, "cr")):
                v = tb.cell(r, col).value
                if _num(v):
                    out.append({"sheet": TB, "cell": "%s%d" % (get_column_letter(col), r),
                                "factor": 1 if side == natural else -1})
        return out

    def line(cells, ws_map):
        val = 0
        for c in cells:
            val += c["factor"] * ws_map[c["sheet"]][c["cell"]].value
        return {"value": val, "cells": cells}

    wsmap = {TB: tb, CF: cf}
    fields = {}
    for key, (nat, codes) in LINES.items():
        fields[key] = {y: line(tb_cells(codes, nat, y), wsmap) for y in ("cy", "py")}

    # --- cash flow sheet
    cf_rows = {}
    for r in range(1, cf.max_row + 1):
        a = cf.cell(r, 1).value
        if isinstance(a, str):
            cf_rows[a.strip().lower()] = r

    def cf_row(prefix):
        for k, r in cf_rows.items():
            if k.startswith(prefix):
                return r
        return None

    cf_col = {"cy": "B", "py": "C"}
    for key, prefix in CF_LINES.items():
        r = cf_row(prefix)
        fields[key] = {}
        for y in ("cy", "py"):
            cells = []
            if r is not None and _num(cf[cf_col[y] + str(r)].value):
                cells = [{"sheet": CF, "cell": cf_col[y] + str(r), "factor": 1}]
            fields[key][y] = line(cells, wsmap)

    # --- source totals
    source_totals = []
    if total_row is not None:
        for y in ("cy", "py"):
            for col in cols[y]:
                L = get_column_letter(col)
                cells = []
                for r in range(hdr + 1, last_row + 1):
                    if _num(tb.cell(r, 1).value) and _num(tb.cell(r, col).value):
                        cells.append({"sheet": TB, "cell": "%s%d" % (L, r), "factor": 1})
                if _num(tb.cell(total_row, col).value):
                    source_totals.append({"sheet": TB, "cell": "%s%d" % (L, total_row),
                                          "year": y, "factor": 1, "sum_of": cells})
    r_net = cf_row("net increase")
    r_end = cf_row("cash at end")
    for y in ("cy", "py"):
        if r_net:
            source_totals.append({"sheet": CF, "cell": cf_col[y] + str(r_net), "year": y,
                                  "factor": 1, "formula": "net_change_in_cash"})
        if r_end:
            source_totals.append({"sheet": CF, "cell": cf_col[y] + str(r_end), "year": y,
                                  "factor": 1, "formula": "cash_at_end"})

    ignored = [{
        "sheet": TB,
        "range": "A%d:A%d" % (hdr + 1, last_row),
        "reason": "General ledger account codes, identifiers rather than amounts",
    }]

    title = str(tb.cell(1, 1).value or "")
    name = title.split(" - ")[0].strip()
    cy_end = _date(tb.cell(hdr - 1, cols["cy"][0]).value, "2025-12-31")
    py_end = _date(tb.cell(hdr - 1, cols["py"][0]).value, "2024-12-31")

    return {
        "entity_code": "TW-SG01",
        "source_entity_name": name,
        "currency": "SGD",
        "unit": 1,
        "period_end": {"cy": cy_end, "py": py_end},
        "fields": fields,
        "source_totals": source_totals,
        "ignored": ignored,
        "source_exceptions": [],
        "notes": (
            "Source is a pre-closing trial balance, so each amount is debit minus credit with the sign set per line. "
            "Retained earnings = retained earnings b/f + the year's P&L accounts - dividends, since the TB has not been closed; "
            "this rolls forward from the prior-year column. The 1-month fixed deposit is treated as a cash equivalent "
            "(ties to the cash flow closing cash). Rental deposits go to other current assets. The group loan is treated "
            "as non-current as no current portion is shown. Account 6500 has a debit balance and is treated as a net expense. "
            "The prior-year translation reserve sits in the debit column, so it is a deficit and negative. "
            "Allowance for doubtful debts is netted in trade receivables; accumulated depreciation is netted in PPE and "
            "right-of-use assets. Deferred tax expense is included in income tax expense."
        ),
    }
