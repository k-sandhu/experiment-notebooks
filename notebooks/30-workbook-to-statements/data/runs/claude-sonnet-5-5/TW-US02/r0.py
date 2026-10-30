import calendar
import datetime
import re

import openpyxl
from openpyxl.utils import get_column_letter as L

KEYS = [
    "revenue", "cost_of_sales", "distribution_costs", "administrative_expenses",
    "other_operating_income", "finance_income", "finance_costs", "income_tax_expense",
    "property_plant_equipment", "right_of_use_assets", "intangible_assets", "investments",
    "deferred_tax_assets", "inventories", "trade_receivables", "other_current_assets",
    "cash_and_cash_equivalents", "trade_payables", "other_current_liabilities",
    "current_borrowings", "current_tax_payable", "non_current_borrowings", "provisions",
    "deferred_tax_liabilities", "share_capital", "retained_earnings", "other_reserves",
    "net_cash_operating", "net_cash_investing", "net_cash_financing", "fx_effect_on_cash",
    "cash_at_beginning", "dividends",
]


def _norm(s):
    s = str(s).replace("\u2019", "'")
    return re.sub(r"\s+", " ", s).strip().lower()


def _index(ws):
    idx = {}
    for r in range(1, ws.max_row + 1):
        v = ws.cell(r, 1).value
        if isinstance(v, str):
            idx.setdefault(_norm(v), []).append(r)
    return idx


def _row(idx, label, nth=0):
    return idx[_norm(label)][nth]


def _bsdate(v):
    try:
        return datetime.datetime.strptime(str(v).strip(), "%b %d, %Y").date()
    except ValueError:
        return None


def _totals(v):
    t = dict(v)
    t["gross_profit"] = t["revenue"] - t["cost_of_sales"]
    t["operating_profit"] = (t["gross_profit"] - t["distribution_costs"]
                             - t["administrative_expenses"] + t["other_operating_income"])
    t["profit_before_tax"] = t["operating_profit"] + t["finance_income"] - t["finance_costs"]
    t["profit_for_year"] = t["profit_before_tax"] - t["income_tax_expense"]
    t["total_non_current_assets"] = (t["property_plant_equipment"] + t["right_of_use_assets"]
                                     + t["intangible_assets"] + t["investments"]
                                     + t["deferred_tax_assets"])
    t["total_current_assets"] = (t["inventories"] + t["trade_receivables"]
                                 + t["other_current_assets"] + t["cash_and_cash_equivalents"])
    t["total_assets"] = t["total_non_current_assets"] + t["total_current_assets"]
    t["total_current_liabilities"] = (t["trade_payables"] + t["other_current_liabilities"]
                                      + t["current_borrowings"] + t["current_tax_payable"])
    t["total_non_current_liabilities"] = (t["non_current_borrowings"] + t["provisions"]
                                          + t["deferred_tax_liabilities"])
    t["total_liabilities"] = t["total_current_liabilities"] + t["total_non_current_liabilities"]
    t["total_equity"] = t["share_capital"] + t["retained_earnings"] + t["other_reserves"]
    t["total_equity_and_liabilities"] = t["total_equity"] + t["total_liabilities"]
    t["net_change_in_cash"] = (t["net_cash_operating"] + t["net_cash_investing"]
                               + t["net_cash_financing"])
    t["cash_at_end"] = t["cash_at_beginning"] + t["net_change_in_cash"] + t["fx_effect_on_cash"]
    return t


def _ev(expr, t):
    return sum((-1 if s == "-" else 1) * t[n] for s, n in re.findall(r"([+-]?)\s*(\w+)", expr))


def extract(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    pl, bs, cf = wb["P&L"], wb["Balance Sheet"], wb["Cash Flow"]
    ipl, ibs, icf = _index(pl), _index(bs), _index(cf)

    # ---- P&L columns: twelve monthly columns = current year, "FY2025 Total" = prior year
    hdr = next(r for r in range(1, 10)
               if sum(isinstance(pl.cell(r, c).value, datetime.date)
                      for c in range(2, pl.max_column + 1)) >= 12)
    cy_cols = [c for c in range(2, pl.max_column + 1)
               if isinstance(pl.cell(hdr, c).value, datetime.date)]
    py_col = next(c for c in range(2, pl.max_column + 1)
                  if isinstance(pl.cell(hdr, c).value, str)
                  and "total" in pl.cell(hdr, c).value.lower())

    # ---- balance sheet columns by date
    bhdr = next(r for r in range(1, 10)
                if _bsdate(bs.cell(r, 2).value) and _bsdate(bs.cell(r, 3).value))
    bcols = sorted((2, 3), key=lambda c: _bsdate(bs.cell(bhdr, c).value), reverse=True)
    bcol = {"cy": bcols[0], "py": bcols[1]}
    period_end = {y: _bsdate(bs.cell(bhdr, bcol[y]).value).isoformat() for y in ("cy", "py")}

    # ---- cash flow columns by fiscal-year header
    chdr = next(r for r in range(1, 10)
                if re.search(r"FY\s*\d{4}", str(cf.cell(r, 2).value or "")))
    ccols = sorted((2, 3), key=lambda c: int(re.search(r"\d{4}", str(cf.cell(chdr, c).value)).group()),
                   reverse=True)
    ccol = {"cy": ccols[0], "py": ccols[1]}

    cells = {k: {"cy": [], "py": []} for k in KEYS}

    def pl_add(key, label, factor=1):
        r = _row(ipl, label)
        cells[key]["cy"] = [(pl, f"{L(c)}{r}", factor) for c in cy_cols]
        cells[key]["py"] = [(pl, f"{L(py_col)}{r}", factor)]

    def col_add(key, ws, idx, cols, labels):
        for y in ("cy", "py"):
            cells[key][y] = [(ws, f"{L(cols[y])}{_row(idx, lab, n)}", f) for lab, n, f in labels]

    pl_add("revenue", "Net sales")
    pl_add("cost_of_sales", "Cost of goods sold")
    pl_add("distribution_costs", "Selling expenses")
    pl_add("administrative_expenses", "General and administrative expenses")
    pl_add("other_operating_income", "Other operating (income) expense, net", -1)
    pl_add("finance_income", "Interest income")
    pl_add("finance_costs", "Interest expense")
    pl_add("income_tax_expense", "Provision for income taxes")

    bsmap = {
        "property_plant_equipment": [("Property and equipment, net", 0, 1)],
        "right_of_use_assets": [("Operating lease right-of-use assets", 0, 1)],
        "intangible_assets": [("Goodwill", 0, 1), ("Intangible assets, net", 0, 1)],
        "deferred_tax_assets": [("Deferred income taxes", 0, 1)],
        "trade_receivables": [("Accounts receivable, net", 0, 1)],
        "other_current_assets": [("Prepaid expenses and other current assets", 0, 1)],
        "cash_and_cash_equivalents": [("Cash and cash equivalents", 0, 1)],
        "trade_payables": [("Accounts payable", 0, 1)],
        "other_current_liabilities": [("Accrued expenses", 0, 1), ("Deferred revenue", 0, 1)],
        "current_borrowings": [("Current portion of long-term debt", 0, 1)],
        "current_tax_payable": [("Income taxes payable", 0, 1)],
        "non_current_borrowings": [("Long-term debt", 0, 1)],
        "provisions": [("Other long-term liabilities", 0, 1)],
        "deferred_tax_liabilities": [("Deferred income taxes", 1, 1)],
        "share_capital": [("Members' capital", 0, 1)],
        "retained_earnings": [("Retained earnings", 0, 1)],
    }
    for k, labs in bsmap.items():
        col_add(k, bs, ibs, bcol, labs)

    cfmap = {
        "net_cash_operating": [("Net cash provided by operating activities", 0, 1)],
        "net_cash_investing": [("Net cash used in investing activities", 0, 1)],
        "net_cash_financing": [("Net cash provided by (used in) financing activities", 0, 1)],
        "fx_effect_on_cash": [("Effect of exchange rate changes on cash", 0, 1)],
        "cash_at_beginning": [("Cash and cash equivalents, beginning of year", 0, 1)],
        "dividends": [("Dividends declared", 0, -1)],
    }
    for k, labs in cfmap.items():
        col_add(k, cf, icf, ccol, labs)

    fields = {}
    vals = {"cy": {}, "py": {}}
    for k in KEYS:
        fields[k] = {}
        for y in ("cy", "py"):
            cl = cells[k][y]
            value = sum(f * ws[c].value for ws, c, f in cl)
            vals[y][k] = value
            fields[k][y] = {
                "value": value,
                "cells": [{"sheet": ws.title, "cell": c, "factor": f} for ws, c, f in cl],
            }

    # ---- printed totals
    source_totals = []

    def tot(ws, cell, year, formula=None, sum_of=None, factor=1):
        d = {"sheet": ws.title, "cell": cell, "year": year, "factor": factor}
        if formula is not None:
            d["formula"] = formula
        else:
            d["sum_of"] = [{"sheet": w.title, "cell": c, "factor": f} for w, c, f in sum_of]
        source_totals.append(d)

    # P&L prior-year total column
    for lab, formula in [
        ("Gross profit", "gross_profit"),
        ("Total SG&A", "distribution_costs + administrative_expenses"),
        ("Operating income", "operating_profit"),
        ("Income before income taxes", "profit_before_tax"),
        ("Net income", "profit_for_year"),
    ]:
        tot(pl, f"{L(py_col)}{_row(ipl, lab)}", "py", formula)

    # P&L monthly subtotals: checked against the month's own line items
    r_rev, r_cogs = _row(ipl, "Net sales"), _row(ipl, "Cost of goods sold")
    r_gp, r_sell = _row(ipl, "Gross profit"), _row(ipl, "Selling expenses")
    r_ga, r_sga = _row(ipl, "General and administrative expenses"), _row(ipl, "Total SG&A")
    r_oth, r_op = _row(ipl, "Other operating (income) expense, net"), _row(ipl, "Operating income")
    r_ii, r_ie = _row(ipl, "Interest income"), _row(ipl, "Interest expense")
    r_pbt, r_tax = _row(ipl, "Income before income taxes"), _row(ipl, "Provision for income taxes")
    r_ni = _row(ipl, "Net income")
    for c in cy_cols:
        a = L(c)
        tot(pl, f"{a}{r_gp}", "cy", sum_of=[(pl, f"{a}{r_rev}", 1), (pl, f"{a}{r_cogs}", -1)])
        tot(pl, f"{a}{r_sga}", "cy", sum_of=[(pl, f"{a}{r_sell}", 1), (pl, f"{a}{r_ga}", 1)])
        tot(pl, f"{a}{r_op}", "cy", sum_of=[(pl, f"{a}{r_gp}", 1), (pl, f"{a}{r_sga}", -1),
                                           (pl, f"{a}{r_oth}", -1)])
        tot(pl, f"{a}{r_pbt}", "cy", sum_of=[(pl, f"{a}{r_op}", 1), (pl, f"{a}{r_ii}", 1),
                                            (pl, f"{a}{r_ie}", -1)])
        tot(pl, f"{a}{r_ni}", "cy", sum_of=[(pl, f"{a}{r_pbt}", 1), (pl, f"{a}{r_tax}", -1)])

    for lab, formula in [
        ("Total current assets", "total_current_assets"),
        ("Total assets", "total_assets"),
        ("Total current liabilities", "total_current_liabilities"),
        ("Total liabilities", "total_liabilities"),
        ("Total stockholders' equity", "total_equity"),
        ("Total liabilities and stockholders' equity", "total_equity_and_liabilities"),
    ]:
        r = _row(ibs, lab)
        for y in ("cy", "py"):
            tot(bs, f"{L(bcol[y])}{r}", y, formula)

    for lab, formula in [
        ("Net increase (decrease) in cash", "net_change_in_cash"),
        ("Cash and cash equivalents, end of year", "cash_at_end"),
        ("Net income", "profit_for_year"),
        ("Retained earnings, end of year", "retained_earnings"),
    ]:
        r = _row(icf, lab)
        for y in ("cy", "py"):
            tot(cf, f"{L(ccol[y])}{r}", y, formula)

    # opening retained earnings of the current year is last year's closing balance
    r_reb = _row(icf, "Retained earnings, beginning of year")
    tot(cf, f"{L(ccol['cy'])}{r_reb}", "py", "retained_earnings")

    ignored = [{
        "sheet": cf.title,
        "cell": f"{L(ccol['py'])}{r_reb}",
        "reason": "Opening retained earnings of FY2025 (balance at the end of FY2024); "
                  "two years back, so it has no place in the template",
    }]

    # ---- verify printed totals and template cross-checks; declare any workbook inconsistency
    T = {y: _totals(vals[y]) for y in ("cy", "py")}
    exceptions = []

    def add_exc(check, diff, text):
        if abs(diff) > 0.5:
            exceptions.append({"check": check, "difference": round(diff, 2), "explanation": text})

    sheets = {w.title: w for w in (pl, bs, cf)}
    for st in source_totals:
        ws = sheets[st["sheet"]]
        printed = st["factor"] * ws[st["cell"]].value
        if "formula" in st:
            expected = _ev(st["formula"], T[st["year"]])
        else:
            expected = sum(c["factor"] * sheets[c["sheet"]][c["cell"]].value for c in st["sum_of"])
        add_exc(f"T {st['sheet']}!{st['cell']} {st['year']}", printed - expected,
                "Printed total in the workbook does not agree with its components; "
                "workbook numbers kept as they are.")

    for y in ("cy", "py"):
        add_exc(f"R1 balance {y}", T[y]["total_assets"] - T[y]["total_equity_and_liabilities"],
                "Balance sheet in the workbook does not balance; numbers kept as printed.")
        add_exc(f"R2 cash_at_end {y}", T[y]["cash_at_end"] - T[y]["cash_and_cash_equivalents"],
                "Cash flow does not end at balance-sheet cash in the workbook; numbers kept as printed.")
    add_exc("R3 opening cash cy", vals["cy"]["cash_at_beginning"] - vals["py"]["cash_and_cash_equivalents"],
            "Opening cash differs from prior-year balance-sheet cash in the workbook.")
    add_exc("R4 retained earnings cy",
            vals["py"]["retained_earnings"] + T["cy"]["profit_for_year"] - vals["cy"]["dividends"]
            - vals["cy"]["retained_earnings"],
            "Retained earnings do not roll forward in the workbook; numbers kept as printed.")

    name = str(bs["A1"].value).strip()
    notes = (
        "Entity matched to TW-US02 (Tarnwell Software, LLC, USD, workbook unit US$ so unit = 1). "
        "The P&L is monthly: current-year amounts are the sum of the twelve monthly cells October 2025 to "
        "September 2026 (each cell cited with factor 1). The column headed 'FY2025 Total' holds the prior-year "
        "annual figures; its net income equals FY2025 net income on the Cash Flow sheet, so it is used for py. "
        "Other operating (income) expense is shown with income as negative, so it is taken with factor -1. "
        "Selling expenses map to distribution costs and G&A to administrative expenses. Goodwill and intangible "
        "assets are combined; accrued expenses and deferred revenue go to other current liabilities; other "
        "long-term liabilities go to provisions; 'Deferred income taxes' appears twice and is split by section "
        "(assets vs liabilities). Dividends are shown negative in the source and taken with factor -1. "
        "Monthly subtotals are verified against their own line items via sum_of; annual and balance-sheet totals "
        "via template formulas. Numbers are never adjusted; any mismatch is declared in source_exceptions."
    )

    return {
        "entity_code": "TW-US02",
        "source_entity_name": name,
        "currency": "USD",
        "unit": 1,
        "period_end": period_end,
        "fields": fields,
        "source_totals": source_totals,
        "ignored": ignored,
        "source_exceptions": exceptions,
        "notes": notes,
    }
