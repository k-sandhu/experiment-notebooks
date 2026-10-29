import openpyxl

ENTITY_CODE = "TW-PL01"
CURRENCY = "PLN"
UNIT = 1000

# template key -> (row label in column A, factor)
LINE_MAP = {
    "revenue": ("Sales", 1),
    "cost_of_sales": ("Cost of sales", 1),
    "administrative_expenses": ("Overheads", 1),
    "other_operating_income": ("Other operating income/(expense), net", 1),
    "finance_income": ("Interest income", 1),
    "finance_costs": ("Interest expense", 1),
    "income_tax_expense": ("Taxation", 1),
    "property_plant_equipment": ("Tangible assets", 1),
    "right_of_use_assets": ("Leased assets (right of use)", 1),
    "intangible_assets": ("Goodwill and intangibles", 1),
    "deferred_tax_assets": ("Deferred tax asset", 1),
    "trade_receivables": ("Accounts receivable", 1),
    "other_current_assets": ("Other current assets", 1),
    "cash_and_cash_equivalents": ("Cash and cash equivalents", 1),
    "trade_payables": ("Accounts payable", 1),
    "other_current_liabilities": ("Other payables and accruals", 1),
    "current_borrowings": ("Borrowings (current)", 1),
    "current_tax_payable": ("Current tax payable", 1),
    "non_current_borrowings": ("Borrowings (non-current)", 1),
    "provisions": ("Employee benefit obligations", 1),
    "share_capital": ("Share capital", 1),
    "retained_earnings": ("Retained earnings", 1),
    "net_cash_operating": ("Cash generated from operations (net)", 1),
    "net_cash_investing": ("Net cash used in investing activities", 1),
    "net_cash_financing": ("Net cash used in financing activities", 1),
    "fx_effect_on_cash": ("Foreign exchange differences on cash", 1),
    "cash_at_beginning": ("Cash at start of year", 1),
    "dividends": ("Dividends declared", -1),  # shown negative in the workbook
}

ABSENT_KEYS = [
    "distribution_costs",
    "inventories",
    "investments",
    "deferred_tax_liabilities",
    "other_reserves",
]

# printed totals: (label, formula, factor)
TOTAL_MAP = [
    ("Gross margin", "revenue - cost_of_sales", 1),
    ("Operating profit", "operating_profit", 1),
    ("Profit before tax", "profit_before_tax", 1),
    ("Net profit", "profit_for_year", 1),
    ("Total non-current assets", "total_non_current_assets", 1),
    ("Total current assets", "total_current_assets", 1),
    ("Total assets", "total_assets", 1),
    ("Total equity", "total_equity", 1),
    ("Total non-current liabilities", "total_non_current_liabilities", 1),
    ("Total current liabilities", "total_current_liabilities", 1),
    ("Total liabilities", "total_liabilities", 1),
    ("Total equity and liabilities", "total_equity_and_liabilities", 1),
    ("Net change in cash and cash equivalents", "net_change_in_cash", 1),
    ("Closing cash and cash equivalents", "cash_at_end", 1),
    ("Profit for the year", "profit_for_year", 1),
    ("Closing retained earnings", "retained_earnings", 1),
]


def _label_rows(ws):
    rows = {}
    for r in range(1, ws.max_row + 1):
        v = ws.cell(row=r, column=1).value
        if isinstance(v, str):
            rows.setdefault(v.strip(), r)
    return rows


def extract(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    years = sorted(wb.sheetnames, key=lambda s: int(s))
    sheets = {"py": years[-2], "cy": years[-1]}

    rows = {y: _label_rows(wb[sheets[y]]) for y in sheets}

    def cell_ref(y, label):
        r = rows[y][label]
        return sheets[y], "B%d" % r

    def num(y, label):
        s, c = cell_ref(y, label)
        return wb[s][c].value

    fields = {}
    for key, (label, factor) in LINE_MAP.items():
        fields[key] = {}
        for y in ("cy", "py"):
            s, c = cell_ref(y, label)
            fields[key][y] = {
                "value": factor * num(y, label),
                "cells": [{"sheet": s, "cell": c, "factor": factor}],
            }
    for key in ABSENT_KEYS:
        fields[key] = {y: {"value": 0, "cells": []} for y in ("cy", "py")}

    source_totals = []
    for label, formula, factor in TOTAL_MAP:
        for y in ("cy", "py"):
            if label not in rows[y]:
                continue
            s, c = cell_ref(y, label)
            source_totals.append({
                "sheet": s, "cell": c, "year": y,
                "factor": factor, "formula": formula,
            })

    ignored = []
    for y in ("cy", "py"):
        s, c = cell_ref(y, "Opening retained earnings")
        ignored.append({
            "sheet": s, "cell": c,
            "reason": "Opening retained earnings in the reconciliation; equals the prior-year "
                      "closing retained earnings, which is taken from the prior-year sheet",
        })

    title = wb[sheets["cy"]]["A1"].value or ""
    source_name = title.split(" - ")[0].strip()

    return {
        "entity_code": ENTITY_CODE,
        "source_entity_name": source_name,
        "currency": CURRENCY,
        "unit": UNIT,
        "period_end": {
            "cy": "%s-12-31" % sheets["cy"],
            "py": "%s-12-31" % sheets["py"],
        },
        "fields": fields,
        "source_totals": source_totals,
        "ignored": ignored,
        "source_exceptions": [],
        "notes": (
            "Workbook is in PLN thousands (tys. PLN), kept as unit 1000. One sheet per year; "
            "the later sheet is cy and the earlier is py. Period-end dates are not printed in "
            "the workbook; 31 December of each sheet year is assumed. 'Overheads' is a single "
            "line with no selling split, so it is mapped wholly to administrative expenses and "
            "distribution costs are 0. Employee benefit obligations are mapped to provisions. "
            "Dividends are shown as negatives in the reconciliation and are sign-flipped with "
            "factor -1. No inventories, investments, deferred tax liabilities or other reserves "
            "appear, so those lines are 0. All printed totals, the balance sheet, the cash flow "
            "and the retained earnings roll-forward tie, so no source exceptions are declared."
        ),
    }
