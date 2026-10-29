import openpyxl

COLS = {"L": {"cy": "B", "py": "C"}, "R": {"cy": "F", "py": "G"}}
LABEL_COL = {"L": "A", "R": "E"}

# (sheet, side, label, factor)
FIELD_SPECS = {
    "revenue": [("GuV", "L", "Umsatzerlöse", 1)],
    "cost_of_sales": [("GuV", "L", "Herstellungskosten der zur Erzielung der Umsatzerlöse erbrachten Leistungen", -1)],
    "distribution_costs": [("GuV", "L", "Vertriebskosten", -1)],
    "administrative_expenses": [("GuV", "L", "Allgemeine Verwaltungskosten", -1)],
    "other_operating_income": [("GuV", "L", "Sonstige betriebliche Erträge (netto)", 1)],
    "finance_income": [("GuV", "L", "Zinsen und ähnliche Erträge", 1)],
    "finance_costs": [("GuV", "L", "Zinsen und ähnliche Aufwendungen", -1)],
    "income_tax_expense": [("GuV", "L", "Steuern vom Einkommen und vom Ertrag", -1)],
    "property_plant_equipment": [
        ("Bilanz", "L", "Grundstücke und Bauten", 1),
        ("Bilanz", "L", "Technische Anlagen und Maschinen", 1),
        ("Bilanz", "L", "Betriebs- und Geschäftsausstattung", 1),
    ],
    "right_of_use_assets": [("Bilanz", "L", "Nutzungsrechte aus Leasingverhältnissen", 1)],
    "intangible_assets": [("Bilanz", "L", "Immaterielle Vermögensgegenstände", 1)],
    "investments": [],
    "deferred_tax_assets": [("Bilanz", "L", "D. Aktive latente Steuern", 1)],
    "inventories": [
        ("Bilanz", "L", "Roh-, Hilfs- und Betriebsstoffe", 1),
        ("Bilanz", "L", "Fertige Erzeugnisse und Waren", 1),
    ],
    "trade_receivables": [("Bilanz", "L", "Forderungen aus Lieferungen und Leistungen", 1)],
    "other_current_assets": [
        ("Bilanz", "L", "Sonstige Vermögensgegenstände", 1),
        ("Bilanz", "L", "C. Rechnungsabgrenzungsposten", 1),
    ],
    "cash_and_cash_equivalents": [("Bilanz", "L", "Kassenbestand, Guthaben bei Kreditinstituten", 1)],
    "trade_payables": [("Bilanz", "R", "Verbindlichkeiten aus Lieferungen und Leistungen", 1)],
    "other_current_liabilities": [("Bilanz", "R", "Sonstige Verbindlichkeiten", 1)],
    "current_borrowings": [
        ("Bilanz", "R", "Verbindlichkeiten gegenüber Kreditinstituten (Restlaufzeit bis 1 Jahr)", 1),
        ("Bilanz", "R", "Leasingverbindlichkeiten (bis 1 Jahr)", 1),
    ],
    "current_tax_payable": [("Bilanz", "R", "Steuerrückstellungen", 1)],
    "non_current_borrowings": [
        ("Bilanz", "R", "Verbindlichkeiten gegenüber Kreditinstituten (Restlaufzeit über 1 Jahr)", 1),
        ("Bilanz", "R", "Leasingverbindlichkeiten (über 1 Jahr)", 1),
    ],
    "provisions": [("Bilanz", "R", "Rückstellungen für Pensionen und ähnliche Verpflichtungen", 1)],
    "deferred_tax_liabilities": [("Bilanz", "R", "D. Passive latente Steuern", 1)],
    "share_capital": [
        ("Bilanz", "R", "Gezeichnetes Kapital", 1),
        ("Bilanz", "R", "Kapitalrücklage", 1),
    ],
    "retained_earnings": [("Bilanz", "R", "Gewinnrücklagen und Bilanzgewinn", 1)],
    "other_reserves": [("Bilanz", "R", "Sonstige Rücklagen", 1)],
    "net_cash_operating": [("KFR", "L", "Cashflow aus laufender Geschäftstätigkeit", 1)],
    "net_cash_investing": [("KFR", "L", "Cashflow aus der Investitionstätigkeit", 1)],
    "net_cash_financing": [("KFR", "L", "Cashflow aus der Finanzierungstätigkeit", 1)],
    "fx_effect_on_cash": [("KFR", "L", "Wechselkursbedingte Änderungen des Finanzmittelfonds", 1)],
    "cash_at_beginning": [("KFR", "L", "Finanzmittelfonds am Anfang der Periode", 1)],
    "dividends": [("KFR", "L", "Ausschüttung an die Gesellschafter", -1)],
}

# (sheet, side, label, formula string or list of (sheet, side, label) for sum_of)
TOTAL_SPECS = [
    ("GuV", "L", "Bruttoergebnis vom Umsatz", "gross_profit"),
    ("GuV", "L", "Betriebsergebnis", "operating_profit"),
    ("GuV", "L", "Ergebnis vor Steuern", "profit_before_tax"),
    ("GuV", "L", "Jahresüberschuss", "profit_for_year"),
    ("Bilanz", "L", "Summe Anlagevermögen",
     "property_plant_equipment + right_of_use_assets + intangible_assets + investments"),
    ("Bilanz", "L", "Summe Umlaufvermögen", [
        ("Bilanz", "L", "Roh-, Hilfs- und Betriebsstoffe"),
        ("Bilanz", "L", "Fertige Erzeugnisse und Waren"),
        ("Bilanz", "L", "Forderungen aus Lieferungen und Leistungen"),
        ("Bilanz", "L", "Sonstige Vermögensgegenstände"),
        ("Bilanz", "L", "Kassenbestand, Guthaben bei Kreditinstituten"),
    ]),
    ("Bilanz", "L", "Summe Aktiva", "total_assets"),
    ("Bilanz", "R", "Summe Eigenkapital", "total_equity"),
    ("Bilanz", "R", "Summe Rückstellungen", "provisions + current_tax_payable"),
    ("Bilanz", "R", "Summe Verbindlichkeiten",
     "current_borrowings + non_current_borrowings + trade_payables + other_current_liabilities"),
    ("Bilanz", "R", "Summe Passiva", "total_equity_and_liabilities"),
    ("KFR", "L", "Zahlungswirksame Veränderung des Finanzmittelfonds", "net_change_in_cash"),
    ("KFR", "L", "Finanzmittelfonds am Ende der Periode", "cash_at_end"),
    ("KFR", "L", "Jahresüberschuss", "profit_for_year"),
    ("KFR", "L", "Stand 31. Dezember", "retained_earnings"),
]


def extract(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    row_cache = {}

    def find_row(sheet, side, label):
        key = (sheet, side)
        if key not in row_cache:
            ws = wb[sheet]
            col = LABEL_COL[side]
            rows = {}
            for r in range(1, ws.max_row + 1):
                v = ws[f"{col}{r}"].value
                if isinstance(v, str):
                    rows.setdefault(v.strip(), r)
            row_cache[key] = rows
        return row_cache[key][label]

    def addr(sheet, side, label, year):
        return f"{COLS[side][year]}{find_row(sheet, side, label)}"

    def num(sheet, cell):
        v = wb[sheet][cell].value
        return v if isinstance(v, (int, float)) else 0

    fields = {}
    for key, items in FIELD_SPECS.items():
        fields[key] = {}
        for year in ("cy", "py"):
            cells = [{"sheet": s, "cell": addr(s, side, lab, year), "factor": f}
                     for (s, side, lab, f) in items]
            value = sum(c["factor"] * num(c["sheet"], c["cell"]) for c in cells)
            fields[key][year] = {"value": value, "cells": cells}

    source_totals = []
    for sheet, side, label, spec in TOTAL_SPECS:
        for year in ("cy", "py"):
            entry = {"sheet": sheet, "cell": addr(sheet, side, label, year),
                     "year": year, "factor": 1}
            if isinstance(spec, str):
                entry["formula"] = spec
            else:
                entry["sum_of"] = [{"sheet": s, "cell": addr(s, sd, lb, year), "factor": 1}
                                   for (s, sd, lb) in spec]
            source_totals.append(entry)

    open_row = find_row("KFR", "L", "Stand 1. Januar")
    ignored = [{
        "sheet": "KFR",
        "range": f"B{open_row}:C{open_row}",
        "reason": "Opening balance of the retained-earnings roll-forward; equals the prior-year "
                  "closing balance already taken from the balance sheet (checked by the R4 roll-forward).",
    }]

    name = wb["GuV"]["A1"].value
    return {
        "entity_code": "TW-DE01",
        "source_entity_name": name.strip() if isinstance(name, str) else name,
        "currency": "EUR",
        "unit": 1000,
        "period_end": {"cy": "2025-12-31", "py": "2024-12-31"},
        "fields": fields,
        "source_totals": source_totals,
        "ignored": ignored,
        "source_exceptions": [],
        "notes": (
            "German HGB workbook in TEUR (unit 1000), cost-of-sales method. Expenses and tax charge are "
            "shown negative in the source, so they are taken with factor -1; dividends likewise. "
            "Property, plant and equipment is the sum of the three Sachanlagen lines. Prepayments "
            "(Rechnungsabgrenzungsposten) are mapped to other current assets. Bank loans and lease "
            "liabilities are split into current and non-current by their stated maturities. Tax provisions "
            "(Steuerrückstellungen) are mapped to current income tax payable, and pension provisions to "
            "provisions. Gezeichnetes Kapital plus Kapitalrücklage form share capital and premium. "
            "Printed subtotals were checked against the template and all agree; no source exceptions."
        ),
    }
