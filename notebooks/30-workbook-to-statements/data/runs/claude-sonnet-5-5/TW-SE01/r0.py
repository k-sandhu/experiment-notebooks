import re
import openpyxl
from openpyxl.utils import get_column_letter

MINUS = "\u2212\u2012\u2013\u2014"


def _num(v):
    if isinstance(v, bool) or v is None:
        raise ValueError("not a number")
    if isinstance(v, (int, float)):
        return v
    s = str(v).strip()
    for m in MINUS:
        s = s.replace(m, "-")
    s = re.sub(r"[\s\u00a0\u2009\u202f]", "", s)
    neg = s.startswith("(") and s.endswith(")")
    s = s.strip("()")
    x = float(s)
    x = -x if neg else x
    return int(x) if x == int(x) else x


def _norm(s):
    return re.sub(r"\s+", " ", str(s)).strip().casefold()


def _find_row(ws, label, start=1):
    target = _norm(label)
    for r in range(start, ws.max_row + 1):
        v = ws.cell(r, 1).value
        if v is not None and _norm(v) == target:
            return r
    raise KeyError("%s: label not found: %s" % (ws.title, label))


def _year_row(ws, start=1):
    """First row at or after start whose columns B and C hold two 4-digit years."""
    for r in range(start, ws.max_row + 1):
        try:
            b, c = _num(ws.cell(r, 2).value), _num(ws.cell(r, 3).value)
        except (ValueError, TypeError):
            continue
        if 1900 <= b <= 2200 and 1900 <= c <= 2200:
            return r, b, c
    raise KeyError("%s: year header not found" % ws.title)


def extract(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    rr, br, kf = wb["RR"], wb["BR"], wb["KF"]

    # year columns: the later year is the current year
    _, y1, y2 = _year_row(rr)
    cols = {"cy": 2, "py": 3} if y1 >= y2 else {"cy": 3, "py": 2}
    cy_year = max(y1, y2)
    py_year = min(y1, y2)

    name = str(rr.cell(_find_row(rr, "Tarnwell Nordic Aktiebolag"), 1).value).strip()
    unit_text = _norm(rr.cell(4, 1).value or "")
    if "million" in unit_text:
        unit = 1000000
    elif "thousand" in unit_text or unit_text.startswith("belopp i t"):
        unit = 1000
    else:
        unit = 1

    def cell_ref(ws, label, year, factor=1, start=1):
        r = _find_row(ws, label, start)
        coord = "%s%d" % (get_column_letter(cols[year]), r)
        return {"sheet": ws.title, "cell": coord, "factor": factor}

    def value_of(ws, c):
        return c["factor"] * _num(ws[c["cell"]].value)

    spec = {
        "revenue": [(rr, "Turnover", 1)],
        "cost_of_sales": [(rr, "Direct costs", -1)],
        "distribution_costs": [(rr, "Selling expenses", -1)],
        "administrative_expenses": [(rr, "Overheads", -1)],
        "other_operating_income": [(rr, "Other operating income/(expense), net", 1)],
        "finance_income": [(rr, "Finance income", 1)],
        "finance_costs": [(rr, "Interest paid on borrowings", -1)],
        "income_tax_expense": [(rr, "Income tax expense", -1)],
        "property_plant_equipment": [(br, "Property, plant and equipment", 1)],
        "right_of_use_assets": [(br, "Right-of-use assets", 1)],
        "intangible_assets": [(br, "Goodwill and intangibles", 1)],
        "investments": [],
        "deferred_tax_assets": [(br, "Deferred tax asset", 1)],
        "inventories": [(br, "Stock", 1)],
        "trade_receivables": [(br, "Accounts receivable", 1)],
        "other_current_assets": [(br, "Other current assets", 1)],
        "cash_and_cash_equivalents": [
            (br, "Cash at bank", 1),
            (br, "Money market deposits (up to 3 months)", 1),
        ],
        "trade_payables": [(br, "Accounts payable", 1)],
        "other_current_liabilities": [(br, "Accruals and other payables", 1)],
        "current_borrowings": [(br, "Short-term borrowings", 1)],
        "current_tax_payable": [(br, "Corporation tax", 1)],
        "non_current_borrowings": [(br, "Loans due after one year", 1)],
        "provisions": [(br, "Provisions", 1)],
        "deferred_tax_liabilities": [(br, "Deferred tax liabilities", 1)],
        "share_capital": [(br, "Share capital", 1)],
        "retained_earnings": [(br, "Retained earnings", 1)],
        "other_reserves": [(br, "Translation reserve", 1)],
        "net_cash_operating": [(kf, "Net cash from operating activities", 1)],
        "net_cash_investing": [(kf, "Net cash from investing activities", 1)],
        "net_cash_financing": [(kf, "Net cash from financing activities", 1)],
        "fx_effect_on_cash": [(kf, "Foreign exchange differences on cash", 1)],
        "cash_at_beginning": [(kf, "Cash at start of year", 1)],
        "dividends": [(br, "Dividends paid to shareholders", -1)],
    }

    fields = {}
    for key, parts in spec.items():
        fields[key] = {}
        for yr in ("cy", "py"):
            cells = [cell_ref(ws, lab, yr, f) for ws, lab, f in parts]
            val = sum(value_of(ws, c) for (ws, _, _), c in zip(parts, cells))
            fields[key][yr] = {"value": val, "cells": cells}

    # printed totals and subtotals
    totals_spec = [
        (rr, "Gross profit", "revenue - cost_of_sales"),
        (rr, "Profit from operations", "operating_profit"),
        (rr, "Profit before tax", "profit_before_tax"),
        (rr, "Profit after tax", "profit_for_year"),
        (br, "Total non-current assets", "total_non_current_assets"),
        (br, "Total current assets", "total_current_assets"),
        (br, "TOTAL ASSETS", "total_assets"),
        (br, "Total shareholders' funds", "total_equity"),
        (br, "Total non-current liabilities", "total_non_current_liabilities"),
        (br, "Total current liabilities", "total_current_liabilities"),
        (br, "Total liabilities", "total_liabilities"),
        (br, "Total equity and liabilities", "total_equity_and_liabilities"),
        (br, "Profit for the year", "profit_for_year"),
        (br, "Closing retained earnings", "retained_earnings"),
        (kf, "Net increase/(decrease) in cash", "net_change_in_cash"),
        (kf, "Cash at end of year", "cash_at_end"),
    ]
    source_totals = []
    for ws, lab, formula in totals_spec:
        for yr in ("cy", "py"):
            c = cell_ref(ws, lab, yr)
            source_totals.append({
                "sheet": c["sheet"], "cell": c["cell"], "year": yr,
                "factor": 1, "formula": formula,
            })

    # ignored numbers
    ignored = []
    hdr_r, _, _ = _year_row(rr)
    for ws in (rr, br, kf):
        r, _, _ = _year_row(ws)
        ignored.append({
            "sheet": ws.title, "range": "B%d:C%d" % (r, r),
            "reason": "column year headings, not amounts",
        })
    mov = _find_row(br, "Movement in retained earnings")
    r2, _, _ = _year_row(br, mov)
    ignored.append({
        "sheet": "BR", "range": "B%d:C%d" % (r2, r2),
        "reason": "year headings of the retained earnings movement table",
    })
    ro = _find_row(br, "Opening retained earnings")
    ignored.append({
        "sheet": "BR", "range": "B%d:C%d" % (ro, ro),
        "reason": "opening retained earnings is the prior-year closing balance, "
                  "a roll-forward component already covered by the retained earnings line",
    })

    notes = (
        "Entity matched to TW-SE01 (Tarnwell Nordic Aktiebolag, SEK, workbook unit TSEK so unit 1000). "
        "Amounts are stored as text with space thousands separators and Unicode minus signs; they are parsed in code. "
        "Costs and the tax charge are shown negative in the source and are flipped with factor -1; dividends likewise. "
        "Cash and cash equivalents is cash at bank plus money market deposits of up to 3 months, which equals the "
        "cash-flow closing cash. Investments are not present and are 0. Stock maps to inventories, accounts receivable "
        "to trade receivables, and the translation reserve to other reserves. "
        "All printed totals, the balance sheet, the cash flow and the retained earnings roll-forward tie, "
        "so no source exceptions are declared."
    )

    return {
        "entity_code": "TW-SE01",
        "source_entity_name": name,
        "currency": "SEK",
        "unit": unit,
        "period_end": {"cy": "%d-12-31" % cy_year, "py": "%d-12-31" % py_year},
        "fields": fields,
        "source_totals": source_totals,
        "ignored": ignored,
        "source_exceptions": [],
        "notes": notes,
    }
