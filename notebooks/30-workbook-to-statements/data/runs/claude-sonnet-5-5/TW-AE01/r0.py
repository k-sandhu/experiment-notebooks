import re
import datetime
import openpyxl


def _norm(s):
    return re.sub(r"\s+", " ", str(s)).strip().lower()


def _find_sheet(wb, *words):
    for name in wb.sheetnames:
        if all(w in name.lower() for w in words):
            return wb[name]
    raise KeyError(words)


def _row_of(ws, label):
    want = _norm(label)
    for r in range(1, ws.max_row + 1):
        v = ws.cell(row=r, column=1).value
        if v is not None and _norm(v) == want:
            return r
    raise KeyError("%s: %s" % (ws.title, label))


def _year_cols(ws):
    found = {}
    for r in range(1, 10):
        for c in range(2, ws.max_column + 1):
            v = ws.cell(row=r, column=c).value
            if isinstance(v, str):
                m = re.search(r"FY\s*(\d{4})", v)
                if m and int(m.group(1)) not in found:
                    found[int(m.group(1))] = openpyxl.utils.get_column_letter(c)
    ys = sorted(found, reverse=True)
    return {"cy": found[ys[0]], "py": found[ys[1]]}


def _num(x):
    if isinstance(x, float) and x == int(x):
        return int(x)
    return x


def extract(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    pl = _find_sheet(wb, "pl")
    bs = _find_sheet(wb, "bs", "final")
    draft = _find_sheet(wb, "draft")
    cf = _find_sheet(wb, "cf")
    cover = _find_sheet(wb, "cover")

    cols = {id(s): _year_cols(s) for s in (pl, bs, cf)}

    def cell_ref(ws, label, yr):
        return "%s%d" % (cols[id(ws)][yr], _row_of(ws, label))

    def line(ws, label, sign=1):
        out = {}
        for yr in ("cy", "py"):
            ref = cell_ref(ws, label, yr)
            v = ws[ref].value
            out[yr] = {
                "value": _num(sign * v),
                "cells": [{"sheet": ws.title, "cell": ref, "factor": sign}],
            }
        return out

    spec = {
        "revenue": (pl, "Net sales", 1),
        "cost_of_sales": (pl, "Cost of sales", 1),
        "distribution_costs": (pl, "Selling expenses", 1),
        "administrative_expenses": (pl, "Overheads", 1),
        "other_operating_income": (pl, "Other operating income/(expense), net", 1),
        "finance_income": (pl, "Finance income", 1),
        "finance_costs": (pl, "Finance costs", 1),
        "income_tax_expense": (pl, "Taxation", 1),
        "property_plant_equipment": (bs, "Property, plant and equipment", 1),
        "right_of_use_assets": (bs, "Leased assets (right of use)", 1),
        "intangible_assets": (bs, "Intangible assets", 1),
        "deferred_tax_assets": (bs, "Deferred tax asset", 1),
        "inventories": (bs, "Stock", 1),
        "trade_receivables": (bs, "Trade debtors", 1),
        "other_current_assets": (bs, "Other receivables and prepayments", 1),
        "cash_and_cash_equivalents": (bs, "Cash at bank and in hand", 1),
        "trade_payables": (bs, "Accounts payable", 1),
        "other_current_liabilities": (bs, "Other current liabilities", 1),
        "current_borrowings": (bs, "Short-term borrowings", 1),
        "current_tax_payable": (bs, "Corporation tax", 1),
        "non_current_borrowings": (bs, "Borrowings (non-current)", 1),
        "provisions": (bs, "Long-term provisions", 1),
        "deferred_tax_liabilities": (bs, "Deferred taxation", 1),
        "share_capital": (bs, "Share capital", 1),
        "retained_earnings": (bs, "Retained earnings", 1),
        "other_reserves": (bs, "Other reserves", 1),
        "net_cash_operating": (cf, "Net cash from operating activities", 1),
        "net_cash_investing": (cf, "Net cash from investing activities", 1),
        "net_cash_financing": (cf, "Net cash used in financing activities", 1),
        "fx_effect_on_cash": (cf, "Foreign exchange differences on cash", 1),
        "cash_at_beginning": (cf, "Cash at start of year", 1),
        "dividends": (cf, "Dividends", -1),
    }

    fields = {}
    for key, (ws, label, sign) in spec.items():
        fields[key] = line(ws, label, sign)
    # Not shown in this workbook
    fields["investments"] = {
        yr: {"value": 0, "cells": []} for yr in ("cy", "py")
    }

    source_totals = []

    def tot(ws, label, formula):
        for yr in ("cy", "py"):
            source_totals.append({
                "sheet": ws.title,
                "cell": cell_ref(ws, label, yr),
                "year": yr,
                "factor": 1,
                "formula": formula,
            })

    tot(pl, "Gross profit", "gross_profit")
    tot(pl, "Operating profit", "operating_profit")
    tot(pl, "Profit before tax", "profit_before_tax")
    tot(pl, "Profit for the year", "profit_for_year")

    tot(bs, "Total non-current assets", "total_non_current_assets")
    tot(bs, "Total current assets", "total_current_assets")
    tot(bs, "Total assets", "total_assets")
    tot(bs, "Total shareholders' funds", "total_equity")
    tot(bs, "Total non-current liabilities", "total_non_current_liabilities")
    tot(bs, "Total current liabilities", "total_current_liabilities")
    tot(bs, "Total liabilities", "total_liabilities")
    tot(bs, "Total equity and liabilities", "total_equity_and_liabilities")

    tot(cf, "Net change in cash and cash equivalents", "net_change_in_cash")
    tot(cf, "Cash at end of year", "cash_at_end")
    tot(cf, "Profit for the year", "profit_for_year")
    tot(cf, "Closing retained earnings", "retained_earnings")

    # Ignored ranges
    ignored = []
    dcols = _year_cols(draft)
    first = None
    for r in range(1, draft.max_row + 1):
        if isinstance(draft["%s%d" % (dcols["cy"], r)].value, (int, float)):
            first = r
            break
    ignored.append({
        "sheet": draft.title,
        "range": "%s%d:%s%d" % (dcols["cy"], first, dcols["py"], draft.max_row),
        "reason": "Superseded draft v1 balance sheet (19 January 2026); BS_final is the approved "
                  "statement and differs in stock, trade debtors, other current liabilities, "
                  "corporation tax and retained earnings.",
    })
    r_open = _row_of(cf, "Retained earnings at start of year")
    ignored.append({
        "sheet": cf.title,
        "range": "%s%d:%s%d" % (cols[id(cf)]["cy"], r_open, cols[id(cf)]["py"], r_open),
        "reason": "Opening retained earnings in the movement table; equals the prior-year balance "
                  "sheet retained earnings (cy) or is a prior-year opening balance with no template "
                  "line (py). The roll-forward is checked by the template.",
    })

    # Metadata
    name = pl["A1"].value
    cur = None
    for r in range(1, 10):
        v = pl["%s%d" % (cols[id(pl)]["cy"], r)].value
        if isinstance(v, str) and re.fullmatch(r"[A-Z]{3}", v.strip()):
            cur = v.strip()
            break
    unit = 1
    for r in range(1, cover.max_row + 1):
        t = _norm(cover.cell(row=r, column=1).value or "")
        if "currency" in t:
            if "million" in t:
                unit = 1000000
            elif "thousand" in t:
                unit = 1000

    title = str(bs["A2"].value)
    m = re.search(r"(\d{1,2})\s+([A-Za-z]+)\s+(\d{4})", title)
    d_cy = datetime.datetime.strptime(
        "%s %s %s" % (m.group(1), m.group(2), m.group(3)), "%d %B %Y").date()
    try:
        d_py = d_cy.replace(year=d_cy.year - 1)
    except ValueError:
        d_py = d_cy.replace(year=d_cy.year - 1, day=28)

    return {
        "entity_code": "TW-AE01",
        "source_entity_name": name,
        "currency": cur or "AED",
        "unit": unit,
        "period_end": {"cy": d_cy.isoformat(), "py": d_py.isoformat()},
        "fields": fields,
        "source_totals": source_totals,
        "ignored": ignored,
        "source_exceptions": [],
        "notes": (
            "Workbook is 'Tarnwell ME FZE (Jebel Ali)', matched to TW-AE01 Tarnwell Middle East FZE "
            "(AED, full dirhams, unit 1). Balance sheet taken from BS_final; BS_draft_v1 is a "
            "superseded draft and is ignored in full. Dividends come from the retained earnings "
            "movement on the CF sheet (shown negative, factor -1). Investments are not reported, so 0. "
            "Short-term and non-current borrowings are mapped as shown; the workbook does not "
            "separate lease liabilities. All printed totals and cross-statement ties were consistent, "
            "so no source exceptions are declared."
        ),
    }
