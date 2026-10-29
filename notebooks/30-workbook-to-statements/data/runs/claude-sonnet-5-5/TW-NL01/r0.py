import re
import datetime
import openpyxl

SHEET = "FS"


def _norm(s):
    return re.sub(r"\s+", " ", str(s)).strip().lower()


def extract(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[SHEET]

    # label -> row (first occurrence of each label)
    rows = {}
    for r in range(1, ws.max_row + 1):
        v = ws.cell(r, 1).value
        if isinstance(v, str) and v.strip():
            rows.setdefault(_norm(v), r)

    def row_of(label):
        key = _norm(label)
        if key not in rows:
            raise KeyError("label not found: " + label)
        return rows[key]

    # year columns from the first date header row (earlier date = prior year)
    hdr = None
    for r in range(1, ws.max_row + 1):
        dated = [c for c in range(2, ws.max_column + 1)
                 if isinstance(ws.cell(r, c).value, (datetime.date, datetime.datetime))]
        if len(dated) >= 2:
            hdr = r
            cols_dates = sorted((ws.cell(r, c).value, c) for c in dated)
            break
    py_col = cols_dates[0][1]
    cy_col = cols_dates[-1][1]
    col_of = {"cy": cy_col, "py": py_col}

    def iso(d):
        if isinstance(d, datetime.datetime):
            d = d.date()
        return d.isoformat()

    period_end = {"cy": iso(cols_dates[-1][0]), "py": iso(cols_dates[0][0])}

    # unit and currency
    unit = 1
    currency = "EUR"
    for r in range(1, 5):
        v = ws.cell(r, 1).value
        if isinstance(v, str):
            m = re.match(r"\s*([A-Z]{3})\s*x\s*([\d,\.]+)", v)
            if m:
                currency = m.group(1)
                unit = int(m.group(2).replace(",", ""))
    title = str(ws.cell(1, 1).value)
    source_name = re.split(r"\s+-\s+", title)[0].strip()

    def cell_ref(label, year):
        r = row_of(label)
        c = col_of[year]
        coord = ws.cell(r, c).coordinate
        return coord, ws.cell(r, c).value

    def line(parts, year):
        cells = []
        total = 0
        for label, factor in parts:
            coord, val = cell_ref(label, year)
            cells.append({"sheet": SHEET, "cell": coord, "factor": factor})
            total += factor * val
        return {"value": total, "cells": cells}

    spec = {
        "revenue": [("Sales revenue", 1)],
        "cost_of_sales": [("Direct costs", 1)],
        "distribution_costs": [("Distribution costs", 1)],
        "administrative_expenses": [("General and administrative expenses", 1)],
        "other_operating_income": [("Other gains and losses (net)", 1)],
        "finance_income": [("Finance income", 1)],
        "finance_costs": [("Interest expense", 1)],
        "income_tax_expense": [("Income tax expense", 1)],
        "property_plant_equipment": [("Land and buildings", 1), ("Plant and machinery", 1)],
        "right_of_use_assets": [("Right-of-use assets", 1)],
        "intangible_assets": [("Goodwill and intangibles", 1)],
        "investments": [],
        "deferred_tax_assets": [("Deferred tax assets", 1)],
        "inventories": [],
        "trade_receivables": [("Accounts receivable", 1)],
        "other_current_assets": [("Prepayments and other debtors", 1)],
        "cash_and_cash_equivalents": [("Cash at bank and in hand", 1)],
        "trade_payables": [("Accounts payable", 1)],
        "other_current_liabilities": [("Other payables and accruals", 1)],
        "current_borrowings": [("Bank loans - current portion", 1), ("Lease liabilities - current", 1)],
        "current_tax_payable": [("Income tax payable", 1)],
        "non_current_borrowings": [("Bank loans - non-current", 1), ("Lease liabilities - non-current", 1)],
        "provisions": [("Provisions", 1)],
        "deferred_tax_liabilities": [("Deferred tax liabilities", 1)],
        "share_capital": [("Share capital", 1)],
        "retained_earnings": [("Retained earnings", 1)],
        "other_reserves": [],
        "net_cash_operating": [("Net cash from operating activities", 1)],
        "net_cash_investing": [("Net cash from investing activities", 1)],
        "net_cash_financing": [("Net cash used in financing activities", 1)],
        "fx_effect_on_cash": [("Effect of exchange rate changes", 1)],
        "cash_at_beginning": [("Opening cash and cash equivalents", 1)],
        "dividends": [("Dividends", -1)],
    }

    fields = {}
    for key, parts in spec.items():
        fields[key] = {y: line(parts, y) for y in ("cy", "py")}

    totals_spec = [
        ("Gross profit", "revenue - cost_of_sales"),
        ("Operating profit", "operating_profit"),
        ("Profit before taxation", "profit_before_tax"),
        ("Net profit", "profit_for_year"),
        ("Total non-current assets", "total_non_current_assets"),
        ("Total current assets", "total_current_assets"),
        ("TOTAL ASSETS", "total_assets"),
        ("Total equity", "total_equity"),
        ("Total non-current liabilities", "total_non_current_liabilities"),
        ("Total current liabilities", "total_current_liabilities"),
        ("Total liabilities", "total_liabilities"),
        ("TOTAL EQUITY AND LIABILITIES", "total_equity_and_liabilities"),
        ("Net increase/(decrease) in cash", "net_change_in_cash"),
        ("Cash at end of year", "cash_at_end"),
        ("Profit for the year", "profit_for_year"),
        ("Closing retained earnings", "retained_earnings"),
    ]
    source_totals = []
    for label, formula in totals_spec:
        for y in ("cy", "py"):
            coord, _ = cell_ref(label, y)
            source_totals.append({"sheet": SHEET, "cell": coord, "year": y,
                                  "factor": 1, "formula": formula})

    r_open = row_of("Opening retained earnings")
    lo, hi = sorted((py_col, cy_col))
    rng = "%s:%s" % (ws.cell(r_open, lo).coordinate, ws.cell(r_open, hi).coordinate)
    ignored = [{
        "sheet": SHEET,
        "range": rng,
        "reason": "Opening retained earnings in the roll-forward note; equals prior-year closing "
                  "retained earnings already taken from the balance sheet, so not a separate input.",
    }]

    notes = (
        "Workbook is Tarnwell Logistics BV (TW-NL01), EUR in thousands, kept as is. Column B is 2024 (py) and "
        "column C is 2025 (cy), identified from the date headers. Land and buildings are combined with plant and "
        "machinery for PP&E; bank loans and lease liabilities are combined into current and non-current "
        "borrowings. Other gains and losses (net) is added in operating profit, so it is taken as net income. "
        "Dividends are shown negative in the retained earnings movement and are sign-flipped to positive. "
        "Inventories, investments and other reserves do not appear and are 0. All printed totals tie to the "
        "template, so no source exceptions are declared."
    )

    return {
        "entity_code": "TW-NL01",
        "source_entity_name": source_name,
        "currency": currency,
        "unit": unit,
        "period_end": period_end,
        "fields": fields,
        "source_totals": source_totals,
        "ignored": ignored,
        "source_exceptions": [],
        "notes": notes,
    }
