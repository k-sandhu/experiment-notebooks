import re
from openpyxl import load_workbook

REGISTER = {
    "Tarnwell Holdings plc": ("TW-UK01", "GBP"),
    "Tarnwell Engineering Limited": ("TW-UK02", "GBP"),
    "Tarnwell Engineering Services Limited": ("TW-UK03", "GBP"),
    "Tarnwell Antriebstechnik GmbH": ("TW-DE01", "EUR"),
    "Tarnwell Service GmbH": ("TW-DE02", "EUR"),
    "Tarnwell France SAS": ("TW-FR01", "EUR"),
    "Tarnwell Logistics B.V.": ("TW-NL01", "EUR"),
    "Tarnwell Ireland Limited": ("TW-IE01", "EUR"),
    "Tarnwell Inc.": ("TW-US01", "USD"),
    "Tarnwell Software LLC": ("TW-US02", "USD"),
    "Tarnwell Holdings (US) Inc.": ("TW-US03", "USD"),
    "Tarnwell Canada Ltd.": ("TW-CA01", "CAD"),
    "Tarnwell México, S.A. de C.V.": ("TW-MX01", "MXN"),
    "Tarnwell do Brasil Ltda.": ("TW-BR01", "BRL"),
    "Tarnwell Japan K.K.": ("TW-JP01", "JPY"),
    "Tarnwell India Private Limited": ("TW-IN01", "INR"),
    "Tarnwell Asia Pacific Pte. Ltd.": ("TW-SG01", "SGD"),
    "Tarnwell Australia Pty Ltd": ("TW-AU01", "AUD"),
    "Tarnwell New Zealand Limited": ("TW-NZ01", "NZD"),
    "Tarnwell South Africa (Pty) Ltd": ("TW-ZA01", "ZAR"),
    "Tarnwell Middle East FZE": ("TW-AE01", "AED"),
    "Tarnwell Trading L.L.C.": ("TW-AE02", "AED"),
    "Tarnwell Polska sp. z o.o.": ("TW-PL01", "PLN"),
    "Tarnwell Nordic AB": ("TW-SE01", "SEK"),
    "Tarnwell Schweiz AG": ("TW-CH01", "CHF"),
    "Tarnwell East Africa Limited": ("TW-KE01", "KES"),
}

ER, BS, CF = "Erfolgsrechnung", "Bilanz", "Geldflussrechnung"

# template key -> list of (sheet, label, factor)
SPEC = {
    "revenue": [(ER, "Umsatzerlöse", 1)],
    "cost_of_sales": [],
    "distribution_costs": [],
    "administrative_expenses": [(ER, "Allgemeine Verwaltungskosten", -1)],
    "other_operating_income": [],
    "finance_income": [
        (ER, "Zinserträge aus Ausleihungen an verbundene Unternehmen", 1),
        (ER, "Sonstige Zinsen und ähnliche Erträge", 1),
    ],
    "finance_costs": [(ER, "Zinsen und ähnliche Aufwendungen", -1)],
    "income_tax_expense": [(ER, "Steuern vom Einkommen und vom Ertrag", -1)],
    "property_plant_equipment": [(BS, "Sachanlagen", 1)],
    "right_of_use_assets": [],
    "intangible_assets": [],
    "investments": [(BS, "Finanzanlagen (Ausleihungen an verbundene Unternehmen)", 1)],
    "deferred_tax_assets": [],
    "inventories": [],
    "trade_receivables": [],
    "other_current_assets": [(BS, "Sonstige Vermögensgegenstände", 1)],
    "cash_and_cash_equivalents": [(BS, "Kassenbestand, Guthaben bei Kreditinstituten", 1)],
    "trade_payables": [(BS, "Verbindlichkeiten aus Lieferungen und Leistungen", 1)],
    "other_current_liabilities": [(BS, "Sonstige Verbindlichkeiten", 1)],
    "current_borrowings": [(BS, "Kurzfristige Finanzverbindlichkeiten", 1)],
    "current_tax_payable": [(BS, "Steuerrückstellungen", 1)],
    "non_current_borrowings": [(BS, "Langfristige Finanzverbindlichkeiten", 1)],
    "provisions": [],
    "deferred_tax_liabilities": [],
    "share_capital": [(BS, "Aktienkapital", 1)],
    "retained_earnings": [(BS, "Gewinnrücklagen und Bilanzgewinn", 1)],
    "other_reserves": [],
    "net_cash_operating": [(CF, "Cashflow aus laufender Geschäftstätigkeit", 1)],
    "net_cash_investing": [(CF, "Cashflow aus der Investitionstätigkeit", 1)],
    "net_cash_financing": [(CF, "Cashflow aus der Finanzierungstätigkeit", 1)],
    "fx_effect_on_cash": [(CF, "Wechselkursbedingte Änderungen des Finanzmittelfonds", 1)],
    "cash_at_beginning": [(CF, "Finanzmittelfonds am Anfang der Periode", 1)],
    "dividends": [(BS, "Ausschüttung an die Gesellschafter", -1)],
}

# (sheet, label, formula) printed totals, factor 1
TOTALS = [
    (ER, "Bruttoergebnis vom Umsatz", "gross_profit"),
    (ER, "Betriebsergebnis", "operating_profit"),
    (ER, "Ergebnis vor Steuern", "profit_before_tax"),
    (ER, "Jahresüberschuss", "profit_for_year"),
    (BS, "Summe Anlagevermögen", "total_non_current_assets"),
    (BS, "Summe Umlaufvermögen", "total_current_assets"),
    (BS, "Summe Aktiva", "total_assets"),
    (BS, "Summe Eigenkapital", "total_equity"),
    (BS, "Summe langfristiges Fremdkapital", "total_non_current_liabilities"),
    (BS, "Summe kurzfristiges Fremdkapital", "total_current_liabilities"),
    (BS, "Summe Fremdkapital", "total_liabilities"),
    (BS, "Summe Passiva", "total_equity_and_liabilities"),
    (BS, "Jahresüberschuss", "profit_for_year"),
    (BS, "Stand 31. Dezember", "retained_earnings"),
    (CF, "Zahlungswirksame Veränderung des Finanzmittelfonds", "net_change_in_cash"),
    (CF, "Finanzmittelfonds am Ende der Periode", "cash_at_end"),
]


def _norm(s):
    return re.sub(r"\s+", " ", str(s)).strip()


def _num(v):
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return v
    if isinstance(v, str):
        t = v.strip()
        for ch in ("'", "’", "\u00a0", " ", ","):
            t = t.replace(ch, "")
        t = t.replace("−", "-")
        neg = t.startswith("(") and t.endswith(")")
        t = t.strip("()")
        try:
            x = float(t)
        except ValueError:
            return None
        x = -x if neg else x
        return int(x) if x == int(x) else x
    return None


def extract(path):
    wb = load_workbook(path, data_only=True)

    def find_row(ws, label, after=0):
        want = _norm(label)
        for r in range(after + 1, ws.max_row + 1):
            v = ws.cell(r, 2).value
            if v is not None and _norm(v) == want:
                return r
        raise KeyError(label)

    # year header rows and columns
    def header_rows(ws):
        out = []
        for r in range(1, ws.max_row + 1):
            cols = []
            for c in range(3, ws.max_column + 1):
                n = _num(ws.cell(r, c).value)
                if isinstance(n, int) and 1990 <= n <= 2100 and ws.cell(r, 2).value is None:
                    cols.append((n, c))
            if len(cols) >= 2:
                out.append((r, sorted(cols, reverse=True)))
        return out

    heads = {n: header_rows(wb[n]) for n in (ER, BS, CF)}
    colmap = {}
    for n in heads:
        cols = heads[n][0][1]
        colmap[n] = {"cy": cols[0][1], "py": cols[1][1]}
    cy_year = heads[BS][0][1][0][0]

    def letter(c):
        return chr(64 + c)

    def ref(sheet, row, yr):
        return f"{letter(colmap[sheet][yr])}{row}"

    def val(sheet, row, yr):
        return _num(wb[sheet].cell(row, colmap[sheet][yr]).value)

    # entity
    name = _norm(wb[BS].cell(1, 2).value)
    code = cur = None
    for legal, (c, k) in sorted(REGISTER.items(), key=lambda kv: -len(kv[0])):
        if _norm(legal).lower() in name.lower():
            code, cur = c, k
            break

    unit_txt = _norm(wb[BS].cell(4, 2).value)
    unit = 1
    if "1'000'000" in unit_txt or "Mio" in unit_txt:
        unit = 1000000
    elif "1'000" in unit_txt or unit_txt.upper().startswith("IN T"):
        unit = 1000

    m = re.search(r"(\d{1,2})\.(\d{1,2})\.(\d{4})", str(wb[BS].cell(3, 2).value))
    if m:
        d, mo, y = int(m.group(1)), int(m.group(2)), int(m.group(3))
    else:
        d, mo, y = 31, 12, cy_year
    period_end = {
        "cy": f"{y:04d}-{mo:02d}-{d:02d}",
        "py": f"{y - 1:04d}-{mo:02d}-{d:02d}",
    }

    # fields
    fields = {}
    for key, items in SPEC.items():
        fields[key] = {}
        for yr in ("cy", "py"):
            cells, total = [], 0
            for sheet, label, factor in items:
                row = find_row(wb[sheet], label)
                cells.append({"sheet": sheet, "cell": ref(sheet, row, yr), "factor": factor})
                total += factor * val(sheet, row, yr)
            fields[key][yr] = {"value": total, "cells": cells}

    # printed totals
    source_totals = []
    for sheet, label, formula in TOTALS:
        row = find_row(wb[sheet], label)
        for yr in ("cy", "py"):
            source_totals.append({
                "sheet": sheet, "cell": ref(sheet, row, yr), "year": yr,
                "factor": 1, "formula": formula,
            })

    # ignored
    ignored = []
    r_open = find_row(wb[BS], "Stand 1. Januar")
    ignored.append({
        "sheet": BS,
        "range": f"{ref(BS, r_open, 'cy')}:{ref(BS, r_open, 'py')}",
        "reason": "Opening retained earnings in the roll-forward; equals the prior-year "
                  "closing balance already taken from the balance sheet line.",
    })
    for r, cols in heads[BS][1:]:
        cs = sorted(c for _, c in cols)
        ignored.append({
            "sheet": BS,
            "range": f"{letter(cs[0])}{r}:{letter(cs[-1])}{r}",
            "reason": "Year column headers of the retained earnings roll-forward table, not amounts.",
        })

    # workbook consistency checks
    def f(key, yr):
        return fields[key][yr]["value"]

    exceptions = []
    for yr in ("cy", "py"):
        assets = (f("property_plant_equipment", yr) + f("right_of_use_assets", yr)
                  + f("intangible_assets", yr) + f("investments", yr)
                  + f("deferred_tax_assets", yr) + f("inventories", yr)
                  + f("trade_receivables", yr) + f("other_current_assets", yr)
                  + f("cash_and_cash_equivalents", yr))
        eq_liab = (f("share_capital", yr) + f("retained_earnings", yr) + f("other_reserves", yr)
                   + f("trade_payables", yr) + f("other_current_liabilities", yr)
                   + f("current_borrowings", yr) + f("current_tax_payable", yr)
                   + f("non_current_borrowings", yr) + f("provisions", yr)
                   + f("deferred_tax_liabilities", yr))
        if assets != eq_liab:
            exceptions.append({
                "check": f"R1 balance {yr}", "difference": assets - eq_liab,
                "explanation": "Total assets differ from total equity and liabilities in the workbook.",
            })
        end = (f("cash_at_beginning", yr) + f("net_cash_operating", yr)
               + f("net_cash_investing", yr) + f("net_cash_financing", yr)
               + f("fx_effect_on_cash", yr))
        diff = end - f("cash_and_cash_equivalents", yr)
        if diff != 0:
            exceptions.append({
                "check": f"R2 cash_at_end {yr}", "difference": diff,
                "explanation": "The cash flow statement ends at a different amount than the cash "
                               "line on the balance sheet (cash flow end minus balance sheet cash). "
                               "Both are kept exactly as printed; the cash flow statement itself "
                               "adds up and its printed end-of-period figure matches its components.",
            })

    notes = (
        "Figures are in TCHF (unit 1000) and kept as printed; text values with apostrophe thousand "
        "separators are parsed in code. Expenses and tax are printed negative in the source and are "
        "cited with factor -1; the tax line is a charge (profit before tax less profit for the year). "
        "Finance income combines interest on loans to group companies and other interest income. "
        "Steuerrückstellungen are mapped to current income tax payable. The entity is a treasury "
        "centre with no cost of sales. Dividends come from the retained earnings roll-forward "
        "(Ausschüttung an die Gesellschafter). The profit line in the roll-forward is registered "
        "as a checked source total (profit for the year). The cash flow statement ends in 2025 at a "
        "figure that differs from balance sheet cash; this is declared as a source exception and "
        "not adjusted."
    )

    return {
        "entity_code": code,
        "source_entity_name": name,
        "currency": cur,
        "unit": unit,
        "period_end": period_end,
        "fields": fields,
        "source_totals": source_totals,
        "ignored": ignored,
        "source_exceptions": exceptions,
        "notes": notes,
    }
