import re
from datetime import date

import openpyxl
from openpyxl.utils import get_column_letter

KEYS = [
    "revenue", "cost_of_sales", "distribution_costs", "administrative_expenses",
    "other_operating_income", "finance_income", "finance_costs", "income_tax_expense",
    "property_plant_equipment", "right_of_use_assets", "intangible_assets", "investments",
    "deferred_tax_assets", "inventories", "trade_receivables", "other_current_assets",
    "cash_and_cash_equivalents", "trade_payables", "other_current_liabilities",
    "current_borrowings", "current_tax_payable", "non_current_borrowings", "provisions",
    "deferred_tax_liabilities", "share_capital", "retained_earnings", "other_reserves",
    "net_cash_operating", "net_cash_investing", "net_cash_financing", "fx_effect_on_cash",
    "cash_at_beginning", "dividends",
]

MONTHS = {m: i + 1 for i, m in enumerate(
    ["january", "february", "march", "april", "may", "june", "july",
     "august", "september", "october", "november", "december"])}


def _norm(v):
    return re.sub(r"\s+", " ", str(v)).strip().lower() if v is not None else ""


def _find(ws, label, prefix=False, start=1, end=None):
    end = end or ws.max_row
    label = _norm(label)
    for r in range(start, end + 1):
        t = _norm(ws.cell(r, 1).value)
        if (t.startswith(label) if prefix else t == label):
            return r
    raise KeyError(f"{ws.title}: row '{label}' not found")


def _parse_date(text):
    m = re.search(r"(\d{1,2})\s+([A-Za-z]+)\s+(\d{4})", str(text or ""))
    if not m or m.group(2).lower() not in MONTHS:
        return None
    return date(int(m.group(3)), MONTHS[m.group(2).lower()], int(m.group(1)))


def _year_columns(ws):
    hdr = _find(ws, "Particulars")
    cols = []
    for c in range(2, ws.max_column + 1):
        d = _parse_date(ws.cell(hdr, c).value)
        if d:
            cols.append((d, c))
    cols.sort(reverse=True)
    return {"cy": cols[0][1], "py": cols[1][1]}, {"cy": cols[0][0], "py": cols[1][0]}


def extract(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    pl, bs, cf, se = wb["P&L"], wb["BS"], wb["CF"], wb["SOCE"]

    name = str(pl["A1"].value).strip()
    unit_text = _norm(pl["A3"].value)
    if "lakh" in unit_text:
        unit, scale = 1000, 100      # 1 lakh = 100 thousand
    elif "million" in unit_text:
        unit, scale = 1000000, 1
    elif "thousand" in unit_text:
        unit, scale = 1000, 1
    else:
        unit, scale = 1, 1

    cols, dates = {}, {}
    for ws in (pl, bs, cf, se):
        cols[ws.title], d = _year_columns(ws)
        dates[ws.title] = d

    def cell(ws, row, yr):
        return f"{get_column_letter(cols[ws.title][yr])}{row}"

    def num(ws, row, yr):
        return ws[cell(ws, row, yr)].value

    def line(items, yr):
        cells = [{"sheet": ws.title, "cell": cell(ws, r, yr), "factor": sign * scale}
                 for ws, r, sign in items]
        value = sum(c["factor"] * wb[c["sheet"]][c["cell"]].value for c in cells)
        return {"value": value, "cells": cells}

    # ---- locate rows by label ----
    r = {}
    r["rev"] = _find(pl, "Revenue from operations")
    r["cos"] = _find(pl, "Cost of services")
    r["gp"] = _find(pl, "Gross profit")
    r["sell"] = _find(pl, "Selling and marketing expenses")
    r["ga"] = _find(pl, "General and administrative expenses")
    r["oth"] = _find(pl, "Other income - miscellaneous")
    r["int"] = _find(pl, "Other income - interest on deposits")
    r["fin"] = _find(pl, "Finance costs")
    r["pbt"] = _find(pl, "Profit before tax")
    r["tax"] = _find(pl, "Tax expense")
    r["pfy"] = _find(pl, "Profit for the year")

    b = {}
    b["ppe"] = _find(bs, "Property, plant and equipment")
    b["rou"] = _find(bs, "Right-of-use assets")
    b["intg"] = _find(bs, "Other intangible assets")
    b["dta"] = _find(bs, "Deferred tax assets", prefix=True)
    b["tnca"] = _find(bs, "Total non-current assets")
    b["rec"] = _find(bs, "Trade receivables")
    b["bank"] = _find(bs, "Bank balances other than cash and cash equivalents")
    b["oca"] = _find(bs, "Other current assets")
    b["cash"] = _find(bs, "Cash and cash equivalents")
    b["tca"] = _find(bs, "Total current assets")
    b["ta"] = _find(bs, "TOTAL ASSETS")
    b["sc"] = _find(bs, "Equity share capital")
    b["sp"] = _find(bs, "Securities premium")
    b["re"] = _find(bs, "Retained earnings")
    b["toe"] = _find(bs, "Total other equity")
    b["te"] = _find(bs, "Total equity")
    ncl_hdr = _find(bs, "Non-current liabilities")
    cl_hdr = _find(bs, "Current liabilities")
    b["ncl_lease"] = _find(bs, "Lease liabilities", start=ncl_hdr, end=cl_hdr - 1)
    b["prov"] = _find(bs, "Provisions", start=ncl_hdr, end=cl_hdr - 1)
    b["dtl"] = _find(bs, "Deferred tax liabilities", prefix=True, start=ncl_hdr, end=cl_hdr - 1)
    b["tncl"] = _find(bs, "Total non-current liabilities")
    b["tp_a"] = _find(bs, "(a) total outstanding dues of micro", prefix=True, start=cl_hdr)
    b["tp_b"] = _find(bs, "(b) total outstanding dues of creditors", prefix=True, start=cl_hdr)
    b["ocl"] = _find(bs, "Other current liabilities", start=cl_hdr)
    b["cl_lease"] = _find(bs, "Lease liabilities", start=cl_hdr)
    b["ctax"] = _find(bs, "Current tax liabilities", prefix=True, start=cl_hdr)
    b["tcl"] = _find(bs, "Total current liabilities")
    b["tel"] = _find(bs, "TOTAL EQUITY AND LIABILITIES")

    c = {}
    c["op"] = _find(cf, "Net cash generated from operating", prefix=True)
    c["inv"] = _find(cf, "Net cash used in investing", prefix=True)
    c["fin"] = _find(cf, "Net cash used in financing", prefix=True)
    c["net"] = _find(cf, "Net increase/(decrease)", prefix=True)
    c["fx"] = _find(cf, "Effect of exchange", prefix=True)
    c["beg"] = _find(cf, "Cash and cash equivalents at the beginning", prefix=True)
    c["end"] = _find(cf, "Cash and cash equivalents at the end", prefix=True)

    s = {}
    s["open"] = _find(se, "Balance at the beginning", prefix=True)
    s["pfy"] = _find(se, "Profit for the year")
    s["div"] = _find(se, "Final dividend paid", prefix=True)
    s["close"] = _find(se, "Balance at the end", prefix=True)

    spec = {
        "revenue": [(pl, r["rev"], 1)],
        "cost_of_sales": [(pl, r["cos"], 1)],
        "distribution_costs": [(pl, r["sell"], 1)],
        "administrative_expenses": [(pl, r["ga"], 1)],
        "other_operating_income": [(pl, r["oth"], 1)],
        "finance_income": [(pl, r["int"], 1)],
        "finance_costs": [(pl, r["fin"], 1)],
        "income_tax_expense": [(pl, r["tax"], 1)],
        "property_plant_equipment": [(bs, b["ppe"], 1)],
        "right_of_use_assets": [(bs, b["rou"], 1)],
        "intangible_assets": [(bs, b["intg"], 1)],
        "deferred_tax_assets": [(bs, b["dta"], 1)],
        "trade_receivables": [(bs, b["rec"], 1)],
        # bank balances outside cash equivalents (over three months) go to other current assets
        "other_current_assets": [(bs, b["bank"], 1), (bs, b["oca"], 1)],
        "cash_and_cash_equivalents": [(bs, b["cash"], 1)],
        "trade_payables": [(bs, b["tp_a"], 1), (bs, b["tp_b"], 1)],
        "other_current_liabilities": [(bs, b["ocl"], 1)],
        "current_borrowings": [(bs, b["cl_lease"], 1)],
        "current_tax_payable": [(bs, b["ctax"], 1)],
        "non_current_borrowings": [(bs, b["ncl_lease"], 1)],
        "provisions": [(bs, b["prov"], 1)],
        "deferred_tax_liabilities": [(bs, b["dtl"], 1)],
        "share_capital": [(bs, b["sc"], 1), (bs, b["sp"], 1)],
        "retained_earnings": [(bs, b["re"], 1)],
        "net_cash_operating": [(cf, c["op"], 1)],
        "net_cash_investing": [(cf, c["inv"], 1)],
        "net_cash_financing": [(cf, c["fin"], 1)],
        "fx_effect_on_cash": [(cf, c["fx"], 1)],
        "cash_at_beginning": [(cf, c["beg"], 1)],
        "dividends": [(se, s["div"], -1)],   # shown as a negative movement
    }

    fields = {}
    for k in KEYS:
        if k in spec:
            fields[k] = {yr: line(spec[k], yr) for yr in ("cy", "py")}
        else:
            fields[k] = {yr: {"value": 0, "cells": []} for yr in ("cy", "py")}

    # ---- printed totals ----
    totals = []

    def add(ws, row, formula, years=("cy", "py")):
        for yr in years:
            totals.append({"sheet": ws.title, "cell": cell(ws, row, yr), "year": yr,
                           "factor": scale, "formula": formula})

    add(pl, r["gp"], "revenue - cost_of_sales")
    add(pl, r["pbt"], "revenue - cost_of_sales - distribution_costs - administrative_expenses"
                      " + other_operating_income + finance_income - finance_costs")
    add(pl, r["pfy"], "profit_for_year")
    add(bs, b["tnca"], "total_non_current_assets")
    add(bs, b["tca"], "total_current_assets")
    add(bs, b["ta"], "total_assets")
    add(bs, b["te"], "total_equity")
    add(bs, b["tncl"], "total_non_current_liabilities")
    add(bs, b["tcl"], "total_current_liabilities")
    add(bs, b["tel"], "total_equity_and_liabilities")
    add(cf, c["net"], "net_change_in_cash")
    add(cf, c["end"], "cash_at_end")
    add(se, s["pfy"], "profit_for_year")
    add(se, s["close"], "retained_earnings")
    # opening retained earnings of the current year = prior-year closing balance
    totals.append({"sheet": se.title, "cell": cell(se, s["open"], "cy"), "year": "py",
                   "factor": scale, "formula": "retained_earnings"})
    # subtotal with no template equivalent (securities premium + retained earnings)
    for yr in ("cy", "py"):
        totals.append({
            "sheet": bs.title, "cell": cell(bs, b["toe"], yr), "year": yr, "factor": scale,
            "sum_of": [{"sheet": bs.title, "cell": cell(bs, b["sp"], yr), "factor": scale},
                       {"sheet": bs.title, "cell": cell(bs, b["re"], yr), "factor": scale}]})

    # ---- ignored ----
    ignored = []
    for ws in (pl, bs, cf, se):
        rows = [rr for rr in range(1, ws.max_row + 1)
                if isinstance(ws.cell(rr, 2).value, (int, float))]
        if rows:
            ignored.append({"sheet": ws.title, "range": f"B{min(rows)}:B{max(rows)}",
                            "reason": "Note number references, not amounts"})
    ignored.append({
        "sheet": se.title, "cell": cell(se, s["open"], "py"),
        "reason": "Opening retained earnings of the prior year; the template has no "
                  "year before the prior year to compare it with"})

    notes = (
        "Workbook is in INR lakhs; values are converted to thousands (factor 100 on every cell), "
        "so unit is 1000. Interest on deposits is reported as finance income and miscellaneous "
        "other income as other operating income. Bank balances other than cash equivalents "
        "(over three months) are added to other current assets. Securities premium is added to "
        "equity share capital. Lease liabilities map to non-current and current borrowings. "
        "Both trade payable categories (MSME and other creditors) are combined. Dividends are "
        "taken from the SOCE with the sign flipped. All printed totals reconcile, so no "
        "source exceptions are declared."
    )

    return {
        "entity_code": "TW-IN01",
        "source_entity_name": name,
        "currency": "INR",
        "unit": unit,
        "period_end": {yr: dates["BS"][yr].isoformat() for yr in ("cy", "py")},
        "fields": fields,
        "source_totals": totals,
        "ignored": ignored,
        "source_exceptions": [],
        "notes": notes,
    }
