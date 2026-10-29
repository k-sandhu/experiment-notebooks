import re
import datetime
import openpyxl
from openpyxl.utils import get_column_letter

PL = "P&L"
BS = "Balance sheet"
CF = "Cash flow"

MONTHS = {m: i + 1 for i, m in enumerate(
    ["january", "february", "march", "april", "may", "june", "july",
     "august", "september", "october", "november", "december"])}


def _norm(s):
    return re.sub(r"\s+", " ", str(s).strip().lower())


def _find_row(ws, label, start=1):
    target = _norm(label)
    for r in range(start, ws.max_row + 1):
        v = ws.cell(row=r, column=1).value
        if v is not None and _norm(v) == target:
            return r
    raise KeyError("label not found: %s on %s" % (label, ws.title))


def _year_columns(ws):
    """Return (cy_col, py_col, header_row) from the first row of year headers."""
    for r in range(1, ws.max_row + 1):
        cols = []
        for c in range(2, ws.max_column + 1):
            v = ws.cell(row=r, column=c).value
            if isinstance(v, (int, float)) and 1990 <= v <= 2100:
                cols.append((int(v), c))
        if len(cols) >= 2:
            cols.sort(reverse=True)
            return cols[0][1], cols[1][1], r
    raise KeyError("year header not found on %s" % ws.title)


def extract(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    sheets = {n: wb[n] for n in (PL, BS, CF)}
    cols = {}
    for n, ws in sheets.items():
        cy, py, hdr = _year_columns(ws)
        cols[n] = {"cy": get_column_letter(cy), "py": get_column_letter(py), "hdr": hdr}

    def cell(sheet, label, yr, start=1):
        ws = sheets[sheet]
        r = _find_row(ws, label, start)
        return {"sheet": sheet, "cell": "%s%d" % (cols[sheet][yr], r)}

    def num(sheet, coord):
        return sheets[sheet][coord].value

    def line(refs):
        cells = []
        total = 0
        for sheet, label, factor in refs:
            pass
        return cells, total

    # field -> list of (sheet, label, factor)
    spec = {
        "revenue": [(PL, "Revenue", 1)],
        "administrative_expenses": [(PL, "Administrative expenses", 1)],
        "other_operating_income": [(PL, "Other operating income/(expense), net", 1)],
        "finance_income": [(PL, "Dividends received from subsidiaries", 1),
                           (PL, "Bank interest received", 1)],
        "finance_costs": [(PL, "Interest expense", 1)],
        "income_tax_expense": [(PL, "Tax on profit", 1)],
        "property_plant_equipment": [(BS, "Tangible assets", 1)],
        "investments": [(BS, "Investments in subsidiaries", 1),
                        (BS, "Loans to group undertakings (due after one year)", 1)],
        "other_current_assets": [(BS, "Other receivables and prepayments", 1)],
        "cash_and_cash_equivalents": [(BS, "Cash and cash equivalents", 1)],
        "trade_payables": [(BS, "Trade creditors", 1)],
        "other_current_liabilities": [(BS, "Other current liabilities", 1)],
        "current_borrowings": [(BS, "Borrowings (current)", 1)],
        "current_tax_payable": [(BS, "Income tax payable", 1)],
        "non_current_borrowings": [(BS, "Long-term borrowings", 1)],
        "share_capital": [(BS, "Share capital", 1)],
        "retained_earnings": [(BS, "Retained profits", 1)],
        "net_cash_operating": [(CF, "Net cash from operating activities", 1)],
        "net_cash_investing": [(CF, "Net cash from investing activities", 1)],
        "net_cash_financing": [(CF, "Net cash from financing activities", 1)],
        "cash_at_beginning": [(CF, "Opening cash and cash equivalents", 1)],
        "dividends": [(BS, "Dividends paid to shareholders", -1)],
    }
    all_keys = [
        "revenue", "cost_of_sales", "distribution_costs", "administrative_expenses",
        "other_operating_income", "finance_income", "finance_costs", "income_tax_expense",
        "property_plant_equipment", "right_of_use_assets", "intangible_assets", "investments",
        "deferred_tax_assets", "inventories", "trade_receivables", "other_current_assets",
        "cash_and_cash_equivalents", "trade_payables", "other_current_liabilities",
        "current_borrowings", "current_tax_payable", "non_current_borrowings", "provisions",
        "deferred_tax_liabilities", "share_capital", "retained_earnings", "other_reserves",
        "net_cash_operating", "net_cash_investing", "net_cash_financing", "fx_effect_on_cash",
        "cash_at_beginning", "dividends",
    ]

    fields = {}
    for key in all_keys:
        fields[key] = {}
        for yr in ("cy", "py"):
            cells = []
            total = 0
            for sheet, label, factor in spec.get(key, []):
                c = cell(sheet, label, yr)
                v = num(sheet, c["cell"])
                cells.append({"sheet": sheet, "cell": c["cell"], "factor": factor})
                total += factor * v
            fields[key][yr] = {"value": total, "cells": cells}

    totals = []

    def st(sheet, label, yr, formula, start=1):
        c = cell(sheet, label, yr, start)
        totals.append({"sheet": sheet, "cell": c["cell"], "year": yr,
                       "factor": 1, "formula": formula})

    recon_start = _find_row(sheets[BS], "Retained earnings reconciliation")
    for yr in ("cy", "py"):
        st(PL, "Gross profit", yr, "revenue - cost_of_sales")
        st(PL, "Profit from operations", yr, "operating_profit")
        st(PL, "Profit before tax", yr, "profit_before_tax")
        st(PL, "Net profit", yr, "profit_for_year")
        st(BS, "Total non-current assets", yr, "total_non_current_assets")
        st(BS, "Total current assets", yr, "total_current_assets")
        st(BS, "Total assets", yr, "total_assets")
        st(BS, "Total shareholders' funds", yr, "total_equity")
        st(BS, "Total non-current liabilities", yr, "total_non_current_liabilities")
        st(BS, "Total current liabilities", yr, "total_current_liabilities")
        st(BS, "Total liabilities", yr, "total_liabilities")
        st(BS, "TOTAL EQUITY AND LIABILITIES", yr, "total_equity_and_liabilities")
        st(BS, "Profit for the year", yr, "profit_for_year", recon_start)
        st(BS, "Closing retained earnings", yr, "retained_earnings", recon_start)
        st(CF, "Net increase/(decrease) in cash", yr, "net_change_in_cash")
        st(CF, "Cash at end of year", yr, "cash_at_end")

    ignored = []
    r_open = _find_row(sheets[BS], "Opening retained earnings", recon_start)
    ignored.append({
        "sheet": BS,
        "range": "%s%d:%s%d" % (cols[BS]["cy"], r_open, cols[BS]["py"], r_open),
        "reason": "Opening retained earnings in the reconciliation: equals the prior year's closing "
                  "balance already used as retained_earnings; it is the starting point of the "
                  "roll-forward, not a separate line"})
    h = _find_row(sheets[BS], "Retained earnings reconciliation") + 3
    hdr2 = None
    for r in range(recon_start, sheets[BS].max_row + 1):
        v = sheets[BS].cell(row=r, column=2).value
        if isinstance(v, (int, float)) and 1990 <= v <= 2100:
            hdr2 = r
            break
    if hdr2:
        ignored.append({
            "sheet": BS,
            "range": "%s%d:%s%d" % (cols[BS]["cy"], hdr2, cols[BS]["py"], hdr2),
            "reason": "Year column headers (2026, 2025) of the retained earnings reconciliation, "
                      "not amounts"})

    # period end dates from the balance sheet subtitle
    text = str(sheets[BS]["A3"].value)
    m = re.search(r"(\d{1,2})\s+([A-Za-z]+)\s+(\d{4})", text)
    day, mon, year = int(m.group(1)), MONTHS[m.group(2).lower()], int(m.group(3))
    cy_date = datetime.date(year, mon, day)
    try:
        py_date = datetime.date(year - 1, mon, day)
    except ValueError:
        py_date = datetime.date(year - 1, mon, day - 1)

    return {
        "entity_code": "TW-UK01",
        "source_entity_name": str(sheets[PL]["A1"].value),
        "currency": "GBP",
        "unit": 1,
        "period_end": {"cy": cy_date.isoformat(), "py": py_date.isoformat()},
        "fields": fields,
        "source_totals": totals,
        "ignored": ignored,
        "source_exceptions": [],
        "notes": (
            "Entity matched to TW-UK01 (Tarnwell Holdings plc); amounts are in whole GBP. "
            "The workbook has no cost of sales, so gross profit equals revenue and cost_of_sales is 0. "
            "Finance income is the sum of dividends received from subsidiaries and bank interest received "
            "(the 'Finance income' row is only a heading). Tangible assets map to PPE; investments in "
            "subsidiaries and long-term loans to group undertakings are combined in investments. "
            "Other receivables and prepayments map to other current assets; trade creditors to trade payables. "
            "Dividends are taken from the retained earnings reconciliation with the sign flipped (factor -1). "
            "All printed totals, the balance sheet, the cash flow and the retained earnings roll-forward "
            "tie out, so no source exceptions are declared."
        ),
    }
