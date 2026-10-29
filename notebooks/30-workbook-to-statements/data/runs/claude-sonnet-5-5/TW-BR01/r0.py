import re
import unicodedata
import datetime
import openpyxl

SHEET = "DFs"


def _norm(s):
    s = unicodedata.normalize("NFKD", str(s)).encode("ascii", "ignore").decode()
    return re.sub(r"\s+", " ", s).strip().lower()


def _num(v):
    if isinstance(v, bool) or v is None:
        return None
    if isinstance(v, (int, float)):
        return v
    s = str(v).strip()
    neg = s.startswith("(") and s.endswith(")")
    s = s.strip("()").replace(".", "").replace(",", ".")
    try:
        x = float(s)
    except ValueError:
        return None
    x = -x if neg else x
    return int(x) if x == int(x) else x


def _date(v):
    if isinstance(v, (datetime.datetime, datetime.date)):
        return v.strftime("%Y-%m-%d")
    m = re.match(r"\s*(\d{1,2})/(\d{1,2})/(\d{4})", str(v))
    return f"{m.group(3)}-{int(m.group(2)):02d}-{int(m.group(1)):02d}"


def extract(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[SHEET]
    L = openpyxl.utils.get_column_letter
    CI = openpyxl.utils.column_index_from_string

    def val(col, row):
        return _num(ws[f"{col}{row}"].value)

    # Locate rows by label. Column A: income statement; E: balance sheet; I: cash flow / equity.
    def index(labcol, sectioned=False):
        idx = {}
        sec = None
        cy = L(CI(labcol) + 1)
        for r in range(1, ws.max_row + 1):
            lab = ws[f"{labcol}{r}"].value
            if lab is None:
                continue
            n = _norm(lab)
            if sectioned and val(cy, r) is None:
                if n == "ativo circulante":
                    sec = "ac"
                elif n == "ativo nao circulante":
                    sec = "anc"
                elif n == "passivo circulante":
                    sec = "pc"
                elif n == "passivo nao circulante":
                    sec = "pnc"
                elif n.startswith("patrimonio liquido"):
                    sec = "pl"
                continue
            idx.setdefault((sec, n), r)
        return idx

    IS = index("A")
    BS = index("E", True)
    CF = index("I")
    blocks = {"IS": ("A", IS), "BS": ("E", BS), "CF": ("I", CF)}

    def ref(block, sec, label, year):
        lc, idx = blocks[block]
        r = idx[(sec, _norm(label))]
        col = L(CI(lc) + (1 if year == "cy" else 2))
        return col + str(r)

    # (key, [(block, section, label, factor)])
    spec = [
        ("revenue", [("IS", None, "Receita operacional líquida", 1)]),
        ("cost_of_sales", [("IS", None, "Custo das mercadorias vendidas", -1)]),
        ("distribution_costs", [("IS", None, "Despesas com vendas", -1)]),
        ("administrative_expenses", [("IS", None, "Despesas gerais e administrativas", -1)]),
        ("other_operating_income", [("IS", None, "Outras receitas (despesas) operacionais, líquidas", 1)]),
        ("finance_income", [("IS", None, "Receitas financeiras", 1)]),
        ("finance_costs", [("IS", None, "Despesas financeiras", -1)]),
        ("income_tax_expense", [("IS", None, "IR e CSLL correntes", -1),
                                ("IS", None, "IR e CSLL diferidos", -1)]),
        ("property_plant_equipment", [("BS", "anc", "Imobilizado", 1)]),
        ("right_of_use_assets", [("BS", "anc", "Direito de uso", 1)]),
        ("intangible_assets", [("BS", "anc", "Intangível", 1)]),
        ("investments", []),
        ("deferred_tax_assets", [("BS", "anc", "IR e CS diferidos", 1)]),
        ("inventories", [("BS", "ac", "Estoques", 1)]),
        ("trade_receivables", [("BS", "ac", "Contas a receber de clientes", 1)]),
        ("other_current_assets", [("BS", "ac", "Impostos a recuperar", 1),
                                  ("BS", "ac", "Outros créditos", 1)]),
        ("cash_and_cash_equivalents", [("BS", "ac", "Caixa e equivalentes de caixa", 1)]),
        ("trade_payables", [("BS", "pc", "Fornecedores", 1)]),
        ("other_current_liabilities", [("BS", "pc", "Obrigações trabalhistas e outras contas a pagar", 1)]),
        ("current_borrowings", [("BS", "pc", "Empréstimos, financiamentos e arrendamentos", 1)]),
        ("current_tax_payable", [("BS", "pc", "IR e CS a pagar", 1)]),
        ("non_current_borrowings", [("BS", "pnc", "Empréstimos, financiamentos e arrendamentos", 1)]),
        ("provisions", [("BS", "pnc", "Provisões", 1)]),
        ("deferred_tax_liabilities", [("BS", "pnc", "IR e CS diferidos", 1)]),
        ("share_capital", [("BS", "pl", "Capital social", 1)]),
        ("retained_earnings", [("BS", "pl", "Reservas de lucros", 1)]),
        ("other_reserves", [("BS", "pl", "Ajustes de avaliação patrimonial", 1)]),
        ("net_cash_operating", [("CF", None, "Caixa líquido das atividades operacionais", 1)]),
        ("net_cash_investing", [("CF", None, "Caixa líquido das atividades de investimento", 1)]),
        ("net_cash_financing", [("CF", None, "Caixa líquido das atividades de financiamento", 1)]),
        ("fx_effect_on_cash", [("CF", None, "Efeito da variação cambial sobre o caixa", 1)]),
        ("cash_at_beginning", [("CF", None, "Caixa no início do exercício", 1)]),
        ("dividends", [("CF", None, "Dividendos declarados", -1)]),
    ]

    fields = {}
    for key, parts in spec:
        fields[key] = {}
        for year in ("cy", "py"):
            cells = []
            total = 0
            for block, sec, label, f in parts:
                c = ref(block, sec, label, year)
                v = val(re.match(r"[A-Z]+", c).group(), int(re.search(r"\d+", c).group()))
                total += f * v
                cells.append({"sheet": SHEET, "cell": c, "factor": f})
            fields[key][year] = {"value": total, "cells": cells}

    # Printed totals
    totals = []

    def st(block, sec, label, formula, factor=1, years=("cy", "py")):
        for y in years:
            totals.append({"sheet": SHEET, "cell": ref(block, sec, label, y), "year": y,
                           "factor": factor, "formula": formula})

    st("IS", None, "Lucro bruto", "gross_profit")
    st("IS", None, "Resultado operacional", "operating_profit")
    st("IS", None, "Resultado antes do IR e CSLL", "profit_before_tax")
    st("IS", None, "Lucro líquido do exercício", "profit_for_year")
    st("BS", "ac", "Total do ativo circulante", "total_current_assets")
    st("BS", "anc", "Total do ativo não circulante", "total_non_current_assets")
    st("BS", "anc", "Total do ativo", "total_assets")
    st("BS", "pc", "Total do passivo circulante", "total_current_liabilities")
    st("BS", "pnc", "Total do passivo não circulante", "total_non_current_liabilities")
    st("BS", "pnc", "Total liabilities", "total_liabilities")
    st("BS", "pl", "Total do patrimônio líquido", "total_equity")
    st("BS", "pl", "Total do passivo e patrimônio líquido", "total_equity_and_liabilities")
    st("CF", None, "Aumento (redução) de caixa", "net_change_in_cash")
    st("CF", None, "Caixa no fim do exercício", "cash_at_end")
    # Retained-earnings roll-forward block (cash-flow side of the sheet)
    totals.append({"sheet": SHEET, "cell": ref("CF", None, "Saldo inicial", "cy"), "year": "py",
                   "factor": 1, "formula": "retained_earnings"})
    st("CF", None, "Saldo final", "retained_earnings")
    # Net profit repeated in the roll-forward (second "Lucro líquido do exercício" is in column I)
    ni_row = CF[(None, _norm("Lucro líquido do exercício"))]
    for y, col in (("cy", "J"), ("py", "K")):
        totals.append({"sheet": SHEET, "cell": f"{col}{ni_row}", "year": y,
                       "factor": 1, "formula": "profit_for_year"})

    # The prior-year opening balance of profit reserves has no balance sheet to compare with
    ignored = [{
        "sheet": SHEET,
        "cell": ref("CF", None, "Saldo inicial", "py"),
        "reason": "Opening profit reserves at 1 Jan 2024 (start of prior year); no balance sheet "
                  "for that date in the template, so it is not a template line.",
    }]

    # Header data
    title = str(ws["A1"].value)
    name = re.split(r"\s+-\s+", title)[0].strip()
    hdr_row = None
    for r in range(1, 10):
        if re.match(r"\s*\d{1,2}/\d{1,2}/\d{4}", str(ws[f"F{r}"].value or "")):
            hdr_row = r
            break
    period_end = {"cy": _date(ws[f"F{hdr_row}"].value), "py": _date(ws[f"G{hdr_row}"].value)}

    return {
        "entity_code": "TW-BR01",
        "source_entity_name": name,
        "currency": "BRL",
        "unit": 1000,
        "period_end": period_end,
        "fields": fields,
        "source_totals": totals,
        "ignored": ignored,
        "source_exceptions": [],
        "notes": (
            "Brazilian GAAP workbook in R$ thousands, single sheet with income statement, balance sheet "
            "and cash flow side by side. Expenses and the dividend are shown in brackets and were flipped "
            "with factor -1; other operating income (expense) is kept as printed (net expense, negative). "
            "Income tax is current plus deferred IR/CSLL (a credit in 2024 reduces the charge). "
            "Taxes recoverable and other receivables are combined into other current assets. Deferred tax "
            "assets and liabilities are taken from the lines labelled IR e CS diferidos in the matching "
            "section. Profit reserves map to retained earnings and the equity valuation adjustments to "
            "other reserves. The roll-forward block under the cash flow ties to retained earnings and "
            "profit and is used as source totals. The workbook is internally consistent."
        ),
    }
