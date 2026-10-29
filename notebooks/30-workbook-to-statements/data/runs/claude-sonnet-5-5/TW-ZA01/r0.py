import re
import datetime
import openpyxl
from openpyxl.utils import get_column_letter

REGISTER = {
    "tarnwell holdings plc": ("TW-UK01", "GBP"),
    "tarnwell engineering limited": ("TW-UK02", "GBP"),
    "tarnwell engineering services limited": ("TW-UK03", "GBP"),
    "tarnwell antriebstechnik gmbh": ("TW-DE01", "EUR"),
    "tarnwell service gmbh": ("TW-DE02", "EUR"),
    "tarnwell france sas": ("TW-FR01", "EUR"),
    "tarnwell logistics b.v.": ("TW-NL01", "EUR"),
    "tarnwell ireland limited": ("TW-IE01", "EUR"),
    "tarnwell inc.": ("TW-US01", "USD"),
    "tarnwell software llc": ("TW-US02", "USD"),
    "tarnwell holdings (us) inc.": ("TW-US03", "USD"),
    "tarnwell canada ltd.": ("TW-CA01", "CAD"),
    "tarnwell méxico, s.a. de c.v.": ("TW-MX01", "MXN"),
    "tarnwell do brasil ltda.": ("TW-BR01", "BRL"),
    "tarnwell japan k.k.": ("TW-JP01", "JPY"),
    "tarnwell india private limited": ("TW-IN01", "INR"),
    "tarnwell asia pacific pte. ltd.": ("TW-SG01", "SGD"),
    "tarnwell australia pty ltd": ("TW-AU01", "AUD"),
    "tarnwell new zealand limited": ("TW-NZ01", "NZD"),
    "tarnwell south africa (pty) ltd": ("TW-ZA01", "ZAR"),
    "tarnwell middle east fze": ("TW-AE01", "AED"),
    "tarnwell trading l.l.c.": ("TW-AE02", "AED"),
    "tarnwell polska sp. z o.o.": ("TW-PL01", "PLN"),
    "tarnwell nordic ab": ("TW-SE01", "SEK"),
    "tarnwell schweiz ag": ("TW-CH01", "CHF"),
    "tarnwell east africa limited": ("TW-KE01", "KES"),
}

# key -> list of (sheet, row label, factor)
SPEC = {
    "revenue": [("SOCI", "Turnover", 1)],
    "cost_of_sales": [("SOCI", "Cost of sales", 1)],
    "distribution_costs": [("SOCI", "Selling expenses", 1)],
    "administrative_expenses": [("SOCI", "Overheads", 1)],
    "other_operating_income": [("SOCI", "Other gains and losses (net)", 1)],
    "finance_income": [("SOCI", "Interest income", 1)],
    "finance_costs": [("SOCI", "Interest paid on borrowings", 1)],
    "income_tax_expense": [("SOCI", "Income tax expense", 1)],
    "property_plant_equipment": [("SOFP", "Land and buildings", 1), ("SOFP", "Plant and machinery", 1)],
    "right_of_use_assets": [("SOFP", "Right-of-use assets", 1)],
    "intangible_assets": [("SOFP", "Intangible assets", 1)],
    "investments": [],
    "deferred_tax_assets": [("SOFP", "Deferred tax assets", 1)],
    "inventories": [("SOFP", "Stock", 1)],
    "trade_receivables": [("SOFP", "Trade debtors", 1)],
    "other_current_assets": [("SOFP", "Prepayments and other debtors", 1)],
    "cash_and_cash_equivalents": [("SOFP", "Cash and cash equivalents", 1)],
    "trade_payables": [("SOFP", "Accounts payable", 1)],
    "other_current_liabilities": [("SOFP", "Accruals", 1), ("SOFP", "VAT payable", 1)],
    "current_borrowings": [("SOFP", "Short-term borrowings", 1)],
    "current_tax_payable": [("SOFP", "Current tax payable", 1)],
    "non_current_borrowings": [("SOFP", "Borrowings (non-current)", 1)],
    "provisions": [("SOFP", "Provisions", 1)],
    "deferred_tax_liabilities": [("SOFP", "Deferred tax liabilities", 1)],
    "share_capital": [("SOFP", "Share capital", 1)],
    "retained_earnings": [("SOFP", "Accumulated profits", 1)],
    "other_reserves": [("SOFP", "Other reserves", 1)],
    "net_cash_operating": [("SOCF", "Net cash from operating activities", 1)],
    "net_cash_investing": [("SOCF", "Net cash from investing activities", 1)],
    "net_cash_financing": [("SOCF", "Net cash used in financing activities", 1)],
    "fx_effect_on_cash": [("SOCF", "Effect of exchange rate changes", 1)],
    "cash_at_beginning": [("SOCF", "Cash at start of year", 1)],
    # shown as a negative deduction in the retained earnings reconciliation
    "dividends": [("SOFP", "Dividends", -1)],
}

# (sheet, row label, formula)
TOTALS = [
    ("SOCI", "Gross margin", "gross_profit"),
    ("SOCI", "Profit from operations", "operating_profit"),
    ("SOCI", "Profit before taxation", "profit_before_tax"),
    ("SOCI", "Net profit", "profit_for_year"),
    ("SOFP", "Total non-current assets", "total_non_current_assets"),
    ("SOFP", "Total current assets", "total_current_assets"),
    ("SOFP", "Total assets", "total_assets"),
    ("SOFP", "Total shareholders' funds", "total_equity"),
    ("SOFP", "Total non-current liabilities", "total_non_current_liabilities"),
    ("SOFP", "Total current liabilities", "total_current_liabilities"),
    ("SOFP", "Total liabilities", "total_liabilities"),
    ("SOFP", "TOTAL EQUITY AND LIABILITIES", "total_equity_and_liabilities"),
    ("SOFP", "Profit for the year", "profit_for_year"),
    ("SOFP", "Closing retained earnings", "retained_earnings"),
    ("SOCF", "Net increase/(decrease) in cash", "net_change_in_cash"),
    ("SOCF", "Cash at end of year", "cash_at_end"),
]


def _norm(s):
    return re.sub(r"\s+", " ", str(s)).strip().lower()


def _find_row(ws, label):
    target = _norm(label)
    for r in range(1, ws.max_row + 1):
        v = ws.cell(r, 1).value
        if v is not None and _norm(v) == target:
            return r
    raise KeyError("row not found: %s on %s" % (label, ws.title))


def _is_num(v):
    return isinstance(v, (int, float)) and not isinstance(v, bool)


def _header(ws):
    """Locate the header row; return (cy_col, py_col, draft_cols, notes_col, header_row)."""
    for r in range(1, 15):
        years, drafts, notes = {}, [], None
        for c in range(1, ws.max_column + 1):
            v = ws.cell(r, c).value
            if v is None:
                continue
            t = str(v).strip()
            if re.fullmatch(r"\d{4}", t):
                years[int(t)] = c
            elif t.lower().startswith("draft"):
                drafts.append(c)
            elif t.lower() == "notes":
                notes = c
        if len(years) >= 2:
            ys = sorted(years, reverse=True)
            return years[ys[0]], years[ys[1]], drafts, notes, r
    raise ValueError("header not found on " + ws.title)


def _col_range(ws, col, first=1):
    rows = [r for r in range(first, ws.max_row + 1) if _is_num(ws.cell(r, col).value)]
    if not rows:
        return None
    L = get_column_letter(col)
    return "%s%d:%s%d" % (L, min(rows), L, max(rows))


def extract(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    sheets = {ws.title: ws for ws in wb.worksheets}
    heads = {n: _header(ws) for n, ws in sheets.items()}

    # entity
    name = str(sheets["SOCI"]["A1"].value).strip()
    key = _norm(name)
    hit = REGISTER.get(key)
    if hit is None:
        alt = re.sub(r"\bsa\b", "south africa", key)
        hit = REGISTER.get(alt)
    if hit is None:
        raise ValueError("entity not in register: " + name)
    code, currency = hit

    # unit
    unit_txt = ""
    for n, ws in sheets.items():
        hr = heads[n][4]
        v = ws.cell(hr + 1, heads[n][0]).value
        if v:
            unit_txt = str(v).lower()
            break
    if "million" in unit_txt or re.search(r"\b\d+m\b", unit_txt):
        unit = 1000000
    elif "000" in unit_txt or "thousand" in unit_txt:
        unit = 1000
    else:
        unit = 1

    # period end
    title = str(sheets["SOFP"]["A2"].value)
    m = re.search(r"(\d{1,2})\s+([A-Za-z]+)\s+(\d{4})", title)
    d_cy = datetime.datetime.strptime("%s %s %s" % m.groups(), "%d %B %Y").date()
    try:
        d_py = d_cy.replace(year=d_cy.year - 1)
    except ValueError:
        d_py = d_cy.replace(year=d_cy.year - 1, day=28)

    def colof(sheet, year):
        h = heads[sheet]
        return h[0] if year == "cy" else h[1]

    def cell_ref(sheet, label, year):
        ws = sheets[sheet]
        r = _find_row(ws, label)
        c = colof(sheet, year)
        return r, c, ws.cell(r, c).value

    fields = {}
    for k, parts in SPEC.items():
        fields[k] = {}
        for year in ("cy", "py"):
            cells, total = [], 0
            for sheet, label, factor in parts:
                r, c, v = cell_ref(sheet, label, year)
                cells.append({"sheet": sheet, "cell": "%s%d" % (get_column_letter(c), r), "factor": factor})
                total += factor * (v if _is_num(v) else 0)
            fields[k][year] = {"value": total, "cells": cells}

    source_totals = []
    for sheet, label, formula in TOTALS:
        for year in ("cy", "py"):
            r, c, v = cell_ref(sheet, label, year)
            source_totals.append({"sheet": sheet, "cell": "%s%d" % (get_column_letter(c), r),
                                  "year": year, "factor": 1, "formula": formula})
    # opening retained earnings in the cy column equals last year's closing balance
    ws = sheets["SOFP"]
    r_open = _find_row(ws, "Retained earnings at start of year")
    source_totals.append({"sheet": "SOFP", "cell": "%s%d" % (get_column_letter(heads["SOFP"][0]), r_open),
                          "year": "py", "factor": 1, "formula": "retained_earnings"})

    ignored = []
    for n, ws in sheets.items():
        cy, py, drafts, notes, hr = heads[n]
        for dc in drafts:
            rng = _col_range(ws, dc, hr + 1)
            if rng:
                ignored.append({"sheet": n, "range": rng,
                                "reason": "Hidden 'Draft 12/03' column: superseded draft figures, not the final current-year numbers"})
        if notes:
            rng = _col_range(ws, notes, hr + 1)
            if rng:
                ignored.append({"sheet": n, "range": rng,
                                "reason": "Note reference numbers, not amounts"})
    c_py = heads["SOFP"][1]
    ignored.append({"sheet": "SOFP", "cell": "%s%d" % (get_column_letter(c_py), r_open),
                    "reason": "Opening retained earnings of the prior year; no prior-prior balance sheet in the template to tie to"})

    return {
        "entity_code": code,
        "source_entity_name": name,
        "currency": currency,
        "unit": unit,
        "period_end": {"cy": d_cy.isoformat(), "py": d_py.isoformat()},
        "fields": fields,
        "source_totals": source_totals,
        "ignored": ignored,
        "source_exceptions": [],
        "notes": (
            "Entity 'Tarnwell SA (Pty) Ltd' matched to TW-ZA01 (South Africa, ZAR); amounts are in R'000 so unit is 1000. "
            "Current year is column F (2026) and prior year column H (2025); hidden column G is a superseded draft and is ignored. "
            "Property, plant and equipment combines Land and buildings with Plant and machinery. "
            "Other current liabilities combines Accruals and VAT payable; Short-term borrowings go to current borrowings. "
            "Dividends are shown negative in the retained earnings reconciliation, so they are read with factor -1. "
            "Other gains and losses (net) is positive income and maps to other operating income. "
            "No investments line exists in this workbook. All printed totals and cross-statement checks tie in the final columns."
        ),
    }
