import calendar
import math
import re
from datetime import date, datetime
from numbers import Real

import openpyxl


KEYS = (
    "revenue", "cost_of_sales", "distribution_costs",
    "administrative_expenses", "other_operating_income", "finance_income",
    "finance_costs", "income_tax_expense", "property_plant_equipment",
    "right_of_use_assets", "intangible_assets", "investments",
    "deferred_tax_assets", "inventories", "trade_receivables",
    "other_current_assets", "cash_and_cash_equivalents", "trade_payables",
    "other_current_liabilities", "current_borrowings", "current_tax_payable",
    "non_current_borrowings", "provisions", "deferred_tax_liabilities",
    "share_capital", "retained_earnings", "other_reserves",
    "net_cash_operating", "net_cash_investing", "net_cash_financing",
    "fx_effect_on_cash", "cash_at_beginning", "dividends",
)

TOTAL_FORMULAS = {
    "gross_profit": "revenue - cost_of_sales",
    "operating_profit": (
        "gross_profit - distribution_costs - administrative_expenses"
        " + other_operating_income"
    ),
    "profit_before_tax": "operating_profit + finance_income - finance_costs",
    "profit_for_year": "profit_before_tax - income_tax_expense",
    "total_non_current_assets": (
        "property_plant_equipment + right_of_use_assets + intangible_assets"
        " + investments + deferred_tax_assets"
    ),
    "total_current_assets": (
        "inventories + trade_receivables + other_current_assets"
        " + cash_and_cash_equivalents"
    ),
    "total_assets": "total_non_current_assets + total_current_assets",
    "total_current_liabilities": (
        "trade_payables + other_current_liabilities + current_borrowings"
        " + current_tax_payable"
    ),
    "total_non_current_liabilities": (
        "non_current_borrowings + provisions + deferred_tax_liabilities"
    ),
    "total_liabilities": (
        "total_current_liabilities + total_non_current_liabilities"
    ),
    "total_equity": "share_capital + retained_earnings + other_reserves",
    "total_equity_and_liabilities": "total_equity + total_liabilities",
    "net_change_in_cash": (
        "net_cash_operating + net_cash_investing + net_cash_financing"
    ),
    "cash_at_end": (
        "cash_at_beginning + net_change_in_cash + fx_effect_on_cash"
    ),
}


def extract(path) -> dict:
    wb = openpyxl.load_workbook(path, data_only=True)
    try:
        return _extract(wb)
    finally:
        wb.close()


def _extract(wb):
    pl = wb["P&L"]
    bs = wb["Balance Sheet"]
    cf = wb["Cash Flow"]

    def normalise(value):
        return re.sub(r"\s+", " ", str(value or "").strip()).casefold()

    def rows(sheet, label):
        wanted = normalise(label)
        return [
            row for row in range(1, sheet.max_row + 1)
            if normalise(sheet.cell(row, 1).value) == wanted
        ]

    def row(sheet, label):
        matches = rows(sheet, label)
        if len(matches) != 1:
            raise ValueError(
                f"Expected one {label!r} row on {sheet.title!r}; "
                f"found {len(matches)}."
            )
        return matches[0]

    def number(sheet_name, coordinate):
        value = wb[sheet_name][coordinate].value
        if (
            isinstance(value, bool)
            or not isinstance(value, Real)
            or not math.isfinite(value)
        ):
            raise ValueError(
                f"{sheet_name}!{coordinate} must contain a finite number "
                "(formula cells require cached results)."
            )
        return int(value) if isinstance(value, int) else float(value)

    def reference(sheet, row_number, column, factor=1):
        coordinate = sheet.cell(row_number, column).coordinate
        number(sheet.title, coordinate)
        return {
            "sheet": sheet.title,
            "cell": coordinate,
            "factor": factor,
        }

    def labelled_reference(sheet, label, column, factor=1):
        return reference(sheet, row(sheet, label), column, factor)

    def sum_references(references):
        return sum(
            item["factor"] * number(item["sheet"], item["cell"])
            for item in references
        )

    # Match the legal name independently of punctuation and spacing.
    entity_names = [
        cell.value
        for cells in bs.iter_rows()
        for cell in cells
        if isinstance(cell.value, str)
        and re.sub(r"[^a-z0-9]", "", cell.value.casefold())
        == "tarnwellsoftwarellc"
    ]
    if len(entity_names) != 1:
        raise ValueError("Workbook does not uniquely identify Tarnwell Software LLC.")

    for sheet in (pl, bs, cf):
        if not any(
            isinstance(cell.value, str) and "in us$" in cell.value.casefold()
            for cells in sheet.iter_rows()
            for cell in cells
        ):
            raise ValueError(f"Missing USD/unit declaration on {sheet.title!r}.")

    def parse_date(value):
        if isinstance(value, datetime):
            return value.date()
        if isinstance(value, date):
            return value
        if isinstance(value, str):
            for fmt in ("%b %d, %Y", "%B %d, %Y", "%Y-%m-%d"):
                try:
                    return datetime.strptime(value.strip(), fmt).date()
                except ValueError:
                    pass
        return None

    bs_headers = []
    for cells in bs.iter_rows():
        dates = [
            (cell.column, parse_date(cell.value))
            for cell in cells if cell.column > 1
        ]
        dates = [(column, value) for column, value in dates if value]
        if len(dates) == 2:
            bs_headers.append(dates)
    if len(bs_headers) != 1:
        raise ValueError("Cannot uniquely locate the balance-sheet date headers.")

    ordered_dates = sorted(bs_headers[0], key=lambda item: item[1], reverse=True)
    bs_columns = {"cy": ordered_dates[0][0], "py": ordered_dates[1][0]}
    period_dates = {"cy": ordered_dates[0][1], "py": ordered_dates[1][1]}

    cy_end = period_dates["cy"]
    expected_py = date(
        cy_end.year - 1,
        cy_end.month,
        min(cy_end.day, calendar.monthrange(cy_end.year - 1, cy_end.month)[1]),
    )
    if period_dates["py"] != expected_py:
        raise ValueError("The balance-sheet dates are not one year apart.")

    monthly_headers = []
    for cells in pl.iter_rows():
        months = [
            (cell.column, cell.value.date()
             if isinstance(cell.value, datetime) else cell.value)
            for cell in cells
            if isinstance(cell.value, (date, datetime))
        ]
        if len(months) == 12:
            monthly_headers.append(months)
    if len(monthly_headers) != 1:
        raise ValueError("Expected one complete 12-month P&L header.")

    months = sorted(monthly_headers[0], key=lambda item: item[1])
    end_month_index = cy_end.year * 12 + cy_end.month - 1
    expected_months = [
        date(index // 12, index % 12 + 1, 1)
        for index in range(end_month_index - 11, end_month_index + 1)
    ]
    if [value for _, value in months] != expected_months:
        raise ValueError("P&L months do not match the current fiscal year.")
    monthly_columns = [column for column, _ in months]

    def find_header_column(sheet, text):
        matches = [
            cell.column
            for cells in sheet.iter_rows()
            for cell in cells
            if normalise(cell.value) == normalise(text)
        ]
        if len(matches) != 1:
            raise ValueError(f"Cannot uniquely locate {text!r} on {sheet.title!r}.")
        return matches[0]

    pl_prior_column = find_header_column(
        pl, f"FY{period_dates['py'].year} Total"
    )
    cf_columns = {
        year: find_header_column(cf, f"FY{period_dates[year].year}")
        for year in ("cy", "py")
    }

    fields = {
        key: {year: {"value": 0, "cells": []} for year in ("cy", "py")}
        for key in KEYS
    }

    def assign(key, year, references):
        fields[key][year] = {
            "value": sum_references(references),
            "cells": references,
        }

    pl_mapping = {
        "revenue": ("Net sales", 1),
        "cost_of_sales": ("Cost of goods sold", 1),
        "distribution_costs": ("Selling expenses", 1),
        "administrative_expenses": ("General and administrative expenses", 1),
        "other_operating_income": ("Other operating (income) expense, net", -1),
        "finance_income": ("Interest income", 1),
        "finance_costs": ("Interest expense", 1),
        "income_tax_expense": ("Provision for income taxes", 1),
    }
    for key, (label, factor) in pl_mapping.items():
        for year, columns in (
            ("cy", monthly_columns),
            ("py", [pl_prior_column]),
        ):
            assign(key, year, [
                labelled_reference(pl, label, column, factor)
                for column in columns
            ])

    bs_mapping = {
        "property_plant_equipment": ["Property and equipment, net"],
        "right_of_use_assets": ["Operating lease right-of-use assets"],
        "intangible_assets": ["Goodwill", "Intangible assets, net"],
        "trade_receivables": ["Accounts receivable, net"],
        "other_current_assets": ["Prepaid expenses and other current assets"],
        "cash_and_cash_equivalents": ["Cash and cash equivalents"],
        "trade_payables": ["Accounts payable"],
        "other_current_liabilities": ["Accrued expenses", "Deferred revenue"],
        "current_borrowings": ["Current portion of long-term debt"],
        "current_tax_payable": ["Income taxes payable"],
        "non_current_borrowings": ["Long-term debt"],
        "provisions": ["Other long-term liabilities"],
        "share_capital": ["Members' capital"],
        "retained_earnings": ["Retained earnings"],
    }
    for key, labels in bs_mapping.items():
        for year, column in bs_columns.items():
            assign(key, year, [
                labelled_reference(bs, label, column) for label in labels
            ])

    # The same label occurs in the asset and liability sections.
    asset_total_row = row(bs, "Total assets")
    liability_total_row = row(bs, "Total liabilities")
    tax_rows = rows(bs, "Deferred income taxes")
    asset_tax_rows = [r for r in tax_rows if r < asset_total_row]
    liability_tax_rows = [
        r for r in tax_rows if asset_total_row < r < liability_total_row
    ]
    if len(asset_tax_rows) != 1 or len(liability_tax_rows) != 1:
        raise ValueError("Cannot distinguish deferred tax assets and liabilities.")
    for year, column in bs_columns.items():
        assign("deferred_tax_assets", year, [
            reference(bs, asset_tax_rows[0], column)
        ])
        assign("deferred_tax_liabilities", year, [
            reference(bs, liability_tax_rows[0], column)
        ])

    cf_mapping = {
        "net_cash_operating": "Net cash provided by operating activities",
        "net_cash_investing": "Net cash used in investing activities",
        "net_cash_financing":
            "Net cash provided by (used in) financing activities",
        "fx_effect_on_cash": "Effect of exchange rate changes on cash",
        "cash_at_beginning": "Cash and cash equivalents, beginning of year",
        "dividends": "Dividends declared",
    }
    for key, label in cf_mapping.items():
        for year, column in cf_columns.items():
            assign(key, year, [
                labelled_reference(
                    cf, label, column, -1 if key == "dividends" else 1
                )
            ])

    source_totals = []

    def add_total(sheet, label, column, year, formula=None, components=None):
        item = labelled_reference(sheet, label, column)
        item["year"] = year
        if formula is not None:
            item["formula"] = formula
        else:
            item["sum_of"] = components
        source_totals.append(item)

    # Monthly subtotals are checked against their monthly source components;
    # an individual month cannot be compared to an annual template total.
    pl_subtotals = {
        "Gross profit": (
            "gross_profit",
            [("Net sales", 1), ("Cost of goods sold", -1)],
        ),
        "Total SG&A": (
            "distribution_costs + administrative_expenses",
            [("Selling expenses", 1),
             ("General and administrative expenses", 1)],
        ),
        "Operating income": (
            "operating_profit",
            [("Gross profit", 1), ("Total SG&A", -1),
             ("Other operating (income) expense, net", -1)],
        ),
        "Income before income taxes": (
            "profit_before_tax",
            [("Operating income", 1), ("Interest income", 1),
             ("Interest expense", -1)],
        ),
        "Net income": (
            "profit_for_year",
            [("Income before income taxes", 1),
             ("Provision for income taxes", -1)],
        ),
    }
    for label, (formula, components) in pl_subtotals.items():
        for column in monthly_columns:
            add_total(
                pl, label, column, "cy",
                components=[
                    labelled_reference(pl, component, column, factor)
                    for component, factor in components
                ],
            )
        add_total(pl, label, pl_prior_column, "py", formula=formula)

    bs_subtotals = {
        "Total current assets": "total_current_assets",
        "Total assets": "total_assets",
        "Total current liabilities": "total_current_liabilities",
        "Total liabilities": "total_liabilities",
        "Total stockholders' equity": "total_equity",
        "Total liabilities and stockholders' equity":
            "total_equity_and_liabilities",
    }
    for label, formula in bs_subtotals.items():
        for year, column in bs_columns.items():
            add_total(bs, label, column, year, formula=formula)

    cf_subtotals = {
        "Net increase (decrease) in cash": "net_change_in_cash",
        "Cash and cash equivalents, end of year": "cash_at_end",
        "Net income": "profit_for_year",
        "Retained earnings, end of year": "retained_earnings",
    }
    for label, formula in cf_subtotals.items():
        for year, column in cf_columns.items():
            add_total(cf, label, column, year, formula=formula)

    add_total(
        cf, "Retained earnings, beginning of year", cf_columns["cy"], "cy",
        components=[
            labelled_reference(
                bs, "Retained earnings", bs_columns["py"]
            )
        ],
    )
    # Prior opening retained earnings has no template input. Check it against
    # the prior-year movement statement, preserving signed source dividends.
    add_total(
        cf, "Retained earnings, beginning of year", cf_columns["py"], "py",
        components=[
            labelled_reference(
                cf, "Retained earnings, end of year", cf_columns["py"]
            ),
            labelled_reference(cf, "Net income", cf_columns["py"], -1),
            labelled_reference(cf, "Dividends declared", cf_columns["py"], -1),
        ],
    )

    def expression_value(expression, values):
        tokens = expression.split()
        if not tokens or len(tokens) % 2 != 1:
            raise ValueError(f"Invalid formula: {expression!r}")
        result = values[tokens[0]]
        for index in range(1, len(tokens), 2):
            operator, name = tokens[index:index + 2]
            if operator not in ("+", "-"):
                raise ValueError(f"Invalid operator: {operator!r}")
            result += values[name] if operator == "+" else -values[name]
        return result

    calculated = {}
    for year in ("cy", "py"):
        values = {key: fields[key][year]["value"] for key in KEYS}
        for name, formula in TOTAL_FORMULAS.items():
            values[name] = expression_value(formula, values)
        calculated[year] = values

    exceptions = []

    def check(check_id, actual, expected, explanation):
        difference = actual - expected
        if not math.isclose(difference, 0, rel_tol=0, abs_tol=1e-7):
            exceptions.append({
                "check": check_id,
                "difference": difference,
                "explanation": explanation,
            })

    for item in source_totals:
        actual = item["factor"] * number(item["sheet"], item["cell"])
        expected = (
            expression_value(item["formula"], calculated[item["year"]])
            if "formula" in item
            else sum_references(item["sum_of"])
        )
        check(
            f"T {item['sheet']}!{item['cell']} {item['year']}",
            actual, expected,
            "The printed source amount differs from its stated components "
            "or the corresponding template expression. Difference is "
            "normalised printed amount minus expected amount; source values "
            "have been preserved.",
        )

    for year in ("cy", "py"):
        values = calculated[year]
        check(
            f"R1 balance {year}",
            values["total_assets"], values["total_equity_and_liabilities"],
            "Mapped assets minus mapped equity and liabilities; source "
            "values have been preserved.",
        )
        check(
            f"R2 cash_at_end {year}",
            values["cash_at_end"], values["cash_and_cash_equivalents"],
            "Calculated cash-flow closing cash minus balance-sheet cash; "
            "source values have been preserved.",
        )

    check(
        "R3 opening cash cy",
        calculated["cy"]["cash_at_beginning"],
        calculated["py"]["cash_and_cash_equivalents"],
        "Current opening cash minus prior balance-sheet cash; source "
        "values have been preserved.",
    )
    check(
        "R4 retained earnings cy",
        calculated["cy"]["retained_earnings"],
        calculated["py"]["retained_earnings"]
        + calculated["cy"]["profit_for_year"]
        - calculated["cy"]["dividends"],
        "Current retained earnings minus prior retained earnings plus "
        "current profit less dividends; source values have been preserved.",
    )

    # Ensure every number in each amount column is explicitly accounted for.
    covered = set()
    for years in fields.values():
        for line in years.values():
            covered.update(
                (item["sheet"], item["cell"]) for item in line["cells"]
            )
    for item in source_totals:
        covered.add((item["sheet"], item["cell"]))
        covered.update(
            (component["sheet"], component["cell"])
            for component in item.get("sum_of", [])
        )

    amount_columns = {}
    for sheet_name, coordinate in covered:
        cell = wb[sheet_name][coordinate]
        key = (sheet_name, cell.column)
        amount_columns[key] = min(amount_columns.get(key, cell.row), cell.row)

    for (sheet_name, column), first_row in amount_columns.items():
        sheet = wb[sheet_name]
        for row_number in range(first_row, sheet.max_row + 1):
            cell = sheet.cell(row_number, column)
            if (
                isinstance(cell.value, Real)
                and not isinstance(cell.value, bool)
                and (sheet_name, cell.coordinate) not in covered
            ):
                raise ValueError(
                    f"Unclassified numeric cell {sheet_name}!{cell.coordinate}."
                )

    return {
        "entity_code": "TW-US02",
        "source_entity_name": entity_names[0],
        "currency": "USD",
        "unit": 1,
        "period_end": {
            year: value.isoformat() for year, value in period_dates.items()
        },
        "fields": fields,
        "source_totals": source_totals,
        "ignored": [],
        "source_exceptions": exceptions,
        "notes": (
            "Current-year profit or loss aggregates the twelve fiscal months "
            "October through September; the comparative uses the FY total "
            "column. Amounts remain in US dollars. Other operating income "
            "reverses the source expense convention, and dividends reverse "
            "the negative equity-movement convention. Goodwill and other "
            "intangibles are combined; deferred revenue is included in other "
            "current liabilities. Other long-term liabilities are provisionally "
            "classified as provisions/employee benefit obligations: the source "
            "does not disclose their nature, so this classification requires "
            "reviewer confirmation. Inputs absent from the workbook are zero. "
            "Exception differences are actual minus expected."
        ),
    }
