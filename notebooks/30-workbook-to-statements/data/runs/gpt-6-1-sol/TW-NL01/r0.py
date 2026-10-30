import math
import re
from datetime import date, datetime
from numbers import Real

import openpyxl


def extract(path) -> dict:
    """Extract Tarnwell Logistics B.V.'s financial statements, in EUR thousands."""
    wb = openpyxl.load_workbook(path, data_only=True)
    try:
        ws = wb["FS"]

        def normalize(text):
            return re.sub(r"[^a-z0-9]", "", str(text).casefold())

        def row_for(label):
            matches = [
                row
                for row in range(1, ws.max_row + 1)
                if normalize(ws.cell(row, 1).value) == normalize(label)
            ]
            if len(matches) != 1:
                raise ValueError(
                    f"Expected exactly one row labelled {label!r}; found {matches}"
                )
            return matches[0]

        def number(address):
            value = ws[address].value
            if (
                isinstance(value, bool)
                or not isinstance(value, Real)
                or not math.isfinite(value)
            ):
                raise ValueError(f"{ws.title}!{address} is not a finite number")
            return int(value) if isinstance(value, int) else float(value)

        def reference(address, factor=1):
            number(address)
            return {"sheet": ws.title, "cell": address, "factor": factor}

        def ref_sum(refs):
            return sum(item["factor"] * number(item["cell"]) for item in refs)

        # Identify this subsidiary from the name actually present in the workbook.
        titles = [
            cell.value
            for row in ws.iter_rows()
            for cell in row
            if isinstance(cell.value, str)
            and re.search(r"\bfinancial statements\b", cell.value, re.I)
        ]
        if len(titles) != 1:
            raise ValueError("Cannot uniquely identify the financial statement title")
        source_entity_name = re.split(
            r"\s*[-–—:]\s*financial statements\b", titles[0], maxsplit=1,
            flags=re.I,
        )[0].strip()
        if normalize(source_entity_name) != normalize("Tarnwell Logistics B.V."):
            raise ValueError(f"Unexpected entity: {source_entity_name!r}")

        unit_matches = []
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str):
                    match = re.fullmatch(
                        r"\s*([A-Z]{3})\s*[x×]\s*([\d,]+)\s*", cell.value
                    )
                    if match:
                        unit_matches.append(
                            (match.group(1), int(match.group(2).replace(",", "")))
                        )
        if len(unit_matches) != 1 or unit_matches[0] != ("EUR", 1000):
            raise ValueError("Expected one unambiguous EUR x 1,000 unit declaration")
        currency, unit = unit_matches[0]

        def statement_dates(section):
            header_row = row_for(section) + 1
            result = {}
            for col in range(2, ws.max_column + 1):
                value = ws.cell(header_row, col).value
                if isinstance(value, datetime):
                    value = value.date()
                if isinstance(value, date):
                    result[col] = value
            if len(result) != 2 or len(set(result.values())) != 2:
                raise ValueError(f"Expected two distinct dates under {section!r}")
            return result

        dates_by_col = statement_dates("Statement of financial position")
        for section in (
            "Profit and loss statement",
            "Cash flow statement",
            "Movement in retained earnings",
        ):
            if statement_dates(section) != dates_by_col:
                raise ValueError(f"Year columns differ in {section!r}")

        ordered = sorted(dates_by_col, key=dates_by_col.get)
        columns = {"cy": ordered[1], "py": ordered[0]}
        period_end = {
            year: dates_by_col[col].isoformat() for year, col in columns.items()
        }

        def address(label, year):
            return ws.cell(row_for(label), columns[year]).coordinate

        # Each source entry specifies its label and its provenance multiplier.
        # An empty list means the line is absent, rather than inferred or plugged.
        mapping = {
            "revenue": [("Sales revenue", 1)],
            "cost_of_sales": [("Direct costs", 1)],
            "distribution_costs": [("Distribution costs", 1)],
            "administrative_expenses": [
                ("General and administrative expenses", 1)
            ],
            "other_operating_income": [("Other gains and losses (net)", 1)],
            "finance_income": [("Finance income", 1)],
            "finance_costs": [("Interest expense", 1)],
            "income_tax_expense": [("Income tax expense", 1)],
            "property_plant_equipment": [
                ("Land and buildings", 1),
                ("Plant and machinery", 1),
            ],
            "right_of_use_assets": [("Right-of-use assets", 1)],
            "intangible_assets": [("Goodwill and intangibles", 1)],
            "investments": [],
            "deferred_tax_assets": [("Deferred tax assets", 1)],
            "inventories": [],
            "trade_receivables": [("Accounts receivable", 1)],
            "other_current_assets": [("Prepayments and other debtors", 1)],
            "cash_and_cash_equivalents": [("Cash at bank and in hand", 1)],
            "trade_payables": [("Accounts payable", 1)],
            "other_current_liabilities": [("Other payables and accruals", 1)],
            "current_borrowings": [
                ("Bank loans - current portion", 1),
                ("Lease liabilities - current", 1),
            ],
            "current_tax_payable": [("Income tax payable", 1)],
            "non_current_borrowings": [
                ("Bank loans - non-current", 1),
                ("Lease liabilities - non-current", 1),
            ],
            "provisions": [("Provisions", 1)],
            "deferred_tax_liabilities": [("Deferred tax liabilities", 1)],
            "share_capital": [("Share capital", 1)],
            "retained_earnings": [("Retained earnings", 1)],
            "other_reserves": [],
            "net_cash_operating": [("Net cash from operating activities", 1)],
            "net_cash_investing": [("Net cash from investing activities", 1)],
            "net_cash_financing": [("Net cash used in financing activities", 1)],
            "fx_effect_on_cash": [("Effect of exchange rate changes", 1)],
            "cash_at_beginning": [("Opening cash and cash equivalents", 1)],
            "dividends": [("Dividends", -1)],
        }

        fields = {}
        for key, sources in mapping.items():
            fields[key] = {}
            for year in columns:
                refs = [
                    reference(address(label, year), factor)
                    for label, factor in sources
                ]
                fields[key][year] = {"value": ref_sum(refs), "cells": refs}

        total_formulas = {
            "gross_profit": "revenue - cost_of_sales",
            "operating_profit": (
                "gross_profit - distribution_costs - administrative_expenses"
                " + other_operating_income"
            ),
            "profit_before_tax": (
                "operating_profit + finance_income - finance_costs"
            ),
            "profit_for_year": "profit_before_tax - income_tax_expense",
            "total_non_current_assets": (
                "property_plant_equipment + right_of_use_assets"
                " + intangible_assets + investments + deferred_tax_assets"
            ),
            "total_current_assets": (
                "inventories + trade_receivables + other_current_assets"
                " + cash_and_cash_equivalents"
            ),
            "total_assets": "total_non_current_assets + total_current_assets",
            "total_current_liabilities": (
                "trade_payables + other_current_liabilities"
                " + current_borrowings + current_tax_payable"
            ),
            "total_non_current_liabilities": (
                "non_current_borrowings + provisions + deferred_tax_liabilities"
            ),
            "total_liabilities": (
                "total_current_liabilities + total_non_current_liabilities"
            ),
            "total_equity": "share_capital + retained_earnings + other_reserves",
            "total_equity_and_liabilities": "total_equity + total_liabilities",
            "net_change_in_cash": (
                "net_cash_operating + net_cash_investing + net_cash_financing"
            ),
            "cash_at_end": (
                "cash_at_beginning + net_change_in_cash + fx_effect_on_cash"
            ),
        }

        def evaluate(expression, values):
            tokens = expression.split()
            if not tokens or len(tokens) % 2 != 1:
                raise ValueError(f"Invalid formula: {expression}")
            result = values[tokens[0]]
            for index in range(1, len(tokens), 2):
                operator, key = tokens[index:index + 2]
                if operator not in ("+", "-"):
                    raise ValueError(f"Unsupported operator: {operator}")
                result += values[key] if operator == "+" else -values[key]
            return result

        calculated = {}
        for year in columns:
            values = {key: periods[year]["value"] for key, periods in fields.items()}
            for key, formula in total_formulas.items():
                values[key] = evaluate(formula, values)
            calculated[year] = values

        printed_totals = {
            "Gross profit": "gross_profit",
            "Operating profit": "operating_profit",
            "Profit before taxation": "profit_before_tax",
            "Net profit": "profit_for_year",
            "Total non-current assets": "total_non_current_assets",
            "Total current assets": "total_current_assets",
            "TOTAL ASSETS": "total_assets",
            "Total equity": "total_equity",
            "Total non-current liabilities": "total_non_current_liabilities",
            "Total current liabilities": "total_current_liabilities",
            "Total liabilities": "total_liabilities",
            "TOTAL EQUITY AND LIABILITIES": "total_equity_and_liabilities",
            "Net increase/(decrease) in cash": "net_change_in_cash",
            "Cash at end of year": "cash_at_end",
            "Profit for the year": "profit_for_year",
            "Closing retained earnings": "retained_earnings",
        }

        source_totals = []
        for label, formula in printed_totals.items():
            for year in columns:
                source_totals.append({
                    **reference(address(label, year)),
                    "year": year,
                    "formula": formula,
                })

        # This local roll-forward also checks the prior year's opening balance,
        # for which no earlier balance sheet is supplied.
        for year in columns:
            source_totals.append({
                **reference(address("Opening retained earnings", year)),
                "year": year,
                "sum_of": [
                    reference(address("Closing retained earnings", year), 1),
                    reference(address("Profit for the year", year), -1),
                    reference(address("Dividends", year), -1),
                ],
            })

        exceptions = []

        def check(check_id, actual, expected, explanation):
            difference = actual - expected
            if not math.isclose(actual, expected, rel_tol=0, abs_tol=1e-8):
                exceptions.append({
                    "check": check_id,
                    "difference": difference,
                    "explanation": explanation,
                })

        for total in source_totals:
            expected = (
                evaluate(total["formula"], calculated[total["year"]])
                if "formula" in total
                else ref_sum(total["sum_of"])
            )
            check(
                f"T {total['sheet']}!{total['cell']} {total['year']}",
                total["factor"] * number(total["cell"]),
                expected,
                "Printed amount minus the amount calculated from its "
                "documented formula or source-cell reconciliation; "
                "all source amounts are preserved.",
            )

        for year in columns:
            values = calculated[year]
            check(
                f"R1 balance {year}",
                values["total_assets"],
                values["total_equity_and_liabilities"],
                "Assets minus equity and liabilities, using mapped input lines.",
            )
            check(
                f"R2 cash_at_end {year}",
                values["cash_at_end"],
                values["cash_and_cash_equivalents"],
                "Calculated cash-flow closing cash minus balance-sheet cash.",
            )

        check(
            "R3 opening cash cy",
            calculated["cy"]["cash_at_beginning"],
            calculated["py"]["cash_and_cash_equivalents"],
            "Current-year opening cash minus prior-year balance-sheet cash.",
        )
        check(
            "R4 retained earnings cy",
            calculated["py"]["retained_earnings"]
            + calculated["cy"]["profit_for_year"]
            - calculated["cy"]["dividends"],
            calculated["cy"]["retained_earnings"],
            "Prior retained earnings plus current profit less dividends "
            "minus current retained earnings.",
        )

        # Fail explicitly on an unexpected numeric row rather than silently
        # omitting it or inventing a reason to ignore it.
        covered = set()
        for periods in fields.values():
            for line in periods.values():
                covered.update(item["cell"] for item in line["cells"])
        for total in source_totals:
            covered.add(total["cell"])
            covered.update(item["cell"] for item in total.get("sum_of", []))

        for year, col in columns.items():
            first_row = min(
                row_for(label)
                for sources in mapping.values()
                for label, _ in sources
            )
            for row in range(first_row, ws.max_row + 1):
                cell = ws.cell(row, col)
                if (
                    isinstance(cell.value, Real)
                    and not isinstance(cell.value, bool)
                    and cell.coordinate not in covered
                ):
                    raise ValueError(
                        f"Unclassified numeric cell: {ws.title}!{cell.coordinate}"
                    )

        return {
            "entity_code": "TW-NL01",
            "source_entity_name": source_entity_name,
            "currency": currency,
            "unit": unit,
            "period_end": period_end,
            "fields": fields,
            "source_totals": source_totals,
            "ignored": [],
            "source_exceptions": exceptions,
            "notes": (
                "The workbook name 'Tarnwell Logistics BV' matches register "
                "entity TW-NL01, Tarnwell Logistics B.V. Amounts remain in EUR "
                "thousands; year columns are determined from their dates. "
                "Land and buildings and plant and machinery are combined as "
                "property, plant and equipment; bank and lease liabilities "
                "are combined by maturity. Signed cash flows and other gains "
                "are preserved; negative dividend movements are converted to "
                "positive dividends declared. Absent input lines are zero. "
                "Repeated profit and retained earnings figures and both "
                "opening retained earnings balances are reconciled as "
                "source totals."
            ),
        }
    finally:
        wb.close()
