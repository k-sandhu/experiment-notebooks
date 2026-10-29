import math
import re
from datetime import datetime
from numbers import Real

import openpyxl


def extract(path) -> dict:
    # This adapter publishes the Australian subsidiary from the mixed ANZ export.
    target = "AU01"
    entity_code = "TW-AU01"
    currency = "AUD"

    mappings = {
        ("IS", "Sales revenue - stores"): ("revenue", -1),
        ("IS", "Sales revenue - online"): ("revenue", -1),
        ("IS", "Cost of goods sold"): ("cost_of_sales", 1),
        ("IS", "Store operating costs"): ("distribution_costs", 1),
        ("IS", "Marketing"): ("distribution_costs", 1),
        ("IS", "Support office costs"): ("administrative_expenses", 1),
        ("IS", "Other operating income/expense"): ("other_operating_income", -1),
        ("IS", "Finance income"): ("finance_income", -1),
        ("IS", "Finance costs - leases"): ("finance_costs", 1),
        ("IS", "Finance costs - bank"): ("finance_costs", 1),
        ("IS", "Income tax expense"): ("income_tax_expense", 1),
        ("BS", "Cash on hand and at bank"): ("cash_and_cash_equivalents", 1),
        ("BS", "Trade receivables"): ("trade_receivables", 1),
        ("BS", "Inventory - merchandise"): ("inventories", 1),
        ("BS", "Prepayments and other receivables"): ("other_current_assets", 1),
        ("BS", "Property, plant and equipment"): ("property_plant_equipment", 1),
        ("BS", "Right-of-use assets - stores"): ("right_of_use_assets", 1),
        ("BS", "Intangible assets"): ("intangible_assets", 1),
        ("BS", "Deferred tax assets"): ("deferred_tax_assets", 1),
        ("BS", "Trade payables"): ("trade_payables", -1),
        ("BS", "Accruals and employee entitlements"): ("other_current_liabilities", -1),
        ("BS", "Lease liabilities - current"): ("current_borrowings", -1),
        ("BS", "Income tax payable"): ("current_tax_payable", -1),
        ("BS", "Lease liabilities - non-current"): ("non_current_borrowings", -1),
        ("BS", "Bank facility"): ("non_current_borrowings", -1),
        ("BS", "Provisions - make good"): ("provisions", -1),
        ("BS", "Deferred tax liabilities"): ("deferred_tax_liabilities", -1),
        ("BS", "Contributed equity"): ("share_capital", -1),
        ("BS", "Retained earnings"): ("retained_earnings", -1),
        ("CF", "Net cash from operating activities"): ("net_cash_operating", 1),
        ("CF", "Net cash from investing activities"): ("net_cash_investing", 1),
        ("CF", "Net cash from financing activities"): ("net_cash_financing", 1),
        ("CF", "FX on cash held"): ("fx_effect_on_cash", 1),
        ("CF", "Cash at beginning of period"): ("cash_at_beginning", 1),
        ("EQ", "Dividends declared"): ("dividends", 1),
    }

    subtotal_mappings = {
        ("IS", "GROSS PROFIT"): ("gross_profit", -1),
        ("IS", "NET PROFIT AFTER TAX"): ("profit_for_year", -1),
        ("BS", "TOTAL ASSETS"): ("total_assets", 1),
        ("BS", "TOTAL LIABILITIES AND EQUITY"): (
            "total_equity_and_liabilities", -1
        ),
        ("CF", "CASH AT END OF PERIOD"): ("cash_at_end", 1),
    }

    keys = (
        "revenue cost_of_sales distribution_costs administrative_expenses "
        "other_operating_income finance_income finance_costs income_tax_expense "
        "property_plant_equipment right_of_use_assets intangible_assets "
        "investments deferred_tax_assets inventories trade_receivables "
        "other_current_assets cash_and_cash_equivalents trade_payables "
        "other_current_liabilities current_borrowings current_tax_payable "
        "non_current_borrowings provisions deferred_tax_liabilities "
        "share_capital retained_earnings other_reserves net_cash_operating "
        "net_cash_investing net_cash_financing fx_effect_on_cash "
        "cash_at_beginning dividends"
    ).split()

    formulas = {
        "gross_profit": "revenue - cost_of_sales",
        "operating_profit": (
            "gross_profit - distribution_costs - administrative_expenses "
            "+ other_operating_income"
        ),
        "profit_before_tax": "operating_profit + finance_income - finance_costs",
        "profit_for_year": "profit_before_tax - income_tax_expense",
        "total_non_current_assets": (
            "property_plant_equipment + right_of_use_assets + intangible_assets "
            "+ investments + deferred_tax_assets"
        ),
        "total_current_assets": (
            "inventories + trade_receivables + other_current_assets "
            "+ cash_and_cash_equivalents"
        ),
        "total_assets": "total_non_current_assets + total_current_assets",
        "total_current_liabilities": (
            "trade_payables + other_current_liabilities "
            "+ current_borrowings + current_tax_payable"
        ),
        "total_non_current_liabilities": (
            "non_current_borrowings + provisions + deferred_tax_liabilities"
        ),
        "total_liabilities": (
            "total_current_liabilities + total_non_current_liabilities"
        ),
        "total_equity": "share_capital + retained_earnings + other_reserves",
        "total_equity_and_liabilities": "total_equity + total_liabilities",
        "net_change_in_cash": (
            "net_cash_operating + net_cash_investing + net_cash_financing"
        ),
        "cash_at_end": "cash_at_beginning + net_change_in_cash + fx_effect_on_cash",
    }

    def text(value):
        return " ".join(str(value).split()) if value is not None else ""

    def number(cell):
        value = cell.value
        if (
            isinstance(value, bool)
            or not isinstance(value, Real)
            or not math.isfinite(value)
        ):
            raise ValueError(f"Expected a finite number at {cell.parent.title}!{cell.coordinate}")
        return int(value) if int(value) == value else float(value)

    def evaluate(expression, values):
        result = 0
        sign = 1
        for token in expression.split():
            if token == "+":
                sign = 1
            elif token == "-":
                sign = -1
            else:
                result += sign * values[token]
        return result

    wb = openpyxl.load_workbook(path, data_only=True)
    try:
        ws = wb["Export"]
        parameters = [
            text(cell.value)
            for row in wb["Parameters"].iter_rows()
            for cell in row
            if cell.value is not None
        ]
        fiscal_matches = []
        for parameter in parameters:
            match = re.fullmatch(
                r"Fiscal year (FY\d{2,4}) = .+ to (\d{1,2} [A-Za-z]+ \d{4})",
                parameter,
                flags=re.IGNORECASE,
            )
            if match:
                fiscal_matches.append(match)
        if len(fiscal_matches) != 1:
            raise ValueError("Expected exactly one fiscal-year definition in Parameters.")

        fiscal_match = fiscal_matches[0]
        current_fy = fiscal_match.group(1).upper()
        current_end = datetime.strptime(fiscal_match.group(2), "%d %B %Y").date()
        fiscal_digits = current_fy[2:]
        if current_end.year % (10 ** len(fiscal_digits)) != int(fiscal_digits):
            raise ValueError("Fiscal-year label and period-end year disagree.")
        prior_fy = "FY" + str(int(fiscal_digits) - 1).zfill(len(fiscal_digits))
        years = {current_fy: "cy", prior_fy: "py"}
        prior_end = current_end.replace(year=current_end.year - 1)

        required_headers = {
            "Entity", "Entity name", "Fiscal year", "Statement",
            "Line code", "Line description", "Amount",
        }
        header_candidates = []
        for row in ws.iter_rows():
            columns = {text(cell.value): cell.column for cell in row if cell.value is not None}
            if required_headers <= columns.keys():
                header_candidates.append((row[0].row, columns))
        if len(header_candidates) != 1:
            raise ValueError("Expected exactly one Export header row.")
        header_row, columns = header_candidates[0]

        fields = {
            key: {
                year: {"value": 0, "cells": []}
                for year in ("cy", "py")
            }
            for key in keys
        }
        source_totals = []
        ignored = []
        names = set()
        seen = set()
        selected_years = set()

        for row_index in range(header_row + 1, ws.max_row + 1):
            amount_cell = ws.cell(row_index, columns["Amount"])
            if amount_cell.value is None:
                continue
            amount = number(amount_cell)

            def column_value(header):
                return text(ws.cell(row_index, columns[header]).value)

            local_entity = column_value("Entity")
            source_name = column_value("Entity name")
            fiscal_year = column_value("Fiscal year").upper()

            if local_entity != target:
                ignored.append({
                    "sheet": ws.title,
                    "cell": amount_cell.coordinate,
                    "reason": (
                        f"Another entity: {local_entity} ({source_name}); "
                        f"this adapter selects {target} / {entity_code}."
                    ),
                })
                continue

            names.add(source_name)
            if fiscal_year not in years:
                ignored.append({
                    "sheet": ws.title,
                    "cell": amount_cell.coordinate,
                    "reason": (
                        f"Fiscal year {fiscal_year} is outside the selected "
                        f"{current_fy}/{prior_fy} comparison."
                    ),
                })
                continue

            year = years[fiscal_year]
            selected_years.add(year)
            statement = column_value("Statement")
            label = column_value("Line description")
            identity = (year, statement, label)
            if identity in seen:
                raise ValueError(f"Duplicate selected line: {identity!r}")
            seen.add(identity)
            lookup = (statement, label)

            if lookup in mappings:
                key, factor = mappings[lookup]
                line = fields[key][year]
                line["value"] += factor * amount
                line["cells"].append({
                    "sheet": ws.title,
                    "cell": amount_cell.coordinate,
                    "factor": factor,
                })
            elif lookup in subtotal_mappings:
                formula, factor = subtotal_mappings[lookup]
                source_totals.append({
                    "sheet": ws.title,
                    "cell": amount_cell.coordinate,
                    "year": year,
                    "factor": factor,
                    "formula": formula,
                })
            else:
                raise ValueError(
                    f"Unmapped selected amount at {ws.title}!{amount_cell.coordinate}: "
                    f"{statement} / {label}"
                )

        if len(names) != 1 or not next(iter(names), ""):
            raise ValueError("Expected exactly one nonempty Australian entity name.")
        source_entity_name = next(iter(names))
        normalized_name = re.sub(r"[^A-Z0-9]", "", source_entity_name.upper())
        if normalized_name not in {
            "TARNWELLAUSTRALIAPTYLIMITED",
            "TARNWELLAUSTRALIAPTYLTD",
        }:
            raise ValueError("AU01 entity name does not match the Australian register entry.")
        if selected_years != {"cy", "py"}:
            raise ValueError("Both current and prior Australian fiscal years are required.")

        values = {}
        for year in ("cy", "py"):
            values[year] = {key: fields[key][year]["value"] for key in keys}
            for total, expression in formulas.items():
                values[year][total] = evaluate(expression, values[year])

        exceptions = []

        def check(check_id, difference, explanation):
            if not math.isclose(difference, 0, rel_tol=0, abs_tol=1e-7):
                exceptions.append({
                    "check": check_id,
                    "difference": difference,
                    "explanation": explanation,
                })

        for total in source_totals:
            printed = total["factor"] * number(wb[total["sheet"]][total["cell"]])
            calculated = evaluate(total["formula"], values[total["year"]])
            check(
                f"T {total['sheet']}!{total['cell']} {total['year']}",
                printed - calculated,
                f"The sign-adjusted printed {total['formula']} is {printed}, "
                f"but the mapped input lines calculate {calculated}. "
                "Difference is printed minus calculated; source amounts are preserved.",
            )

        for year in ("cy", "py"):
            v = values[year]
            check(
                f"R1 balance {year}",
                v["total_assets"] - v["total_equity_and_liabilities"],
                "Mapped assets minus mapped equity and liabilities is nonzero. "
                "All source balances are preserved.",
            )
            check(
                f"R2 cash_at_end {year}",
                v["cash_at_end"] - v["cash_and_cash_equivalents"],
                "Opening cash plus the three cash flows and FX differs from "
                "balance-sheet cash. Difference is calculated closing cash minus "
                "balance-sheet cash; source amounts are preserved.",
            )

        check(
            "R3 opening cash cy",
            values["cy"]["cash_at_beginning"]
            - values["py"]["cash_and_cash_equivalents"],
            "Current opening cash differs from prior balance-sheet cash. "
            "Difference is current opening minus prior closing; source amounts are preserved.",
        )
        check(
            "R4 retained earnings cy",
            values["py"]["retained_earnings"]
            + values["cy"]["profit_for_year"]
            - values["cy"]["dividends"]
            - values["cy"]["retained_earnings"],
            "Prior retained earnings plus calculated current profit less declared "
            "dividends differs from current retained earnings. Difference is "
            "calculated closing minus reported closing; source amounts are preserved.",
        )

        return {
            "entity_code": entity_code,
            "source_entity_name": source_entity_name,
            "currency": currency,
            "unit": 1,
            "period_end": {
                "cy": current_end.isoformat(),
                "py": prior_end.isoformat(),
            },
            "fields": fields,
            "source_totals": source_totals,
            "ignored": ignored,
            "source_exceptions": exceptions,
            "notes": (
                "The workbook contains Australia and New Zealand; this adapter explicitly "
                "selects AU01, matched to TW-AU01, and excludes every other entity's amount. "
                "AUD is inferred from the entity register and the local-currency parameter; "
                "unscaled export amounts are retained in whole currency units. "
                "Debit-positive/credit-negative signs are reversed for income, liabilities "
                "and equity; cash-flow signs and dividends are retained. Store operating "
                "costs and marketing are selling costs. The bank facility is classified "
                "as non-current because it appears in the non-current liability section "
                "between non-current leases and make-good provisions. Investments and "
                "other reserves are absent and remain zero. Fiscal dates come from "
                "Parameters; any reconciliation discrepancies are reported without plugs."
            ),
        }
    finally:
        wb.close()
