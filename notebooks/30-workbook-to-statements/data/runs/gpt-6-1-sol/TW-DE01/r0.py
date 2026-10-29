import datetime as dt
import math
import re

import openpyxl


def extract(path) -> dict:
    wb = openpyxl.load_workbook(path, data_only=True)

    def normalize(value):
        return " ".join(str(value or "").split())

    def find_row(sheet, label_column, label):
        ws = wb[sheet]
        matches = [
            row
            for row in range(1, ws.max_row + 1)
            if normalize(ws.cell(row, label_column).value) == normalize(label)
        ]
        if len(matches) != 1:
            raise ValueError(
                f"{sheet}: expected exactly one label {label!r}; "
                f"found {len(matches)}"
            )
        return matches[0]

    def number(sheet, address):
        value = wb[sheet][address].value
        if (
            isinstance(value, bool)
            or not isinstance(value, (int, float))
            or not math.isfinite(value)
        ):
            raise ValueError(f"{sheet}!{address} must contain a finite number")
        return value

    def reference(sheet, label_column, amount_column, label, factor=1):
        row = find_row(sheet, label_column, label)
        address = wb[sheet].cell(row, amount_column).coordinate
        number(sheet, address)
        return {"sheet": sheet, "cell": address, "factor": factor}

    def contribution(ref):
        return ref["factor"] * number(ref["sheet"], ref["cell"])

    required_sheets = ("GuV", "Bilanz", "KFR")
    for sheet in required_sheets:
        if sheet not in wb.sheetnames:
            raise ValueError(f"Missing required sheet: {sheet}")

    legal_name = "Tarnwell Antriebstechnik GmbH"
    entity_names = []
    for sheet in ("GuV", "Bilanz"):
        names = [
            cell.value
            for row in wb[sheet].iter_rows()
            for cell in row
            if isinstance(cell.value, str)
            and normalize(cell.value).split(",", 1)[0] == legal_name
        ]
        if len(names) != 1:
            raise ValueError(f"Cannot uniquely identify the entity in {sheet}")
        entity_names.append(names[0])

    if normalize(entity_names[0]) != normalize(entity_names[1]):
        raise ValueError("Entity names differ between GuV and Bilanz")

    # TEUR means EUR thousands. Validate each statement's stated unit.
    for sheet in required_sheets:
        if not any(
            isinstance(cell.value, str)
            and re.search(r"\bTEUR\b", cell.value, re.IGNORECASE)
            for row in wb[sheet].iter_rows()
            for cell in row
        ):
            raise ValueError(f"{sheet}: expected amounts stated in TEUR")

    def parse_balance_date(value):
        if isinstance(value, dt.datetime):
            return value.date()
        if isinstance(value, dt.date):
            return value
        if isinstance(value, str):
            match = re.fullmatch(
                r"\s*(\d{1,2})\.(\d{1,2})\.(\d{4})\s*", value
            )
            if match:
                day, month, year = map(int, match.groups())
                return dt.date(year, month, day)
        return None

    date_headers = []
    for row in wb["Bilanz"].iter_rows():
        for cell in row:
            parsed = parse_balance_date(cell.value)
            if parsed is not None:
                date_headers.append((cell.row, cell.column, parsed))

    dates = sorted({entry[2] for entry in date_headers}, reverse=True)
    if len(dates) != 2 or dates[0].year != dates[1].year + 1:
        raise ValueError("Expected two consecutive balance-sheet dates")

    period_dates = {"cy": dates[0], "py": dates[1]}
    balance_columns = {}
    for year, date in period_dates.items():
        matches = sorted(
            (row, column)
            for row, column, parsed in date_headers
            if parsed == date
        )
        if len(matches) != 2 or matches[0][0] != matches[1][0]:
            raise ValueError(f"Expected asset and liability date headers for {year}")
        balance_columns[year] = {
            "assets": matches[0][1],
            "liabilities": matches[1][1],
        }

    def statement_year_columns(sheet, prefix):
        columns = {}
        for year, date in period_dates.items():
            expected = (
                f"{prefix[year]} {date.year}" if prefix else str(date.year)
            )
            matches = [
                cell.column
                for row in wb[sheet].iter_rows()
                for cell in row
                if normalize(cell.value) == expected
            ]
            if len(matches) != 1:
                raise ValueError(
                    f"{sheet}: expected exactly one year header {expected!r}"
                )
            columns[year] = matches[0]
        return columns

    guv_columns = statement_year_columns(
        "GuV", {"cy": "GJ", "py": "VJ"}
    )
    kfr_columns = statement_year_columns("KFR", None)

    keys = """
        revenue cost_of_sales distribution_costs administrative_expenses
        other_operating_income finance_income finance_costs income_tax_expense
        property_plant_equipment right_of_use_assets intangible_assets
        investments deferred_tax_assets inventories trade_receivables
        other_current_assets cash_and_cash_equivalents trade_payables
        other_current_liabilities current_borrowings current_tax_payable
        non_current_borrowings provisions deferred_tax_liabilities
        share_capital retained_earnings other_reserves net_cash_operating
        net_cash_investing net_cash_financing fx_effect_on_cash
        cash_at_beginning dividends
    """.split()

    fields = {
        key: {year: {"value": 0, "cells": []} for year in period_dates}
        for key in keys
    }

    def assign(key, year, refs):
        fields[key][year] = {
            "value": sum(contribution(ref) for ref in refs),
            "cells": refs,
        }

    guv_mapping = {
        "revenue": ("Umsatzerlöse", 1),
        "cost_of_sales": (
            "Herstellungskosten der zur Erzielung der Umsatzerlöse "
            "erbrachten Leistungen",
            -1,
        ),
        "distribution_costs": ("Vertriebskosten", -1),
        "administrative_expenses": ("Allgemeine Verwaltungskosten", -1),
        "other_operating_income": ("Sonstige betriebliche Erträge (netto)", 1),
        "finance_income": ("Zinsen und ähnliche Erträge", 1),
        "finance_costs": ("Zinsen und ähnliche Aufwendungen", -1),
        "income_tax_expense": ("Steuern vom Einkommen und vom Ertrag", -1),
    }

    asset_mapping = {
        "property_plant_equipment": [
            "Grundstücke und Bauten",
            "Technische Anlagen und Maschinen",
            "Betriebs- und Geschäftsausstattung",
        ],
        "right_of_use_assets": ["Nutzungsrechte aus Leasingverhältnissen"],
        "intangible_assets": ["Immaterielle Vermögensgegenstände"],
        "deferred_tax_assets": ["D. Aktive latente Steuern"],
        "inventories": [
            "Roh-, Hilfs- und Betriebsstoffe",
            "Fertige Erzeugnisse und Waren",
        ],
        "trade_receivables": ["Forderungen aus Lieferungen und Leistungen"],
        "other_current_assets": [
            "Sonstige Vermögensgegenstände",
            "C. Rechnungsabgrenzungsposten",
        ],
        "cash_and_cash_equivalents": [
            "Kassenbestand, Guthaben bei Kreditinstituten"
        ],
    }

    liability_mapping = {
        "share_capital": ["Gezeichnetes Kapital", "Kapitalrücklage"],
        "retained_earnings": ["Gewinnrücklagen und Bilanzgewinn"],
        "other_reserves": ["Sonstige Rücklagen"],
        "provisions": [
            "Rückstellungen für Pensionen und ähnliche Verpflichtungen"
        ],
        "current_tax_payable": ["Steuerrückstellungen"],
        "current_borrowings": [
            "Verbindlichkeiten gegenüber Kreditinstituten "
            "(Restlaufzeit bis 1 Jahr)",
            "Leasingverbindlichkeiten (bis 1 Jahr)",
        ],
        "non_current_borrowings": [
            "Verbindlichkeiten gegenüber Kreditinstituten "
            "(Restlaufzeit über 1 Jahr)",
            "Leasingverbindlichkeiten (über 1 Jahr)",
        ],
        "trade_payables": ["Verbindlichkeiten aus Lieferungen und Leistungen"],
        "other_current_liabilities": ["Sonstige Verbindlichkeiten"],
        "deferred_tax_liabilities": ["D. Passive latente Steuern"],
    }

    kfr_mapping = {
        "net_cash_operating": (
            "Cashflow aus laufender Geschäftstätigkeit", 1
        ),
        "net_cash_investing": (
            "Cashflow aus der Investitionstätigkeit", 1
        ),
        "net_cash_financing": (
            "Cashflow aus der Finanzierungstätigkeit", 1
        ),
        "fx_effect_on_cash": (
            "Wechselkursbedingte Änderungen des Finanzmittelfonds", 1
        ),
        "cash_at_beginning": (
            "Finanzmittelfonds am Anfang der Periode", 1
        ),
        "dividends": ("Ausschüttung an die Gesellschafter", -1),
    }

    source_totals = []

    def add_total(sheet, label_column, amount_column, label, year,
                  formula=None, sum_of=None):
        ref = reference(sheet, label_column, amount_column, label)
        entry = {**ref, "year": year}
        if (formula is None) == (sum_of is None):
            raise ValueError("A total needs exactly one of formula or sum_of")
        if formula is not None:
            entry["formula"] = formula
        else:
            entry["sum_of"] = sum_of
        source_totals.append(entry)

    for year in period_dates:
        gc = guv_columns[year]
        ac = balance_columns[year]["assets"]
        lc = balance_columns[year]["liabilities"]
        kc = kfr_columns[year]

        for key, (label, factor) in guv_mapping.items():
            assign(key, year, [reference("GuV", 1, gc, label, factor)])
        for key, labels in asset_mapping.items():
            assign(key, year, [
                reference("Bilanz", 1, ac, label) for label in labels
            ])
        for key, labels in liability_mapping.items():
            assign(key, year, [
                reference("Bilanz", 5, lc, label) for label in labels
            ])
        for key, (label, factor) in kfr_mapping.items():
            assign(key, year, [reference("KFR", 1, kc, label, factor)])

        for label, formula in [
            ("Bruttoergebnis vom Umsatz", "gross_profit"),
            ("Betriebsergebnis", "operating_profit"),
            ("Ergebnis vor Steuern", "profit_before_tax"),
            ("Jahresüberschuss", "profit_for_year"),
        ]:
            add_total("GuV", 1, gc, label, year, formula=formula)

        add_total(
            "Bilanz", 1, ac, "Summe Anlagevermögen", year,
            formula="total_non_current_assets - deferred_tax_assets",
        )

        current_asset_labels = (
            asset_mapping["inventories"]
            + asset_mapping["trade_receivables"]
            + ["Sonstige Vermögensgegenstände"]
            + asset_mapping["cash_and_cash_equivalents"]
        )
        add_total(
            "Bilanz", 1, ac, "Summe Umlaufvermögen", year,
            sum_of=[
                reference("Bilanz", 1, ac, label)
                for label in current_asset_labels
            ],
        )

        for label, formula in [
            ("Summe Eigenkapital", "total_equity"),
            ("Summe Rückstellungen", "provisions + current_tax_payable"),
            (
                "Summe Verbindlichkeiten",
                "trade_payables + other_current_liabilities "
                "+ current_borrowings + non_current_borrowings",
            ),
            ("Summe Passiva", "total_equity_and_liabilities"),
        ]:
            add_total("Bilanz", 5, lc, label, year, formula=formula)

        add_total(
            "Bilanz", 1, ac, "Summe Aktiva", year, formula="total_assets"
        )

        for label, formula in [
            (
                "Zahlungswirksame Veränderung des Finanzmittelfonds",
                "net_change_in_cash",
            ),
            ("Finanzmittelfonds am Ende der Periode", "cash_at_end"),
            (
                "Stand 1. Januar",
                "retained_earnings - profit_for_year + dividends",
            ),
            ("Jahresüberschuss", "profit_for_year"),
            ("Stand 31. Dezember", "retained_earnings"),
        ]:
            add_total("KFR", 1, kc, label, year, formula=formula)

    total_formulas = {
        "gross_profit": "revenue - cost_of_sales",
        "operating_profit": (
            "gross_profit - distribution_costs - administrative_expenses "
            "+ other_operating_income"
        ),
        "profit_before_tax": "operating_profit + finance_income - finance_costs",
        "profit_for_year": "profit_before_tax - income_tax_expense",
        "total_non_current_assets": (
            "property_plant_equipment + right_of_use_assets "
            "+ intangible_assets + investments + deferred_tax_assets"
        ),
        "total_current_assets": (
            "inventories + trade_receivables + other_current_assets "
            "+ cash_and_cash_equivalents"
        ),
        "total_assets": "total_non_current_assets + total_current_assets",
        "total_current_liabilities": (
            "trade_payables + other_current_liabilities "
            "+ current_borrowings + current_tax_payable"
        ),
        "total_non_current_liabilities": (
            "non_current_borrowings + provisions + deferred_tax_liabilities"
        ),
        "total_liabilities": (
            "total_current_liabilities + total_non_current_liabilities"
        ),
        "total_equity": "share_capital + retained_earnings + other_reserves",
        "total_equity_and_liabilities": "total_equity + total_liabilities",
        "net_change_in_cash": (
            "net_cash_operating + net_cash_investing + net_cash_financing"
        ),
        "cash_at_end": (
            "cash_at_beginning + net_change_in_cash + fx_effect_on_cash"
        ),
    }

    def evaluate(expression, values):
        tokens = expression.split()
        if not tokens or len(tokens) % 2 == 0:
            raise ValueError(f"Invalid expression: {expression}")
        result = values[tokens[0]]
        for index in range(1, len(tokens), 2):
            operator, name = tokens[index:index + 2]
            if operator not in ("+", "-"):
                raise ValueError(f"Invalid operator: {operator}")
            result += values[name] * (1 if operator == "+" else -1)
        return result

    computed = {}
    for year in period_dates:
        values = {key: fields[key][year]["value"] for key in keys}
        for name, expression in total_formulas.items():
            values[name] = evaluate(expression, values)
        computed[year] = values

    source_exceptions = []

    def report_difference(check, actual, expected, explanation):
        difference = actual - expected
        if not math.isclose(actual, expected, rel_tol=0, abs_tol=1e-8):
            source_exceptions.append({
                "check": check,
                "difference": difference,
                "explanation": explanation,
            })

    for entry in source_totals:
        actual = entry["factor"] * number(entry["sheet"], entry["cell"])
        if "formula" in entry:
            expected = evaluate(entry["formula"], computed[entry["year"]])
        else:
            expected = sum(contribution(ref) for ref in entry["sum_of"])
        report_difference(
            f"T {entry['sheet']}!{entry['cell']} {entry['year']}",
            actual,
            expected,
            "The printed amount differs from its mapped formula or "
            "component cells; all source amounts have been preserved.",
        )

    for year, values in computed.items():
        report_difference(
            f"R1 balance {year}",
            values["total_assets"],
            values["total_equity_and_liabilities"],
            "Mapped assets differ from mapped equity and liabilities.",
        )
        report_difference(
            f"R2 cash_at_end {year}",
            values["cash_at_end"],
            values["cash_and_cash_equivalents"],
            "Cash calculated from opening cash and cash flows differs "
            "from balance-sheet cash.",
        )

    report_difference(
        "R3 opening cash cy",
        computed["cy"]["cash_at_beginning"],
        computed["py"]["cash_and_cash_equivalents"],
        "Current-year opening cash differs from prior-year balance-sheet cash.",
    )
    report_difference(
        "R4 retained earnings cy",
        computed["cy"]["retained_earnings"],
        computed["py"]["retained_earnings"]
        + computed["cy"]["profit_for_year"]
        - computed["cy"]["dividends"],
        "Current retained earnings differ from prior retained earnings "
        "plus current profit less dividends declared.",
    )

    # All numeric amounts in the selected columns must be accounted for.
    accounted = set()
    for periods in fields.values():
        for line in periods.values():
            for ref in line["cells"]:
                accounted.add((ref["sheet"], ref["cell"]))
    for entry in source_totals:
        accounted.add((entry["sheet"], entry["cell"]))
        for ref in entry.get("sum_of", []):
            accounted.add((ref["sheet"], ref["cell"]))

    starts = {}
    for sheet, address in accounted:
        cell = wb[sheet][address]
        key = (sheet, cell.column)
        starts[key] = min(starts.get(key, cell.row), cell.row)

    for (sheet, column), first_row in starts.items():
        ws = wb[sheet]
        for row in range(first_row, ws.max_row + 1):
            cell = ws.cell(row, column)
            if (
                isinstance(cell.value, (int, float))
                and not isinstance(cell.value, bool)
                and (sheet, cell.coordinate) not in accounted
            ):
                raise ValueError(
                    f"Unmapped numeric amount: {sheet}!{cell.coordinate}"
                )

    return {
        "entity_code": "TW-DE01",
        "source_entity_name": entity_names[0],
        "currency": "EUR",
        "unit": 1000,
        "period_end": {
            year: date.isoformat() for year, date in period_dates.items()
        },
        "fields": fields,
        "source_totals": source_totals,
        "ignored": [],
        "source_exceptions": source_exceptions,
        "notes": (
            "Amounts remain in EUR thousands (TEUR). Negative GuV expenses "
            "and negative shareholder distributions are sign-reversed. "
            "Steuerrückstellungen are treated as current income tax payable; "
            "the workbook provides no further tax breakdown. Prepayments "
            "are included in other current assets, and deferred taxes are "
            "classified separately. Capital reserve is included in share "
            "capital, profit reserves in retained earnings, and pension "
            "obligations in non-current provisions. Summe Umlaufvermögen "
            "is checked against its component cells because the source "
            "excludes separately presented prepayments."
        ),
    }
