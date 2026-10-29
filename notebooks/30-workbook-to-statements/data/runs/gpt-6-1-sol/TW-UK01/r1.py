import math
import re
from datetime import date
from numbers import Real

import openpyxl


FIELD_KEYS = (
    "revenue cost_of_sales distribution_costs administrative_expenses "
    "other_operating_income finance_income finance_costs income_tax_expense "
    "property_plant_equipment right_of_use_assets intangible_assets investments "
    "deferred_tax_assets inventories trade_receivables other_current_assets "
    "cash_and_cash_equivalents trade_payables other_current_liabilities "
    "current_borrowings current_tax_payable non_current_borrowings provisions "
    "deferred_tax_liabilities share_capital retained_earnings other_reserves "
    "net_cash_operating net_cash_investing net_cash_financing fx_effect_on_cash "
    "cash_at_beginning dividends"
).split()

TOTAL_FORMULAS = {
    "gross_profit": "revenue - cost_of_sales",
    "operating_profit": (
        "gross_profit - distribution_costs - administrative_expenses "
        "+ other_operating_income"
    ),
    "profit_before_tax": "operating_profit + finance_income - finance_costs",
    "profit_for_year": "profit_before_tax - income_tax_expense",
    "total_non_current_assets": (
        "property_plant_equipment + right_of_use_assets + intangible_assets "
        "+ investments + deferred_tax_assets"
    ),
    "total_current_assets": (
        "inventories + trade_receivables + other_current_assets "
        "+ cash_and_cash_equivalents"
    ),
    "total_assets": "total_non_current_assets + total_current_assets",
    "total_current_liabilities": (
        "trade_payables + other_current_liabilities + current_borrowings "
        "+ current_tax_payable"
    ),
    "total_non_current_liabilities": (
        "non_current_borrowings + provisions + deferred_tax_liabilities"
    ),
    "total_liabilities": "total_current_liabilities + total_non_current_liabilities",
    "total_equity": "share_capital + retained_earnings + other_reserves",
    "total_equity_and_liabilities": "total_equity + total_liabilities",
    "net_change_in_cash": (
        "net_cash_operating + net_cash_investing + net_cash_financing"
    ),
    "cash_at_end": "cash_at_beginning + net_change_in_cash + fx_effect_on_cash",
}


def extract(path) -> dict:
    wb = openpyxl.load_workbook(path, data_only=True)
    try:
        def normalize(text):
            return " ".join(str(text).split()).casefold()

        sheet_names = ("P&L", "Balance sheet", "Cash flow")
        rows = {}
        for name in sheet_names:
            index = {}
            for row in wb[name].iter_rows():
                label = row[0].value
                if isinstance(label, str):
                    index.setdefault(normalize(label), []).append(row[0].row)
            rows[name] = index

        def row_for(sheet, label):
            matches = rows[sheet].get(normalize(label), [])
            if len(matches) != 1:
                raise ValueError(
                    f"Expected exactly one {label!r} row in {sheet!r}; "
                    f"found {len(matches)}."
                )
            return matches[0]

        expected_name = "Tarnwell Holdings plc"
        entity_names = []
        for name in sheet_names:
            matches = rows[name].get(normalize(expected_name), [])
            if len(matches) != 1:
                raise ValueError(f"Cannot identify TW-UK01 uniquely in {name!r}.")
            entity_names.append(wb[name].cell(matches[0], 1).value)
        source_entity_name = entity_names[0]

        date_matches = []
        for label, matching_rows in rows["Balance sheet"].items():
            match = re.fullmatch(r"as at (\d{1,2}) june (\d{4})", label)
            if match:
                date_matches.extend(
                    (int(match[1]), int(match[2]))
                    for _ in matching_rows
                )
        if len(date_matches) != 1:
            raise ValueError("Cannot identify the balance-sheet reporting date.")
        day, cy = date_matches[0]
        py = cy - 1
        period_end = {
            "cy": date(cy, 6, day).isoformat(),
            "py": date(py, 6, day).isoformat(),
        }

        # Unit headings may legitimately repeat for separate sections.
        # Require their presence, not a unique row.
        expected_units = normalize("All amounts in £ (GBP)")
        for name in sheet_names:
            unit_labels = [
                label for label in rows[name]
                if label.startswith("all amounts in ")
            ]
            if not unit_labels or any(
                label != expected_units for label in unit_labels
            ):
                raise ValueError(
                    f"Missing or unsupported currency/unit heading in {name!r}."
                )
        currency, unit = "GBP", 1

        columns = {}
        year_headers = []
        for name in sheet_names:
            matches = []
            for row in wb[name].iter_rows():
                current = [c.column for c in row if c.value == cy]
                prior = [c.column for c in row if c.value == py]
                if len(current) == 1 and len(prior) == 1:
                    matches.append((row[0].row, current[0], prior[0]))
            if not matches:
                raise ValueError(f"Missing comparative year headers in {name!r}.")
            column_pairs = {(c, p) for _, c, p in matches}
            if len(column_pairs) != 1:
                raise ValueError(f"Ambiguous year columns in {name!r}.")
            current_col, prior_col = next(iter(column_pairs))
            columns[name] = {"cy": current_col, "py": prior_col}
            for row_number, c, p in matches:
                year_headers.extend(
                    (name, wb[name].cell(row_number, col).coordinate)
                    for col in (c, p)
                )

        def number(sheet, coordinate):
            value = wb[sheet][coordinate].value
            if (
                isinstance(value, bool)
                or not isinstance(value, Real)
                or not math.isfinite(value)
            ):
                raise ValueError(f"{sheet}!{coordinate} is not a finite number.")
            return int(value) if isinstance(value, int) else float(value)

        def reference(sheet, label, year, factor=1):
            coordinate = wb[sheet].cell(
                row_for(sheet, label), columns[sheet][year]
            ).coordinate
            number(sheet, coordinate)
            return {"sheet": sheet, "cell": coordinate, "factor": factor}

        def sum_cells(cells):
            return sum(
                item["factor"] * number(item["sheet"], item["cell"])
                for item in cells
            )

        fields = {
            key: {
                year: {"value": 0, "cells": []}
                for year in ("cy", "py")
            }
            for key in FIELD_KEYS
        }

        mappings = {
            "P&L": {
                "revenue": [("Revenue", 1)],
                "administrative_expenses": [("Administrative expenses", 1)],
                "other_operating_income": [
                    ("Other operating income/(expense), net", 1)
                ],
                "finance_income": [
                    ("Dividends received from subsidiaries", 1),
                    ("Bank interest received", 1),
                ],
                "finance_costs": [("Interest expense", 1)],
                "income_tax_expense": [("Tax on profit", 1)],
            },
            "Balance sheet": {
                "property_plant_equipment": [("Tangible assets", 1)],
                "investments": [
                    ("Investments in subsidiaries", 1),
                    ("Loans to group undertakings (due after one year)", 1),
                ],
                "other_current_assets": [
                    ("Other receivables and prepayments", 1)
                ],
                "cash_and_cash_equivalents": [("Cash and cash equivalents", 1)],
                "share_capital": [("Share capital", 1)],
                "retained_earnings": [("Retained profits", 1)],
                "non_current_borrowings": [("Long-term borrowings", 1)],
                "trade_payables": [("Trade creditors", 1)],
                "other_current_liabilities": [("Other current liabilities", 1)],
                "current_borrowings": [("Borrowings (current)", 1)],
                "current_tax_payable": [("Income tax payable", 1)],
                "dividends": [("Dividends paid to shareholders", -1)],
            },
            "Cash flow": {
                "net_cash_operating": [
                    ("Net cash from operating activities", 1)
                ],
                "net_cash_investing": [
                    ("Net cash from investing activities", 1)
                ],
                "net_cash_financing": [
                    ("Net cash from financing activities", 1)
                ],
                "cash_at_beginning": [
                    ("Opening cash and cash equivalents", 1)
                ],
            },
        }
        for sheet, mapping in mappings.items():
            for key, labels in mapping.items():
                for year in ("cy", "py"):
                    cells = [
                        reference(sheet, label, year, factor)
                        for label, factor in labels
                    ]
                    fields[key][year] = {
                        "value": sum_cells(cells),
                        "cells": cells,
                    }

        def evaluate(expression, values):
            tokens = expression.split()
            if (
                not tokens
                or len(tokens) % 2 == 0
                or any(t not in ("+", "-") for t in tokens[1::2])
            ):
                raise ValueError(f"Unsupported formula: {expression!r}")
            result = values[tokens[0]]
            for operator, key in zip(tokens[1::2], tokens[2::2]):
                result += values[key] if operator == "+" else -values[key]
            return result

        values = {}
        for year in ("cy", "py"):
            values[year] = {
                key: fields[key][year]["value"] for key in FIELD_KEYS
            }
            for key, expression in TOTAL_FORMULAS.items():
                values[year][key] = evaluate(expression, values[year])

        printed_totals = {
            "P&L": {
                "Gross profit": "gross_profit",
                "Profit from operations": "operating_profit",
                "Profit before tax": "profit_before_tax",
                "Net profit": "profit_for_year",
            },
            "Balance sheet": {
                "Total non-current assets": "total_non_current_assets",
                "Total current assets": "total_current_assets",
                "Total assets": "total_assets",
                "Total shareholders' funds": "total_equity",
                "Total non-current liabilities": "total_non_current_liabilities",
                "Total current liabilities": "total_current_liabilities",
                "Total liabilities": "total_liabilities",
                "TOTAL EQUITY AND LIABILITIES": "total_equity_and_liabilities",
                "Profit for the year": "profit_for_year",
                "Closing retained earnings": "retained_earnings",
            },
            "Cash flow": {
                "Net increase/(decrease) in cash": "net_change_in_cash",
                "Cash at end of year": "cash_at_end",
            },
        }

        source_totals = []
        for sheet, mapping in printed_totals.items():
            for label, formula in mapping.items():
                for year in ("cy", "py"):
                    source_totals.append({
                        **reference(sheet, label, year),
                        "year": year,
                        "formula": formula,
                    })

        # Cross-year opening equity requires explicit cell references.
        source_totals.append({
            **reference("Balance sheet", "Opening retained earnings", "cy"),
            "year": "cy",
            "sum_of": [
                reference("Balance sheet", "Retained profits", "py")
            ],
        })
        # No earlier balance sheet is supplied: verify comparative opening
        # equity against its own printed reconciliation.
        source_totals.append({
            **reference("Balance sheet", "Opening retained earnings", "py"),
            "year": "py",
            "sum_of": [
                reference("Balance sheet", "Closing retained earnings", "py"),
                reference("Balance sheet", "Profit for the year", "py", -1),
                reference(
                    "Balance sheet", "Dividends paid to shareholders", "py", -1
                ),
            ],
        })

        ignored = [
            {
                "sheet": sheet,
                "cell": cell,
                "reason": "Comparative year heading; not a monetary amount.",
            }
            for sheet, cell in year_headers
        ]

        exceptions = []

        def check(identifier, difference, explanation):
            if difference != 0:
                exceptions.append({
                    "check": identifier,
                    "difference": difference,
                    "explanation": explanation,
                })

        for total in source_totals:
            expected = (
                evaluate(total["formula"], values[total["year"]])
                if "formula" in total
                else sum_cells(total["sum_of"])
            )
            actual = total["factor"] * number(total["sheet"], total["cell"])
            check(
                f"T {total['sheet']}!{total['cell']} {total['year']}",
                actual - expected,
                "The signed printed amount differs from its mapped formula "
                "or cited reconciliation. Source amounts are preserved.",
            )

        for year in ("cy", "py"):
            v = values[year]
            check(
                f"R1 balance {year}",
                v["total_assets"] - v["total_equity_and_liabilities"],
                "Mapped assets less mapped equity and liabilities; "
                "source amounts are preserved.",
            )
            check(
                f"R2 cash_at_end {year}",
                v["cash_at_end"] - v["cash_and_cash_equivalents"],
                "Cash-flow closing cash less balance-sheet cash; "
                "source amounts are preserved.",
            )
        check(
            "R3 opening cash cy",
            values["cy"]["cash_at_beginning"]
            - values["py"]["cash_and_cash_equivalents"],
            "Current opening cash less prior balance-sheet cash; "
            "source amounts are preserved.",
        )
        check(
            "R4 retained earnings cy",
            values["cy"]["retained_earnings"]
            - (
                values["py"]["retained_earnings"]
                + values["cy"]["profit_for_year"]
                - values["cy"]["dividends"]
            ),
            "Current retained earnings less (prior retained earnings plus "
            "current profit less dividends); source amounts are preserved.",
        )

        covered = set()
        for years in fields.values():
            for line in years.values():
                covered.update((r["sheet"], r["cell"]) for r in line["cells"])
        for total in source_totals:
            covered.add((total["sheet"], total["cell"]))
            covered.update(
                (r["sheet"], r["cell"]) for r in total.get("sum_of", [])
            )
        covered.update((item["sheet"], item["cell"]) for item in ignored)
        header_cells = set(year_headers)

        for sheet in sheet_names:
            ws = wb[sheet]
            for column in columns[sheet].values():
                used_rows = [
                    ws[cell].row
                    for name, cell in covered
                    if name == sheet
                    and ws[cell].column == column
                    and (name, cell) not in header_cells
                ]
                if not used_rows:
                    continue
                for row in range(min(used_rows), ws.max_row + 1):
                    cell = ws.cell(row, column)
                    if (
                        isinstance(cell.value, Real)
                        and not isinstance(cell.value, bool)
                        and (sheet, cell.coordinate) not in covered
                    ):
                        raise ValueError(
                            f"Unclassified number at {sheet}!{cell.coordinate}."
                        )

        return {
            "entity_code": "TW-UK01",
            "source_entity_name": source_entity_name,
            "currency": currency,
            "unit": unit,
            "period_end": period_end,
            "fields": fields,
            "source_totals": source_totals,
            "ignored": ignored,
            "source_exceptions": exceptions,
            "notes": (
                "The entity headings identify Tarnwell Holdings plc "
                "(TW-UK01); amounts remain in unscaled GBP. Repeated unit "
                "headings apply to separate sections. Finance income combines "
                "subsidiary dividends and bank interest; investments combine "
                "subsidiary investments and long-term group loans. Dividends "
                "paid in the retained-earnings reconciliation are used as "
                "dividends declared because no separate declaration figure is "
                "provided; their negative source signs are reversed. Lines "
                "absent from the workbook are zero with no cell references."
            ),
        }
    finally:
        wb.close()
