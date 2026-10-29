import math
import re
from numbers import Real

import openpyxl


def extract(path) -> dict:
    field_labels = {
        "revenue": "Sales",
        "cost_of_sales": "Cost of sales",
        "distribution_costs": None,
        "administrative_expenses": "Overheads",
        "other_operating_income": "Other operating income/(expense), net",
        "finance_income": "Interest income",
        "finance_costs": "Interest expense",
        "income_tax_expense": "Taxation",
        "property_plant_equipment": "Tangible assets",
        "right_of_use_assets": "Leased assets (right of use)",
        "intangible_assets": "Goodwill and intangibles",
        "investments": None,
        "deferred_tax_assets": "Deferred tax asset",
        "inventories": None,
        "trade_receivables": "Accounts receivable",
        "other_current_assets": "Other current assets",
        "cash_and_cash_equivalents": "Cash and cash equivalents",
        "trade_payables": "Accounts payable",
        "other_current_liabilities": "Other payables and accruals",
        "current_borrowings": "Borrowings (current)",
        "current_tax_payable": "Current tax payable",
        "non_current_borrowings": "Borrowings (non-current)",
        "provisions": "Employee benefit obligations",
        "deferred_tax_liabilities": None,
        "share_capital": "Share capital",
        "retained_earnings": "Retained earnings",
        "other_reserves": None,
        "net_cash_operating": "Cash generated from operations (net)",
        "net_cash_investing": "Net cash used in investing activities",
        "net_cash_financing": "Net cash used in financing activities",
        "fx_effect_on_cash": "Foreign exchange differences on cash",
        "cash_at_beginning": "Cash at start of year",
        "dividends": "Dividends declared",
    }

    formulas = {
        "gross_profit": "revenue - cost_of_sales",
        "operating_profit": (
            "gross_profit - distribution_costs - administrative_expenses"
            " + other_operating_income"
        ),
        "profit_before_tax": "operating_profit + finance_income - finance_costs",
        "profit_for_year": "profit_before_tax - income_tax_expense",
        "total_non_current_assets": (
            "property_plant_equipment + right_of_use_assets + intangible_assets"
            " + investments + deferred_tax_assets"
        ),
        "total_current_assets": (
            "inventories + trade_receivables + other_current_assets"
            " + cash_and_cash_equivalents"
        ),
        "total_assets": "total_non_current_assets + total_current_assets",
        "total_current_liabilities": (
            "trade_payables + other_current_liabilities"
            " + current_borrowings + current_tax_payable"
        ),
        "total_non_current_liabilities": (
            "non_current_borrowings + provisions + deferred_tax_liabilities"
        ),
        "total_liabilities": (
            "total_current_liabilities + total_non_current_liabilities"
        ),
        "total_equity": "share_capital + retained_earnings + other_reserves",
        "total_equity_and_liabilities": "total_equity + total_liabilities",
        "net_change_in_cash": (
            "net_cash_operating + net_cash_investing + net_cash_financing"
        ),
        "cash_at_end": "cash_at_beginning + net_change_in_cash + fx_effect_on_cash",
    }

    total_labels = {
        "Gross margin": "gross_profit",
        "Operating profit": "operating_profit",
        "Profit before tax": "profit_before_tax",
        "Net profit": "profit_for_year",
        "Total non-current assets": "total_non_current_assets",
        "Total current assets": "total_current_assets",
        "Total assets": "total_assets",
        "Total equity": "total_equity",
        "Total non-current liabilities": "total_non_current_liabilities",
        "Total current liabilities": "total_current_liabilities",
        "Total liabilities": "total_liabilities",
        "Total equity and liabilities": "total_equity_and_liabilities",
        "Net change in cash and cash equivalents": "net_change_in_cash",
        "Closing cash and cash equivalents": "cash_at_end",
        "Profit for the year": "profit_for_year",
        "Closing retained earnings": "retained_earnings",
    }

    def normalize(value):
        return " ".join(str(value).split()).casefold()

    def numeric(value):
        return (
            isinstance(value, Real)
            and not isinstance(value, bool)
            and math.isfinite(value)
        )

    def evaluate(expression, values):
        tokens = expression.split()
        result = values[tokens[0]]
        for index in range(1, len(tokens), 2):
            operator, name = tokens[index:index + 2]
            if operator == "+":
                result += values[name]
            elif operator == "-":
                result -= values[name]
            else:
                raise ValueError("Unsupported formula operator")
        return result

    wb = openpyxl.load_workbook(path, data_only=True)
    try:
        if set(wb.sheetnames) != {"2025", "2024"}:
            raise ValueError("Expected the 2025 and 2024 worksheets")

        sheets = {"cy": wb["2025"], "py": wb["2024"]}
        indexes = {}
        names = {}
        periods = {}

        for year, ws in sheets.items():
            index = {}
            titles = []
            for row in ws.iter_rows():
                label = row[0].value
                if not isinstance(label, str):
                    continue
                key = normalize(label)
                if key in index:
                    raise ValueError(f"Duplicate label in {ws.title}: {label}")
                index[key] = row[0].row
                if " - sprawozdanie finansowe / financial statements " in key:
                    titles.append(label)

            indexes[year] = index
            if len(titles) != 1:
                raise ValueError(f"Expected one entity title in {ws.title}")

            match = re.fullmatch(
                r"\s*(.*?)\s+-\s+sprawozdanie finansowe\s*/\s*"
                r"financial statements\s+(\d{4})\s*",
                titles[0],
                flags=re.IGNORECASE,
            )
            if not match or match.group(2) != ws.title:
                raise ValueError(f"Invalid financial statement title in {ws.title}")

            names[year] = match.group(1)
            if normalize(names[year]) != normalize("Tarnwell Polska sp. z o.o."):
                raise ValueError("Workbook does not identify TW-PL01")

            if normalize("tys. PLN (PLN thousands)") not in index:
                raise ValueError(f"Expected amounts in PLN thousands in {ws.title}")

            for section in (
                "Income statement",
                "Statement of financial position",
                "Cash flow statement",
                "Retained earnings reconciliation",
            ):
                if normalize(f"{section} {ws.title}") not in index:
                    raise ValueError(f"Missing {section} in {ws.title}")

            periods[year] = f"{ws.title}-12-31"

        if normalize(names["cy"]) != normalize(names["py"]):
            raise ValueError("The worksheets identify different entities")

        used = set()

        def reference(year, label, factor=1):
            ws = sheets[year]
            key = normalize(label)
            if key not in indexes[year]:
                raise ValueError(f"Missing label in {ws.title}: {label}")
            cell = ws.cell(indexes[year][key], 2)
            if not numeric(cell.value):
                raise ValueError(f"Expected a number in {ws.title}!{cell.coordinate}")
            used.add((ws.title, cell.coordinate))
            return {"sheet": ws.title, "cell": cell.coordinate, "factor": factor}

        def amount(ref):
            return ref["factor"] * wb[ref["sheet"]][ref["cell"]].value

        fields = {}
        for key, label in field_labels.items():
            fields[key] = {}
            for year in sheets:
                refs = (
                    [reference(year, label, -1 if key == "dividends" else 1)]
                    if label is not None else []
                )
                fields[key][year] = {
                    "value": sum(amount(ref) for ref in refs),
                    "cells": refs,
                }

        values = {}
        for year in sheets:
            values[year] = {
                key: fields[key][year]["value"] for key in fields
            }
            for key, expression in formulas.items():
                values[year][key] = evaluate(expression, values[year])

        source_totals = []
        for year in sheets:
            for label, expression in total_labels.items():
                ref = reference(year, label)
                source_totals.append({**ref, "year": year, "formula": expression})

            opening = reference(year, "Opening retained earnings")
            if year == "cy":
                components = [reference("py", "Retained earnings")]
            else:
                # Opening = closing - profit - the signed dividend movement.
                components = [
                    reference(year, "Closing retained earnings"),
                    reference(year, "Profit for the year", -1),
                    reference(year, "Dividends declared", -1),
                ]
            source_totals.append({
                **opening,
                "year": year,
                "sum_of": components,
            })

        exceptions = []

        def check(check_id, difference, explanation):
            if not math.isclose(difference, 0, rel_tol=0, abs_tol=1e-9):
                exceptions.append({
                    "check": check_id,
                    "difference": difference,
                    "explanation": explanation,
                })

        for total in source_totals:
            expected = (
                evaluate(total["formula"], values[total["year"]])
                if "formula" in total
                else sum(amount(ref) for ref in total["sum_of"])
            )
            check(
                f"T {total['sheet']}!{total['cell']} {total['year']}",
                amount(total) - expected,
                "Printed source amount minus its mapped formula or component sum; "
                "source numbers are preserved.",
            )

        for year in sheets:
            v = values[year]
            check(
                f"R1 balance {year}",
                v["total_assets"] - v["total_equity_and_liabilities"],
                "Mapped assets minus mapped equity and liabilities.",
            )
            check(
                f"R2 cash_at_end {year}",
                v["cash_at_end"] - v["cash_and_cash_equivalents"],
                "Cash flow closing cash minus balance-sheet cash.",
            )

        cy, py = values["cy"], values["py"]
        check(
            "R3 opening cash cy",
            cy["cash_at_beginning"] - py["cash_and_cash_equivalents"],
            "Current-year opening cash minus prior-year balance-sheet cash.",
        )
        check(
            "R4 retained earnings cy",
            cy["retained_earnings"]
            - (py["retained_earnings"] + cy["profit_for_year"] - cy["dividends"]),
            "Current retained earnings minus prior retained earnings plus "
            "current profit less dividends.",
        )

        # Refuse to silently omit newly introduced numeric rows.
        for year, ws in sheets.items():
            first_row = min(
                wb[sheet][coordinate].row
                for sheet, coordinate in used
                if sheet == ws.title
            )
            for row in ws.iter_rows(
                min_row=first_row, min_col=2, max_col=2
            ):
                cell = row[0]
                if numeric(cell.value) and (ws.title, cell.coordinate) not in used:
                    raise ValueError(
                        f"Unmapped numeric cell: {ws.title}!{cell.coordinate}"
                    )

        return {
            "entity_code": "TW-PL01",
            "source_entity_name": names["cy"],
            "currency": "PLN",
            "unit": 1000,
            "period_end": periods,
            "fields": fields,
            "source_totals": source_totals,
            "ignored": [],
            "source_exceptions": exceptions,
            "notes": (
                "Amounts remain in the workbook's PLN thousands. Calendar "
                "year-end dates are inferred from the annual statement headings. "
                "Unsplit Overheads are classified as administrative expenses; "
                "no separate selling expense is identifiable. Lines absent from "
                "the workbook are zero. Negative dividend movements are reversed "
                "to positive declared dividends; cash flows retain their signs. "
                "Printed totals and retained earnings reconciliation amounts "
                "are retained as source checks."
            ),
        }
    finally:
        wb.close()
