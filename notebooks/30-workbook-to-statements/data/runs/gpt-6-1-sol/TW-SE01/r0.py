import math
import re
from datetime import date
from numbers import Real

import openpyxl


def extract(path) -> dict:
    wb = openpyxl.load_workbook(path, data_only=True)
    try:
        def normalize(value):
            return " ".join(str(value).split()).casefold()

        def number(cell):
            value = cell.value
            if isinstance(value, bool):
                raise ValueError(f"Boolean amount at {cell.parent.title}!{cell.coordinate}")
            if isinstance(value, Real):
                result = float(value)
            elif isinstance(value, str):
                text = re.sub(r"\s+", "", value).replace("\u2212", "-")
                if text.startswith("(") and text.endswith(")"):
                    text = "-" + text[1:-1]
                if not re.fullmatch(r"[+-]?\d+(?:[.,]\d+)?", text):
                    raise ValueError(
                        f"Invalid amount at {cell.parent.title}!{cell.coordinate}"
                    )
                result = float(text.replace(",", "."))
            else:
                raise ValueError(
                    f"Missing amount at {cell.parent.title}!{cell.coordinate}"
                )
            if not math.isfinite(result):
                raise ValueError("Non-finite amount")
            return int(result) if result.is_integer() else result

        rows = {}
        year_columns = {}
        period_end = None
        source_entity_name = None
        ignored = []

        for sheet_name in ("RR", "BR", "KF"):
            ws = wb[sheet_name]
            labels = {}
            for row in ws.iter_rows():
                value = row[0].value
                if value is not None:
                    labels.setdefault(normalize(value), []).append(row[0].row)
            rows[sheet_name] = labels

            entity = ws["A1"].value
            if not isinstance(entity, str):
                raise ValueError("Missing entity name")
            if source_entity_name is None:
                source_entity_name = entity
            elif normalize(entity) != normalize(source_entity_name):
                raise ValueError("Sheets belong to different entities")

            date_pairs = []
            for row in ws.iter_rows():
                for cell in row:
                    if isinstance(cell.value, str):
                        dates = re.findall(r"\b\d{4}-\d{2}-\d{2}\b", cell.value)
                        if len(dates) == 2 and "räkenskapsår" in normalize(cell.value):
                            date_pairs.append(dates)

            if len(date_pairs) != 1:
                raise ValueError(f"Ambiguous reporting period in {sheet_name}")
            start, end = map(date.fromisoformat, date_pairs[0])
            if start > end:
                raise ValueError("Invalid reporting period")
            dates = {
                "cy": end.isoformat(),
                "py": date(end.year - 1, end.month, end.day).isoformat(),
            }
            if period_end is None:
                period_end = dates
            elif period_end != dates:
                raise ValueError("Reporting periods differ between sheets")

            units = [
                value
                for value in labels
                if "belopp i tsek" in value and "thousands of sek" in value
            ]
            if not units:
                raise ValueError(f"Expected TSEK amounts in {sheet_name}")

            headers = []
            for row in ws.iter_rows():
                matches = {}
                for cell in row:
                    if cell.value is None:
                        continue
                    text = str(cell.value).strip()
                    if text == str(end.year):
                        matches["cy"] = cell.column
                    elif text == str(end.year - 1):
                        matches["py"] = cell.column
                if set(matches) == {"cy", "py"}:
                    headers.append((row[0].row, matches))

            if not headers:
                raise ValueError(f"Missing year columns in {sheet_name}")
            year_columns[sheet_name] = headers[0][1]
            for header_row, columns in headers:
                if columns != year_columns[sheet_name]:
                    raise ValueError("Year columns differ within a sheet")
                for column in columns.values():
                    ignored.append({
                        "sheet": sheet_name,
                        "cell": ws.cell(header_row, column).coordinate,
                        "reason": "Year heading; not a financial amount.",
                    })

        if normalize(source_entity_name) not in {
            normalize("Tarnwell Nordic Aktiebolag"),
            normalize("Tarnwell Nordic AB"),
        }:
            raise ValueError("Workbook does not identify Tarnwell Nordic AB")

        def find_row(sheet, label):
            matches = rows[sheet].get(normalize(label), [])
            if len(matches) != 1:
                raise ValueError(f"Expected one row labelled {label!r} in {sheet}")
            return matches[0]

        def reference(sheet, label, year, factor=1):
            ws = wb[sheet]
            cell = ws.cell(find_row(sheet, label), year_columns[sheet][year])
            number(cell)
            return {"sheet": sheet, "cell": cell.coordinate, "factor": factor}

        def reference_value(ref):
            return ref["factor"] * number(wb[ref["sheet"]][ref["cell"]])

        mappings = {
            "revenue": ("RR", [("Turnover", 1)]),
            "cost_of_sales": ("RR", [("Direct costs", -1)]),
            "distribution_costs": ("RR", [("Selling expenses", -1)]),
            "administrative_expenses": ("RR", [("Overheads", -1)]),
            "other_operating_income": (
                "RR", [("Other operating income/(expense), net", 1)]
            ),
            "finance_income": ("RR", [("Finance income", 1)]),
            "finance_costs": ("RR", [("Interest paid on borrowings", -1)]),
            "income_tax_expense": ("RR", [("Income tax expense", -1)]),
            "property_plant_equipment": (
                "BR", [("Property, plant and equipment", 1)]
            ),
            "right_of_use_assets": ("BR", [("Right-of-use assets", 1)]),
            "intangible_assets": ("BR", [("Goodwill and intangibles", 1)]),
            "investments": ("BR", []),
            "deferred_tax_assets": ("BR", [("Deferred tax asset", 1)]),
            "inventories": ("BR", [("Stock", 1)]),
            "trade_receivables": ("BR", [("Accounts receivable", 1)]),
            "other_current_assets": ("BR", [("Other current assets", 1)]),
            "cash_and_cash_equivalents": (
                "BR",
                [
                    ("Cash at bank", 1),
                    ("Money market deposits (up to 3 months)", 1),
                ],
            ),
            "trade_payables": ("BR", [("Accounts payable", 1)]),
            "other_current_liabilities": (
                "BR", [("Accruals and other payables", 1)]
            ),
            "current_borrowings": ("BR", [("Short-term borrowings", 1)]),
            "current_tax_payable": ("BR", [("Corporation tax", 1)]),
            "non_current_borrowings": ("BR", [("Loans due after one year", 1)]),
            "provisions": ("BR", [("Provisions", 1)]),
            "deferred_tax_liabilities": ("BR", [("Deferred tax liabilities", 1)]),
            "share_capital": ("BR", [("Share capital", 1)]),
            "retained_earnings": ("BR", [("Retained earnings", 1)]),
            "other_reserves": ("BR", [("Translation reserve", 1)]),
            "net_cash_operating": (
                "KF", [("Net cash from operating activities", 1)]
            ),
            "net_cash_investing": (
                "KF", [("Net cash from investing activities", 1)]
            ),
            "net_cash_financing": (
                "KF", [("Net cash from financing activities", 1)]
            ),
            "fx_effect_on_cash": (
                "KF", [("Foreign exchange differences on cash", 1)]
            ),
            "cash_at_beginning": ("KF", [("Cash at start of year", 1)]),
            "dividends": ("BR", [("Dividends paid to shareholders", -1)]),
        }

        fields = {}
        for key, (sheet, components) in mappings.items():
            fields[key] = {}
            for year in ("cy", "py"):
                cells = [
                    reference(sheet, label, year, factor)
                    for label, factor in components
                ]
                fields[key][year] = {
                    "value": sum(reference_value(ref) for ref in cells),
                    "cells": cells,
                }

        total_formulas = {
            "gross_profit": "revenue - cost_of_sales",
            "operating_profit":
                "gross_profit - distribution_costs - administrative_expenses"
                " + other_operating_income",
            "profit_before_tax":
                "operating_profit + finance_income - finance_costs",
            "profit_for_year": "profit_before_tax - income_tax_expense",
            "total_non_current_assets":
                "property_plant_equipment + right_of_use_assets"
                " + intangible_assets + investments + deferred_tax_assets",
            "total_current_assets":
                "inventories + trade_receivables + other_current_assets"
                " + cash_and_cash_equivalents",
            "total_assets": "total_non_current_assets + total_current_assets",
            "total_current_liabilities":
                "trade_payables + other_current_liabilities"
                " + current_borrowings + current_tax_payable",
            "total_non_current_liabilities":
                "non_current_borrowings + provisions + deferred_tax_liabilities",
            "total_liabilities":
                "total_current_liabilities + total_non_current_liabilities",
            "total_equity": "share_capital + retained_earnings + other_reserves",
            "total_equity_and_liabilities": "total_equity + total_liabilities",
            "net_change_in_cash":
                "net_cash_operating + net_cash_investing + net_cash_financing",
            "cash_at_end":
                "cash_at_beginning + net_change_in_cash + fx_effect_on_cash",
        }

        def expression_value(expression, values):
            tokens = expression.split()
            if not tokens or len(tokens) % 2 == 0:
                raise ValueError("Invalid internal formula")
            result = values[tokens[0]]
            for index in range(1, len(tokens), 2):
                operator, key = tokens[index:index + 2]
                if operator == "+":
                    result += values[key]
                elif operator == "-":
                    result -= values[key]
                else:
                    raise ValueError("Invalid internal operator")
            return result

        calculated = {}
        for year in ("cy", "py"):
            values = {
                key: by_year[year]["value"] for key, by_year in fields.items()
            }
            for key, formula in total_formulas.items():
                values[key] = expression_value(formula, values)
            calculated[year] = values

        printed_totals = [
            ("RR", "Gross profit", "gross_profit"),
            ("RR", "Profit from operations", "operating_profit"),
            ("RR", "Profit before tax", "profit_before_tax"),
            ("RR", "Profit after tax", "profit_for_year"),
            ("BR", "Total non-current assets", "total_non_current_assets"),
            ("BR", "Total current assets", "total_current_assets"),
            ("BR", "TOTAL ASSETS", "total_assets"),
            ("BR", "Total shareholders' funds", "total_equity"),
            ("BR", "Total non-current liabilities", "total_non_current_liabilities"),
            ("BR", "Total current liabilities", "total_current_liabilities"),
            ("BR", "Total liabilities", "total_liabilities"),
            ("BR", "Total equity and liabilities", "total_equity_and_liabilities"),
            ("BR", "Profit for the year", "profit_for_year"),
            ("BR", "Closing retained earnings", "retained_earnings"),
            ("KF", "Net increase/(decrease) in cash", "net_change_in_cash"),
            ("KF", "Cash at end of year", "cash_at_end"),
        ]

        source_totals = []
        for sheet, label, formula in printed_totals:
            for year in ("cy", "py"):
                source_totals.append({
                    **reference(sheet, label, year),
                    "year": year,
                    "formula": formula,
                })

        # Check the movement schedule's own arithmetic as well as its closing
        # balance against the balance sheet.
        for year in ("cy", "py"):
            source_totals.append({
                **reference("BR", "Closing retained earnings", year),
                "year": year,
                "sum_of": [
                    reference("BR", "Opening retained earnings", year),
                    reference("BR", "Profit for the year", year),
                    reference("BR", "Dividends paid to shareholders", year),
                ],
            })

        source_exceptions = []

        def check(check_id, difference, explanation):
            if not math.isclose(difference, 0, rel_tol=0, abs_tol=1e-7):
                source_exceptions.append({
                    "check": check_id,
                    "difference": difference,
                    "explanation": explanation,
                })

        # Multiple validations of one printed cell share the prescribed ID.
        # Emit at most one exception per ID, retaining all distinct failures
        # in the explanation.
        total_failures = {}
        for item in source_totals:
            actual = reference_value(item)
            if "formula" in item:
                expected = expression_value(
                    item["formula"], calculated[item["year"]]
                )
                description = f"template expression {item['formula']}"
            else:
                expected = sum(reference_value(ref) for ref in item["sum_of"])
                description = "opening retained earnings plus profit less dividends"
            difference = actual - expected
            if not math.isclose(difference, 0, rel_tol=0, abs_tol=1e-7):
                check_id = f"T {item['sheet']}!{item['cell']} {item['year']}"
                explanation = (
                    f"Printed amount less {description} is {difference}; "
                    "source figures are preserved."
                )
                if check_id not in total_failures:
                    total_failures[check_id] = {
                        "check": check_id,
                        "difference": difference,
                        "explanation": explanation,
                    }
                else:
                    total_failures[check_id]["explanation"] += " " + explanation
        source_exceptions.extend(total_failures.values())

        for year in ("cy", "py"):
            values = calculated[year]
            check(
                f"R1 balance {year}",
                values["total_assets"] - values["total_equity_and_liabilities"],
                "Assets less equity and liabilities; source figures are preserved.",
            )
            check(
                f"R2 cash_at_end {year}",
                values["cash_at_end"] - values["cash_and_cash_equivalents"],
                "Cash-flow closing cash less balance-sheet cash equivalents; "
                "source figures are preserved.",
            )

        check(
            "R3 opening cash cy",
            calculated["cy"]["cash_at_beginning"]
            - calculated["py"]["cash_and_cash_equivalents"],
            "Current opening cash less prior balance-sheet cash equivalents; "
            "source figures are preserved.",
        )
        check(
            "R4 retained earnings cy",
            calculated["cy"]["retained_earnings"]
            - (
                calculated["py"]["retained_earnings"]
                + calculated["cy"]["profit_for_year"]
                - calculated["cy"]["dividends"]
            ),
            "Current retained earnings less prior retained earnings plus "
            "current profit less dividends; source figures are preserved.",
        )

        return {
            "entity_code": "TW-SE01",
            "source_entity_name": source_entity_name,
            "currency": "SEK",
            "unit": 1000,
            "period_end": period_end,
            "fields": fields,
            "source_totals": source_totals,
            "ignored": ignored,
            "source_exceptions": source_exceptions,
            "notes": (
                "Aktiebolag is the expanded form of AB, identifying TW-SE01. "
                "Amounts remain in the workbook's TSEK unit; nonbreaking-space "
                "grouping and Unicode minus signs are parsed. Negative expense "
                "presentation is reversed, while cash flows and reserves retain "
                "their signs. Money market deposits up to three months are "
                "classified as cash equivalents. Investments are absent and "
                "reported as zero. Dividends paid in the retained earnings "
                "movement are treated as the year's declared distributions "
                "because that schedule deducts them from retained earnings."
            ),
        }
    finally:
        wb.close()
