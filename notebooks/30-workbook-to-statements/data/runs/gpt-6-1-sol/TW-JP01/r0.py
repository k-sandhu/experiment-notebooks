import calendar
import math
import re
import unicodedata

import openpyxl


def extract(path) -> dict:
    wb = openpyxl.load_workbook(path, data_only=True)

    def normalise(value):
        return " ".join(unicodedata.normalize("NFKC", str(value)).split())

    def number(sheet, address):
        value = wb[sheet][address].value
        if (
            isinstance(value, bool)
            or not isinstance(value, (int, float))
            or not math.isfinite(value)
        ):
            raise ValueError(f"Expected a finite number at {sheet}!{address}")
        return value

    def row_for(sheet, label, after=0, before=None):
        ws = wb[sheet]
        stop = before if before is not None else ws.max_row + 1
        matches = [
            row
            for row in range(after + 1, stop)
            if normalise(ws.cell(row, 1).value) == normalise(label)
        ]
        if len(matches) != 1:
            raise ValueError(
                f"Expected one row labelled {label!r} in {sheet}; found {matches}"
            )
        return matches[0]

    sheets = ("損益計算書", "貸借対照表", "CF計算書", "株主資本等変動")
    for sheet in sheets:
        if sheet not in wb.sheetnames:
            raise ValueError(f"Missing sheet: {sheet}")

    try:
        source_name = wb[sheets[0]]["A1"].value
        expected_name = "ターンウェル・ジャパン株式会社 (Tarnwell Japan KK)"
        if not isinstance(source_name, str):
            raise ValueError("Missing source entity name")
        for sheet in sheets:
            if normalise(wb[sheet]["A1"].value) != normalise(expected_name):
                raise ValueError(f"Unexpected entity on {sheet}")

        # The legal entity identifies the register's JPY entity.
        # Each statement must explicitly declare millions of yen.
        for sheet in sheets:
            declarations = [
                normalise(cell.value)
                for row in wb[sheet].iter_rows()
                for cell in row
                if isinstance(cell.value, str) and "単位" in cell.value
            ]
            if declarations != [normalise("（単位：百万円）")]:
                raise ValueError(f"Unexpected or ambiguous unit on {sheet}")

        columns = {}
        dates_by_sheet = {}
        for sheet in sheets:
            found = {}
            for row in wb[sheet].iter_rows():
                for cell in row:
                    if not isinstance(cell.value, str):
                        continue
                    text = normalise(cell.value)
                    match = re.fullmatch(
                        r"(前期|当期)\s*\((\d{4})年(\d{1,2})月期\)", text
                    )
                    if match:
                        year_key = "cy" if match.group(1) == "当期" else "py"
                        if year_key in found:
                            raise ValueError(f"Duplicate year header on {sheet}")
                        year, month = int(match.group(2)), int(match.group(3))
                        day = calendar.monthrange(year, month)[1]
                        found[year_key] = (
                            cell.column_letter,
                            f"{year:04d}-{month:02d}-{day:02d}",
                        )
            if set(found) != {"cy", "py"}:
                raise ValueError(f"Missing year headers on {sheet}")
            columns[sheet] = {key: item[0] for key, item in found.items()}
            dates_by_sheet[sheet] = {key: item[1] for key, item in found.items()}

        period_end = dates_by_sheet["貸借対照表"]
        if any(dates != period_end for dates in dates_by_sheet.values()):
            raise ValueError("Statement dates disagree")

        keys = (
            "revenue cost_of_sales distribution_costs administrative_expenses "
            "other_operating_income finance_income finance_costs income_tax_expense "
            "property_plant_equipment right_of_use_assets intangible_assets investments "
            "deferred_tax_assets inventories trade_receivables other_current_assets "
            "cash_and_cash_equivalents trade_payables other_current_liabilities "
            "current_borrowings current_tax_payable non_current_borrowings provisions "
            "deferred_tax_liabilities share_capital retained_earnings other_reserves "
            "net_cash_operating net_cash_investing net_cash_financing fx_effect_on_cash "
            "cash_at_beginning dividends"
        ).split()
        fields = {
            key: {year: {"value": 0, "cells": []} for year in ("cy", "py")}
            for key in keys
        }
        source_totals = []
        accounted = set()

        def reference(sheet, row, year, factor=1):
            address = f"{columns[sheet][year]}{row}"
            number(sheet, address)
            accounted.add((sheet, address))
            return {"sheet": sheet, "cell": address, "factor": factor}

        def add_field(key, sheet, label, factor=1, after=0, before=None):
            row = row_for(sheet, label, after, before)
            for year in ("cy", "py"):
                ref = reference(sheet, row, year, factor)
                line = fields[key][year]
                line["cells"].append(ref)
                line["value"] += factor * number(sheet, ref["cell"])

        def add_total(sheet, label, formula):
            row = row_for(sheet, label)
            for year in ("cy", "py"):
                ref = reference(sheet, row, year)
                source_totals.append({**ref, "year": year, "formula": formula})

        income = "損益計算書"
        for key, label in (
            ("revenue", "売上高 Net sales"),
            ("cost_of_sales", "売上原価 Cost of sales"),
            ("distribution_costs", "販売費 Selling expenses"),
            ("administrative_expenses", "一般管理費 General and administrative"),
            ("other_operating_income", "雑収入 Miscellaneous income"),
            ("finance_income", "受取利息 Interest income"),
            ("finance_costs", "支払利息 Interest expense"),
            ("income_tax_expense", "法人税、住民税及び事業税 Current income taxes"),
            ("income_tax_expense", "法人税等調整額 Deferred income taxes"),
        ):
            add_field(key, income, label)

        for label, formula in (
            ("売上総利益 Gross profit", "gross_profit"),
            (
                "販売費及び一般管理費合計 Total SG&A",
                "distribution_costs + administrative_expenses",
            ),
            (
                "営業利益 Operating income",
                "operating_profit - other_operating_income",
            ),
            (
                "営業外収益合計 Total non-operating income",
                "finance_income + other_operating_income",
            ),
            ("経常利益 Ordinary income", "profit_before_tax"),
            ("税引前当期純利益 Income before income taxes", "profit_before_tax"),
            ("当期純利益 Net income", "profit_for_year"),
        ):
            add_total(income, label, formula)

        balance = "貸借対照表"
        for key, label in (
            ("inventories", "棚卸資産 Inventories"),
            ("trade_receivables", "売掛金 Accounts receivable"),
            ("cash_and_cash_equivalents", "現金及び預金 Cash and deposits"),
            ("property_plant_equipment", "有形固定資産 Property and equipment"),
            ("right_of_use_assets", "使用権資産 Right-of-use assets"),
            ("intangible_assets", "無形固定資産 Intangible assets"),
            ("deferred_tax_assets", "繰延税金資産 Deferred tax assets"),
            ("trade_payables", "買掛金 Accounts payable"),
            ("current_borrowings", "短期借入金 Short-term loans"),
            ("current_tax_payable", "未払法人税等 Income taxes payable"),
            ("non_current_borrowings", "長期借入金 Long-term loans"),
            ("provisions", "退職給付に係る負債 Retirement benefit liability"),
            ("deferred_tax_liabilities", "繰延税金負債 Deferred tax liabilities"),
            ("share_capital", "資本金 Capital stock"),
            ("share_capital", "資本剰余金 Capital surplus"),
            ("retained_earnings", "利益剰余金 Retained earnings"),
            (
                "other_reserves",
                "その他の包括利益累計額 Accumulated other comprehensive income",
            ),
        ):
            add_field(key, balance, label)

        for key, start_label, end_label in (
            (
                "other_current_assets",
                "流動資産 Current assets",
                "流動資産合計 Total current assets",
            ),
            (
                "other_current_liabilities",
                "流動負債 Current liabilities",
                "流動負債合計 Total current liabilities",
            ),
        ):
            add_field(
                key,
                balance,
                "その他 Other",
                after=row_for(balance, start_label),
                before=row_for(balance, end_label),
            )

        for label, formula in (
            ("流動資産合計 Total current assets", "total_current_assets"),
            ("固定資産合計 Total non-current assets", "total_non_current_assets"),
            ("資産合計 Total assets", "total_assets"),
            ("流動負債合計 Total current liabilities", "total_current_liabilities"),
            ("負債合計 Total liabilities", "total_liabilities"),
            ("純資産合計 Total net assets", "total_equity"),
            (
                "負債純資産合計 Total liabilities and net assets",
                "total_equity_and_liabilities",
            ),
        ):
            add_total(balance, label, formula)

        cashflow = "CF計算書"
        for key, label in (
            ("net_cash_operating", "営業活動によるキャッシュ・フロー Operating activities"),
            ("net_cash_investing", "投資活動によるキャッシュ・フロー Investing activities"),
            ("net_cash_financing", "財務活動によるキャッシュ・フロー Financing activities"),
            ("fx_effect_on_cash", "現金及び現金同等物に係る換算差額 FX effect"),
            ("cash_at_beginning", "現金及び現金同等物の期首残高 Beginning balance"),
        ):
            add_field(key, cashflow, label)
        add_total(
            cashflow, "現金及び現金同等物の増減額 Net change", "net_change_in_cash"
        )
        add_total(
            cashflow, "現金及び現金同等物の期末残高 Ending balance", "cash_at_end"
        )

        equity = "株主資本等変動"
        add_field("dividends", equity, "剰余金の配当 Dividends", factor=-1)
        for label, formula in (
            (
                "当期首残高 Beginning balance",
                "retained_earnings - profit_for_year + dividends",
            ),
            ("当期純利益 Net income", "profit_for_year"),
            ("当期末残高 Ending balance", "retained_earnings"),
        ):
            add_total(equity, label, formula)

        total_formulas = {
            "gross_profit": "revenue - cost_of_sales",
            "operating_profit": (
                "gross_profit - distribution_costs - administrative_expenses "
                "+ other_operating_income"
            ),
            "profit_before_tax": "operating_profit + finance_income - finance_costs",
            "profit_for_year": "profit_before_tax - income_tax_expense",
            "total_non_current_assets": (
                "property_plant_equipment + right_of_use_assets + intangible_assets "
                "+ investments + deferred_tax_assets"
            ),
            "total_current_assets": (
                "inventories + trade_receivables + other_current_assets "
                "+ cash_and_cash_equivalents"
            ),
            "total_assets": "total_non_current_assets + total_current_assets",
            "total_current_liabilities": (
                "trade_payables + other_current_liabilities + current_borrowings "
                "+ current_tax_payable"
            ),
            "total_non_current_liabilities": (
                "non_current_borrowings + provisions + deferred_tax_liabilities"
            ),
            "total_liabilities": (
                "total_current_liabilities + total_non_current_liabilities"
            ),
            "total_equity": "share_capital + retained_earnings + other_reserves",
            "total_equity_and_liabilities": "total_equity + total_liabilities",
            "net_change_in_cash": (
                "net_cash_operating + net_cash_investing + net_cash_financing"
            ),
            "cash_at_end": "cash_at_beginning + net_change_in_cash + fx_effect_on_cash",
        }

        def evaluate(expression, values):
            tokens = expression.split()
            if len(tokens) % 2 != 1 or tokens[0] not in values:
                raise ValueError(f"Invalid formula: {expression}")
            result = values[tokens[0]]
            for index in range(1, len(tokens), 2):
                operator, name = tokens[index:index + 2]
                if operator not in ("+", "-") or name not in values:
                    raise ValueError(f"Invalid formula: {expression}")
                result += values[name] if operator == "+" else -values[name]
            return result

        computed = {}
        for year in ("cy", "py"):
            values = {key: fields[key][year]["value"] for key in keys}
            for key, expression in total_formulas.items():
                values[key] = evaluate(expression, values)
            computed[year] = values

        exceptions = []

        def check(check_id, difference, explanation):
            if not math.isclose(difference, 0, rel_tol=0, abs_tol=1e-9):
                exceptions.append({
                    "check": check_id,
                    "difference": difference,
                    "explanation": explanation,
                })

        for total in source_totals:
            actual = total["factor"] * number(total["sheet"], total["cell"])
            expected = evaluate(total["formula"], computed[total["year"]])
            check(
                f"T {total['sheet']}!{total['cell']} {total['year']}",
                actual - expected,
                "Printed amount minus the mapped template expression "
                f"({total['formula']}); source amounts are preserved.",
            )

        for year in ("cy", "py"):
            values = computed[year]
            check(
                f"R1 balance {year}",
                values["total_assets"] - values["total_equity_and_liabilities"],
                "Total assets minus total equity and liabilities.",
            )
            check(
                f"R2 cash_at_end {year}",
                values["cash_at_end"] - values["cash_and_cash_equivalents"],
                "Cash-flow closing cash minus balance-sheet cash.",
            )

        check(
            "R3 opening cash cy",
            computed["cy"]["cash_at_beginning"]
            - computed["py"]["cash_and_cash_equivalents"],
            "Current-year opening cash minus prior-year balance-sheet cash.",
        )
        check(
            "R4 retained earnings cy",
            computed["py"]["retained_earnings"]
            + computed["cy"]["profit_for_year"]
            - computed["cy"]["dividends"]
            - computed["cy"]["retained_earnings"],
            "Prior retained earnings plus current profit less dividends "
            "minus current retained earnings.",
        )

        # Fail explicitly if a changed workbook introduces an unclassified amount.
        for sheet in sheets:
            ws = wb[sheet]
            for column in set(columns[sheet].values()):
                used_rows = [
                    ws[address].row
                    for used_sheet, address in accounted
                    if used_sheet == sheet
                    and ws[address].column_letter == column
                ]
                for row in range(min(used_rows), ws.max_row + 1):
                    cell = ws[f"{column}{row}"]
                    if (
                        isinstance(cell.value, (int, float))
                        and not isinstance(cell.value, bool)
                        and (sheet, cell.coordinate) not in accounted
                    ):
                        raise ValueError(
                            f"Unclassified numeric amount: {sheet}!{cell.coordinate}"
                        )

        return {
            "entity_code": "TW-JP01",
            "source_entity_name": source_name,
            "currency": "JPY",
            "unit": 1000000,
            "period_end": period_end,
            "fields": fields,
            "source_totals": source_totals,
            "ignored": [],
            "source_exceptions": exceptions,
            "notes": (
                "The workbook identifies Tarnwell Japan K.K.; amounts remain in "
                "millions of JPY. Miscellaneous income is mapped to other operating "
                "income, so Japanese operating income is checked against group "
                "operating profit less that income. Capital surplus is combined "
                "with capital stock; accumulated OCI retains its sign. Cash and "
                "deposits is treated as cash equivalents because both years agree "
                "with the cash-flow closing balances and no restricted or longer-term "
                "deposits are identified. Dividend outflows are sign-reversed. "
                "Investments and other absent input lines are zero with no citations. "
                "The retained-earnings schedule's opening balances are checked against "
                "closing retained earnings less profit plus dividends."
            ),
        }
    finally:
        wb.close()
