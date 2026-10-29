import calendar
import datetime as dt
import math
import re
from decimal import Decimal
from numbers import Real

import openpyxl


def extract(path) -> dict:
    wb = openpyxl.load_workbook(path, data_only=True)
    try:
        income = "Erfolgsrechnung"
        balance = "Bilanz"
        cash = "Geldflussrechnung"
        sheet_names = (income, balance, cash)
        periods = ("cy", "py")

        for sheet in sheet_names:
            if sheet not in wb.sheetnames:
                raise ValueError(f"Missing sheet: {sheet}")

        def normalise(value):
            return " ".join(str(value).split())

        def number(sheet, address):
            value = wb[sheet][address].value
            if isinstance(value, bool):
                raise ValueError(f"Boolean amount at {sheet}!{address}")
            if isinstance(value, Decimal):
                result = value
            elif isinstance(value, Real):
                if not math.isfinite(value):
                    raise ValueError(f"Non-finite amount at {sheet}!{address}")
                result = Decimal(str(value))
            elif isinstance(value, str):
                text = value.strip().replace("\u2212", "-")
                text = re.sub(r"['’\s]", "", text)
                if not re.fullmatch(r"[+-]?\d+(?:\.\d+)?", text):
                    raise ValueError(f"Non-numeric amount at {sheet}!{address}")
                result = Decimal(text)
            else:
                raise ValueError(f"Missing numeric amount at {sheet}!{address}")
            if not result.is_finite():
                raise ValueError(f"Non-finite amount at {sheet}!{address}")
            return result

        def json_number(value):
            return int(value) if value == value.to_integral_value() else float(value)

        # Repeated headings are legitimate. Require uniqueness only when
        # resolving a label used for an amount.
        rows = {}
        for sheet in sheet_names:
            index = {}
            for row in wb[sheet].iter_rows(min_col=2, max_col=2):
                cell = row[0]
                if isinstance(cell.value, str):
                    index.setdefault(normalise(cell.value), []).append(cell.row)
            rows[sheet] = index

        def row_for(sheet, label):
            matches = rows[sheet].get(normalise(label), [])
            if not matches:
                raise ValueError(f"Missing label in {sheet}: {label}")
            if len(matches) != 1:
                raise ValueError(f"Ambiguous amount label in {sheet}: {label}")
            return matches[0]

        names = []
        for sheet in sheet_names:
            matches = [
                cell.value.strip()
                for row in wb[sheet].iter_rows()
                for cell in row
                if isinstance(cell.value, str)
                and re.match(r"^\s*Tarnwell Schweiz AG(?:\s|$)", cell.value)
            ]
            if len(matches) != 1:
                raise ValueError(f"Cannot identify exactly one entity in {sheet}")
            names.append(matches[0])

            unit_texts = [
                normalise(cell.value)
                for row in wb[sheet].iter_rows()
                for cell in row
                if isinstance(cell.value, str) and "TCHF" in cell.value
            ]
            if not unit_texts or any(
                not re.fullmatch(r"in TCHF \(CHF 1['’]000\)", text)
                for text in unit_texts
            ):
                raise ValueError(f"Cannot confirm CHF thousands in {sheet}")

        if len(set(names)) != 1:
            raise ValueError("Entity names differ between statements")

        date_matches = []
        for label, occurrences in rows[balance].items():
            match = re.fullmatch(r"per\s+(\d{2})\.(\d{2})\.(\d{4})", label)
            if match:
                date_matches.extend([match] * len(occurrences))
        if len(date_matches) != 1:
            raise ValueError("Cannot identify the balance-sheet date")

        day, month, year = map(int, date_matches[0].groups())
        cy_date = dt.date(year, month, day)
        py_date = dt.date(
            year - 1, month,
            min(day, calendar.monthrange(year - 1, month)[1]),
        )

        columns = {}
        header_rows = {}
        for sheet in sheet_names:
            matches = []
            for row in wb[sheet].iter_rows():
                found = {}
                for period, expected in (("cy", year), ("py", year - 1)):
                    candidates = [
                        cell.column_letter
                        for cell in row
                        if str(cell.value).strip()
                        in (str(expected), f"{expected}.0")
                    ]
                    if len(candidates) == 1:
                        found[period] = candidates[0]
                if set(found) == set(periods):
                    matches.append((row[0].row, found))
            if not matches:
                raise ValueError(f"Missing year columns in {sheet}")
            if any(mapping != matches[0][1] for _, mapping in matches):
                raise ValueError(f"Inconsistent year columns in {sheet}")
            columns[sheet] = matches[0][1]
            header_rows[sheet] = [row for row, _ in matches]

        def address(sheet, label, period):
            return f"{columns[sheet][period]}{row_for(sheet, label)}"

        keys = (
            "revenue cost_of_sales distribution_costs administrative_expenses "
            "other_operating_income finance_income finance_costs income_tax_expense "
            "property_plant_equipment right_of_use_assets intangible_assets "
            "investments deferred_tax_assets inventories trade_receivables "
            "other_current_assets cash_and_cash_equivalents trade_payables "
            "other_current_liabilities current_borrowings current_tax_payable "
            "non_current_borrowings provisions deferred_tax_liabilities "
            "share_capital retained_earnings other_reserves net_cash_operating "
            "net_cash_investing net_cash_financing fx_effect_on_cash "
            "cash_at_beginning dividends"
        ).split()
        fields = {
            key: {period: {"value": 0, "cells": []} for period in periods}
            for key in keys
        }
        values = {
            period: {key: Decimal(0) for key in keys}
            for period in periods
        }

        def map_line(key, sheet, *components):
            for period in periods:
                cells = []
                total = Decimal(0)
                for label, factor in components:
                    cell = address(sheet, label, period)
                    total += Decimal(str(factor)) * number(sheet, cell)
                    cells.append({"sheet": sheet, "cell": cell, "factor": factor})
                values[period][key] = total
                fields[key][period] = {
                    "value": json_number(total),
                    "cells": cells,
                }

        for key, label, factor in (
            ("revenue", "Umsatzerlöse", 1),
            ("administrative_expenses", "Allgemeine Verwaltungskosten", -1),
            ("finance_costs", "Zinsen und ähnliche Aufwendungen", -1),
            ("income_tax_expense", "Steuern vom Einkommen und vom Ertrag", -1),
        ):
            map_line(key, income, (label, factor))

        map_line(
            "finance_income", income,
            ("Zinserträge aus Ausleihungen an verbundene Unternehmen", 1),
            ("Sonstige Zinsen und ähnliche Erträge", 1),
        )

        for key, label in (
            ("property_plant_equipment", "Sachanlagen"),
            ("investments", "Finanzanlagen (Ausleihungen an verbundene Unternehmen)"),
            ("other_current_assets", "Sonstige Vermögensgegenstände"),
            ("cash_and_cash_equivalents",
             "Kassenbestand, Guthaben bei Kreditinstituten"),
            ("share_capital", "Aktienkapital"),
            ("retained_earnings", "Gewinnrücklagen und Bilanzgewinn"),
            ("non_current_borrowings", "Langfristige Finanzverbindlichkeiten"),
            ("trade_payables", "Verbindlichkeiten aus Lieferungen und Leistungen"),
            ("other_current_liabilities", "Sonstige Verbindlichkeiten"),
            ("current_borrowings", "Kurzfristige Finanzverbindlichkeiten"),
            ("current_tax_payable", "Steuerrückstellungen"),
        ):
            map_line(key, balance, (label, 1))

        map_line(
            "dividends", balance,
            ("Ausschüttung an die Gesellschafter", -1),
        )

        for key, label in (
            ("net_cash_operating", "Cashflow aus laufender Geschäftstätigkeit"),
            ("net_cash_investing", "Cashflow aus der Investitionstätigkeit"),
            ("net_cash_financing", "Cashflow aus der Finanzierungstätigkeit"),
            ("fx_effect_on_cash",
             "Wechselkursbedingte Änderungen des Finanzmittelfonds"),
            ("cash_at_beginning", "Finanzmittelfonds am Anfang der Periode"),
        ):
            map_line(key, cash, (label, 1))

        formulas = {
            "gross_profit": "revenue - cost_of_sales",
            "operating_profit":
                "gross_profit - distribution_costs - administrative_expenses "
                "+ other_operating_income",
            "profit_before_tax": "operating_profit + finance_income - finance_costs",
            "profit_for_year": "profit_before_tax - income_tax_expense",
            "total_non_current_assets":
                "property_plant_equipment + right_of_use_assets + intangible_assets "
                "+ investments + deferred_tax_assets",
            "total_current_assets":
                "inventories + trade_receivables + other_current_assets "
                "+ cash_and_cash_equivalents",
            "total_assets": "total_non_current_assets + total_current_assets",
            "total_current_liabilities":
                "trade_payables + other_current_liabilities + current_borrowings "
                "+ current_tax_payable",
            "total_non_current_liabilities":
                "non_current_borrowings + provisions + deferred_tax_liabilities",
            "total_liabilities":
                "total_current_liabilities + total_non_current_liabilities",
            "total_equity": "share_capital + retained_earnings + other_reserves",
            "total_equity_and_liabilities": "total_equity + total_liabilities",
            "net_change_in_cash":
                "net_cash_operating + net_cash_investing + net_cash_financing",
            "cash_at_end": "cash_at_beginning + net_change_in_cash + fx_effect_on_cash",
        }

        def evaluate(expression, available):
            total = Decimal(0)
            sign = 1
            for token in expression.split():
                if token == "+":
                    sign = 1
                elif token == "-":
                    sign = -1
                else:
                    total += sign * available[token]
            return total

        for period in periods:
            for key, expression in formulas.items():
                values[period][key] = evaluate(expression, values[period])

        source_totals = []
        exceptions = []

        def check(check_id, difference, explanation):
            if difference != 0:
                exceptions.append({
                    "check": check_id,
                    "difference": json_number(difference),
                    "explanation": explanation,
                })

        def add_total(sheet, label, formula):
            for period in periods:
                cell = address(sheet, label, period)
                source_totals.append({
                    "sheet": sheet,
                    "cell": cell,
                    "year": period,
                    "factor": 1,
                    "formula": formula,
                })
                check(
                    f"T {sheet}!{cell} {period}",
                    number(sheet, cell) - evaluate(formula, values[period]),
                    f"Printed '{label}' less the mapped expression "
                    f"'{formula}', in TCHF. Source amounts are preserved.",
                )

        for sheet, label, formula in (
            (income, "Bruttoergebnis vom Umsatz", "gross_profit"),
            (income, "Betriebsergebnis", "operating_profit"),
            (income, "Ergebnis vor Steuern", "profit_before_tax"),
            (income, "Jahresüberschuss", "profit_for_year"),
            (balance, "Summe Anlagevermögen", "total_non_current_assets"),
            (balance, "Summe Umlaufvermögen", "total_current_assets"),
            (balance, "Summe Aktiva", "total_assets"),
            (balance, "Summe Eigenkapital", "total_equity"),
            (balance, "Summe langfristiges Fremdkapital",
             "total_non_current_liabilities"),
            (balance, "Summe kurzfristiges Fremdkapital",
             "total_current_liabilities"),
            (balance, "Summe Fremdkapital", "total_liabilities"),
            (balance, "Summe Passiva", "total_equity_and_liabilities"),
            (balance, "Jahresüberschuss", "profit_for_year"),
            (balance, "Stand 31. Dezember", "retained_earnings"),
            (cash, "Zahlungswirksame Veränderung des Finanzmittelfonds",
             "net_change_in_cash"),
            (cash, "Finanzmittelfonds am Ende der Periode", "cash_at_end"),
        ):
            add_total(sheet, label, formula)

        opening_cy = address(balance, "Stand 1. Januar", "cy")
        prior_retained = address(balance, "Gewinnrücklagen und Bilanzgewinn", "py")
        source_totals.append({
            "sheet": balance,
            "cell": opening_cy,
            "year": "cy",
            "factor": 1,
            "sum_of": [
                {"sheet": balance, "cell": prior_retained, "factor": 1},
            ],
        })
        check(
            f"T {balance}!{opening_cy} cy",
            number(balance, opening_cy) - number(balance, prior_retained),
            "Opening current-year retained earnings less comparative "
            "balance-sheet retained earnings, in TCHF.",
        )

        for period in periods:
            v = values[period]
            check(
                f"R1 balance {period}",
                v["total_assets"] - v["total_equity_and_liabilities"],
                "Total assets less total equity and liabilities, in TCHF.",
            )
            check(
                f"R2 cash_at_end {period}",
                v["cash_at_end"] - v["cash_and_cash_equivalents"],
                "Closing cash calculated from cash-flow inputs less "
                "balance-sheet cash, in TCHF. Both source amounts are preserved.",
            )

        check(
            "R3 opening cash cy",
            values["cy"]["cash_at_beginning"]
            - values["py"]["cash_and_cash_equivalents"],
            "Current-year opening cash less comparative balance-sheet cash, in TCHF.",
        )
        check(
            "R4 retained earnings cy",
            values["py"]["retained_earnings"] + values["cy"]["profit_for_year"]
            - values["cy"]["dividends"] - values["cy"]["retained_earnings"],
            "Prior retained earnings plus current profit less dividends "
            "and closing retained earnings, in TCHF.",
        )

        opening_py = address(balance, "Stand 1. Januar", "py")
        check(
            "R4 retained earnings py",
            number(balance, opening_py) + values["py"]["profit_for_year"]
            - values["py"]["dividends"] - values["py"]["retained_earnings"],
            "Comparative opening retained earnings plus comparative profit "
            "less dividends and closing retained earnings, in TCHF.",
        )

        ignored = []
        for sheet in sheet_names:
            for row in header_rows[sheet]:
                for period in periods:
                    ignored.append({
                        "sheet": sheet,
                        "cell": f"{columns[sheet][period]}{row}",
                        "reason": "Year-column heading; period metadata, not an amount.",
                    })
        ignored.append({
            "sheet": balance,
            "cell": opening_py,
            "reason": "Opening comparative retained earnings precede both "
                      "reported balance-sheet dates; no template input. "
                      "Used only to validate the comparative equity roll-forward.",
        })

        return {
            "entity_code": "TW-CH01",
            "source_entity_name": names[0],
            "currency": "CHF",
            "unit": 1000,
            "period_end": {
                "cy": cy_date.isoformat(),
                "py": py_date.isoformat(),
            },
            "fields": fields,
            "source_totals": source_totals,
            "ignored": ignored,
            "source_exceptions": exceptions,
            "notes": (
                "The Treasury Centre designation and Zug location identify "
                "Tarnwell Schweiz AG, TW-CH01; amounts remain in TCHF. "
                "Group-loan interest is finance income and long-term group "
                "loans are investments; short-term Steuerrückstellungen are "
                "current income tax payable. Negative expense and dividend "
                "entries are reversed to the required convention, and absent "
                "input lines are zero. Cash-flow and balance-sheet cash are "
                "preserved independently, with reconciliation differences "
                "declared as source exceptions."
            ),
        }
    finally:
        wb.close()
