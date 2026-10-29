import math
import re
from datetime import date, datetime
from numbers import Real

import openpyxl
from openpyxl.utils import get_column_letter


def extract(path) -> dict:
    wb = openpyxl.load_workbook(path, data_only=True)
    try:
        def normalized(value):
            return re.sub(r"[^a-z0-9]", "", str(value).casefold())

        def find_row(sheet, label, start=1, end=None):
            ws = wb[sheet]
            matches = [
                row
                for row in range(start, (end or ws.max_row) + 1)
                if normalized(ws.cell(row, 1).value) == normalized(label)
            ]
            if len(matches) != 1:
                raise ValueError(
                    f"{sheet}: expected one occurrence of {label!r}; "
                    f"found {len(matches)}"
                )
            return matches[0]

        def number(sheet, address):
            value = wb[sheet][address].value
            if (
                isinstance(value, bool)
                or not isinstance(value, Real)
                or not math.isfinite(value)
            ):
                raise ValueError(f"{sheet}!{address} is not a finite number")
            return value

        def reference(sheet, address, factor=1):
            number(sheet, address)
            return {"sheet": sheet, "cell": address, "factor": factor}

        def contribution(ref):
            return ref["factor"] * number(ref["sheet"], ref["cell"])

        required_sheets = (
            "Income Statement",
            "Balance Sheet",
            "Cash Flow",
            "Retained Earnings",
        )
        for sheet in required_sheets:
            if sheet not in wb.sheetnames:
                raise ValueError(f"Missing sheet: {sheet}")

        source_name = wb["Income Statement"]["A1"].value
        if not isinstance(source_name, str):
            raise ValueError("Missing source entity name")
        # Punctuation differs from the register: "Tarnwell, Inc."
        if normalized(source_name) != normalized("Tarnwell Inc."):
            raise ValueError(f"Unexpected entity: {source_name!r}")
        for sheet in required_sheets:
            if normalized(wb[sheet]["A1"].value) != normalized(source_name):
                raise ValueError(f"Entity name differs on {sheet}")

        # The dollar currency is USD because this is the US registered entity.
        for sheet in required_sheets:
            header = " ".join(
                str(wb[sheet].cell(row, 1).value or "")
                for row in range(1, 4)
            )
            if not re.search(r"\$\s+in\s+thousands", header, re.I):
                raise ValueError(f"Unexpected currency/unit heading on {sheet}")

        def parse_date(value):
            if isinstance(value, datetime):
                return value.date()
            if isinstance(value, date):
                return value
            if isinstance(value, str):
                for fmt in ("%b %d, %Y", "%B %d, %Y", "%Y-%m-%d"):
                    try:
                        return datetime.strptime(value.strip(), fmt).date()
                    except ValueError:
                        pass
            return None

        first_bs_amount = find_row("Balance Sheet", "Inventories")
        date_headers = []
        for row in wb["Balance Sheet"].iter_rows(
            min_row=1, max_row=first_bs_amount - 1
        ):
            dated = [
                (cell.column, parse_date(cell.value))
                for cell in row
                if cell.column > 1 and parse_date(cell.value) is not None
            ]
            if len(dated) == 2:
                date_headers.append(dated)
        if len(date_headers) != 1:
            raise ValueError("Expected one balance-sheet header with two dates")
        dated = sorted(date_headers[0], key=lambda item: item[1], reverse=True)
        if dated[0][1] == dated[1][1]:
            raise ValueError("Balance-sheet dates must differ")

        dates = {"cy": dated[0][1], "py": dated[1][1]}
        columns = {
            "Balance Sheet": {"cy": dated[0][0], "py": dated[1][0]}
        }

        first_labels = {
            "Income Statement": "Net sales",
            "Cash Flow": "Net cash provided by operating activities",
            "Retained Earnings": "Retained earnings, beginning of year",
        }
        for sheet, first_label in first_labels.items():
            first = find_row(sheet, first_label)
            columns[sheet] = {}
            for year, period in dates.items():
                matches = [
                    cell
                    for row in wb[sheet].iter_rows(
                        min_row=1, max_row=first - 1
                    )
                    for cell in row
                    if str(cell.value).strip().upper() == f"FY{period.year}"
                ]
                if len(matches) != 1:
                    raise ValueError(
                        f"{sheet}: expected one FY{period.year} header"
                    )
                columns[sheet][year] = matches[0].column

        def cell_for(sheet, label, year, start=1, end=None):
            row = find_row(sheet, label, start, end)
            return f"{get_column_letter(columns[sheet][year])}{row}"

        def refs_for(sheet, labels, year, factor=1, start=1, end=None):
            return [
                reference(
                    sheet, cell_for(sheet, label, year, start, end), factor
                )
                for label in labels
            ]

        keys = """
            revenue cost_of_sales distribution_costs administrative_expenses
            other_operating_income finance_income finance_costs
            income_tax_expense property_plant_equipment right_of_use_assets
            intangible_assets investments deferred_tax_assets inventories
            trade_receivables other_current_assets cash_and_cash_equivalents
            trade_payables other_current_liabilities current_borrowings
            current_tax_payable non_current_borrowings provisions
            deferred_tax_liabilities share_capital retained_earnings
            other_reserves net_cash_operating net_cash_investing
            net_cash_financing fx_effect_on_cash cash_at_beginning dividends
        """.split()
        fields = {
            key: {year: {"value": 0, "cells": []} for year in dates}
            for key in keys
        }

        def map_field(key, sheet, labels, factor=1, start=1, end=None):
            if isinstance(labels, str):
                labels = [labels]
            for year in dates:
                refs = refs_for(sheet, labels, year, factor, start, end)
                fields[key][year] = {
                    "value": sum(contribution(ref) for ref in refs),
                    "cells": refs,
                }

        income_mapping = {
            "revenue": "Net sales",
            "cost_of_sales": "Cost of goods sold",
            "distribution_costs": "Selling expenses",
            "administrative_expenses": "General and administrative expenses",
            "finance_income": "Interest income",
            "finance_costs": "Interest expense",
            "income_tax_expense": "Provision for income taxes",
        }
        for key, label in income_mapping.items():
            map_field(key, "Income Statement", label)
        map_field(
            "other_operating_income",
            "Income Statement",
            "Other operating (income) expense, net",
            factor=-1,
        )

        balance_mapping = {
            "inventories": "Inventories",
            "trade_receivables": "Accounts receivable, net",
            "other_current_assets": "Prepaid expenses and other current assets",
            "cash_and_cash_equivalents": "Cash and cash equivalents",
            "property_plant_equipment": "Property and equipment, net",
            "right_of_use_assets": "Operating lease right-of-use assets",
            "intangible_assets": ["Goodwill", "Intangible assets, net"],
            "trade_payables": "Accounts payable",
            "other_current_liabilities":
                "Accrued expenses and other current liabilities",
            "current_borrowings": [
                "Current portion of long-term debt",
                "Current operating lease liabilities",
            ],
            "current_tax_payable": "Income taxes payable",
            "non_current_borrowings": [
                "Long-term debt, net of current portion",
                "Operating lease liabilities, non-current",
            ],
            "provisions": "Other long-term liabilities",
            "share_capital": ["Common stock", "Additional paid-in capital"],
            "retained_earnings": "Retained earnings",
            "other_reserves": "Accumulated other comprehensive income (loss)",
        }
        for key, labels in balance_mapping.items():
            map_field(key, "Balance Sheet", labels)

        # Identical labels are distinguished by their statement sections.
        total_assets_row = find_row("Balance Sheet", "Total assets")
        current_liabilities_total = find_row(
            "Balance Sheet", "Total current liabilities"
        )
        liabilities_total = find_row("Balance Sheet", "Total liabilities")
        map_field(
            "deferred_tax_assets",
            "Balance Sheet",
            "Deferred income taxes",
            end=total_assets_row - 1,
        )
        map_field(
            "deferred_tax_liabilities",
            "Balance Sheet",
            "Deferred income taxes",
            start=current_liabilities_total + 1,
            end=liabilities_total - 1,
        )

        cash_mapping = {
            "net_cash_operating": "Net cash provided by operating activities",
            "net_cash_investing": "Net cash used in investing activities",
            "net_cash_financing":
                "Net cash provided by (used in) financing activities",
            "fx_effect_on_cash": "Effect of exchange rate changes on cash",
            "cash_at_beginning":
                "Cash and cash equivalents, beginning of year",
        }
        for key, label in cash_mapping.items():
            map_field(key, "Cash Flow", label)
        map_field(
            "dividends", "Retained Earnings", "Dividends declared", factor=-1
        )

        total_formulas = {
            "gross_profit": "revenue - cost_of_sales",
            "operating_profit":
                "gross_profit - distribution_costs - administrative_expenses "
                "+ other_operating_income",
            "profit_before_tax":
                "operating_profit + finance_income - finance_costs",
            "profit_for_year": "profit_before_tax - income_tax_expense",
            "total_non_current_assets":
                "property_plant_equipment + right_of_use_assets "
                "+ intangible_assets + investments + deferred_tax_assets",
            "total_current_assets":
                "inventories + trade_receivables + other_current_assets "
                "+ cash_and_cash_equivalents",
            "total_assets":
                "total_non_current_assets + total_current_assets",
            "total_current_liabilities":
                "trade_payables + other_current_liabilities "
                "+ current_borrowings + current_tax_payable",
            "total_non_current_liabilities":
                "non_current_borrowings + provisions + deferred_tax_liabilities",
            "total_liabilities":
                "total_current_liabilities + total_non_current_liabilities",
            "total_equity": "share_capital + retained_earnings + other_reserves",
            "total_equity_and_liabilities":
                "total_equity + total_liabilities",
            "net_change_in_cash":
                "net_cash_operating + net_cash_investing + net_cash_financing",
            "cash_at_end":
                "cash_at_beginning + net_change_in_cash + fx_effect_on_cash",
        }

        def evaluate(expression, values):
            tokens = expression.split()
            result = values[tokens[0]]
            for index in range(1, len(tokens), 2):
                operator, name = tokens[index:index + 2]
                if operator not in ("+", "-"):
                    raise ValueError(f"Unsupported operator: {operator}")
                result += values[name] * (1 if operator == "+" else -1)
            return result

        values = {}
        for year in dates:
            values[year] = {
                key: fields[key][year]["value"] for key in fields
            }
            for key, formula in total_formulas.items():
                values[year][key] = evaluate(formula, values[year])

        source_totals = []

        def add_total(sheet, label, formula):
            for year in dates:
                source_totals.append({
                    "sheet": sheet,
                    "cell": cell_for(sheet, label, year),
                    "year": year,
                    "factor": 1,
                    "formula": formula,
                })

        for label, formula in {
            "Gross profit": "gross_profit",
            "Total SG&A": "distribution_costs + administrative_expenses",
            "Operating income": "operating_profit",
            "Income before income taxes": "profit_before_tax",
            "Net income": "profit_for_year",
        }.items():
            add_total("Income Statement", label, formula)

        for label, formula in {
            "Total current assets": "total_current_assets",
            "Total assets": "total_assets",
            "Total current liabilities": "total_current_liabilities",
            "Total liabilities": "total_liabilities",
            "Total stockholders' equity": "total_equity",
            "Total liabilities and stockholders' equity":
                "total_equity_and_liabilities",
        }.items():
            add_total("Balance Sheet", label, formula)

        add_total(
            "Cash Flow", "Net increase (decrease) in cash", "net_change_in_cash"
        )
        add_total(
            "Cash Flow", "Cash and cash equivalents, end of year", "cash_at_end"
        )
        add_total("Retained Earnings", "Net income", "profit_for_year")
        add_total(
            "Retained Earnings", "Retained earnings, end of year",
            "retained_earnings",
        )

        # Cross-year opening equity cannot be expressed using same-year keys.
        for year in dates:
            if year == "cy":
                refs = refs_for(
                    "Balance Sheet", ["Retained earnings"], "py"
                )
            else:
                # Opening = closing - income - signed dividend movement.
                refs = (
                    refs_for(
                        "Retained Earnings",
                        ["Retained earnings, end of year"], year,
                    )
                    + refs_for(
                        "Retained Earnings",
                        ["Net income", "Dividends declared"], year, factor=-1,
                    )
                )
            source_totals.append({
                "sheet": "Retained Earnings",
                "cell": cell_for(
                    "Retained Earnings",
                    "Retained earnings, beginning of year",
                    year,
                ),
                "year": year,
                "factor": 1,
                "sum_of": refs,
            })

        ignored = []
        first_income_row = find_row("Income Statement", "Net sales")
        for column, reason in (
            ("C", "Budget amounts; group inputs use actual results."),
            ("D", "Actual-minus-budget monetary variances."),
            ("E", "Actual-minus-budget percentage variances."),
        ):
            ignored.append({
                "sheet": "Income Statement",
                "range": (
                    f"{column}{first_income_row}:"
                    f"{column}{wb['Income Statement'].max_row}"
                ),
                "reason": reason,
            })

        source_exceptions = []

        def check(check_id, difference, explanation):
            if not math.isclose(difference, 0, rel_tol=0, abs_tol=1e-8):
                source_exceptions.append({
                    "check": check_id,
                    "difference": difference,
                    "explanation": explanation,
                })

        for total in source_totals:
            actual = total["factor"] * number(total["sheet"], total["cell"])
            expected = (
                evaluate(total["formula"], values[total["year"]])
                if "formula" in total
                else sum(contribution(ref) for ref in total["sum_of"])
            )
            check(
                f"T {total['sheet']}!{total['cell']} {total['year']}",
                actual - expected,
                "Printed amount minus its mapped expression or component "
                "sum. Source amounts are preserved.",
            )

        for year in dates:
            v = values[year]
            check(
                f"R1 balance {year}",
                v["total_assets"] - v["total_equity_and_liabilities"],
                "Mapped assets minus mapped equity and liabilities.",
            )
            check(
                f"R2 cash_at_end {year}",
                v["cash_at_end"] - v["cash_and_cash_equivalents"],
                "Computed cash-flow closing cash minus balance-sheet cash.",
            )
        check(
            "R3 opening cash cy",
            values["cy"]["cash_at_beginning"]
            - values["py"]["cash_and_cash_equivalents"],
            "Current opening cash minus prior balance-sheet cash.",
        )
        check(
            "R4 retained earnings cy",
            values["py"]["retained_earnings"]
            + values["cy"]["profit_for_year"]
            - values["cy"]["dividends"]
            - values["cy"]["retained_earnings"],
            "Prior retained earnings plus current profit less dividends "
            "and current closing retained earnings.",
        )

        return {
            "entity_code": "TW-US01",
            "source_entity_name": source_name,
            "currency": "USD",
            "unit": 1000,
            "period_end": {
                year: period.isoformat() for year, period in dates.items()
            },
            "fields": fields,
            "source_totals": source_totals,
            "ignored": ignored,
            "source_exceptions": source_exceptions,
            "notes": (
                "Tarnwell, Inc. matches register entity TW-US01 after "
                "punctuation normalization; dollar amounts are USD thousands. "
                "Income Statement actuals use the FY headers, including the "
                "merged current-year header. Other operating expense is "
                "negated, and signed dividend movements become positive "
                "dividends declared. Goodwill and intangibles, debt and lease "
                "maturities, and common stock and additional paid-in capital "
                "are combined. Deferred taxes are classified by their asset "
                "or liability section. Other long-term liabilities are "
                "assigned to provisions, the template's available residual "
                "non-current obligation category; the workbook provides no "
                "further breakdown, so this classification merits review. "
                "Investments are absent and therefore zero. All reconciliation "
                "differences are reported without changing source amounts."
            ),
        }
    finally:
        wb.close()
