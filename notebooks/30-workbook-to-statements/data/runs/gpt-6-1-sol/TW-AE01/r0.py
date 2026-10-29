import math
import re
from datetime import datetime
from numbers import Real

import openpyxl
from openpyxl.utils import get_column_letter


def extract(path) -> dict:
    wb = openpyxl.load_workbook(path, data_only=True)
    try:
        def normalise(value):
            return re.sub(r"\s+", " ", str(value or "").strip()).casefold()

        def find_row(sheet, label):
            matches = [
                row
                for row in range(1, wb[sheet].max_row + 1)
                if normalise(wb[sheet].cell(row, 1).value) == normalise(label)
            ]
            if len(matches) != 1:
                raise ValueError(
                    f"{sheet}: expected exactly one row labelled {label!r}"
                )
            return matches[0]

        def number(sheet, address):
            value = wb[sheet][address].value
            if (
                isinstance(value, bool)
                or not isinstance(value, Real)
                or not math.isfinite(value)
            ):
                raise ValueError(f"{sheet}!{address} is not a finite number")
            return value

        source_name = wb["Cover"]["A1"].value
        if normalise(source_name) != normalise("Tarnwell ME FZE (Jebel Ali)"):
            raise ValueError(f"Unrecognised source entity: {source_name!r}")

        cover_text = "\n".join(
            str(cell.value)
            for row in wb["Cover"].iter_rows()
            for cell in row
            if cell.value is not None
        )
        currency_match = re.search(
            r"Currency:\s*([A-Z]{3})\s*\(full dirhams\)",
            cover_text,
            re.IGNORECASE,
        )
        if not currency_match or currency_match.group(1).upper() != "AED":
            raise ValueError("Expected amounts in AED, full dirhams")
        if "status: final" not in cover_text.casefold():
            raise ValueError("Reporting package is not marked final")

        date_match = re.search(
            r"as at\s+(\d{1,2}\s+[A-Za-z]+\s+\d{4})",
            str(wb["BS_final"]["A2"].value),
            re.IGNORECASE,
        )
        if not date_match:
            raise ValueError("Cannot identify the final balance-sheet date")
        cy_date = datetime.strptime(date_match.group(1), "%d %B %Y").date()

        columns = {}
        for sheet in ("PL", "BS_final", "CF"):
            headers = {}
            for row in wb[sheet].iter_rows():
                for cell in row:
                    match = re.fullmatch(
                        r"FY\s+(\d{4})", str(cell.value or "").strip(),
                        re.IGNORECASE,
                    )
                    if match:
                        year = int(match.group(1))
                        if year in headers:
                            raise ValueError(f"{sheet}: duplicate FY {year} header")
                        headers[year] = cell.column
            if set(headers) != {cy_date.year, cy_date.year - 1}:
                raise ValueError(f"{sheet}: unexpected reporting-year headers")
            columns[sheet] = {
                "cy": headers[cy_date.year],
                "py": headers[cy_date.year - 1],
            }

        def ref(sheet, label, year, factor=1):
            row = find_row(sheet, label)
            col = columns[sheet][year]
            address = f"{get_column_letter(col)}{row}"
            number(sheet, address)
            return {"sheet": sheet, "cell": address, "factor": factor}

        def reference_value(reference):
            return (
                reference["factor"]
                * number(reference["sheet"], reference["cell"])
            )

        mappings = {
            "revenue": ("PL", "Net sales", 1),
            "cost_of_sales": ("PL", "Cost of sales", 1),
            "distribution_costs": ("PL", "Selling expenses", 1),
            "administrative_expenses": ("PL", "Overheads", 1),
            "other_operating_income": (
                "PL", "Other operating income/(expense), net", 1
            ),
            "finance_income": ("PL", "Finance income", 1),
            "finance_costs": ("PL", "Finance costs", 1),
            "income_tax_expense": ("PL", "Taxation", 1),
            "property_plant_equipment": (
                "BS_final", "Property, plant and equipment", 1
            ),
            "right_of_use_assets": (
                "BS_final", "Leased assets (right of use)", 1
            ),
            "intangible_assets": ("BS_final", "Intangible assets", 1),
            "investments": None,
            "deferred_tax_assets": ("BS_final", "Deferred tax asset", 1),
            "inventories": ("BS_final", "Stock", 1),
            "trade_receivables": ("BS_final", "Trade debtors", 1),
            "other_current_assets": (
                "BS_final", "Other receivables and prepayments", 1
            ),
            "cash_and_cash_equivalents": (
                "BS_final", "Cash at bank and in hand", 1
            ),
            "trade_payables": ("BS_final", "Accounts payable", 1),
            "other_current_liabilities": (
                "BS_final", "Other current liabilities", 1
            ),
            "current_borrowings": ("BS_final", "Short-term borrowings", 1),
            "current_tax_payable": ("BS_final", "Corporation tax", 1),
            "non_current_borrowings": (
                "BS_final", "Borrowings (non-current)", 1
            ),
            "provisions": ("BS_final", "Long-term provisions", 1),
            "deferred_tax_liabilities": ("BS_final", "Deferred taxation", 1),
            "share_capital": ("BS_final", "Share capital", 1),
            "retained_earnings": ("BS_final", "Retained earnings", 1),
            "other_reserves": ("BS_final", "Other reserves", 1),
            "net_cash_operating": (
                "CF", "Net cash from operating activities", 1
            ),
            "net_cash_investing": (
                "CF", "Net cash from investing activities", 1
            ),
            "net_cash_financing": (
                "CF", "Net cash used in financing activities", 1
            ),
            "fx_effect_on_cash": (
                "CF", "Foreign exchange differences on cash", 1
            ),
            "cash_at_beginning": ("CF", "Cash at start of year", 1),
            "dividends": ("CF", "Dividends", -1),
        }

        fields = {}
        for key, mapping in mappings.items():
            fields[key] = {}
            for year in ("cy", "py"):
                cells = []
                if mapping is not None:
                    sheet, label, factor = mapping
                    cells.append(ref(sheet, label, year, factor))
                fields[key][year] = {
                    "value": sum(reference_value(cell) for cell in cells),
                    "cells": cells,
                }

        total_formulas = {
            "gross_profit": "revenue - cost_of_sales",
            "operating_profit": (
                "gross_profit - distribution_costs "
                "- administrative_expenses + other_operating_income"
            ),
            "profit_before_tax": (
                "operating_profit + finance_income - finance_costs"
            ),
            "profit_for_year": "profit_before_tax - income_tax_expense",
            "total_non_current_assets": (
                "property_plant_equipment + right_of_use_assets "
                "+ intangible_assets + investments + deferred_tax_assets"
            ),
            "total_current_assets": (
                "inventories + trade_receivables + other_current_assets "
                "+ cash_and_cash_equivalents"
            ),
            "total_assets": "total_non_current_assets + total_current_assets",
            "total_current_liabilities": (
                "trade_payables + other_current_liabilities "
                "+ current_borrowings + current_tax_payable"
            ),
            "total_non_current_liabilities": (
                "non_current_borrowings + provisions + deferred_tax_liabilities"
            ),
            "total_liabilities": (
                "total_current_liabilities + total_non_current_liabilities"
            ),
            "total_equity": "share_capital + retained_earnings + other_reserves",
            "total_equity_and_liabilities": "total_equity + total_liabilities",
            "net_change_in_cash": (
                "net_cash_operating + net_cash_investing + net_cash_financing"
            ),
            "cash_at_end": (
                "cash_at_beginning + net_change_in_cash + fx_effect_on_cash"
            ),
        }

        def evaluate(expression, values):
            result = 0
            sign = 1
            for token in expression.split():
                if token == "+":
                    sign = 1
                elif token == "-":
                    sign = -1
                else:
                    result += sign * values[token]
            return result

        computed = {}
        for year in ("cy", "py"):
            values = {key: fields[key][year]["value"] for key in fields}
            for key, expression in total_formulas.items():
                values[key] = evaluate(expression, values)
            computed[year] = values

        printed_totals = {
            "PL": {
                "Gross profit": "gross_profit",
                "Operating profit": "operating_profit",
                "Profit before tax": "profit_before_tax",
                "Profit for the year": "profit_for_year",
            },
            "BS_final": {
                "Total non-current assets": "total_non_current_assets",
                "Total current assets": "total_current_assets",
                "Total assets": "total_assets",
                "Total shareholders' funds": "total_equity",
                "Total non-current liabilities": "total_non_current_liabilities",
                "Total current liabilities": "total_current_liabilities",
                "Total liabilities": "total_liabilities",
                "Total equity and liabilities": "total_equity_and_liabilities",
            },
            "CF": {
                "Net change in cash and cash equivalents": "net_change_in_cash",
                "Cash at end of year": "cash_at_end",
                "Profit for the year": "profit_for_year",
                "Closing retained earnings": "retained_earnings",
            },
        }
        source_totals = []
        for sheet, labels in printed_totals.items():
            for label, formula in labels.items():
                for year in ("cy", "py"):
                    source_totals.append({
                        **ref(sheet, label, year),
                        "year": year,
                        "formula": formula,
                    })

        # Opening retained earnings has no same-year template key.
        # Check the current opening against the prior balance sheet, and
        # the comparative opening against its printed equity movement.
        source_totals.append({
            **ref("CF", "Retained earnings at start of year", "cy"),
            "year": "cy",
            "sum_of": [ref("BS_final", "Retained earnings", "py")],
        })
        source_totals.append({
            **ref("CF", "Retained earnings at start of year", "py"),
            "year": "py",
            "sum_of": [
                ref("CF", "Closing retained earnings", "py"),
                ref("CF", "Profit for the year", "py", -1),
                ref("CF", "Dividends", "py", -1),
            ],
        })

        exceptions = []

        def check(check_id, actual, expected, explanation):
            difference = actual - expected
            if difference != 0:
                exceptions.append({
                    "check": check_id,
                    "difference": difference,
                    "explanation": explanation,
                })

        for total in source_totals:
            if "formula" in total:
                expected = evaluate(
                    total["formula"], computed[total["year"]]
                )
            else:
                expected = sum(
                    reference_value(cell) for cell in total["sum_of"]
                )
            check(
                f"T {total['sheet']}!{total['cell']} {total['year']}",
                reference_value(total),
                expected,
                "The signed printed amount differs from its stated "
                "reconciliation. Difference is printed amount minus "
                "reconciled amount; source values are preserved.",
            )

        for year in ("cy", "py"):
            values = computed[year]
            check(
                f"R1 balance {year}",
                values["total_assets"],
                values["total_equity_and_liabilities"],
                "Assets minus equity and liabilities is nonzero in the "
                "final balance sheet; source values are preserved.",
            )
            check(
                f"R2 cash_at_end {year}",
                values["cash_at_end"],
                values["cash_and_cash_equivalents"],
                "Cash-flow closing cash minus final balance-sheet cash "
                "is nonzero; source values are preserved.",
            )

        check(
            "R3 opening cash cy",
            computed["cy"]["cash_at_beginning"],
            computed["py"]["cash_and_cash_equivalents"],
            "Current opening cash minus prior balance-sheet cash is "
            "nonzero; source values are preserved.",
        )
        check(
            "R4 retained earnings cy",
            computed["py"]["retained_earnings"]
            + computed["cy"]["profit_for_year"]
            - computed["cy"]["dividends"],
            computed["cy"]["retained_earnings"],
            "Prior retained earnings plus current profit less dividends "
            "minus current retained earnings is nonzero; source values "
            "are preserved.",
        )

        ignored = [{
            "sheet": "BS_draft_v1",
            "range": wb["BS_draft_v1"].calculate_dimension(),
            "reason": (
                "Superseded draft dated 19 January 2026. BS_final is the "
                "balance sheet in the final approved reporting package."
            ),
        }]

        # Refuse to silently omit any additional numeric rows introduced
        # into the selected statement amount columns.
        covered = set()
        first_used = {}
        for years in fields.values():
            for line in years.values():
                for cell in line["cells"]:
                    covered.add((cell["sheet"], cell["cell"]))
                    source_cell = wb[cell["sheet"]][cell["cell"]]
                    key = (cell["sheet"], source_cell.column)
                    first_used[key] = min(
                        first_used.get(key, source_cell.row), source_cell.row
                    )
        for total in source_totals:
            covered.add((total["sheet"], total["cell"]))
            for cell in total.get("sum_of", []):
                covered.add((cell["sheet"], cell["cell"]))

        for (sheet, col), first_row in first_used.items():
            for row in range(first_row, wb[sheet].max_row + 1):
                cell = wb[sheet].cell(row, col)
                if (
                    isinstance(cell.value, Real)
                    and not isinstance(cell.value, bool)
                    and (sheet, cell.coordinate) not in covered
                ):
                    raise ValueError(
                        f"Unclassified numeric cell: {sheet}!{cell.coordinate}"
                    )

        return {
            "entity_code": "TW-AE01",
            "source_entity_name": source_name,
            "currency": currency_match.group(1).upper(),
            "unit": 1,
            "period_end": {
                "cy": cy_date.isoformat(),
                "py": cy_date.replace(year=cy_date.year - 1).isoformat(),
            },
            "fields": fields,
            "source_totals": source_totals,
            "ignored": ignored,
            "source_exceptions": exceptions,
            "notes": (
                "The cover name is the abbreviated Jebel Ali name of "
                "Tarnwell Middle East FZE (TW-AE01), not Tarnwell Trading "
                "L.L.C. The approved BS_final supersedes BS_draft_v1. "
                "Amounts remain in full AED; overheads map to administrative "
                "expenses, and no investments line is reported. Cash flows "
                "retain their source signs, while negative dividends in the "
                "equity movement are reversed to positive dividends declared."
            ),
        }
    finally:
        wb.close()
