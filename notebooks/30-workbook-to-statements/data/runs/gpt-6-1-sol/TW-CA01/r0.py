import datetime as dt
import math
import re

import openpyxl


def extract(path) -> dict:
    wb = openpyxl.load_workbook(path, data_only=True)

    def norm(value):
        return re.sub(r"\s+", " ", str(value or "").strip()).casefold()

    def find_row(sheet, label, column):
        matches = [
            row
            for row in range(1, sheet.max_row + 1)
            if norm(sheet.cell(row, column).value) == norm(label)
        ]
        if len(matches) != 1:
            raise ValueError(
                f"{sheet.title}: expected one row labelled {label!r}, "
                f"found {len(matches)}"
            )
        return matches[0]

    def number(sheet_name, address):
        value = wb[sheet_name][address].value
        if (
            isinstance(value, bool)
            or not isinstance(value, (int, float))
            or (isinstance(value, float) and not math.isfinite(value))
        ):
            raise ValueError(f"Expected a finite number at {sheet_name}!{address}")
        return value

    def reference(sheet_name, label, year, factor=1):
        sheet = wb[sheet_name]
        label_column = 2 if sheet_name == "IS" else 1
        row = find_row(sheet, label, label_column)
        cell = sheet.cell(row, columns[sheet_name][year])
        number(sheet_name, cell.coordinate)
        return {"sheet": sheet_name, "cell": cell.coordinate, "factor": factor}

    def ref_value(ref):
        return ref["factor"] * number(ref["sheet"], ref["cell"])

    required_sheets = ("IS", "BS", "CF")
    if any(name not in wb.sheetnames for name in required_sheets):
        raise ValueError("Expected IS, BS and CF worksheets")

    source_name = wb["IS"]["A1"].value
    accepted_names = {"tarnwell canada limited", "tarnwell canada ltd."}
    for name in required_sheets:
        if norm(wb[name]["A1"].value) not in accepted_names:
            raise ValueError(f"Unexpected entity name on {name}")

    # This adapter recognises the workbook's full-dollar presentation.
    if "in canadian dollars" not in norm(wb["IS"]["A2"].value):
        raise ValueError("Expected income statement amounts in Canadian dollars")
    for name in ("BS", "CF"):
        if "(cad)" not in norm(wb[name]["A2"].value):
            raise ValueError(f"Expected CAD currency designation on {name}")

    def parse_date(value):
        if isinstance(value, dt.datetime):
            return value.date()
        if isinstance(value, dt.date):
            return value
        if isinstance(value, str):
            for fmt in ("%b %d, %Y", "%B %d, %Y", "%Y-%m-%d"):
                try:
                    return dt.datetime.strptime(value.strip(), fmt).date()
                except ValueError:
                    pass
        return None

    # Discover the balance-sheet date headers and order them chronologically.
    date_header_rows = []
    for row in wb["BS"].iter_rows():
        dates = [
            (cell.column, parse_date(cell.value))
            for cell in row
            if parse_date(cell.value) is not None
        ]
        if len(dates) == 2:
            date_header_rows.append(dates)
    if len(date_header_rows) != 1:
        raise ValueError("Could not uniquely identify balance-sheet date headers")

    dated_columns = sorted(
        date_header_rows[0], key=lambda item: item[1], reverse=True
    )
    if dated_columns[0][1] == dated_columns[1][1]:
        raise ValueError("Current and prior balance-sheet dates must differ")

    dates = {
        year: item[1]
        for year, item in zip(("cy", "py"), dated_columns)
    }
    columns = {
        "BS": {
            year: item[0]
            for year, item in zip(("cy", "py"), dated_columns)
        }
    }

    def merged_header_value(sheet, row, column):
        for merged in sheet.merged_cells.ranges:
            if (
                merged.min_row <= row <= merged.max_row
                and merged.min_col <= column <= merged.max_col
            ):
                return sheet.cell(merged.min_row, merged.min_col).value
        return sheet.cell(row, column).value

    # Select annual Total columns, excluding the quarterly detail.
    annual_header_rows = []
    for row in wb["IS"].iter_rows():
        total_cells = [cell for cell in row if norm(cell.value) == "total"]
        if len(total_cells) == 2:
            annual_header_rows.append(total_cells)
    if len(annual_header_rows) != 1:
        raise ValueError("Could not uniquely identify annual IS Total columns")

    columns["IS"] = {}
    for cell in annual_header_rows[0]:
        fiscal_years = set()
        for row in range(1, cell.row):
            header = merged_header_value(wb["IS"], row, cell.column)
            match = re.fullmatch(r"fiscal\s+(\d{4})", norm(header))
            if match:
                fiscal_years.add(int(match.group(1)))
        if len(fiscal_years) != 1:
            raise ValueError("Annual IS column has an ambiguous fiscal year")
        fiscal_year = fiscal_years.pop()
        matching_years = [
            year for year in ("cy", "py")
            if dates[year].year == fiscal_year
        ]
        if len(matching_years) != 1:
            raise ValueError("IS fiscal year does not match BS dates")
        year = matching_years[0]
        if year in columns["IS"]:
            raise ValueError("Duplicate annual IS fiscal year")
        columns["IS"][year] = cell.column
    if set(columns["IS"]) != {"cy", "py"}:
        raise ValueError("Missing annual IS year")

    columns["CF"] = {}
    for year in ("cy", "py"):
        header = f"F{dates[year].year % 100:02d}"
        matches = [
            cell
            for row in wb["CF"].iter_rows()
            for cell in row
            if norm(cell.value) == norm(header)
        ]
        if len(matches) != 1:
            raise ValueError(f"Could not uniquely identify CF header {header}")
        columns["CF"][year] = matches[0].column

    keys = """
        revenue cost_of_sales distribution_costs administrative_expenses
        other_operating_income finance_income finance_costs income_tax_expense
        property_plant_equipment right_of_use_assets intangible_assets investments
        deferred_tax_assets inventories trade_receivables other_current_assets
        cash_and_cash_equivalents trade_payables other_current_liabilities
        current_borrowings current_tax_payable non_current_borrowings provisions
        deferred_tax_liabilities share_capital retained_earnings other_reserves
        net_cash_operating net_cash_investing net_cash_financing fx_effect_on_cash
        cash_at_beginning dividends
    """.split()
    fields = {
        key: {
            year: {"value": 0, "cells": []}
            for year in ("cy", "py")
        }
        for key in keys
    }

    mappings = {
        "IS": {
            "revenue": ["Net sales"],
            "cost_of_sales": ["Cost of goods sold"],
            "distribution_costs": ["Selling expenses"],
            "administrative_expenses": [
                "Staff costs (administration)",
                "Rent and building costs",
                "Other overheads",
            ],
            "other_operating_income": ["Other gains and losses (net)"],
            "finance_income": ["Finance income"],
            "finance_costs": [
                "Interest on bank loans",
                "Interest on lease liabilities",
            ],
            "income_tax_expense": ["Income tax expense"],
        },
        "BS": {
            "property_plant_equipment": ["Fixed assets (net book value)"],
            "right_of_use_assets": ["Leased assets (right of use)"],
            "intangible_assets": ["Intangible assets"],
            "deferred_tax_assets": ["Deferred tax asset"],
            "trade_receivables": ["Accounts receivable"],
            "other_current_assets": ["Other receivables and prepayments"],
            "cash_and_cash_equivalents": ["Cash at bank and in hand"],
            "share_capital": ["Share capital"],
            "retained_earnings": ["Retained earnings"],
            "other_reserves": ["Translation reserve"],
            "non_current_borrowings": ["Long-term borrowings"],
            "provisions": ["Long-term provisions"],
            "deferred_tax_liabilities": ["Deferred taxation"],
            "trade_payables": ["Accounts payable"],
            "other_current_liabilities": ["Accruals and other payables"],
            "current_borrowings": ["Borrowings (current)"],
            "current_tax_payable": ["Corporation tax"],
        },
        "CF": {
            "net_cash_operating": ["Net cash from operating activities"],
            "net_cash_investing": ["Net cash from investing activities"],
            "net_cash_financing": ["Net cash from financing activities"],
            "fx_effect_on_cash": ["Foreign exchange differences on cash"],
            "cash_at_beginning": ["Opening cash and cash equivalents"],
            "dividends": ["Dividends paid to shareholders"],
        },
    }

    for sheet_name, mapping in mappings.items():
        for key, labels in mapping.items():
            factor = -1 if key == "dividends" else 1
            for year in ("cy", "py"):
                refs = [
                    reference(sheet_name, label, year, factor)
                    for label in labels
                ]
                fields[key][year] = {
                    "value": sum(ref_value(ref) for ref in refs),
                    "cells": refs,
                }

    total_formulas = {
        "gross_profit": "revenue - cost_of_sales",
        "operating_profit": (
            "gross_profit - distribution_costs - administrative_expenses"
            " + other_operating_income"
        ),
        "profit_before_tax": "operating_profit + finance_income - finance_costs",
        "profit_for_year": "profit_before_tax - income_tax_expense",
        "total_non_current_assets": (
            "property_plant_equipment + right_of_use_assets + intangible_assets"
            " + investments + deferred_tax_assets"
        ),
        "total_current_assets": (
            "inventories + trade_receivables + other_current_assets"
            " + cash_and_cash_equivalents"
        ),
        "total_assets": "total_non_current_assets + total_current_assets",
        "total_current_liabilities": (
            "trade_payables + other_current_liabilities"
            " + current_borrowings + current_tax_payable"
        ),
        "total_non_current_liabilities": (
            "non_current_borrowings + provisions + deferred_tax_liabilities"
        ),
        "total_liabilities": (
            "total_current_liabilities + total_non_current_liabilities"
        ),
        "total_equity": "share_capital + retained_earnings + other_reserves",
        "total_equity_and_liabilities": "total_equity + total_liabilities",
        "net_change_in_cash": (
            "net_cash_operating + net_cash_investing + net_cash_financing"
        ),
        "cash_at_end": (
            "cash_at_beginning + net_change_in_cash + fx_effect_on_cash"
        ),
    }

    def evaluate(expression, values):
        tokens = expression.split()
        if not tokens or len(tokens) % 2 != 1:
            raise ValueError(f"Invalid reconciliation expression: {expression}")
        result = values[tokens[0]]
        for index in range(1, len(tokens), 2):
            operator, key = tokens[index:index + 2]
            if operator == "+":
                result += values[key]
            elif operator == "-":
                result -= values[key]
            else:
                raise ValueError(f"Unsupported operator: {operator}")
        return result

    calculated = {}
    for year in ("cy", "py"):
        values = {key: fields[key][year]["value"] for key in keys}
        for key, expression in total_formulas.items():
            values[key] = evaluate(expression, values)
        calculated[year] = values

    printed_totals = {
        "IS": {
            "Gross profit": "gross_profit",
            "Operating profit": "operating_profit",
            "Profit before tax": "profit_before_tax",
            "Net profit": "profit_for_year",
        },
        "BS": {
            "Total non-current assets": "total_non_current_assets",
            "Total current assets": "total_current_assets",
            "TOTAL ASSETS": "total_assets",
            "Total shareholders' funds": "total_equity",
            "Total non-current liabilities": "total_non_current_liabilities",
            "Total current liabilities": "total_current_liabilities",
            "Total liabilities": "total_liabilities",
            "TOTAL EQUITY AND LIABILITIES": "total_equity_and_liabilities",
        },
        "CF": {
            "Net change in cash and cash equivalents": "net_change_in_cash",
            "Cash at end of year": "cash_at_end",
            "Retained earnings at start of year": (
                "retained_earnings - profit_for_year + dividends"
            ),
            "Profit for the year": "profit_for_year",
            "Retained earnings at end of year": "retained_earnings",
        },
    }

    source_totals = []
    exceptions = []

    def record_difference(check, difference, explanation):
        if difference != 0:
            exceptions.append({
                "check": check,
                "difference": difference,
                "explanation": explanation,
            })

    for sheet_name, mapping in printed_totals.items():
        for label, expression in mapping.items():
            for year in ("cy", "py"):
                ref = reference(sheet_name, label, year)
                source_totals.append({
                    **ref,
                    "year": year,
                    "formula": expression,
                })
                record_difference(
                    f"T {sheet_name}!{ref['cell']} {year}",
                    ref_value(ref) - evaluate(expression, calculated[year]),
                    f"Printed {label!r} less the mapped expression "
                    f"({expression}); source amounts are preserved.",
                )

    for year in ("cy", "py"):
        values = calculated[year]
        record_difference(
            f"R1 balance {year}",
            values["total_assets"] - values["total_equity_and_liabilities"],
            "Mapped assets less mapped equity and liabilities; "
            "source amounts are preserved.",
        )
        record_difference(
            f"R2 cash_at_end {year}",
            values["cash_at_end"] - values["cash_and_cash_equivalents"],
            "Calculated cash-flow closing cash less balance-sheet cash; "
            "source amounts are preserved.",
        )

    record_difference(
        "R3 opening cash cy",
        calculated["cy"]["cash_at_beginning"]
        - calculated["py"]["cash_and_cash_equivalents"],
        "Current-year opening cash less prior-year balance-sheet cash; "
        "source amounts are preserved.",
    )
    record_difference(
        "R4 retained earnings cy",
        calculated["cy"]["retained_earnings"]
        - (
            calculated["py"]["retained_earnings"]
            + calculated["cy"]["profit_for_year"]
            - calculated["cy"]["dividends"]
        ),
        "Current retained earnings less prior retained earnings plus "
        "current profit less dividends; source amounts are preserved.",
    )

    ignored = []
    quarter_headers = [
        cell
        for row in wb["IS"].iter_rows()
        for cell in row
        if norm(cell.value) in {"q1", "q2", "q3", "q4"}
    ]
    if sorted(norm(cell.value) for cell in quarter_headers) != [
        "q1", "q2", "q3", "q4"
    ]:
        raise ValueError("Expected exactly four quarterly detail columns")

    for header in quarter_headers:
        letter = header.column_letter
        ignored.append({
            "sheet": "IS",
            "range": f"{letter}{header.row + 1}:{letter}{wb['IS'].max_row}",
            "reason": (
                f"{header.value} quarterly detail; annual Total columns "
                "provide the reporting amounts."
            ),
        })

    # Ensure no numeric amounts in selected reporting columns are overlooked.
    covered = {
        (ref["sheet"], ref["cell"])
        for field in fields.values()
        for line in field.values()
        for ref in line["cells"]
    }
    covered.update((item["sheet"], item["cell"]) for item in source_totals)
    for sheet_name, year_columns in columns.items():
        sheet = wb[sheet_name]
        for column in set(year_columns.values()):
            used_rows = [
                sheet[address].row
                for name, address in covered
                if name == sheet_name and sheet[address].column == column
            ]
            first_row = min(used_rows)
            for row in range(first_row, sheet.max_row + 1):
                cell = sheet.cell(row, column)
                if (
                    isinstance(cell.value, (int, float))
                    and not isinstance(cell.value, bool)
                    and (sheet_name, cell.coordinate) not in covered
                ):
                    raise ValueError(
                        f"Unclassified reporting amount: "
                        f"{sheet_name}!{cell.coordinate}"
                    )

    return {
        "entity_code": "TW-CA01",
        "source_entity_name": source_name,
        "currency": "CAD",
        "unit": 1,
        "period_end": {
            year: dates[year].isoformat()
            for year in ("cy", "py")
        },
        "fields": fields,
        "source_totals": source_totals,
        "ignored": ignored,
        "source_exceptions": exceptions,
        "notes": (
            "The workbook name Tarnwell Canada Limited is matched to "
            "TW-CA01, registered as Tarnwell Canada Ltd. Amounts remain in "
            "Canadian dollars, with fiscal years determined from the "
            "March balance-sheet dates. Annual income-statement totals "
            "are used; quarterly detail is excluded. Administrative "
            "overheads and borrowing interest components are aggregated. "
            "Dividends paid are treated as dividends declared because "
            "they are the dividend movement in the retained-earnings "
            "reconciliation; their negative cash-flow sign is reversed. "
            "Investments and inventories are absent and reported as zero."
        ),
    }
