import re
from datetime import date
from decimal import Decimal
from numbers import Real

import openpyxl


def extract(path) -> dict:
    wb = openpyxl.load_workbook(path, data_only=True)

    keys = """
        revenue cost_of_sales distribution_costs administrative_expenses
        other_operating_income finance_income finance_costs income_tax_expense
        property_plant_equipment right_of_use_assets intangible_assets investments
        deferred_tax_assets inventories trade_receivables other_current_assets
        cash_and_cash_equivalents trade_payables other_current_liabilities
        current_borrowings current_tax_payable non_current_borrowings provisions
        deferred_tax_liabilities share_capital retained_earnings other_reserves
        net_cash_operating net_cash_investing net_cash_financing fx_effect_on_cash
        cash_at_beginning dividends
    """.split()

    total_formulas = {
        "gross_profit": "revenue - cost_of_sales",
        "operating_profit": (
            "gross_profit - distribution_costs - administrative_expenses"
            " + other_operating_income"
        ),
        "profit_before_tax": "operating_profit + finance_income - finance_costs",
        "profit_for_year": "profit_before_tax - income_tax_expense",
        "total_non_current_assets": (
            "property_plant_equipment + right_of_use_assets + intangible_assets"
            " + investments + deferred_tax_assets"
        ),
        "total_current_assets": (
            "inventories + trade_receivables + other_current_assets"
            " + cash_and_cash_equivalents"
        ),
        "total_assets": "total_non_current_assets + total_current_assets",
        "total_current_liabilities": (
            "trade_payables + other_current_liabilities + current_borrowings"
            " + current_tax_payable"
        ),
        "total_non_current_liabilities": (
            "non_current_borrowings + provisions + deferred_tax_liabilities"
        ),
        "total_liabilities": (
            "total_current_liabilities + total_non_current_liabilities"
        ),
        "total_equity": "share_capital + retained_earnings + other_reserves",
        "total_equity_and_liabilities": "total_equity + total_liabilities",
        "net_change_in_cash": (
            "net_cash_operating + net_cash_investing + net_cash_financing"
        ),
        "cash_at_end": (
            "cash_at_beginning + net_change_in_cash + fx_effect_on_cash"
        ),
    }

    def normalize(value):
        return " ".join(str(value).split()).casefold()

    def number(value):
        if isinstance(value, bool) or not isinstance(value, (Real, Decimal)):
            raise ValueError(f"Expected a numeric amount, received {value!r}")
        result = Decimal(str(value))
        if not result.is_finite():
            raise ValueError("Non-finite workbook amount")
        return result

    def json_number(value):
        return int(value) if value == value.to_integral_value() else float(value)

    def run():
        sheets = ("P&L", "BS", "CF", "SOCE")
        if any(sheet not in wb.sheetnames for sheet in sheets):
            raise ValueError("Required statements are missing")

        source_name = wb["P&L"]["A1"].value
        accepted_names = {
            normalize("Tarnwell India Pvt. Ltd."),
            normalize("Tarnwell India Private Limited"),
        }
        if normalize(source_name) not in accepted_names:
            raise ValueError("Workbook does not identify entity TW-IN01")

        for sheet in sheets:
            ws = wb[sheet]
            if normalize(ws["A1"].value) != normalize(source_name):
                raise ValueError(f"Inconsistent entity name on {sheet}")
            unit_labels = [
                normalize(cell.value)
                for row in ws.iter_rows()
                for cell in row
                if isinstance(cell.value, str)
                and "all amounts" in normalize(cell.value)
            ]
            if (
                len(unit_labels) != 1
                or "₹" not in unit_labels[0]
                or "lakhs" not in unit_labels[0]
            ):
                raise ValueError(f"Expected amounts in INR lakhs on {sheet}")

        # One lakh INR = 100 thousand INR.
        scale = 100
        months = {
            name: index
            for index, name in enumerate(
                (
                    "january", "february", "march", "april", "may", "june",
                    "july", "august", "september", "october", "november",
                    "december",
                ),
                1,
            )
        }

        def header_date(value):
            if not isinstance(value, str):
                return None
            text = normalize(value)
            if not text.startswith(("year ended ", "as at ")):
                return None
            match = re.search(r"(\d{1,2})\s+([a-z]+)\s+(\d{4})", text)
            if not match or match[2] not in months:
                return None
            return date(int(match[3]), months[match[2]], int(match[1]))

        columns = {}
        period_end = None
        for sheet in sheets:
            headers = [
                (cell.column, header_date(cell.value))
                for row in wb[sheet].iter_rows()
                for cell in row
                if header_date(cell.value) is not None
            ]
            if len(headers) != 2 or headers[0][0] == headers[1][0]:
                raise ValueError(f"Expected two year columns on {sheet}")
            headers.sort(key=lambda item: item[1], reverse=True)
            dates = dict(zip(("cy", "py"), (d.isoformat() for _, d in headers)))
            if period_end is None:
                period_end = dates
            elif dates != period_end:
                raise ValueError(f"Inconsistent reporting dates on {sheet}")
            columns[sheet] = dict(
                zip(("cy", "py"), (column for column, _ in headers))
            )

        def row_for(sheet, label, section=None):
            ws = wb[sheet]
            start, stop = 1, ws.max_row + 1
            if section:
                start = row_for(sheet, section[0]) + 1
                stop = row_for(sheet, section[1])
            matches = [
                row
                for row in range(start, stop)
                if normalize(ws.cell(row, 1).value) == normalize(label)
            ]
            if len(matches) != 1:
                raise ValueError(
                    f"Expected one row for {sheet}: {label!r}; found {matches}"
                )
            return matches[0]

        def reference(sheet, label, year, factor=scale, section=None):
            cell = wb[sheet].cell(
                row_for(sheet, label, section), columns[sheet][year]
            )
            number(cell.value)
            return {"sheet": sheet, "cell": cell.coordinate, "factor": factor}

        def reference_value(ref):
            return (
                number(wb[ref["sheet"]][ref["cell"]].value)
                * number(ref["factor"])
            )

        def references_value(refs):
            return sum((reference_value(ref) for ref in refs), Decimal(0))

        fields = {
            key: {
                year: {"value": 0, "cells": []}
                for year in ("cy", "py")
            }
            for key in keys
        }

        def map_line(key, sheet, labels, factor=scale, section=None):
            if isinstance(labels, str):
                labels = [labels]
            for year in ("cy", "py"):
                refs = [
                    reference(sheet, label, year, factor, section)
                    for label in labels
                ]
                fields[key][year] = {
                    "value": json_number(references_value(refs)),
                    "cells": refs,
                }

        pl_mapping = {
            "revenue": "Revenue from operations",
            "cost_of_sales": "Cost of services",
            "distribution_costs": "Selling and marketing expenses",
            "administrative_expenses": "General and administrative expenses",
            "other_operating_income": "Other income - miscellaneous",
            "finance_income": "Other income - interest on deposits",
            "finance_costs": "Finance costs",
            "income_tax_expense": "Tax expense",
        }
        for key, label in pl_mapping.items():
            map_line(key, "P&L", label)

        bs_mapping = {
            "property_plant_equipment": "Property, plant and equipment",
            "right_of_use_assets": "Right-of-use assets",
            "intangible_assets": "Other intangible assets",
            "deferred_tax_assets": "Deferred tax assets (net)",
            "trade_receivables": "Trade receivables",
            "other_current_assets": [
                "Bank balances other than cash and cash equivalents",
                "Other current assets",
            ],
            "cash_and_cash_equivalents": "Cash and cash equivalents",
            "share_capital": ["Equity share capital", "Securities premium"],
            "retained_earnings": "Retained earnings",
            "provisions": "Provisions",
            "deferred_tax_liabilities": "Deferred tax liabilities (net)",
            "trade_payables": [
                "(a) total outstanding dues of micro enterprises and small enterprises",
                "(b) total outstanding dues of creditors other than micro enterprises and small enterprises",
            ],
            "other_current_liabilities": "Other current liabilities",
            "current_tax_payable": "Current tax liabilities (net)",
        }
        for key, labels in bs_mapping.items():
            map_line(key, "BS", labels)

        map_line(
            "non_current_borrowings", "BS", "Lease liabilities",
            section=("Non-current liabilities", "Total non-current liabilities"),
        )
        map_line(
            "current_borrowings", "BS", "Lease liabilities",
            section=("Current liabilities", "Total current liabilities"),
        )

        cf_mapping = {
            "net_cash_operating":
                "Net cash generated from operating activities (A)",
            "net_cash_investing":
                "Net cash used in investing activities (B)",
            "net_cash_financing":
                "Net cash used in financing activities (C)",
            "fx_effect_on_cash":
                "Effect of exchange differences on translation of foreign currency cash",
            "cash_at_beginning":
                "Cash and cash equivalents at the beginning of the year",
        }
        for key, label in cf_mapping.items():
            map_line(key, "CF", label)
        map_line("dividends", "SOCE", "Final dividend paid", factor=-scale)

        source_totals = []

        def add_total(sheet, label, formula):
            for year in ("cy", "py"):
                ref = reference(sheet, label, year)
                source_totals.append(dict(ref, year=year, formula=formula))

        for label, formula in (
            ("Gross profit", "gross_profit"),
            ("Profit before tax", "profit_before_tax"),
            ("Profit for the year", "profit_for_year"),
        ):
            add_total("P&L", label, formula)

        for label, formula in (
            ("Total non-current assets", "total_non_current_assets"),
            ("Total current assets", "total_current_assets"),
            ("TOTAL ASSETS", "total_assets"),
            ("Total equity", "total_equity"),
            ("Total non-current liabilities", "total_non_current_liabilities"),
            ("Total current liabilities", "total_current_liabilities"),
            ("TOTAL EQUITY AND LIABILITIES", "total_equity_and_liabilities"),
        ):
            add_total("BS", label, formula)

        # Other equity includes securities premium, which the group template
        # combines with issued capital rather than other reserves.
        for year in ("cy", "py"):
            source_totals.append(dict(
                reference("BS", "Total other equity", year),
                year=year,
                sum_of=[
                    reference("BS", "Securities premium", year),
                    reference("BS", "Retained earnings", year),
                ],
            ))

        add_total(
            "CF",
            "Net increase/(decrease) in cash and cash equivalents (A+B+C)",
            "net_change_in_cash",
        )
        add_total(
            "CF", "Cash and cash equivalents at the end of the year",
            "cash_at_end",
        )
        add_total("SOCE", "Profit for the year", "profit_for_year")
        add_total("SOCE", "Balance at the end of the year", "retained_earnings")

        # Cross-year comparisons use explicit cell references.
        source_totals.append(dict(
            reference("SOCE", "Balance at the beginning of the year", "cy"),
            year="cy",
            sum_of=[reference("BS", "Retained earnings", "py")],
        ))
        for year in ("cy", "py"):
            source_totals.append(dict(
                reference("SOCE", "Balance at the end of the year", year),
                year=year,
                sum_of=[
                    reference("SOCE", "Balance at the beginning of the year", year),
                    reference("SOCE", "Profit for the year", year),
                    reference("SOCE", "Final dividend paid", year),
                ],
            ))

        values = {
            year: {
                key: references_value(fields[key][year]["cells"])
                for key in keys
            }
            for year in ("cy", "py")
        }

        def expression_value(expression, year):
            tokens = expression.split()
            if not tokens or len(tokens) % 2 != 1:
                raise ValueError(f"Invalid subtotal expression: {expression}")
            result = values[year][tokens[0]]
            for index in range(1, len(tokens), 2):
                operator, key = tokens[index:index + 2]
                if operator not in ("+", "-"):
                    raise ValueError(f"Unsupported operator: {operator}")
                result += values[year][key] * (1 if operator == "+" else -1)
            return result

        for year in ("cy", "py"):
            for key, expression in total_formulas.items():
                values[year][key] = expression_value(expression, year)

        exceptions = []
        exception_ids = set()

        def check(check_id, difference, explanation):
            # Decimal arithmetic avoids binary floating-point discrepancies.
            if difference != 0:
                if check_id in exception_ids:
                    for entry in exceptions:
                        if entry["check"] == check_id:
                            entry["explanation"] += " " + explanation
                            break
                else:
                    exceptions.append({
                        "check": check_id,
                        "difference": json_number(difference),
                        "explanation": explanation,
                    })
                    exception_ids.add(check_id)

        for total in source_totals:
            lhs = reference_value(total)
            rhs = (
                expression_value(total["formula"], total["year"])
                if "formula" in total
                else references_value(total["sum_of"])
            )
            check(
                f'T {total["sheet"]}!{total["cell"]} {total["year"]}',
                lhs - rhs,
                "Printed subtotal minus "
                + (
                    total["formula"]
                    if "formula" in total
                    else "the sum of the cited component cells"
                )
                + ", in INR thousands; source amounts have been preserved.",
            )

        for year in ("cy", "py"):
            v = values[year]
            check(
                f"R1 balance {year}",
                v["total_assets"] - v["total_equity_and_liabilities"],
                "Assets minus equity and liabilities, in INR thousands.",
            )
            check(
                f"R2 cash_at_end {year}",
                v["cash_at_end"] - v["cash_and_cash_equivalents"],
                "Calculated closing cash minus balance-sheet cash, "
                "in INR thousands.",
            )
        check(
            "R3 opening cash cy",
            values["cy"]["cash_at_beginning"]
            - values["py"]["cash_and_cash_equivalents"],
            "Current opening cash minus prior-year balance-sheet cash, "
            "in INR thousands.",
        )
        check(
            "R4 retained earnings cy",
            values["cy"]["retained_earnings"]
            - (
                values["py"]["retained_earnings"]
                + values["cy"]["profit_for_year"]
                - values["cy"]["dividends"]
            ),
            "Current retained earnings minus prior retained earnings plus "
            "current profit less dividends, in INR thousands.",
        )

        ignored = []
        for sheet in sheets:
            ws = wb[sheet]
            note_headers = [
                cell for row in ws.iter_rows() for cell in row
                if normalize(cell.value) == "note no."
            ]
            for header in note_headers:
                for row in range(header.row + 1, ws.max_row + 1):
                    cell = ws.cell(row, header.column)
                    if (
                        not isinstance(cell.value, bool)
                        and isinstance(cell.value, (Real, Decimal))
                    ):
                        ignored.append({
                            "sheet": sheet,
                            "cell": cell.coordinate,
                            "reason": "Note reference number, not a financial amount.",
                        })

        # Refuse to silently discard new or unmapped numeric amount rows.
        covered = set()
        for key in keys:
            for year in ("cy", "py"):
                for ref in fields[key][year]["cells"]:
                    covered.add((ref["sheet"], ref["cell"]))
        for total in source_totals:
            covered.add((total["sheet"], total["cell"]))
            for ref in total.get("sum_of", []):
                covered.add((ref["sheet"], ref["cell"]))

        for sheet in sheets:
            ws = wb[sheet]
            for column in columns[sheet].values():
                used_rows = [
                    ws[coordinate].row
                    for covered_sheet, coordinate in covered
                    if covered_sheet == sheet
                    and ws[coordinate].column == column
                ]
                if not used_rows:
                    raise ValueError(f"No financial amounts mapped on {sheet}")
                for row in range(min(used_rows), ws.max_row + 1):
                    cell = ws.cell(row, column)
                    if (
                        not isinstance(cell.value, bool)
                        and isinstance(cell.value, (Real, Decimal))
                        and (sheet, cell.coordinate) not in covered
                    ):
                        raise ValueError(
                            f"Unmapped numeric amount at {sheet}!{cell.coordinate}"
                        )

        return {
            "entity_code": "TW-IN01",
            "source_entity_name": source_name,
            "currency": "INR",
            "unit": 1000,
            "period_end": period_end,
            "fields": fields,
            "source_totals": source_totals,
            "ignored": ignored,
            "source_exceptions": exceptions,
            "notes": (
                "The workbook's abbreviated legal name identifies Tarnwell "
                "India Private Limited (TW-IN01). INR lakhs are converted to "
                "INR thousands using factor 100. Securities premium is included "
                "in share capital; non-cash-equivalent bank balances are included "
                "in other current assets, and both supplier categories are "
                "combined in trade payables. Lease liabilities are classified "
                "by their current or non-current section. Cash-flow signs are "
                "preserved, including the positive current-year financing "
                "amount despite its 'used in' label. The SOCE final dividend "
                "paid is used as dividends declared because no separate "
                "declaration amount is supplied; its negative movement is "
                "converted to a positive dividend. Unpresented input lines "
                "are zero without cell citations."
            ),
        }

    try:
        return run()
    finally:
        wb.close()
