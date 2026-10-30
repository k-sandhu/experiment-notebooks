import calendar
import math
import numbers
import re
import unicodedata
from datetime import date

import openpyxl


def extract(path) -> dict:
    wb = openpyxl.load_workbook(path, data_only=True)

    def normalize(value):
        text = unicodedata.normalize("NFKD", str(value or ""))
        text = "".join(c for c in text if not unicodedata.combining(c))
        text = text.casefold().replace("’", "'")
        return re.sub(r"\s+", " ", text).strip()

    def find_row(sheet, label):
        target = normalize(label)
        matches = [
            row
            for row in range(1, sheet.max_row + 1)
            if normalize(sheet.cell(row, 1).value) == target
        ]
        if len(matches) != 1:
            raise ValueError(
                f"{sheet.title}: expected exactly one row labelled {label!r}"
            )
        return matches[0]

    def number(sheet_name, coordinate):
        value = wb[sheet_name][coordinate].value
        if isinstance(value, bool) or not isinstance(value, numbers.Real):
            raise ValueError(
                f"{sheet_name}!{coordinate} does not contain a number"
            )
        if not math.isfinite(value):
            raise ValueError(f"{sheet_name}!{coordinate} is not finite")
        return int(value) if value == int(value) else float(value)

    def year_columns(sheet, net=False):
        candidates = []
        for row in sheet.iter_rows():
            found = {}
            for cell in row:
                match = re.fullmatch(
                    r"(?:exercice\s+)?(\d{4})", normalize(cell.value)
                )
                if match:
                    year = int(match.group(1))
                    if year in found:
                        raise ValueError(f"Duplicate year on {sheet.title}")
                    found[year] = cell
            if len(found) == 2:
                candidates.append(found)

        if len(candidates) != 1:
            raise ValueError(f"Ambiguous year headers on {sheet.title}")

        headers = candidates[0]
        years = sorted(headers, reverse=True)
        columns = {}

        for period, year in zip(("cy", "py"), years):
            header = headers[year]
            column = header.column
            if net:
                left = right = column
                for merged in sheet.merged_cells.ranges:
                    if (
                        merged.min_row <= header.row <= merged.max_row
                        and merged.min_col <= column <= merged.max_col
                    ):
                        left, right = merged.min_col, merged.max_col
                        break

                net_columns = [
                    col
                    for col in range(left, right + 1)
                    if normalize(sheet.cell(header.row + 1, col).value) == "net"
                ]
                if len(net_columns) != 1:
                    raise ValueError(
                        f"Cannot identify the net column on {sheet.title}"
                    )
                column = net_columns[0]
            columns[period] = column

        return dict(zip(("cy", "py"), years)), columns

    cr = wb["CR"]
    actif = wb["Bilan actif"]
    passif = wb["Bilan passif"]
    tft = wb["TFT"]

    years, cr_columns = year_columns(cr)
    actif_years, actif_columns = year_columns(actif, net=True)
    passif_years, passif_columns = year_columns(passif)
    tft_years, tft_columns = year_columns(tft)

    if any(y != years for y in (actif_years, passif_years, tft_years)):
        raise ValueError("Statement year headers disagree")

    def entity_identity(value):
        return re.sub(r"[^a-z0-9]", "", normalize(value))

    expected_identity = entity_identity("Tarnwell France SAS")
    names = [
        cell.value
        for row in cr.iter_rows()
        for cell in row
        if isinstance(cell.value, str)
        and entity_identity(cell.value) == expected_identity
    ]
    if len(names) != 1:
        raise ValueError("Workbook does not identify exactly Tarnwell France SAS")
    source_entity_name = names[0]

    for sheet in (actif, passif):
        identities = [
            cell.value
            for row in sheet.iter_rows()
            for cell in row
            if isinstance(cell.value, str)
            and entity_identity(cell.value) == expected_identity
        ]
        if len(identities) != 1:
            raise ValueError(f"Entity identity is missing on {sheet.title}")

    heading_specs = (
        (cr, "compte de resultat"),
        (actif, "bilan actif"),
        (passif, "bilan passif"),
        (tft, "tableau des flux"),
    )
    headings = []
    for sheet, prefix in heading_specs:
        matches = [
            normalize(cell.value)
            for row in sheet.iter_rows()
            for cell in row
            if isinstance(cell.value, str)
            and normalize(cell.value).startswith(prefix)
        ]
        if len(matches) != 1 or "(en euros)" not in matches[0]:
            raise ValueError(
                f"Expected a statement heading specifying euros on {sheet.title}"
            )
        headings.append(matches[0])

    months = {
        normalize(calendar.month_name[i]): i for i in range(1, 13)
    }
    months.update({
        "janvier": 1, "fevrier": 2, "mars": 3, "avril": 4,
        "mai": 5, "juin": 6, "juillet": 7, "aout": 8,
        "septembre": 9, "octobre": 10, "novembre": 11, "decembre": 12,
    })
    dates = []
    for heading in headings[:3]:
        match = re.search(r"\b(\d{1,2})\s+([a-z]+)\s+(\d{4})\b", heading)
        if not match or match.group(2) not in months:
            raise ValueError("Cannot read the balance-sheet date")
        dates.append(date(
            int(match.group(3)), months[match.group(2)], int(match.group(1))
        ))
    if len(set(dates)) != 1 or dates[0].year != years["cy"]:
        raise ValueError("Statement dates disagree")

    period_end = {
        period: date(year, dates[0].month, dates[0].day).isoformat()
        for period, year in years.items()
    }

    keys = (
        "revenue cost_of_sales distribution_costs administrative_expenses "
        "other_operating_income finance_income finance_costs income_tax_expense "
        "property_plant_equipment right_of_use_assets intangible_assets "
        "investments deferred_tax_assets inventories trade_receivables "
        "other_current_assets cash_and_cash_equivalents trade_payables "
        "other_current_liabilities current_borrowings current_tax_payable "
        "non_current_borrowings provisions deferred_tax_liabilities "
        "share_capital retained_earnings other_reserves net_cash_operating "
        "net_cash_investing net_cash_financing fx_effect_on_cash "
        "cash_at_beginning dividends"
    ).split()
    fields = {
        key: {period: {"value": 0, "cells": []} for period in ("cy", "py")}
        for key in keys
    }
    source_totals = []
    ignored = []

    def reference(sheet, columns, label, period, factor=1):
        row = find_row(sheet, label)
        coordinate = sheet.cell(row, columns[period]).coordinate
        number(sheet.title, coordinate)
        return {"sheet": sheet.title, "cell": coordinate, "factor": factor}

    def reference_value(ref):
        return ref["factor"] * number(ref["sheet"], ref["cell"])

    def add_field(key, sheet, columns, terms):
        for period in ("cy", "py"):
            refs = [
                reference(sheet, columns, label, period, factor)
                for label, factor in terms
            ]
            fields[key][period] = {
                "value": sum(reference_value(ref) for ref in refs),
                "cells": refs,
            }

    def add_total(sheet, columns, label, formula):
        for period in ("cy", "py"):
            ref = reference(sheet, columns, label, period)
            source_totals.append({
                **ref, "year": period, "formula": formula
            })

    cr_mapping = {
        "revenue": [("Chiffre d'affaires net", 1)],
        "cost_of_sales": [("Coût des ventes", 1)],
        "distribution_costs": [("Frais commerciaux", 1)],
        "administrative_expenses": [("Frais administratifs", 1)],
        "other_operating_income": [
            ("Autres produits d'exploitation", 1),
            ("Autres charges d'exploitation", -1),
        ],
        "finance_income": [("Produits financiers", 1)],
        "finance_costs": [("Charges financières", 1)],
        "income_tax_expense": [("Impôts sur les bénéfices", 1)],
    }
    for key, terms in cr_mapping.items():
        add_field(key, cr, cr_columns, terms)

    actif_mapping = {
        "intangible_assets": "Immobilisations incorporelles",
        "property_plant_equipment": "Immobilisations corporelles",
        "right_of_use_assets": "Droits d'utilisation (contrats de location)",
        "deferred_tax_assets": "Impôts différés actifs",
        "inventories": "Stocks de marchandises",
        "trade_receivables": "Créances clients et comptes rattachés",
        "other_current_assets": "Autres créances et charges constatées d'avance",
        "cash_and_cash_equivalents": "Disponibilités",
    }
    for key, label in actif_mapping.items():
        add_field(key, actif, actif_columns, [(label, 1)])

    passif_mapping = {
        "share_capital": [("Capital social", 1)],
        "retained_earnings": [
            ("Report à nouveau", 1),
            ("Résultat de l'exercice", 1),
        ],
        "other_reserves": [("Autres réserves", 1)],
        "provisions": [("Provisions pour risques et charges", 1)],
        "non_current_borrowings": [
            ("Emprunts auprès des établissements de crédit (à plus d'un an)", 1),
            ("Dettes locatives (à plus d'un an)", 1),
        ],
        "current_borrowings": [
            ("Emprunts auprès des établissements de crédit (à moins d'un an)", 1),
            ("Dettes locatives (à moins d'un an)", 1),
        ],
        "trade_payables": [
            ("Dettes fournisseurs et comptes rattachés", 1),
        ],
        "other_current_liabilities": [
            ("Dettes fiscales et sociales (hors IS)", 1),
            ("Autres dettes", 1),
            ("Produits constatés d'avance", 1),
        ],
        "current_tax_payable": [("Impôt sur les sociétés à payer", 1)],
        "deferred_tax_liabilities": [("Impôts différés passifs", 1)],
    }
    for key, terms in passif_mapping.items():
        add_field(key, passif, passif_columns, terms)

    tft_mapping = {
        "net_cash_operating": (
            "Flux net de trésorerie généré par l'activité", 1
        ),
        "net_cash_investing": (
            "Flux net de trésorerie lié aux opérations d'investissement", 1
        ),
        "net_cash_financing": (
            "Flux net de trésorerie lié aux opérations de financement", 1
        ),
        "fx_effect_on_cash": (
            "Incidence des variations de cours des devises", 1
        ),
        "cash_at_beginning": ("Trésorerie d'ouverture", 1),
        "dividends": ("Dividendes versés", -1),
    }
    for key, term in tft_mapping.items():
        add_field(key, tft, tft_columns, [term])

    for label, formula in (
        ("Marge brute", "gross_profit"),
        ("Résultat d'exploitation", "operating_profit"),
        ("Résultat financier", "finance_income - finance_costs"),
        ("Résultat courant avant impôts", "profit_before_tax"),
        ("Résultat net de l'exercice", "profit_for_year"),
    ):
        add_total(cr, cr_columns, label, formula)

    for label, formula in (
        ("Total actif immobilisé", "total_non_current_assets"),
        ("Total actif circulant", "total_current_assets"),
        ("TOTAL ACTIF", "total_assets"),
    ):
        add_total(actif, actif_columns, label, formula)

    for label, formula in (
        ("Résultat de l'exercice", "profit_for_year"),
        ("Total capitaux propres", "total_equity"),
        ("Total dettes", "total_liabilities - provisions"),
        ("TOTAL PASSIF", "total_equity_and_liabilities"),
    ):
        add_total(passif, passif_columns, label, formula)

    add_total(tft, tft_columns, "Variation de trésorerie", "net_change_in_cash")
    add_total(tft, tft_columns, "Trésorerie de clôture", "cash_at_end")

    for period in ("cy", "py"):
        ref = reference(
            tft, tft_columns, "Report à nouveau à la clôture", period
        )
        source_totals.append({
            **ref,
            "year": period,
            "sum_of": [
                reference(tft, tft_columns, label, period)
                for label in (
                    "Report à nouveau à l'ouverture",
                    "Résultat de l'exercice précédent affecté",
                    "Dividendes versés",
                )
            ],
        })

    # Gross cost and accumulated depreciation are alternative presentations
    # of the net asset values already extracted.
    for row in actif.iter_rows():
        for cell in row:
            if normalize(cell.value) in ("brut", "amort. et deprec."):
                for r in range(cell.row + 1, actif.max_row + 1):
                    candidate = actif.cell(r, cell.column)
                    if (
                        isinstance(candidate.value, numbers.Real)
                        and not isinstance(candidate.value, bool)
                    ):
                        ignored.append({
                            "sheet": actif.title,
                            "cell": candidate.coordinate,
                            "reason": (
                                "Gross asset cost; the net carrying amount is used."
                                if normalize(cell.value) == "brut"
                                else "Accumulated depreciation or impairment; "
                                     "already reflected in the net carrying amount."
                            ),
                        })

    total_formulas = {
        "gross_profit": "revenue - cost_of_sales",
        "operating_profit": (
            "gross_profit - distribution_costs - administrative_expenses "
            "+ other_operating_income"
        ),
        "profit_before_tax": "operating_profit + finance_income - finance_costs",
        "profit_for_year": "profit_before_tax - income_tax_expense",
        "total_non_current_assets": (
            "property_plant_equipment + right_of_use_assets + intangible_assets "
            "+ investments + deferred_tax_assets"
        ),
        "total_current_assets": (
            "inventories + trade_receivables + other_current_assets "
            "+ cash_and_cash_equivalents"
        ),
        "total_assets": "total_non_current_assets + total_current_assets",
        "total_current_liabilities": (
            "trade_payables + other_current_liabilities + current_borrowings "
            "+ current_tax_payable"
        ),
        "total_non_current_liabilities": (
            "non_current_borrowings + provisions + deferred_tax_liabilities"
        ),
        "total_liabilities": (
            "total_current_liabilities + total_non_current_liabilities"
        ),
        "total_equity": "share_capital + retained_earnings + other_reserves",
        "total_equity_and_liabilities": "total_equity + total_liabilities",
        "net_change_in_cash": (
            "net_cash_operating + net_cash_investing + net_cash_financing"
        ),
        "cash_at_end": "cash_at_beginning + net_change_in_cash + fx_effect_on_cash",
    }

    def expression_value(expression, values):
        tokens = expression.split()
        if not tokens or len(tokens) % 2 != 1:
            raise ValueError(f"Invalid formula: {expression}")
        result = values[tokens[0]]
        for index in range(1, len(tokens), 2):
            operator, name = tokens[index:index + 2]
            if operator not in ("+", "-"):
                raise ValueError(f"Invalid formula: {expression}")
            result += values[name] if operator == "+" else -values[name]
        return result

    computed = {}
    for period in ("cy", "py"):
        values = {key: fields[key][period]["value"] for key in keys}
        for key, formula in total_formulas.items():
            values[key] = expression_value(formula, values)
        computed[period] = values

    exceptions = []

    def check(check_id, left, right, explanation):
        difference = left - right
        if not math.isclose(difference, 0, rel_tol=0, abs_tol=1e-7):
            exceptions.append({
                "check": check_id,
                "difference": difference,
                "explanation": explanation,
            })

    for total in source_totals:
        expected = (
            expression_value(total["formula"], computed[total["year"]])
            if "formula" in total
            else sum(reference_value(ref) for ref in total["sum_of"])
        )
        check(
            f"T {total['sheet']}!{total['cell']} {total['year']}",
            reference_value(total),
            expected,
            "The printed source subtotal differs from its mapped components; "
            "the source amounts have been preserved.",
        )

    for period in ("cy", "py"):
        values = computed[period]
        check(
            f"R1 balance {period}",
            values["total_assets"],
            values["total_equity_and_liabilities"],
            "Source assets differ from source equity and liabilities; "
            "no balancing adjustment has been made.",
        )
        check(
            f"R2 cash_at_end {period}",
            values["cash_at_end"],
            values["cash_and_cash_equivalents"],
            "The source cash-flow closing balance differs from balance-sheet "
            "cash; both sets of source amounts have been preserved.",
        )

    check(
        "R3 opening cash cy",
        computed["cy"]["cash_at_beginning"],
        computed["py"]["cash_and_cash_equivalents"],
        "Current-year opening cash differs from prior-year balance-sheet cash.",
    )
    check(
        "R4 retained earnings cy",
        computed["py"]["retained_earnings"]
        + computed["cy"]["profit_for_year"]
        - computed["cy"]["dividends"],
        computed["cy"]["retained_earnings"],
        "Prior retained earnings plus current profit less dividends differs "
        "from current retained earnings; no adjustment has been made.",
    )

    return {
        "entity_code": "TW-FR01",
        "source_entity_name": source_entity_name,
        "currency": "EUR",
        "unit": 1,
        "period_end": period_end,
        "fields": fields,
        "source_totals": source_totals,
        "ignored": ignored,
        "source_exceptions": exceptions,
        "notes": (
            "Amounts are in euros, using the net asset columns. Other operating "
            "income is presented net of other operating charges. Retained "
            "earnings includes report à nouveau and the separately presented "
            "annual result; other reserves retain their source signs. "
            "Borrowings combine bank debt and lease liabilities by maturity. "
            "Total dettes excludes the separately presented provisions. "
            "Dividendes versés in the profit-appropriation schedule is treated "
            "as the year's dividend distribution and reversed to a positive "
            "amount; the workbook does not separately identify declared but "
            "unpaid dividends. Investments are absent and reported as zero. "
            "Source inconsistencies, if any, are reported without adjustment."
        ),
    }
