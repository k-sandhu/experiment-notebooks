import datetime as dt
import math
import numbers
import re
import unicodedata

import openpyxl


def extract(path) -> dict:
    wb = openpyxl.load_workbook(path, data_only=True)
    try:
        ws = wb["DFs"]

        def normalize(value):
            text = unicodedata.normalize("NFKD", str(value or ""))
            text = "".join(c for c in text if not unicodedata.combining(c))
            return re.sub(r"[^a-z0-9]+", " ", text.lower()).strip()

        def find(label, column, start=1, end=None):
            matches = [
                row
                for row in range(start, (end or ws.max_row) + 1)
                if normalize(ws.cell(row, column).value) == normalize(label)
            ]
            if len(matches) != 1:
                raise ValueError(
                    f"Expected one {label!r} in column {column}; found {matches}"
                )
            return matches[0]

        def number(address):
            value = ws[address].value
            if isinstance(value, bool) or not isinstance(value, numbers.Real):
                raise ValueError(f"DFs!{address} must contain a numeric amount")
            if not math.isfinite(value):
                raise ValueError(f"DFs!{address} contains a non-finite amount")
            return value

        def reference(row, column, factor=1):
            address = ws.cell(row, column).coordinate
            number(address)
            return {"sheet": ws.title, "cell": address, "factor": factor}

        def read_reference(ref):
            return ref["factor"] * number(ref["cell"])

        def parse_date(value):
            if isinstance(value, dt.datetime):
                return value.date().isoformat()
            if isinstance(value, dt.date):
                return value.isoformat()
            if isinstance(value, str):
                try:
                    return dt.datetime.strptime(
                        value.strip(), "%d/%m/%Y"
                    ).date().isoformat()
                except ValueError:
                    pass
            return None

        title = str(ws["A1"].value or "")
        source_name = title.split(" - ", 1)[0].strip()
        if normalize(source_name) != normalize("Tarnwell do Brasil Ltda."):
            raise ValueError(f"Unexpected subsidiary: {source_name!r}")
        if "milhares de reais" not in normalize(title):
            raise ValueError("Expected amounts expressed in thousands of BRL")

        # Locate the shared year-header row without relying on its row number.
        header_rows = [
            row
            for row in range(1, ws.max_row + 1)
            if parse_date(ws.cell(row, 2).value)
            and parse_date(ws.cell(row, 3).value)
        ]
        if len(header_rows) != 1:
            raise ValueError("Cannot uniquely identify the year headers")
        header_row = header_rows[0]
        period_end = {
            "cy": parse_date(ws.cell(header_row, 2).value),
            "py": parse_date(ws.cell(header_row, 3).value),
        }
        if period_end["cy"] <= period_end["py"]:
            raise ValueError("Current-year date must follow prior-year date")
        for columns in ((6, 7), (10, 11)):
            for year, column in zip(("cy", "py"), columns):
                if parse_date(ws.cell(header_row, column).value) != period_end[year]:
                    raise ValueError("Statement year headers disagree")

        keys = (
            "revenue", "cost_of_sales", "distribution_costs",
            "administrative_expenses", "other_operating_income",
            "finance_income", "finance_costs", "income_tax_expense",
            "property_plant_equipment", "right_of_use_assets",
            "intangible_assets", "investments", "deferred_tax_assets",
            "inventories", "trade_receivables", "other_current_assets",
            "cash_and_cash_equivalents", "trade_payables",
            "other_current_liabilities", "current_borrowings",
            "current_tax_payable", "non_current_borrowings", "provisions",
            "deferred_tax_liabilities", "share_capital", "retained_earnings",
            "other_reserves", "net_cash_operating", "net_cash_investing",
            "net_cash_financing", "fx_effect_on_cash", "cash_at_beginning",
            "dividends",
        )
        fields = {
            key: {
                year: {"value": 0, "cells": []}
                for year in ("cy", "py")
            }
            for key in keys
        }

        def assign(key, label_column, amount_columns, components):
            # Each component is (label, sign, first_row, last_row).
            rows = [
                (find(label, label_column, start, end), factor)
                for label, factor, start, end in components
            ]
            for year, column in zip(("cy", "py"), amount_columns):
                cells = [reference(row, column, factor) for row, factor in rows]
                fields[key][year] = {
                    "value": sum(read_reference(ref) for ref in cells),
                    "cells": cells,
                }

        def single(key, label, label_column, columns, factor=1,
                   start=1, end=None):
            assign(
                key, label_column, columns,
                [(label, factor, start, end)],
            )

        is_mapping = (
            ("revenue", "Receita operacional líquida", 1),
            ("cost_of_sales", "Custo das mercadorias vendidas", -1),
            ("distribution_costs", "Despesas com vendas", -1),
            ("administrative_expenses", "Despesas gerais e administrativas", -1),
            ("other_operating_income",
             "Outras receitas (despesas) operacionais, líquidas", 1),
            ("finance_income", "Receitas financeiras", 1),
            ("finance_costs", "Despesas financeiras", -1),
        )
        for key, label, factor in is_mapping:
            single(key, label, 1, (2, 3), factor)

        assign("income_tax_expense", 1, (2, 3), [
            ("IR e CSLL correntes", -1, 1, None),
            ("IR e CSLL diferidos", -1, 1, None),
        ])

        current_assets = find("Ativo circulante", 5)
        noncurrent_assets = find("Ativo não circulante", 5)
        current_liabilities = find("Passivo circulante", 5)
        noncurrent_liabilities = find("Passivo não circulante", 5)
        equity = find("Patrimônio líquido", 5)

        bs_mapping = (
            ("inventories", "Estoques", current_assets, noncurrent_assets - 1),
            ("trade_receivables", "Contas a receber de clientes",
             current_assets, noncurrent_assets - 1),
            ("cash_and_cash_equivalents", "Caixa e equivalentes de caixa",
             current_assets, noncurrent_assets - 1),
            ("property_plant_equipment", "Imobilizado",
             noncurrent_assets, current_liabilities - 1),
            ("right_of_use_assets", "Direito de uso",
             noncurrent_assets, current_liabilities - 1),
            ("intangible_assets", "Intangível",
             noncurrent_assets, current_liabilities - 1),
            ("deferred_tax_assets", "IR e CS diferidos",
             noncurrent_assets, current_liabilities - 1),
            ("trade_payables", "Fornecedores",
             current_liabilities, noncurrent_liabilities - 1),
            ("other_current_liabilities",
             "Obrigações trabalhistas e outras contas a pagar",
             current_liabilities, noncurrent_liabilities - 1),
            ("current_borrowings", "Empréstimos, financiamentos e arrendamentos",
             current_liabilities, noncurrent_liabilities - 1),
            ("current_tax_payable", "IR e CS a pagar",
             current_liabilities, noncurrent_liabilities - 1),
            ("non_current_borrowings",
             "Empréstimos, financiamentos e arrendamentos",
             noncurrent_liabilities, equity - 1),
            ("provisions", "Provisões", noncurrent_liabilities, equity - 1),
            ("deferred_tax_liabilities", "IR e CS diferidos",
             noncurrent_liabilities, equity - 1),
            ("share_capital", "Capital social", equity, ws.max_row),
            ("retained_earnings", "Reservas de lucros", equity, ws.max_row),
            ("other_reserves", "Ajustes de avaliação patrimonial",
             equity, ws.max_row),
        )
        for key, label, start, end in bs_mapping:
            single(key, label, 5, (6, 7), start=start, end=end)

        assign("other_current_assets", 5, (6, 7), [
            ("Impostos a recuperar", 1, current_assets, noncurrent_assets - 1),
            ("Outros créditos", 1, current_assets, noncurrent_assets - 1),
        ])

        cf_mapping = (
            ("net_cash_operating", "Caixa líquido das atividades operacionais", 1),
            ("net_cash_investing", "Caixa líquido das atividades de investimento", 1),
            ("net_cash_financing", "Caixa líquido das atividades de financiamento", 1),
            ("fx_effect_on_cash", "Efeito da variação cambial sobre o caixa", 1),
            ("cash_at_beginning", "Caixa no início do exercício", 1),
            ("dividends", "Dividendos declarados", -1),
        )
        for key, label, factor in cf_mapping:
            single(key, label, 9, (10, 11), factor)

        total_formulas = {
            "gross_profit": "revenue - cost_of_sales",
            "operating_profit": (
                "gross_profit - distribution_costs - administrative_expenses"
                " + other_operating_income"
            ),
            "profit_before_tax": "operating_profit + finance_income - finance_costs",
            "profit_for_year": "profit_before_tax - income_tax_expense",
            "total_non_current_assets": (
                "property_plant_equipment + right_of_use_assets"
                " + intangible_assets + investments + deferred_tax_assets"
            ),
            "total_current_assets": (
                "inventories + trade_receivables + other_current_assets"
                " + cash_and_cash_equivalents"
            ),
            "total_assets": "total_non_current_assets + total_current_assets",
            "total_current_liabilities": (
                "trade_payables + other_current_liabilities"
                " + current_borrowings + current_tax_payable"
            ),
            "total_non_current_liabilities": (
                "non_current_borrowings + provisions + deferred_tax_liabilities"
            ),
            "total_liabilities": (
                "total_current_liabilities + total_non_current_liabilities"
            ),
            "total_equity": "share_capital + retained_earnings + other_reserves",
            "total_equity_and_liabilities": "total_equity + total_liabilities",
            "net_change_in_cash": (
                "net_cash_operating + net_cash_investing + net_cash_financing"
            ),
            "cash_at_end": "cash_at_beginning + net_change_in_cash + fx_effect_on_cash",
        }

        def evaluate(expression, values):
            tokens = expression.split()
            if not tokens or len(tokens) % 2 != 1:
                raise ValueError(f"Invalid expression: {expression}")
            result = values[tokens[0]]
            for index in range(1, len(tokens), 2):
                operator, name = tokens[index:index + 2]
                if operator not in ("+", "-"):
                    raise ValueError(f"Invalid operator: {operator}")
                result += values[name] * (1 if operator == "+" else -1)
            return result

        calculated = {}
        for year in ("cy", "py"):
            values = {key: fields[key][year]["value"] for key in keys}
            for key, expression in total_formulas.items():
                values[key] = evaluate(expression, values)
            calculated[year] = values

        source_totals = []

        def add_total(label, label_column, columns, formula):
            row = find(label, label_column)
            for year, column in zip(("cy", "py"), columns):
                source_totals.append({
                    **reference(row, column),
                    "year": year,
                    "formula": formula,
                })

        for label, formula in (
            ("Lucro bruto", "gross_profit"),
            ("Resultado operacional", "operating_profit"),
            ("Resultado antes do IR e CSLL", "profit_before_tax"),
            ("Lucro líquido do exercício", "profit_for_year"),
        ):
            add_total(label, 1, (2, 3), formula)

        for label, formula in (
            ("Total do ativo circulante", "total_current_assets"),
            ("Total do ativo não circulante", "total_non_current_assets"),
            ("Total do ativo", "total_assets"),
            ("Total do passivo circulante", "total_current_liabilities"),
            ("Total do passivo não circulante", "total_non_current_liabilities"),
            ("Total liabilities", "total_liabilities"),
            ("Total do patrimônio líquido", "total_equity"),
            ("Total do passivo e patrimônio líquido", "total_equity_and_liabilities"),
        ):
            add_total(label, 5, (6, 7), formula)

        for label, formula in (
            ("Aumento (redução) de caixa", "net_change_in_cash"),
            ("Caixa no fim do exercício", "cash_at_end"),
            ("Lucro líquido do exercício", "profit_for_year"),
            ("Saldo final", "retained_earnings"),
        ):
            add_total(label, 9, (10, 11), formula)

        # Opening retained earnings has no same-year template expression.
        opening_row = find("Saldo inicial", 9)
        closing_row = find("Saldo final", 9)
        profit_row = find("Lucro líquido do exercício", 9)
        dividend_row = find("Dividendos declarados", 9)
        for year, column in zip(("cy", "py"), (10, 11)):
            source_totals.append({
                **reference(opening_row, column),
                "year": year,
                "sum_of": [
                    reference(closing_row, column, 1),
                    reference(profit_row, column, -1),
                    reference(dividend_row, column, -1),
                ],
            })

        source_exceptions = []

        def check(check_id, difference, explanation):
            if not math.isclose(difference, 0, rel_tol=0, abs_tol=1e-7):
                source_exceptions.append({
                    "check": check_id,
                    "difference": difference,
                    "explanation": explanation,
                })

        for total in source_totals:
            expected = (
                evaluate(total["formula"], calculated[total["year"]])
                if "formula" in total
                else sum(read_reference(ref) for ref in total["sum_of"])
            )
            check(
                f"T {total['sheet']}!{total['cell']} {total['year']}",
                read_reference(total) - expected,
                "The signed printed total differs from its component amounts; "
                "all source amounts have been preserved.",
            )

        for year in ("cy", "py"):
            values = calculated[year]
            check(
                f"R1 balance {year}",
                values["total_assets"] - values["total_equity_and_liabilities"],
                "Assets less equity and liabilities; source amounts are unchanged.",
            )
            check(
                f"R2 cash_at_end {year}",
                values["cash_at_end"] - values["cash_and_cash_equivalents"],
                "Calculated closing cash less balance-sheet cash; "
                "source amounts are unchanged.",
            )
        check(
            "R3 opening cash cy",
            calculated["cy"]["cash_at_beginning"]
            - calculated["py"]["cash_and_cash_equivalents"],
            "Current-year opening cash less prior-year balance-sheet cash.",
        )
        check(
            "R4 retained earnings cy",
            calculated["py"]["retained_earnings"]
            + calculated["cy"]["profit_for_year"]
            - calculated["cy"]["dividends"]
            - calculated["cy"]["retained_earnings"],
            "Prior retained earnings plus current profit less dividends "
            "and current retained earnings.",
        )

        # Require explicit accounting for every numeric amount in used columns.
        covered = set()
        for years in fields.values():
            for line in years.values():
                covered.update(ref["cell"] for ref in line["cells"])
        for total in source_totals:
            covered.add(total["cell"])
            covered.update(ref["cell"] for ref in total.get("sum_of", []))

        for column in (2, 3, 6, 7, 10, 11):
            first_rows = [
                ws[address].row
                for address in covered
                if ws[address].column == column
            ]
            if not first_rows:
                continue
            for row in range(min(first_rows), ws.max_row + 1):
                cell = ws.cell(row, column)
                if (
                    isinstance(cell.value, numbers.Real)
                    and not isinstance(cell.value, bool)
                    and cell.coordinate not in covered
                ):
                    raise ValueError(
                        f"Unclassified numeric amount at DFs!{cell.coordinate}"
                    )

        return {
            "entity_code": "TW-BR01",
            "source_entity_name": source_name,
            "currency": "BRL",
            "unit": 1000,
            "period_end": period_end,
            "fields": fields,
            "source_totals": source_totals,
            "ignored": [],
            "source_exceptions": source_exceptions,
            "notes": (
                "Amounts remain in thousands of Brazilian reais. Parenthesized "
                "expense and dividend amounts are reversed using citation factors; "
                "other operating results and cash flows retain their source signs. "
                "Current and deferred IR/CSLL are combined as income tax expense. "
                "Asset-side deferred tax is separated from current assets, and "
                "liability-side deferred tax is separated from provisions. "
                "Reservas de lucros represents retained earnings; Ajustes de "
                "avaliação patrimonial represents other reserves. No investments "
                "line is reported, so investments is zero."
            ),
        }
    finally:
        wb.close()
