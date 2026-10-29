import math
import re
import unicodedata
from datetime import date
from numbers import Real

import openpyxl


def extract(path) -> dict:
    def normalize(value):
        text = unicodedata.normalize("NFKD", str(value or ""))
        text = "".join(c for c in text if not unicodedata.combining(c))
        return " ".join(text.casefold().split())

    wb = openpyxl.load_workbook(path, data_only=True)
    try:
        sheets = {name: wb[name] for name in ("ER", "ESF", "EFE")}
        rows = {}

        for name, ws in sheets.items():
            index = {}
            for row in ws.iter_rows():
                label = normalize(row[0].value)
                if label:
                    index.setdefault(label, []).append(row[0].row)
            rows[name] = index

        def find_row(sheet, label):
            matches = rows[sheet].get(normalize(label), [])
            if len(matches) != 1:
                raise ValueError(
                    f"{sheet}: expected one row labelled {label!r}; "
                    f"found {len(matches)}"
                )
            return matches[0]

        def number(sheet, coordinate):
            value = sheets[sheet][coordinate].value
            if (
                isinstance(value, bool)
                or not isinstance(value, Real)
                or not math.isfinite(value)
            ):
                raise ValueError(f"{sheet}!{coordinate} is not a finite number")
            return int(value) if isinstance(value, int) else float(value)

        # This adapter recognises the workbook's unaccented legal-name variant.
        expected_name = normalize("Tarnwell Mexico SA de CV")
        source_name = None
        for name, ws in sheets.items():
            matches = [
                cell.value
                for row in ws.iter_rows()
                for cell in row
                if isinstance(cell.value, str)
                and normalize(cell.value) == expected_name
            ]
            if len(matches) != 1:
                raise ValueError(f"{name}: cannot uniquely identify TW-MX01")
            if source_name is None:
                source_name = matches[0]

            if not any(
                "cifras en miles de pesos mexicanos" in normalize(cell.value)
                for row in ws.iter_rows()
                for cell in row
                if isinstance(cell.value, str)
            ):
                raise ValueError(f"{name}: expected amounts in MXN thousands")

        # Locate year columns from the primary "Nota" header.
        columns = {}
        header_rows = {}
        workbook_years = None

        for name, ws in sheets.items():
            candidates = []
            for row in ws.iter_rows():
                if normalize(ws.cell(row[0].row, 2).value) != "nota":
                    continue

                year_columns = {}
                for cell in row:
                    text = str(cell.value).strip()
                    if re.fullmatch(r"20\d{2}", text):
                        year_columns[int(text)] = cell.column

                if len(year_columns) == 2:
                    candidates.append((row[0].row, year_columns))

            if not candidates:
                raise ValueError(f"{name}: missing two-year amount header")

            primary_row, primary_columns = candidates[0]
            years = sorted(primary_columns, reverse=True)
            if years[0] != years[1] + 1:
                raise ValueError(f"{name}: expected consecutive reporting years")
            if workbook_years is None:
                workbook_years = years
            elif years != workbook_years:
                raise ValueError("Reporting years differ between statements")

            for _, candidate_columns in candidates:
                if candidate_columns != primary_columns:
                    raise ValueError(f"{name}: inconsistent year columns")

            columns[name] = {
                "cy": primary_columns[years[0]],
                "py": primary_columns[years[1]],
            }
            header_rows[name] = {r for r, _ in candidates}

            titles = [
                normalize(cell.value)
                for row in ws.iter_rows()
                for cell in row
                if isinstance(cell.value, str)
            ]
            date_pattern = (
                rf"31 de diciembre de {years[0]} y {years[1]}\b"
            )
            if not any(re.search(date_pattern, text) for text in titles):
                raise ValueError(f"{name}: unrecognised statement dates")

        mappings = {
            "ER": {
                "revenue": ["Ingresos por ventas netas"],
                "cost_of_sales": ["Costo de ventas"],
                "distribution_costs": ["Gastos de venta"],
                "administrative_expenses": ["Gastos de administración"],
                "other_operating_income": [
                    "Otros ingresos (gastos) de operación, neto"
                ],
                "finance_income": ["Ingresos financieros"],
                "finance_costs": ["Gastos financieros"],
                "income_tax_expense": [
                    "Impuesto sobre la renta causado",
                    "Impuesto sobre la renta diferido",
                ],
            },
            "ESF": {
                "property_plant_equipment": [
                    "Propiedades, planta y equipo, neto"
                ],
                "right_of_use_assets": ["Activos por derecho de uso"],
                "intangible_assets": ["Activos intangibles"],
                "deferred_tax_assets": ["Impuestos diferidos activos"],
                "inventories": ["Inventarios"],
                "trade_receivables": [
                    "Cuentas por cobrar a clientes, neto"
                ],
                "other_current_assets": [
                    "Otras cuentas por cobrar y pagos anticipados"
                ],
                "cash_and_cash_equivalents": [
                    "Efectivo y equivalentes de efectivo"
                ],
                "trade_payables": ["Proveedores"],
                "other_current_liabilities": [
                    "Otras cuentas por pagar y pasivos acumulados",
                    "Participación de los trabajadores en las utilidades "
                    "(PTU) por pagar",
                ],
                "current_borrowings": ["Préstamos bancarios a corto plazo"],
                "current_tax_payable": ["Impuestos a la utilidad por pagar"],
                "non_current_borrowings": [
                    "Préstamos bancarios a largo plazo"
                ],
                "provisions": ["Beneficios a los empleados"],
                "deferred_tax_liabilities": ["Impuestos diferidos pasivos"],
                "share_capital": ["Capital social"],
                "retained_earnings": ["Utilidades retenidas"],
                "other_reserves": ["Otras partidas del capital"],
            },
            "EFE": {
                "net_cash_operating": [
                    "Flujos netos de efectivo de actividades de operación"
                ],
                "net_cash_investing": [
                    "Flujos netos de efectivo de actividades de inversión"
                ],
                "net_cash_financing": [
                    "Flujos netos de efectivo de actividades de financiamiento"
                ],
                "fx_effect_on_cash": [
                    "Efectos por cambios en el valor del efectivo"
                ],
                "cash_at_beginning": ["Efectivo al inicio del periodo"],
                "dividends": ["Dividendos decretados"],
            },
        }

        keys = [
            "revenue", "cost_of_sales", "distribution_costs",
            "administrative_expenses", "other_operating_income",
            "finance_income", "finance_costs", "income_tax_expense",
            "property_plant_equipment", "right_of_use_assets",
            "intangible_assets", "investments", "deferred_tax_assets",
            "inventories", "trade_receivables", "other_current_assets",
            "cash_and_cash_equivalents", "trade_payables",
            "other_current_liabilities", "current_borrowings",
            "current_tax_payable", "non_current_borrowings", "provisions",
            "deferred_tax_liabilities", "share_capital",
            "retained_earnings", "other_reserves", "net_cash_operating",
            "net_cash_investing", "net_cash_financing", "fx_effect_on_cash",
            "cash_at_beginning", "dividends",
        ]
        fields = {
            key: {year: {"value": 0, "cells": []} for year in ("cy", "py")}
            for key in keys
        }
        used = set()
        source_totals = []

        def reference(sheet, label, year, factor=1):
            row = find_row(sheet, label)
            cell = sheets[sheet].cell(row, columns[sheet][year])
            number(sheet, cell.coordinate)
            used.add((sheet, cell.coordinate))
            return {
                "sheet": sheet,
                "cell": cell.coordinate,
                "factor": factor,
            }

        def sum_references(refs):
            return sum(
                ref["factor"] * number(ref["sheet"], ref["cell"])
                for ref in refs
            )

        for sheet, mapping in mappings.items():
            for key, labels in mapping.items():
                factor = -1 if key == "dividends" else 1
                for year in ("cy", "py"):
                    refs = [
                        reference(sheet, label, year, factor)
                        for label in labels
                    ]
                    fields[key][year] = {
                        "value": sum_references(refs),
                        "cells": refs,
                    }

        printed_totals = {
            "ER": {
                "Utilidad bruta": "gross_profit",
                "Utilidad de operación": "operating_profit",
                "Utilidad antes de impuestos a la utilidad":
                    "profit_before_tax",
                "Utilidad neta del ejercicio": "profit_for_year",
            },
            "ESF": {
                "Total activo circulante": "total_current_assets",
                "Total activo no circulante": "total_non_current_assets",
                "Total activo": "total_assets",
                "Total pasivo a corto plazo": "total_current_liabilities",
                "Total pasivo a largo plazo": "total_non_current_liabilities",
                "Total pasivo": "total_liabilities",
                "Total capital contable": "total_equity",
                "Total pasivo y capital contable":
                    "total_equity_and_liabilities",
            },
            "EFE": {
                "Incremento (disminución) neto de efectivo":
                    "net_change_in_cash",
                "Efectivo al final del periodo": "cash_at_end",
                "Utilidad neta del ejercicio": "profit_for_year",
                "Saldo final": "retained_earnings",
            },
        }
        for sheet, mapping in printed_totals.items():
            for label, formula in mapping.items():
                for year in ("cy", "py"):
                    ref = reference(sheet, label, year)
                    source_totals.append({
                        **ref, "year": year, "formula": formula
                    })

        # Check each retained-earnings schedule using its printed cells.
        # The dividend cell is already signed as an equity reduction.
        for year in ("cy", "py"):
            source_totals.append({
                **reference("EFE", "Saldo inicial", year),
                "year": year,
                "sum_of": [
                    reference("EFE", "Saldo final", year),
                    reference("EFE", "Utilidad neta del ejercicio", year, -1),
                    reference("EFE", "Dividendos decretados", year, -1),
                ],
            })

        ignored = []
        for name, ws in sheets.items():
            for row in ws.iter_rows():
                for cell in row:
                    if (
                        not isinstance(cell.value, Real)
                        or isinstance(cell.value, bool)
                        or (name, cell.coordinate) in used
                    ):
                        continue
                    if cell.column == 2:
                        reason = (
                            "Numeric note reference in the Nota column; "
                            "not a financial amount."
                        )
                    elif (
                        cell.row in header_rows[name]
                        and cell.column in columns[name].values()
                    ):
                        reason = (
                            "Reporting-year column heading, including the "
                            "repeated retained-earnings schedule heading."
                        )
                    else:
                        raise ValueError(
                            f"Unclassified numeric cell: "
                            f"{name}!{cell.coordinate}"
                        )
                    ignored.append({
                        "sheet": name,
                        "cell": cell.coordinate,
                        "reason": reason,
                    })

        formulas = {
            "gross_profit": "revenue - cost_of_sales",
            "operating_profit":
                "gross_profit - distribution_costs "
                "- administrative_expenses + other_operating_income",
            "profit_before_tax":
                "operating_profit + finance_income - finance_costs",
            "profit_for_year": "profit_before_tax - income_tax_expense",
            "total_non_current_assets":
                "property_plant_equipment + right_of_use_assets "
                "+ intangible_assets + investments + deferred_tax_assets",
            "total_current_assets":
                "inventories + trade_receivables + other_current_assets "
                "+ cash_and_cash_equivalents",
            "total_assets":
                "total_non_current_assets + total_current_assets",
            "total_current_liabilities":
                "trade_payables + other_current_liabilities "
                "+ current_borrowings + current_tax_payable",
            "total_non_current_liabilities":
                "non_current_borrowings + provisions "
                "+ deferred_tax_liabilities",
            "total_liabilities":
                "total_current_liabilities + total_non_current_liabilities",
            "total_equity":
                "share_capital + retained_earnings + other_reserves",
            "total_equity_and_liabilities":
                "total_equity + total_liabilities",
            "net_change_in_cash":
                "net_cash_operating + net_cash_investing "
                "+ net_cash_financing",
            "cash_at_end":
                "cash_at_beginning + net_change_in_cash + fx_effect_on_cash",
        }

        def evaluate(expression, values):
            tokens = expression.split()
            result = values[tokens[0]]
            for i in range(1, len(tokens), 2):
                if tokens[i] == "+":
                    result += values[tokens[i + 1]]
                elif tokens[i] == "-":
                    result -= values[tokens[i + 1]]
                else:
                    raise ValueError("Unsupported formula operator")
            return result

        calculated = {}
        for year in ("cy", "py"):
            values = {key: fields[key][year]["value"] for key in keys}
            for key, expression in formulas.items():
                values[key] = evaluate(expression, values)
            calculated[year] = values

        exceptions = []

        def check(check_id, actual, expected, explanation):
            difference = actual - expected
            if not math.isclose(difference, 0, rel_tol=0, abs_tol=1e-9):
                exceptions.append({
                    "check": check_id,
                    "difference": difference,
                    "explanation": explanation,
                })

        for total in source_totals:
            actual = total["factor"] * number(
                total["sheet"], total["cell"]
            )
            expected = (
                evaluate(total["formula"], calculated[total["year"]])
                if "formula" in total
                else sum_references(total["sum_of"])
            )
            check(
                f"T {total['sheet']}!{total['cell']} {total['year']}",
                actual,
                expected,
                "The signed printed amount differs from its mapped "
                "components. Difference is printed minus calculated; "
                "all source amounts are preserved.",
            )

        for year in ("cy", "py"):
            values = calculated[year]
            check(
                f"R1 balance {year}",
                values["total_assets"],
                values["total_equity_and_liabilities"],
                "Mapped assets minus mapped equity and liabilities; "
                "source amounts are preserved.",
            )
            check(
                f"R2 cash_at_end {year}",
                values["cash_at_end"],
                values["cash_and_cash_equivalents"],
                "Calculated closing cash minus balance-sheet cash; "
                "source amounts are preserved.",
            )

        check(
            "R3 opening cash cy",
            calculated["cy"]["cash_at_beginning"],
            calculated["py"]["cash_and_cash_equivalents"],
            "Current-year opening cash minus prior-year balance-sheet "
            "cash; source amounts are preserved.",
        )
        check(
            "R4 retained earnings cy",
            calculated["cy"]["retained_earnings"],
            calculated["py"]["retained_earnings"]
            + calculated["cy"]["profit_for_year"]
            - calculated["cy"]["dividends"],
            "Current retained earnings minus prior retained earnings "
            "plus current profit less dividends; source amounts "
            "are preserved.",
        )

        return {
            "entity_code": "TW-MX01",
            "source_entity_name": source_name,
            "currency": "MXN",
            "unit": 1000,
            "period_end": {
                "cy": date(workbook_years[0], 12, 31).isoformat(),
                "py": date(workbook_years[1], 12, 31).isoformat(),
            },
            "fields": fields,
            "source_totals": source_totals,
            "ignored": ignored,
            "source_exceptions": exceptions,
            "notes": (
                "The workbook name Tarnwell Mexico SA de CV identifies "
                "TW-MX01, Tarnwell México, S.A. de C.V. Amounts remain in "
                "MXN thousands. Current and signed deferred income tax "
                "are combined; PTU payable is included in other current "
                "liabilities. Declared dividends are reversed from the "
                "negative equity-movement presentation to positive "
                "dividends. Investments are absent and therefore zero."
            ),
        }
    finally:
        wb.close()
