import math
import re
from datetime import date, datetime

import openpyxl


def extract(path) -> dict:
    wb = openpyxl.load_workbook(path, data_only=True)
    try:
        tb = wb["TB"]
        cf = wb["CF summary"]

        def norm(value):
            return re.sub(r"\s+", " ", str(value or "").strip()).casefold()

        def is_number(value):
            return (
                isinstance(value, (int, float))
                and not isinstance(value, bool)
                and math.isfinite(value)
            )

        def number(sheet, coordinate):
            value = wb[sheet][coordinate].value
            if not is_number(value):
                raise ValueError(f"Expected a numeric amount at {sheet}!{coordinate}")
            return value

        def find_row(ws, column, label):
            matches = [
                row
                for row in range(1, ws.max_row + 1)
                if norm(ws.cell(row, column).value) == norm(label)
            ]
            if len(matches) != 1:
                raise ValueError(
                    f"Expected exactly one {label!r} row in {ws.title}; "
                    f"found {len(matches)}"
                )
            return matches[0]

        def reference(ws, row, column, factor):
            cell = ws.cell(row, column)
            if cell.value is None:
                return []
            number(ws.title, cell.coordinate)
            return [{
                "sheet": ws.title,
                "cell": cell.coordinate,
                "factor": factor,
            }]

        def line(cells):
            return {
                "value": sum(
                    item["factor"] * number(item["sheet"], item["cell"])
                    for item in cells
                ),
                "cells": cells,
            }

        def combine(parts):
            cells = []
            for source_line, multiplier in parts:
                cells.extend({
                    "sheet": item["sheet"],
                    "cell": item["cell"],
                    "factor": multiplier * item["factor"],
                } for item in source_line["cells"])
            return line(cells)

        # Identify the legal entity from the workbook, preserving its spelling.
        title_matches = []
        for row in tb.iter_rows():
            for cell in row:
                if isinstance(cell.value, str):
                    match = re.match(
                        r"^(.*?)\s*-\s*Trial Balance as at\b",
                        cell.value.strip(),
                        flags=re.IGNORECASE,
                    )
                    if match:
                        title_matches.append(match.group(1).strip())
        if len(title_matches) != 1:
            raise ValueError("Cannot uniquely identify the trial-balance title")
        source_name = title_matches[0]

        def legal_name_key(value):
            return re.sub(r"[^a-z0-9]", "", value.casefold())

        if legal_name_key(source_name) != legal_name_key(
            "Tarnwell Asia Pacific Pte. Ltd."
        ):
            raise ValueError(f"Unexpected subsidiary: {source_name!r}")

        currency_matches = []
        for row in tb.iter_rows():
            for cell in row:
                if isinstance(cell.value, str):
                    currency_matches.extend(re.findall(
                        r"\bAmounts in\s+([A-Z]{3})\b", cell.value
                    ))
        if currency_matches != ["SGD"]:
            raise ValueError("Workbook currency must uniquely identify SGD")

        # Locate debit/credit pairs and read their balance-sheet dates.
        header_candidates = []
        for row in tb.iter_rows():
            pairs = []
            for left, right in zip(row, row[1:]):
                if norm(left.value) == "debit" and norm(right.value) == "credit":
                    pairs.append((left.column, right.column))
            if len(pairs) == 2:
                header_candidates.append((row[0].row, pairs))
        if len(header_candidates) != 1:
            raise ValueError("Cannot uniquely locate the debit/credit headers")

        header_row, pairs = header_candidates[0]

        def read_date(value):
            if isinstance(value, datetime):
                return value.date()
            if isinstance(value, date):
                return value
            return datetime.strptime(str(value).strip(), "%d/%m/%Y").date()

        dated_pairs = [
            (read_date(tb.cell(header_row - 1, debit).value), debit, credit)
            for debit, credit in pairs
        ]
        dated_pairs.sort(reverse=True)
        if dated_pairs[0][0] <= dated_pairs[1][0]:
            raise ValueError("Comparative dates must be distinct")

        periods = dict(zip(("cy", "py"), (item[0] for item in dated_pairs)))
        tb_columns = {
            year: (item[1], item[2])
            for year, item in zip(("cy", "py"), dated_pairs)
        }

        keys = (
            "revenue", "cost_of_sales", "distribution_costs",
            "administrative_expenses", "other_operating_income",
            "finance_income", "finance_costs", "income_tax_expense",
            "property_plant_equipment", "right_of_use_assets",
            "intangible_assets", "investments", "deferred_tax_assets",
            "inventories", "trade_receivables", "other_current_assets",
            "cash_and_cash_equivalents", "trade_payables",
            "other_current_liabilities", "current_borrowings",
            "current_tax_payable", "non_current_borrowings", "provisions",
            "deferred_tax_liabilities", "share_capital", "retained_earnings",
            "other_reserves", "net_cash_operating", "net_cash_investing",
            "net_cash_financing", "fx_effect_on_cash", "cash_at_beginning",
            "dividends",
        )
        fields = {
            key: {year: line([]) for year in ("cy", "py")}
            for key in keys
        }

        # A credit-normal balance is credit less debit; all others are
        # debit less credit. Contra accounts therefore retain their signs.
        mapping = {
            "revenue": (True, [
                "Revenue - regional service fees",
                "Revenue - IT recharges to group",
            ]),
            "cost_of_sales": (False, [
                "Cost of services - staff",
                "Cost of services - subcontractors",
            ]),
            "distribution_costs": (False, ["Business development expenses"]),
            "administrative_expenses": (False, [
                "Admin - salaries and CPF",
                "Admin - rent and utilities",
                "Depreciation and amortisation",
                "Admin - professional fees",
            ]),
            "other_operating_income": (
                True, ["Other operating income / (expense)"]
            ),
            "finance_income": (True, ["Interest income"]),
            "finance_costs": (False, [
                "Interest expense - lease liabilities",
                "Interest expense - group loan",
            ]),
            "income_tax_expense": (False, [
                "Income tax expense - current",
                "Deferred tax expense",
            ]),
            "property_plant_equipment": (False, [
                "Office equipment - cost",
                "Office equipment - accumulated depreciation",
            ]),
            "right_of_use_assets": (False, [
                "Right-of-use asset - office lease",
                "Right-of-use asset - accumulated depreciation",
            ]),
            "intangible_assets": (False, ["Software licences (net)"]),
            "deferred_tax_assets": (False, ["Deferred tax asset"]),
            "trade_receivables": (False, [
                "Trade debtors", "Allowance for doubtful debts",
            ]),
            "other_current_assets": (False, [
                "Prepayments", "GST input tax receivable", "Rental deposits paid",
            ]),
            "cash_and_cash_equivalents": (False, [
                "Cash at bank - DBS current account", "Fixed deposit - 1 month",
            ]),
            "trade_payables": (True, ["Trade creditors"]),
            "other_current_liabilities": (True, [
                "Accrued expenses", "GST output tax payable",
                "CPF contributions payable",
            ]),
            "current_borrowings": (True, ["Lease liability - current"]),
            "current_tax_payable": (True, ["Income tax payable"]),
            "non_current_borrowings": (True, [
                "Lease liability - non-current",
                "Loan from Tarnwell Holdings plc",
            ]),
            "provisions": (True, ["Provision for reinstatement cost"]),
            "deferred_tax_liabilities": (True, ["Deferred tax liability"]),
            "share_capital": (True, ["Share capital - ordinary shares"]),
            "other_reserves": (True, ["Foreign currency translation reserve"]),
            "dividends": (False, ["Dividends declared"]),
        }

        consumed_tb_rows = set()

        def tb_line(labels, year, credit_normal):
            debit_col, credit_col = tb_columns[year]
            debit_factor = -1 if credit_normal else 1
            cells = []
            for label in labels:
                row = find_row(tb, 2, label)
                consumed_tb_rows.add(row)
                cells.extend(reference(tb, row, debit_col, debit_factor))
                cells.extend(reference(tb, row, credit_col, -debit_factor))
            return line(cells)

        for key, (credit_normal, labels) in mapping.items():
            for year in ("cy", "py"):
                fields[key][year] = tb_line(labels, year, credit_normal)

        profit_components = (
            ("revenue", 1), ("cost_of_sales", -1),
            ("distribution_costs", -1), ("administrative_expenses", -1),
            ("other_operating_income", 1), ("finance_income", 1),
            ("finance_costs", -1), ("income_tax_expense", -1),
        )
        for year in ("cy", "py"):
            opening_retained = tb_line(
                ["Retained earnings b/f"], year, credit_normal=True
            )
            fields["retained_earnings"][year] = combine(
                [(opening_retained, 1)]
                + [(fields[key][year], sign) for key, sign in profit_components]
                + [(fields["dividends"][year], -1)]
            )

        # Match cash-flow columns to the actual balance-sheet years.
        cf_columns = {}
        for year in ("cy", "py"):
            label = f"FY{periods[year].year}"
            matches = [
                cell.column
                for row in cf.iter_rows()
                for cell in row
                if norm(cell.value) == norm(label)
            ]
            if len(matches) != 1:
                raise ValueError(f"Cannot uniquely locate cash-flow header {label}")
            cf_columns[year] = matches[0]

        cf_mapping = {
            "net_cash_operating": "Net cash from operating activities",
            "net_cash_investing": "Net cash from investing activities",
            "net_cash_financing": "Net cash used in financing activities",
            "fx_effect_on_cash": "Foreign exchange differences on cash",
            "cash_at_beginning": "Cash at start of year",
        }
        for key, label in cf_mapping.items():
            row = find_row(cf, 1, label)
            for year in ("cy", "py"):
                fields[key][year] = line(
                    reference(cf, row, cf_columns[year], 1)
                )

        source_totals = []
        total_row = find_row(tb, 2, "Total")
        amount_columns = {column for pair in pairs for column in pair}

        # Reject additional accounts rather than silently omit their amounts.
        for row in range(header_row + 1, tb.max_row + 1):
            if row == total_row:
                continue
            if any(is_number(tb.cell(row, col).value) for col in amount_columns):
                if row not in consumed_tb_rows:
                    raise ValueError(
                        f"Unmapped trial-balance account at row {row}: "
                        f"{tb.cell(row, 2).value!r}"
                    )

        for year in ("cy", "py"):
            for column in tb_columns[year]:
                components = []
                for row in range(header_row + 1, total_row):
                    if is_number(tb.cell(row, column).value):
                        components.extend(reference(tb, row, column, 1))
                total_cell = tb.cell(total_row, column)
                number(tb.title, total_cell.coordinate)
                source_totals.append({
                    "sheet": tb.title,
                    "cell": total_cell.coordinate,
                    "year": year,
                    "factor": 1,
                    "sum_of": components,
                })

        for label, formula in (
            ("Net increase/(decrease) in cash", "net_change_in_cash"),
            ("Cash at end of year", "cash_at_end"),
        ):
            row = find_row(cf, 1, label)
            for year in ("cy", "py"):
                cell = cf.cell(row, cf_columns[year])
                number(cf.title, cell.coordinate)
                source_totals.append({
                    "sheet": cf.title,
                    "cell": cell.coordinate,
                    "year": year,
                    "factor": 1,
                    "formula": formula,
                })

        formulas = (
            ("gross_profit", "revenue - cost_of_sales"),
            ("operating_profit",
             "gross_profit - distribution_costs - administrative_expenses "
             "+ other_operating_income"),
            ("profit_before_tax",
             "operating_profit + finance_income - finance_costs"),
            ("profit_for_year", "profit_before_tax - income_tax_expense"),
            ("total_non_current_assets",
             "property_plant_equipment + right_of_use_assets + intangible_assets "
             "+ investments + deferred_tax_assets"),
            ("total_current_assets",
             "inventories + trade_receivables + other_current_assets "
             "+ cash_and_cash_equivalents"),
            ("total_assets", "total_non_current_assets + total_current_assets"),
            ("total_current_liabilities",
             "trade_payables + other_current_liabilities + current_borrowings "
             "+ current_tax_payable"),
            ("total_non_current_liabilities",
             "non_current_borrowings + provisions + deferred_tax_liabilities"),
            ("total_liabilities",
             "total_current_liabilities + total_non_current_liabilities"),
            ("total_equity", "share_capital + retained_earnings + other_reserves"),
            ("total_equity_and_liabilities", "total_equity + total_liabilities"),
            ("net_change_in_cash",
             "net_cash_operating + net_cash_investing + net_cash_financing"),
            ("cash_at_end",
             "cash_at_beginning + net_change_in_cash + fx_effect_on_cash"),
        )

        def evaluate(expression, values):
            tokens = expression.split()
            result = values[tokens[0]]
            for index in range(1, len(tokens), 2):
                operator, name = tokens[index:index + 2]
                if operator not in ("+", "-"):
                    raise ValueError("Unsupported formula operator")
                result += values[name] * (1 if operator == "+" else -1)
            return result

        calculated = {}
        for year in ("cy", "py"):
            values = {key: fields[key][year]["value"] for key in keys}
            for name, expression in formulas:
                values[name] = evaluate(expression, values)
            calculated[year] = values

        exceptions = []

        def check(check_id, actual, expected, explanation):
            difference = actual - expected
            if abs(difference) > 1e-7:
                exceptions.append({
                    "check": check_id,
                    "difference": difference,
                    "explanation": explanation,
                })

        for total in source_totals:
            printed = total["factor"] * number(total["sheet"], total["cell"])
            expected = (
                evaluate(total["formula"], calculated[total["year"]])
                if "formula" in total
                else line(total["sum_of"])["value"]
            )
            check(
                f"T {total['sheet']}!{total['cell']} {total['year']}",
                printed,
                expected,
                "The printed total differs from its underlying amounts "
                "(difference = printed total minus calculated total). "
                "Source amounts have been preserved.",
            )

        for year in ("cy", "py"):
            values = calculated[year]
            check(
                f"R1 balance {year}",
                values["total_assets"],
                values["total_equity_and_liabilities"],
                "Assets differ from equity and liabilities "
                "(difference = assets minus equity and liabilities). "
                "Closing retained earnings include unclosed profit and dividends.",
            )
            check(
                f"R2 cash_at_end {year}",
                values["cash_at_end"],
                values["cash_and_cash_equivalents"],
                "Cash-flow closing cash differs from balance-sheet cash "
                "(difference = cash-flow closing cash minus balance-sheet cash). "
                "Source amounts have been preserved.",
            )

        check(
            "R3 opening cash cy",
            calculated["cy"]["cash_at_beginning"],
            calculated["py"]["cash_and_cash_equivalents"],
            "Current opening cash differs from prior closing balance-sheet cash "
            "(difference = opening cash minus prior closing cash).",
        )
        check(
            "R4 retained earnings cy",
            calculated["cy"]["retained_earnings"],
            calculated["py"]["retained_earnings"]
            + calculated["cy"]["profit_for_year"]
            - calculated["cy"]["dividends"],
            "Closing retained earnings differ from prior closing retained "
            "earnings plus profit less dividends (difference = closing balance "
            "minus the roll-forward). The trial-balance opening balance is "
            "preserved.",
        )

        account_rows = sorted(
            row for row in consumed_tb_rows if is_number(tb.cell(row, 1).value)
        )
        ignored = [{
            "sheet": tb.title,
            "cell": tb.cell(row, 1).coordinate,
            "reason": "Account identifier, not a monetary amount.",
        } for row in account_rows]

        return {
            "entity_code": "TW-SG01",
            "source_entity_name": source_name,
            "currency": "SGD",
            "unit": 1,
            "period_end": {
                year: periods[year].isoformat() for year in ("cy", "py")
            },
            "fields": fields,
            "source_totals": source_totals,
            "ignored": ignored,
            "source_exceptions": exceptions,
            "notes": (
                "Amounts remain in whole SGD. The trial balance is before "
                "closing entries, so each year's closing retained earnings "
                "include its opening retained earnings plus profit less "
                "dividends, with all constituent cells cited. The one-month "
                "fixed deposit is a cash equivalent; rental deposits paid are "
                "included in other current assets. Business development is "
                "classified as selling expense, and depreciation and "
                "amortisation as administrative expense. The group loan is "
                "classified as non-current because no current portion is "
                "identified. Debit other operating expense and debit translation "
                "reserves retain negative signs."
            ),
        }
    finally:
        wb.close()
