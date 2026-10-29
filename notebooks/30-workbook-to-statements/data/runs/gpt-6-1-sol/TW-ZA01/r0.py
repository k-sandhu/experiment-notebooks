import calendar
import math
import re
from datetime import date
from numbers import Real

import openpyxl


FIELD_MAP = {
    "revenue": ("SOCI", ("Turnover",)),
    "cost_of_sales": ("SOCI", ("Cost of sales",)),
    "distribution_costs": ("SOCI", ("Selling expenses",)),
    "administrative_expenses": ("SOCI", ("Overheads",)),
    "other_operating_income": ("SOCI", ("Other gains and losses (net)",)),
    "finance_income": ("SOCI", ("Interest income",)),
    "finance_costs": ("SOCI", ("Interest paid on borrowings",)),
    "income_tax_expense": ("SOCI", ("Income tax expense",)),
    "property_plant_equipment": (
        "SOFP", ("Land and buildings", "Plant and machinery")
    ),
    "right_of_use_assets": ("SOFP", ("Right-of-use assets",)),
    "intangible_assets": ("SOFP", ("Intangible assets",)),
    "investments": ("SOFP", ()),
    "deferred_tax_assets": ("SOFP", ("Deferred tax assets",)),
    "inventories": ("SOFP", ("Stock",)),
    "trade_receivables": ("SOFP", ("Trade debtors",)),
    "other_current_assets": ("SOFP", ("Prepayments and other debtors",)),
    "cash_and_cash_equivalents": ("SOFP", ("Cash and cash equivalents",)),
    "trade_payables": ("SOFP", ("Accounts payable",)),
    "other_current_liabilities": ("SOFP", ("Accruals", "VAT payable")),
    "current_borrowings": ("SOFP", ("Short-term borrowings",)),
    "current_tax_payable": ("SOFP", ("Current tax payable",)),
    "non_current_borrowings": ("SOFP", ("Borrowings (non-current)",)),
    "provisions": ("SOFP", ("Provisions",)),
    "deferred_tax_liabilities": ("SOFP", ("Deferred tax liabilities",)),
    "share_capital": ("SOFP", ("Share capital",)),
    "retained_earnings": ("SOFP", ("Accumulated profits",)),
    "other_reserves": ("SOFP", ("Other reserves",)),
    "net_cash_operating": ("SOCF", ("Net cash from operating activities",)),
    "net_cash_investing": ("SOCF", ("Net cash from investing activities",)),
    "net_cash_financing": ("SOCF", ("Net cash used in financing activities",)),
    "fx_effect_on_cash": ("SOCF", ("Effect of exchange rate changes",)),
    "cash_at_beginning": ("SOCF", ("Cash at start of year",)),
    "dividends": ("SOFP", ("Dividends",)),
}

TOTAL_FORMULAS = {
    "gross_profit": "revenue - cost_of_sales",
    "operating_profit": (
        "gross_profit - distribution_costs - administrative_expenses"
        " + other_operating_income"
    ),
    "profit_before_tax": "operating_profit + finance_income - finance_costs",
    "profit_for_year": "profit_before_tax - income_tax_expense",
    "total_non_current_assets": (
        "property_plant_equipment + right_of_use_assets + intangible_assets"
        " + investments + deferred_tax_assets"
    ),
    "total_current_assets": (
        "inventories + trade_receivables + other_current_assets"
        " + cash_and_cash_equivalents"
    ),
    "total_assets": "total_non_current_assets + total_current_assets",
    "total_current_liabilities": (
        "trade_payables + other_current_liabilities + current_borrowings"
        " + current_tax_payable"
    ),
    "total_non_current_liabilities": (
        "non_current_borrowings + provisions + deferred_tax_liabilities"
    ),
    "total_liabilities": (
        "total_current_liabilities + total_non_current_liabilities"
    ),
    "total_equity": "share_capital + retained_earnings + other_reserves",
    "total_equity_and_liabilities": "total_equity + total_liabilities",
    "net_change_in_cash": (
        "net_cash_operating + net_cash_investing + net_cash_financing"
    ),
    "cash_at_end": (
        "cash_at_beginning + net_change_in_cash + fx_effect_on_cash"
    ),
}

PRINTED_TOTALS = {
    "SOCI": {
        "Gross margin": "gross_profit",
        "Profit from operations": "operating_profit",
        "Profit before taxation": "profit_before_tax",
        "Net profit": "profit_for_year",
    },
    "SOFP": {
        "Total non-current assets": "total_non_current_assets",
        "Total current assets": "total_current_assets",
        "Total assets": "total_assets",
        "Total shareholders' funds": "total_equity",
        "Total non-current liabilities": "total_non_current_liabilities",
        "Total current liabilities": "total_current_liabilities",
        "Total liabilities": "total_liabilities",
        "TOTAL EQUITY AND LIABILITIES": "total_equity_and_liabilities",
        "Profit for the year": "profit_for_year",
        "Closing retained earnings": "retained_earnings",
    },
    "SOCF": {
        "Net increase/(decrease) in cash": "net_change_in_cash",
        "Cash at end of year": "cash_at_end",
    },
}


def _normalise(value):
    return re.sub(
        r"\s+", " ",
        str(value).replace("\u2019", "'").replace("\u00a0", " ").strip()
    ).casefold()


def _number(cell):
    value = cell.value
    if (
        isinstance(value, bool)
        or not isinstance(value, Real)
        or not math.isfinite(value)
    ):
        raise ValueError(
            f"Expected a finite number at {cell.parent.title}!{cell.coordinate}"
        )
    return int(value) if isinstance(value, int) else float(value)


def _expression_value(expression, values):
    tokens = expression.split()
    if not tokens or len(tokens) % 2 != 1:
        raise ValueError(f"Invalid expression: {expression}")
    result = values[tokens[0]]
    for index in range(1, len(tokens), 2):
        operator, name = tokens[index:index + 2]
        if operator == "+":
            result += values[name]
        elif operator == "-":
            result -= values[name]
        else:
            raise ValueError(f"Unsupported operator: {operator}")
    return result


def extract(path) -> dict:
    wb = openpyxl.load_workbook(path, data_only=True)
    try:
        sheet_names = ("SOCI", "SOFP", "SOCF")
        for name in sheet_names:
            if name not in wb.sheetnames:
                raise ValueError(f"Required sheet missing: {name}")

        # This adapter recognises the workbook's abbreviated legal-name alias.
        source_entity_name = wb["SOCI"]["A1"].value
        if _normalise(source_entity_name) not in {
            _normalise("Tarnwell SA (Pty) Ltd"),
            _normalise("Tarnwell South Africa (Pty) Ltd"),
        }:
            raise ValueError(f"Unrecognised entity: {source_entity_name!r}")
        for name in sheet_names:
            if _normalise(wb[name]["A1"].value) != _normalise(source_entity_name):
                raise ValueError(f"Conflicting entity name on {name}")

        rows = {}
        for name in sheet_names:
            index = {}
            for row in wb[name].iter_rows(min_col=1, max_col=1):
                cell = row[0]
                if isinstance(cell.value, str) and cell.value.strip():
                    index.setdefault(_normalise(cell.value), []).append(cell.row)
            rows[name] = index

        def row_for(sheet, label):
            matches = rows[sheet].get(_normalise(label), [])
            if len(matches) != 1:
                raise ValueError(
                    f"Expected one row labelled {label!r} on {sheet}; "
                    f"found {len(matches)}"
                )
            return matches[0]

        title_row = row_for(
            "SOFP",
            next(
                (
                    cell.value
                    for row in wb["SOFP"].iter_rows(min_col=1, max_col=1)
                    for cell in row
                    if isinstance(cell.value, str)
                    and "statement of financial position as at"
                    in _normalise(cell.value)
                ),
                "",
            ),
        )
        title = wb["SOFP"].cell(title_row, 1).value
        match = re.search(
            r"\bas at\s+(\d{1,2})\s+([A-Za-z]+)\s+(\d{4})\b",
            title, re.IGNORECASE,
        )
        if not match:
            raise ValueError("Cannot parse the balance-sheet date")
        months = {
            month.casefold(): index
            for index, month in enumerate(calendar.month_name) if month
        }
        day, month_name, year = match.groups()
        month = months.get(month_name.casefold())
        if month is None:
            raise ValueError(f"Unknown month: {month_name}")
        cy_date = date(int(year), month, int(day))
        prior_year = cy_date.year - 1
        # Preserve a month-end reporting convention in leap years.
        is_month_end = (
            cy_date.day == calendar.monthrange(cy_date.year, month)[1]
        )
        py_day = (
            calendar.monthrange(prior_year, month)[1]
            if is_month_end else cy_date.day
        )
        py_date = date(prior_year, month, py_day)

        columns = {}
        header_rows = {}
        for name in sheet_names:
            ws = wb[name]
            candidates = []
            for row in ws.iter_rows():
                cy_cells = [
                    cell for cell in row
                    if str(cell.value).strip() == str(cy_date.year)
                ]
                py_cells = [
                    cell for cell in row
                    if str(cell.value).strip() == str(py_date.year)
                ]
                if len(cy_cells) == len(py_cells) == 1:
                    candidates.append((cy_cells[0], py_cells[0]))
            if len(candidates) != 1:
                raise ValueError(f"Cannot uniquely identify year columns on {name}")
            cy_header, py_header = candidates[0]
            columns[name] = {
                "cy": cy_header.column,
                "py": py_header.column,
            }
            header_rows[name] = cy_header.row
            for column in columns[name].values():
                units = _normalise(ws.cell(cy_header.row + 1, column).value)
                if units != "r'000":
                    raise ValueError(f"Unexpected currency/unit on {name}: {units}")

        def reference(sheet, label, year, factor=1):
            cell = wb[sheet].cell(
                row_for(sheet, label), columns[sheet][year]
            )
            _number(cell)
            return {
                "sheet": sheet,
                "cell": cell.coordinate,
                "factor": factor,
            }

        def reference_value(ref):
            return ref["factor"] * _number(
                wb[ref["sheet"]][ref["cell"]]
            )

        fields = {}
        for key, (sheet, labels) in FIELD_MAP.items():
            fields[key] = {}
            for year in ("cy", "py"):
                refs = [
                    reference(
                        sheet, label, year,
                        factor=-1 if key == "dividends" else 1,
                    )
                    for label in labels
                ]
                fields[key][year] = {
                    "value": sum(reference_value(ref) for ref in refs),
                    "cells": refs,
                }

        source_totals = []
        for sheet, labels in PRINTED_TOTALS.items():
            for label, formula in labels.items():
                for year in ("cy", "py"):
                    ref = reference(sheet, label, year)
                    source_totals.append({
                        **ref, "year": year, "formula": formula
                    })

        # The current-year opening retained earnings should equal the
        # comparative balance-sheet amount.
        opening_ref = reference(
            "SOFP", "Retained earnings at start of year", "cy"
        )
        source_totals.append({
            **opening_ref,
            "year": "cy",
            "sum_of": [reference("SOFP", "Accumulated profits", "py")],
        })

        ignored = []
        historical_opening = reference(
            "SOFP", "Retained earnings at start of year", "py"
        )
        ignored.append({
            "sheet": "SOFP",
            "cell": historical_opening["cell"],
            "reason": (
                "Opening retained earnings for the comparative year precede "
                "both template balance-sheet dates; retained earnings inputs "
                "use the closing balances."
            ),
        })

        # Explicitly identify draft amounts and numeric note references.
        for name in sheet_names:
            ws = wb[name]
            header_row = header_rows[name]
            for cell in ws[header_row]:
                header = _normalise(cell.value)
                if "draft" in header:
                    column_letter = cell.column_letter
                    ignored.append({
                        "sheet": name,
                        "range": (
                            f"{column_letter}{header_row + 1}:"
                            f"{column_letter}{ws.max_row}"
                        ),
                        "reason": (
                            "Superseded draft column; the columns headed "
                            "with the reporting years contain the final amounts."
                        ),
                    })
                elif header == "notes":
                    column_letter = cell.column_letter
                    ignored.append({
                        "sheet": name,
                        "range": (
                            f"{column_letter}{header_row + 1}:"
                            f"{column_letter}{ws.max_row}"
                        ),
                        "reason": "Note-reference numbers, not financial amounts.",
                    })

        values = {}
        for year in ("cy", "py"):
            values[year] = {
                key: periods[year]["value"]
                for key, periods in fields.items()
            }
            for key, formula in TOTAL_FORMULAS.items():
                values[year][key] = _expression_value(
                    formula, values[year]
                )

        source_exceptions = []

        def check(check_id, actual, expected, explanation):
            difference = actual - expected
            if not math.isclose(
                actual, expected, rel_tol=0, abs_tol=1e-8
            ):
                source_exceptions.append({
                    "check": check_id,
                    "difference": difference,
                    "explanation": explanation,
                })

        for total in source_totals:
            actual = reference_value(total)
            if "formula" in total:
                expected = _expression_value(
                    total["formula"], values[total["year"]]
                )
            else:
                expected = sum(
                    reference_value(ref) for ref in total["sum_of"]
                )
            check(
                f"T {total['sheet']}!{total['cell']} {total['year']}",
                actual,
                expected,
                "The printed amount differs from its mapped reconciliation. "
                "Difference is the signed printed amount minus the "
                "recomputed amount; source inputs have been preserved.",
            )

        for year in ("cy", "py"):
            check(
                f"R1 balance {year}",
                values[year]["total_assets"],
                values[year]["total_equity_and_liabilities"],
                "Mapped assets minus mapped equity and liabilities; "
                "source inputs have been preserved.",
            )
            check(
                f"R2 cash_at_end {year}",
                values[year]["cash_at_end"],
                values[year]["cash_and_cash_equivalents"],
                "Calculated closing cash minus balance-sheet cash; "
                "source inputs have been preserved.",
            )

        check(
            "R3 opening cash cy",
            values["cy"]["cash_at_beginning"],
            values["py"]["cash_and_cash_equivalents"],
            "Current opening cash minus prior-year balance-sheet cash; "
            "source inputs have been preserved.",
        )
        check(
            "R4 retained earnings cy",
            values["py"]["retained_earnings"]
            + values["cy"]["profit_for_year"]
            - values["cy"]["dividends"],
            values["cy"]["retained_earnings"],
            "Prior retained earnings plus current profit less dividends "
            "minus current retained earnings; source inputs have been preserved.",
        )

        # Reject unexplained amounts instead of silently ignoring new lines.
        accounted = set()
        for periods in fields.values():
            for line in periods.values():
                for ref in line["cells"]:
                    accounted.add((ref["sheet"], ref["cell"]))
        for total in source_totals:
            accounted.add((total["sheet"], total["cell"]))
            for ref in total.get("sum_of", []):
                accounted.add((ref["sheet"], ref["cell"]))
        accounted.add(("SOFP", historical_opening["cell"]))

        for name in sheet_names:
            ws = wb[name]
            for column in columns[name].values():
                used_rows = [
                    ws[cell].row
                    for sheet, cell in accounted
                    if sheet == name and ws[cell].column == column
                ]
                if not used_rows:
                    raise ValueError(f"No statement amounts found on {name}")
                for row in ws.iter_rows(
                    min_row=min(used_rows),
                    min_col=column,
                    max_col=column,
                ):
                    cell = row[0]
                    if (
                        isinstance(cell.value, Real)
                        and not isinstance(cell.value, bool)
                        and (name, cell.coordinate) not in accounted
                    ):
                        raise ValueError(
                            f"Unmapped amount at {name}!{cell.coordinate}; "
                            "review its classification before extraction."
                        )

        return {
            "entity_code": "TW-ZA01",
            "source_entity_name": source_entity_name,
            "currency": "ZAR",
            "unit": 1000,
            "period_end": {
                "cy": cy_date.isoformat(),
                "py": py_date.isoformat(),
            },
            "fields": fields,
            "source_totals": source_totals,
            "ignored": ignored,
            "source_exceptions": source_exceptions,
            "notes": (
                "The workbook name 'Tarnwell SA (Pty) Ltd' is treated as the "
                "abbreviated name of TW-ZA01, Tarnwell South Africa (Pty) Ltd. "
                "Amounts remain in thousands of South African rand. "
                "Final year columns are used; draft columns are excluded. "
                "Land and buildings and plant and machinery form PPE; "
                "accruals and VAT payable form other current liabilities. "
                "Dividends are sign-reversed from the retained earnings "
                "reconciliation. Investments are absent and reported as zero. "
                "Printed totals are retained for validation only."
            ),
        }
    finally:
        wb.close()
