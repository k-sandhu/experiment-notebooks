import calendar
import math
import re
from datetime import datetime
from numbers import Real

import openpyxl


FIELD_KEYS = (
    "revenue", "cost_of_sales", "distribution_costs",
    "administrative_expenses", "other_operating_income",
    "finance_income", "finance_costs", "income_tax_expense",
    "property_plant_equipment", "right_of_use_assets", "intangible_assets",
    "investments", "deferred_tax_assets", "inventories", "trade_receivables",
    "other_current_assets", "cash_and_cash_equivalents", "trade_payables",
    "other_current_liabilities", "current_borrowings", "current_tax_payable",
    "non_current_borrowings", "provisions", "deferred_tax_liabilities",
    "share_capital", "retained_earnings", "other_reserves",
    "net_cash_operating", "net_cash_investing", "net_cash_financing",
    "fx_effect_on_cash", "cash_at_beginning", "dividends",
)

TOTAL_FORMULAS = {
    "gross_profit": "revenue - cost_of_sales",
    "operating_profit":
        "gross_profit - distribution_costs - administrative_expenses"
        " + other_operating_income",
    "profit_before_tax": "operating_profit + finance_income - finance_costs",
    "profit_for_year": "profit_before_tax - income_tax_expense",
    "total_non_current_assets":
        "property_plant_equipment + right_of_use_assets + intangible_assets"
        " + investments + deferred_tax_assets",
    "total_current_assets":
        "inventories + trade_receivables + other_current_assets"
        " + cash_and_cash_equivalents",
    "total_assets": "total_non_current_assets + total_current_assets",
    "total_current_liabilities":
        "trade_payables + other_current_liabilities + current_borrowings"
        " + current_tax_payable",
    "total_non_current_liabilities":
        "non_current_borrowings + provisions + deferred_tax_liabilities",
    "total_liabilities":
        "total_current_liabilities + total_non_current_liabilities",
    "total_equity": "share_capital + retained_earnings + other_reserves",
    "total_equity_and_liabilities": "total_equity + total_liabilities",
    "net_change_in_cash":
        "net_cash_operating + net_cash_investing + net_cash_financing",
    "cash_at_end": "cash_at_beginning + net_change_in_cash + fx_effect_on_cash",
}


def _normalise(value):
    return " ".join(str(value or "").strip().casefold().split())


def _is_number(value):
    return (
        isinstance(value, Real)
        and not isinstance(value, bool)
        and math.isfinite(value)
    )


def _expression_value(expression, values):
    """Evaluate the adapter's internal expressions without using eval()."""
    tokens = expression.split()
    if not tokens or len(tokens) % 2 != 1:
        raise ValueError("Invalid reconciliation expression")
    result = values[tokens[0]]
    for index in range(1, len(tokens), 2):
        operator, name = tokens[index:index + 2]
        if operator == "+":
            result += values[name]
        elif operator == "-":
            result -= values[name]
        else:
            raise ValueError("Unsupported reconciliation operator")
    return result


def extract(path):
    workbook = openpyxl.load_workbook(path, data_only=True)
    try:
        sheet_names = ("P&L", "BS", "CF", "Reserves")
        for name in sheet_names:
            if name not in workbook.sheetnames:
                raise ValueError(f"Required sheet missing: {name}")

        sheets = {name: workbook[name] for name in sheet_names}
        source_name = sheets["BS"]["B1"].value
        aliases = {
            "tarnwell engineering ltd",
            "tarnwell engineering ltd.",
            "tarnwell engineering limited",
        }
        if _normalise(source_name) not in aliases:
            raise ValueError("Workbook does not identify TW-UK02 unambiguously")
        for sheet in sheets.values():
            if _normalise(sheet["B1"].value) not in aliases:
                raise ValueError(f"Inconsistent entity name on {sheet.title}")

        date_match = re.search(
            r"\b(\d{1,2}\s+[A-Za-z]+\s+\d{4})\b",
            str(sheets["BS"]["B2"].value),
        )
        if not date_match:
            raise ValueError("Cannot identify the balance-sheet date")
        closing_date = datetime.strptime(
            date_match.group(1), "%d %B %Y"
        ).date()

        columns = {}
        header_rows = {}
        for name, sheet in sheets.items():
            headers = []
            for row in sheet.iter_rows():
                if any(_normalise(cell.value) == "note" for cell in row):
                    headers.append(row)
            if len(headers) != 1:
                raise ValueError(f"Cannot identify unique year header on {name}")

            header = headers[0]
            header_rows[name] = header[0].row
            year_columns = {}
            for cell in header:
                text = str(cell.value).strip()
                if re.fullmatch(r"\d{4}", text):
                    year = int(text)
                    if year in year_columns:
                        raise ValueError(f"Duplicate year column on {name}")
                    year_columns[year] = cell.column

            cy_year = closing_date.year
            if set(year_columns) != {cy_year, cy_year - 1}:
                raise ValueError(f"Unexpected reporting years on {name}")
            columns[name] = {
                "cy": year_columns[cy_year],
                "py": year_columns[cy_year - 1],
            }

            for column in columns[name].values():
                unit_text = str(
                    sheet.cell(header_rows[name] + 1, column).value
                ).strip()
                if unit_text != "£'000":
                    raise ValueError(f"Unexpected currency or unit on {name}")

        prior_day = min(
            closing_date.day,
            calendar.monthrange(closing_date.year - 1, closing_date.month)[1],
        )
        prior_date = closing_date.replace(
            year=closing_date.year - 1, day=prior_day
        )

        result = {
            "entity_code": "TW-UK02",
            "source_entity_name": source_name,
            "currency": "GBP",
            "unit": 1000,
            "period_end": {
                "cy": closing_date.isoformat(),
                "py": prior_date.isoformat(),
            },
            "fields": {
                key: {
                    year: {"value": 0, "cells": []}
                    for year in ("cy", "py")
                }
                for key in FIELD_KEYS
            },
            "source_totals": [],
            "ignored": [],
            "source_exceptions": [],
            "notes": (
                "The workbook's abbreviated entity name identifies Tarnwell "
                "Engineering Limited (TW-UK02). Amounts remain in GBP thousands. "
                "Source expenses and dividends are reversed from negative "
                "presentation; other operating income retains its signed value. "
                "The deferred tax asset is reclassified from local debtors to "
                "non-current assets; local subtotals retain their original scope. "
                "Overdrafts are current borrowings, while cash follows the "
                "workbook's cash-flow definition. Dividends paid are used as "
                "the year's dividends because they are the distribution shown "
                "in the retained earnings movement. Unpresented input lines "
                "are zero."
            ),
        }

        accounted = {name: set() for name in sheet_names}

        def find_row(sheet_name, label, section=None):
            sheet = sheets[sheet_name]
            start, stop = 1, sheet.max_row + 1
            if section is not None:
                start = find_row(sheet_name, section) + 1
                # The repeated bank-loan label is bounded by its next section.
                next_sections = {
                    "Creditors: amounts falling due within one year":
                        "Total creditors due within one year",
                    "Creditors: amounts falling due after more than one year":
                        "Provisions for liabilities",
                }
                stop = find_row(sheet_name, next_sections[section])
            matches = [
                row for row in range(start, stop)
                if _normalise(sheet.cell(row, 2).value) == _normalise(label)
            ]
            if len(matches) != 1:
                raise ValueError(
                    f"Expected one row for {sheet_name}: {label!r}"
                )
            return matches[0]

        def reference(sheet_name, label, year, factor=1, section=None):
            row = find_row(sheet_name, label, section)
            cell = sheets[sheet_name].cell(row, columns[sheet_name][year])
            if not _is_number(cell.value):
                raise ValueError(
                    f"Expected numeric amount at {sheet_name}!{cell.coordinate}"
                )
            accounted[sheet_name].add(cell.coordinate)
            return {
                "sheet": sheet_name,
                "cell": cell.coordinate,
                "factor": factor,
            }

        def amount(ref):
            return ref["factor"] * sheets[ref["sheet"]][ref["cell"]].value

        def add_field(key, sheet_name, label, factor=1, section=None):
            for year in ("cy", "py"):
                ref = reference(sheet_name, label, year, factor, section)
                line = result["fields"][key][year]
                line["cells"].append(ref)
                line["value"] += amount(ref)

        mappings = (
            ("revenue", "P&L", "Turnover", 1),
            ("cost_of_sales", "P&L", "Cost of sales", -1),
            ("distribution_costs", "P&L", "Distribution costs", -1),
            ("administrative_expenses", "P&L", "Administrative expenses", -1),
            ("other_operating_income", "P&L", "Other operating income", 1),
            ("finance_income", "P&L",
             "Interest receivable and similar income", 1),
            ("finance_costs", "P&L",
             "Interest payable and similar expenses", -1),
            ("income_tax_expense", "P&L", "Tax on profit", -1),
            ("property_plant_equipment", "BS", "Tangible assets", 1),
            ("intangible_assets", "BS", "Intangible assets", 1),
            ("deferred_tax_assets", "BS", "Deferred tax asset", 1),
            ("inventories", "BS", "Stocks", 1),
            ("trade_receivables", "BS", "Trade debtors", 1),
            ("other_current_assets", "BS", "Prepayments and accrued income", 1),
            ("other_current_assets", "BS", "Other debtors", 1),
            ("cash_and_cash_equivalents", "BS", "Cash at bank and in hand", 1),
            ("trade_payables", "BS", "Trade creditors", 1),
            ("other_current_liabilities", "BS",
             "Accruals and deferred income", 1),
            ("other_current_liabilities", "BS",
             "Other taxation and social security", 1),
            ("current_borrowings", "BS", "Bank overdraft", 1),
            ("current_tax_payable", "BS", "Corporation tax", 1),
            ("provisions", "BS", "Other provisions", 1),
            ("deferred_tax_liabilities", "BS", "Deferred taxation", 1),
            ("share_capital", "BS", "Called up share capital", 1),
            ("share_capital", "BS", "Share premium account", 1),
            ("retained_earnings", "BS", "Profit and loss account", 1),
            ("other_reserves", "BS", "Other reserves", 1),
            ("net_cash_operating", "CF",
             "Net cash from operating activities", 1),
            ("net_cash_investing", "CF",
             "Net cash from investing activities", 1),
            ("net_cash_financing", "CF",
             "Net cash from financing activities", 1),
            ("fx_effect_on_cash", "CF",
             "Effect of foreign exchange rate changes", 1),
            ("cash_at_beginning", "CF", "Cash at beginning of year", 1),
            ("dividends", "Reserves", "Dividends paid on equity shares", -1),
        )
        for mapping in mappings:
            add_field(*mapping)

        add_field(
            "current_borrowings", "BS", "Bank loans",
            section="Creditors: amounts falling due within one year",
        )
        add_field(
            "non_current_borrowings", "BS", "Bank loans",
            section="Creditors: amounts falling due after more than one year",
        )

        def add_total(sheet_name, label, formula):
            for year in ("cy", "py"):
                ref = reference(sheet_name, label, year)
                result["source_totals"].append({
                    **ref, "year": year, "formula": formula,
                })

        totals = (
            ("P&L", "Gross profit", "gross_profit"),
            ("P&L", "Operating profit", "operating_profit"),
            ("P&L", "Profit before taxation", "profit_before_tax"),
            ("P&L", "Profit for the financial year", "profit_for_year"),
            ("BS", "Total fixed assets",
             "property_plant_equipment + intangible_assets"),
            ("BS", "Total debtors",
             "trade_receivables + other_current_assets + deferred_tax_assets"),
            ("BS", "Total current assets",
             "total_current_assets + deferred_tax_assets"),
            ("BS", "Total creditors due within one year",
             "total_current_liabilities"),
            ("BS", "Net current assets",
             "total_current_assets + deferred_tax_assets"
             " - total_current_liabilities"),
            ("BS", "Total assets less current liabilities",
             "total_assets - total_current_liabilities"),
            ("BS", "Net assets", "total_assets - total_liabilities"),
            ("BS", "Shareholders' funds", "total_equity"),
            ("CF", "Net increase/(decrease) in cash", "net_change_in_cash"),
            ("CF", "Cash at end of year", "cash_at_end"),
            ("Reserves", "At beginning of year",
             "retained_earnings - profit_for_year + dividends"),
            ("Reserves", "Profit for the financial year", "profit_for_year"),
            ("Reserves", "At end of year", "retained_earnings"),
        )
        for total in totals:
            add_total(*total)

        # Note references are deliberately excluded from financial amounts.
        for name, sheet in sheets.items():
            for row in sheet.iter_rows():
                for cell in row:
                    if cell.column == 3 and _is_number(cell.value):
                        result["ignored"].append({
                            "sheet": name,
                            "cell": cell.coordinate,
                            "reason": "Financial statement note reference.",
                        })

        # Reject unexpected amounts rather than silently discard new content.
        for name, sheet in sheets.items():
            for column in columns[name].values():
                used_rows = [
                    sheet[coordinate].row
                    for coordinate in accounted[name]
                    if sheet[coordinate].column == column
                ]
                first_used = min(used_rows)
                for row in range(first_used, sheet.max_row + 1):
                    cell = sheet.cell(row, column)
                    if (
                        _is_number(cell.value)
                        and cell.coordinate not in accounted[name]
                    ):
                        raise ValueError(
                            f"Unmapped numeric amount: {name}!{cell.coordinate}"
                        )

        values = {}
        for year in ("cy", "py"):
            values[year] = {
                key: result["fields"][key][year]["value"]
                for key in FIELD_KEYS
            }
            for key, expression in TOTAL_FORMULAS.items():
                values[year][key] = _expression_value(
                    expression, values[year]
                )

        def check(check_id, actual, expected, explanation):
            difference = actual - expected
            if not math.isclose(difference, 0, rel_tol=0, abs_tol=1e-8):
                result["source_exceptions"].append({
                    "check": check_id,
                    "difference": difference,
                    "explanation": explanation,
                })

        for total in result["source_totals"]:
            check(
                f"T {total['sheet']}!{total['cell']} {total['year']}",
                amount(total),
                _expression_value(total["formula"], values[total["year"]]),
                "The printed source amount minus its mapped expression "
                "is nonzero; all source amounts have been preserved.",
            )

        for year in ("cy", "py"):
            v = values[year]
            check(
                f"R1 balance {year}",
                v["total_assets"],
                v["total_equity_and_liabilities"],
                "Source assets minus equity and liabilities is nonzero.",
            )
            check(
                f"R2 cash_at_end {year}",
                v["cash_at_end"],
                v["cash_and_cash_equivalents"],
                "Calculated cash-flow closing cash minus balance-sheet "
                "cash is nonzero.",
            )

        check(
            "R3 opening cash cy",
            values["cy"]["cash_at_beginning"],
            values["py"]["cash_and_cash_equivalents"],
            "Current-year opening cash minus prior-year balance-sheet "
            "cash is nonzero.",
        )
        check(
            "R4 retained earnings cy",
            values["cy"]["retained_earnings"],
            values["py"]["retained_earnings"]
            + values["cy"]["profit_for_year"]
            - values["cy"]["dividends"],
            "Closing retained earnings minus prior retained earnings "
            "plus current profit less dividends is nonzero.",
        )

        return result
    finally:
        workbook.close()
